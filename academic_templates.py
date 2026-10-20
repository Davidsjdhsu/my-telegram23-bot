"""Neutral university-derived covers and deterministic academic Word layout.

Provenance and adaptation boundaries are recorded in assets/academic/sources.json.
The profiles are not a claim of approval by every university.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ASSETS = Path(__file__).resolve().parent / 'assets' / 'academic'
KINDS = frozenset(('referat', 'coursework', 'report', 'essay', 'notes'))
COVER_LABELS = {'referat': 'РЕФЕРАТ', 'coursework': 'КУРСОВАЯ РАБОТА',
                'report': 'ДОКЛАД', 'essay': 'ЭССЕ', 'notes': 'КОНСПЕКТ'}
BLANKS = {
    'referat': 'Реферат — титульный лист',
    'coursework': 'Курсовая — титульный лист',
    'course_project': 'Курсовой проект — титульный лист',
    'lab': 'Лабораторная — титульный лист',
    'practice': 'Практика — титульный лист',
    'bachelor': 'ВКР бакалавра — титульный лист',
    'diploma': 'Дипломная — титульный лист',
    'master': 'Магистерская — титульный лист',
    'assignment': 'Задание на ВКР',
}
NOTICE = ('Оформление адаптировано по материалам ТГУ и МПСУ, без чужих логотипов. '
          'Это базовый вариант: требования вашей кафедры могут отличаться.')
META_LABELS = {'org':'Ведомство', 'school':'Учебное заведение', 'department':'Факультет или кафедра',
               'discipline':'Дисциплина', 'author':'ФИО студента', 'group':'Группа',
               'teacher':'Руководитель', 'teacher_position':'Должность руководителя',
               'specialty':'Направление подготовки', 'profile':'Профиль', 'city':'Город', 'year':'Год'}
ALIASES = {'вуз':'school','университет':'school','учебное заведение':'school','фио':'author',
           'автор':'author','студент':'author','фио студента':'author','преподаватель':'teacher',
           'фио преподавателя':'teacher','фио научного руководителя':'teacher','научный руководитель':'teacher',
           'факультет':'department','кафедра':'department','специальность':'specialty'}

def blank_path(kind):
    if kind not in BLANKS:
        raise ValueError('unknown_academic_blank')
    return ASSETS / (kind + '.docx')


def explicit_meta(meta, source):
    """Explicit labels win over generated metadata, including deliberate clearing."""
    result = dict(meta or {})
    names = {v.casefold(): k for k, v in META_LABELS.items()}
    names.update(ALIASES)
    for line in source.splitlines():
        label, sep, value = line.partition(':')
        key = names.get(' '.join(label.casefold().strip().split())) if sep else None
        if key:
            result[key] = value.strip()
    return result


def guidance(kind):
    if kind not in KINDS:
        return ''
    common = ('\nУчебная структура: титульный лист, оглавление и номера страниц создаёт программа. '
              'НЕ включай титульный лист и оглавление в sections. Заголовки разделов должны описывать их содержание. '
              'Данные в meta бери только из запроса. Не придумывай вуз, ФИО, группу, город, год, '
              'подписи, результаты экспериментов, источники или цитаты. Не требуй определённого числа источников: '
              'в список включай только источники, предоставленные пользователем или проверенные по доступному тексту. '
              'При отсутствии источников не создавай фиктивную библиографию. ')
    return common + {
        'referat':'Разделы: Введение (цель и задачи), 2–3 тематических раздела, Заключение, источники при наличии. Главы обозначай 1, 2, подразделы 1.1, 1.2. Введение и заключение не нумеруй.',
        'coursework':'Разделы: Введение (проблема, цель, задачи, объект и предмет), теоретическая глава, практическая/аналитическая глава, Заключение, источники при наличии. Главы 1, 2; подразделы 1.1, 1.2. Никаких вымышленных результатов исследования; если данных нет, опиши методику предполагаемого исследования.',
        'report':'Доклад: введение с вопросом, основная часть с объяснением и примерами, выводы, источники при наличии. Не превращай его в курсовую.',
        'essay':'Эссе: постановка проблемы и тезис, аргументы и примеры, рассмотрение возражения, собственный вывод. Титульный лист с надписью ЭССЕ добавляет программа. Без оглавления; основной текст связный, с абзацами, без видимых заголовков разделов. sections служит только для внутреннего плана. Не приписывай автору личный опыт, которого нет в запросе. Не начинай с фразы «Этот текст — учебное эссе»: сразу переходи к вопросу темы.',
        'notes':'Конспект: тематические пункты, определения, связи и краткие выводы. Сохрани содержание исходного материала. Титульный лист с надписью КОНСПЕКТ добавляет программа. Без оглавления и обязательного введения.',
    }[kind]


def _text(p, text):
    # Preserve the paragraph and first run formatting from the reference.
    nodes = list(p._p.iter(qn('w:t')))
    if nodes:
        nodes[0].text = str(text)
        nodes[0].set(qn('xml:space'), 'preserve')
        for n in nodes[1:]: n.text = ''
    else:
        p.add_run(str(text))


def _cover(kind, title, meta):
    d = Document(blank_path('coursework' if kind == 'coursework' else 'referat'))
    _ensure_styles(d)
    teacher_seen = False
    for i,p in enumerate(d.paragraphs):
        t = p.text.strip()
        v = None
        if i == 0: v = meta.get('org','')
        elif i == 2: v = meta.get('school','')
        elif i == 3: v = meta.get('department','')
        elif t in ('ТЕМА РАБОТЫ','ТЕМА РЕФЕРАТА'): v = title
        elif t == 'РЕФЕРАТ': v = COVER_LABELS[kind]
        elif t.startswith('по дисциплине'): v = ('по дисциплине «'+meta['discipline']+'»') if meta.get('discipline') else ''
        elif t == 'Фамилия Имя Отчество обучающегося': v = meta.get('author','')
        elif t.startswith('Направление подготовки'): v = ('Направление подготовки: '+meta['specialty']) if meta.get('specialty') else ''
        elif t.startswith('Направленность'): v = ('Профиль: '+meta['profile']) if meta.get('profile') else ''
        elif t == 'ученая степень, звание': v = meta.get('teacher_position','')
        elif t == 'Руководитель работы':
            teacher_seen = True
            v = t if meta.get('teacher') or meta.get('teacher_position') else ''
        elif t == 'Автор работы':
            teacher_seen = False
            v = t if meta.get('author') or meta.get('group') else ''
        elif 'И.О. Фамилия' in t: v = meta.get('teacher' if teacher_seen else 'author','')
        elif t.startswith('студент группы'): v = ('Группа '+meta['group']) if meta.get('group') else ''
        elif 'подпись' in t or '20 ___' in t: v = ''
        elif t == '[Город] — [Год]': v = ' — '.join(str(meta[k]) for k in ('city','year') if meta.get(k))
        if v is not None: _text(p,v)
        # Reserved white space absorbs multiline names without shrinking user text.
        if not p.text.strip():
            p.paragraph_format.line_spacing = Pt(6)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs: r.font.size = Pt(6)
        else:
            p.paragraph_format.keep_together = True
        if v == title:
            p.style = d.styles['Title']
            for r in p.runs:
                r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.bold = True
    return d


def _field(p, instruction, cached=''):
    for typ, text in [('begin',None),('instr',instruction),('separate',None),('text',cached),('end',None)]:
        r = p.add_run()
        if typ in ('begin','separate','end'):
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),typ)
        else:
            e = OxmlElement('w:instrText' if typ == 'instr' else 'w:t'); e.text=text
            e.set(qn('xml:space'),'preserve')
        r._r.append(e)


def _ensure_styles(d):
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'TOC 1', 'TOC 2'):
        if name not in d.styles:
            st = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if name != 'Normal': st.base_style = d.styles['Normal']


def _styles(d, covered):
    _ensure_styles(d)
    s=d.sections[0]
    s.page_width,s.page_height=Cm(21),Cm(29.7)
    s.left_margin,s.right_margin=Cm(3),Cm(1.5)
    s.top_margin=s.bottom_margin=Cm(2)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','TOC 1','TOC 2'):
        if name not in d.styles: continue
        st=d.styles[name]; st.font.name='Times New Roman';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
        for attr in list(fonts.attrib):
            if attr.rsplit('}',1)[-1].lower().endswith('theme'): del fonts.attrib[attr]
        for attr in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+attr),'Times New Roman')
        st.font.italic=False
        for e in list(st.element.iter(qn('w:pBdr'))): e.getparent().remove(e)
        pf=st.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.5
        pf.first_line_indent=Cm(0 if name == 'Title' else 1.25);pf.widow_control=True
        pf.left_indent=Cm(0);pf.right_indent=Cm(0)
    for name in ('Heading 1','Heading 2','Heading 3'):
        st=d.styles[name];st.font.bold=True;st.paragraph_format.keep_with_next=True
        props=st.element.get_or_add_pPr()
        for node in list(props.findall(qn('w:outlineLvl'))):props.remove(node)
        level=OxmlElement('w:outlineLvl');level.set(qn('w:val'),str(int(name[-1])-1));props.append(level)
        st.paragraph_format.space_after=Pt(10)
    s.different_first_page_header_footer=covered
    fp=s.first_page_footer.paragraphs[0];fp.clear()
    p=s.footer.paragraphs[0];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    _field(p,' PAGE ')
    flag=OxmlElement('w:updateFields');flag.set(qn('w:val'),'true');d.settings.element.append(flag)


def _level(title):
    return 2 if re.match(r'^\d+\.\d+',title) else 1


def front_matter_pages(kind):
    return 2 if kind in ('referat', 'coursework') else 1 if kind in KINDS else 0


def volume_label(kind, pages):
    """Selected pages always mean body pages; front matter is additional and free."""
    prefix = 'Титульный + оглавление + ' if kind in ('referat', 'coursework') else 'Титульный + '
    return prefix + f'{pages} стр. текста'


def volume_hint(kind):
    extra = ' и оглавление' if kind in ('referat', 'coursework') else ''
    return (f'Выбери число страниц основного текста. Титульный лист{extra} добавляются отдельно, '
            'без доплаты. Данные для титульного можно указать вместе с темой: '
            'учебное заведение, дисциплина, ФИО, группа, преподаватель, город и год.')


def delivery_notice(kind, requested, actual, total, credits):
    return (f'В документе: {volume_label(kind, actual)}. Всего в файле {total} стр.\n'
            + (f'Было выбрано {requested} стр. текста; фактически получилось {actual}.\n'
               if actual != requested else '')
            + f'Стоимость — {credits} кредитов. Титульный лист и оглавление без доплаты; '
              'дополнительные страницы текста не увеличивают цену.')


def page_guidance(kind, pages):
    front = front_matter_pages(kind)
    body = max(1, pages)
    words_per_page = 250 if kind == 'essay' else 200
    compact = kind == 'referat' and pages <= 10
    return (
        f'\nВыбранный объём — {body} страниц основного текста. Дополнительно программа создаст '
        f'{front} служебных стр. (титульный лист и оглавление по жанру); они НЕ входят в выбранный объём. '
        f'Ожидаемый общий объём — {body + front} стр. На основной текст — {body} страниц, '
        f'ориентир {body * words_per_page} слов, включая таблицы, без реквизитов и оглавления. '
        'Введение и выводы краткие; около 80% текста отведи объяснениям, условиям и примерам. '
        'Не наращивай объём повторениями, общими фразами и описанием границ задания. '
        + ('Это короткий реферат: разделы размещаются последовательно, без обязательного '
           'нового листа для каждого. Сохрани все заданные пользователем темы. ' if compact else '')
    )


def inspect_layout(path, kind):
    from word_quality import word_pdf_layout
    return word_pdf_layout(path, front_matter_pages=front_matter_pages(kind))


def render(path, title, sections, kind, meta=None, add_table=None, requested_pages=None):
    if kind not in KINDS: raise ValueError('unknown_academic_kind')
    meta={k:str(v) for k,v in (meta or {}).items() if isinstance(v,(str,int)) and v}
    covered=kind in KINDS
    d=_cover(kind,title,meta) if covered else Document()
    _styles(d,covered)
    blocks=[dict(b) for b in sections or [] if str(b.get('title','')).strip().casefold() not in
            ('содержание','оглавление','титульный лист','contents','table of contents')]
    if covered:
        d.add_page_break()
    else:
        p=d.add_paragraph(title,style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(12)
        for k in ('school','discipline','author','group'):
            if meta.get(k):
                p=d.add_paragraph(meta[k]);p.paragraph_format.first_line_indent=Cm(0)
    toc=kind in ('referat','coursework') and bool(blocks)
    compact = kind == 'referat' and requested_pages is not None and requested_pages <= 10
    if toc:
        p=d.add_paragraph('ОГЛАВЛЕНИЕ');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(14)
        p.runs[0].bold=True
        for i,b in enumerate(blocks):
            head=str(b.get('title') or '').strip()
            if not head: continue
            p=d.add_paragraph(head+'\t');p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.left_indent=Cm(.6 if _level(head)==2 else 0)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(16.2),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            _field(p,' PAGEREF study_'+str(i)+' \\h ')
    for i,b in enumerate(blocks):
        head=str(b.get('title') or '').strip()
        if head and kind != 'essay':
            level=_level(head)
            p=d.add_paragraph(head,style='Heading '+str(level))
            p.paragraph_format.page_break_before=(toc and level==1 and not compact) or (toc and i==0)
            if compact and i:
                p.paragraph_format.space_before=Pt(12)
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(1000+i));start.set(qn('w:name'),'study_'+str(i))
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(1000+i))
            p._p.insert(0,start);p._p.append(end)
        for line in str(b.get('content') or '').splitlines():
            if not line.strip():continue
            p=d.add_paragraph(line.strip());p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if re.match(r'^(?:[-•*]|\d+[.)])\s',line):
                p.paragraph_format.first_line_indent=Cm(0)
        table=b.get('table')
        if table and add_table:
            add_table(d,table)
            for t in d.tables:
                for row in t.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.line_spacing=1
                            p.paragraph_format.first_line_indent=Cm(0)
                            for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
    d.core_properties.author='';d.core_properties.last_modified_by=''
    d.save(path)


def cache_page_references(docx_path, pdf_path):
    """Cache only academic PAGEREF results from the real export's heading outline.

    Keep every other DOCX package part byte-for-byte; do not save through Writer.
    Word still updates these fields when the document is edited later.
    """
    from zipfile import ZipFile, ZIP_DEFLATED
    from lxml import etree
    from pypdf import PdfReader
    import os
    import tempfile
    with ZipFile(docx_path) as z:
        xml=z.read('word/document.xml')
        if b'PAGEREF study_' not in xml:return
        parts={n:z.read(n) for n in z.namelist()}
    root=etree.fromstring(xml)
    reader=PdfReader(pdf_path)
    entries=[]
    def walk(items):
        for item in items:
            if isinstance(item,list):walk(item)
            elif isinstance(item,dict) and item.get('/Title'):
                entries.append((str(item['/Title']),reader.get_destination_page_number(item)+1))
    walk(reader.outline)
    pages={}; cursor=0
    normalize=lambda value:' '.join(value.split())
    for mark in root.iter(qn('w:bookmarkStart')):
        name=mark.get(qn('w:name'),'')
        if not name.startswith('study_'):continue
        head=''.join(mark.getparent().itertext())
        head=''.join(n.text or '' for n in mark.getparent().iter(qn('w:t')))
        for index in range(cursor,len(entries)):
            if normalize(entries[index][0])==normalize(head):
                pages[name]=entries[index][1];cursor=index+1;break
        else:raise ValueError('academic_heading_missing_from_pdf_outline')
    for instruction in root.iter(qn('w:instrText')):
        m=re.fullmatch(r'\s*PAGEREF (study_\d+)\s+\\h\s*',instruction.text or '')
        if not m:continue
        if m[1] not in pages:raise ValueError('academic_page_reference_missing')
        run=instruction.getparent().getnext();separated=False
        while run is not None:
            field=run.find(qn('w:fldChar'))
            if field is not None:
                if field.get(qn('w:fldCharType'))=='separate':separated=True
                elif field.get(qn('w:fldCharType'))=='end':break
            elif separated:
                node=run.find(qn('w:t'))
                if node is not None:
                    node.text=str(pages[m[1]]);break
            run=run.getnext()
        else:raise ValueError('academic_field_result_missing')
    parts['word/document.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    fd,temp=tempfile.mkstemp(suffix='.docx',dir=Path(docx_path).resolve().parent)
    os.close(fd)
    try:
        with ZipFile(temp,'w',ZIP_DEFLATED)as z:
            for name,content in parts.items():z.writestr(name,content)
        os.replace(temp,docx_path)
    finally:
        if os.path.exists(temp):os.remove(temp)
