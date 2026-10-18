"""Convert the saved presentation, without rebuilding its text or layout.

One LibreOffice process at a time per bot process. Each job has an isolated
profile; timeout/cancellation kills its process group before releasing the slot.
No documents are uploaded to an external conversion service.
"""
import asyncio
import os
from pathlib import Path
import shutil
import signal
import tempfile
import weakref

from pptx import Presentation
from pypdf import PdfReader


class PresentationPDFError(RuntimeError):
    pass


_slots = weakref.WeakKeyDictionary()


def ensure_pdf_converter():
    configured = os.environ.get('PRESENTATION_SOFFICE')
    executable = shutil.which(configured or 'soffice')
    if not executable:
        raise PresentationPDFError('converter_unavailable')
    return executable


def validate_pdf(pdf_path, pptx_path, expected_slides):
    """Reject wrong counts, aspect ratios, blank output and lost slide pictures."""
    if not pdf_path.is_file() or not 1000 <= pdf_path.stat().st_size <= 49_000_000:
        raise PresentationPDFError('invalid_pdf_size')
    try:
        source = Presentation(pptx_path)
        result = PdfReader(pdf_path)
        if result.is_encrypted or len(source.slides) != expected_slides or len(result.pages) != expected_slides:
            raise PresentationPDFError('page_count_mismatch')
        width, height = source.slide_width / 12700, source.slide_height / 12700
        for slide, page in zip(source.slides, result.pages):
            if abs(float(page.mediabox.width) - width) > 2 or abs(float(page.mediabox.height) - height) > 2:
                raise PresentationPDFError('page_geometry_mismatch')
            stream = page.get_contents()
            if stream is None or not any(op in (b'Tj', b'TJ', b"'", b'"', b'Do', b'f', b'f*', b'S', b's', b'B', b'B*', b'b', b'b*', b'sh')
                                         for _, op in stream.operations):
                raise PresentationPDFError('empty_pdf_page')
            # LibreOffice retains photographs as image XObjects (also inside forms).
            if any(s.shape_type == 13 for s in slide.shapes) and not list(page.images):
                raise PresentationPDFError('missing_pdf_images')
    except PresentationPDFError:
        raise
    except Exception as exc:
        raise PresentationPDFError('unreadable_pdf') from exc


async def _stop_process(process):
    if process.returncode is None:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
    await process.wait()


async def convert_presentation_pdf(pptx_path, pdf_path, expected_slides, *, timeout=120, queue_timeout=240):
    return await _convert_office_pdf(pptx_path, pdf_path, expected_slides,
        source_suffix='.pptx', export_filter='pdf:impress_pdf_Export', validator=validate_pdf,
        timeout=timeout, queue_timeout=queue_timeout)


async def _convert_office_pdf(pptx_path, pdf_path, expected_slides, *, source_suffix,
                              export_filter, validator, timeout=120, queue_timeout=240):
    executable = ensure_pdf_converter()
    source, target = Path(pptx_path).resolve(), Path(pdf_path).resolve()
    if not source.is_file() or source.suffix.lower() != source_suffix or target.suffix.lower() != '.pdf':
        raise PresentationPDFError('invalid_conversion_paths')
    slot = _slots.setdefault(asyncio.get_running_loop(), asyncio.Semaphore(1))
    try:
        await asyncio.wait_for(slot.acquire(), timeout=queue_timeout)
    except asyncio.TimeoutError as exc:
        raise PresentationPDFError('converter_busy') from exc
    try:
        # Keep output beside the final destination so the successful replace is atomic.
        with tempfile.TemporaryDirectory(prefix='pdf-convert-', dir=target.parent) as td:
            work = Path(td)
            profile = work / 'profile'
            export = work / 'export'
            export.mkdir()
            process = None
            env = dict(os.environ, SAL_USE_VCLPLUGIN='svp', OMP_NUM_THREADS='1')
            try:
                process = await asyncio.create_subprocess_exec(
                    executable, '-env:UserInstallation=' + profile.as_uri(),
                    '--headless', '--nologo', '--nodefault', '--nofirststartwizard',
                    '--convert-to', export_filter, '--outdir', str(export), str(source),
                    stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL,
                    start_new_session=True, env=env,
                )
                await asyncio.wait_for(process.wait(), timeout=timeout)
                if process.returncode != 0:
                    raise PresentationPDFError('converter_exit_' + str(process.returncode))
                generated = export / (source.stem + '.pdf')
                await asyncio.to_thread(validator, generated, source, expected_slides)
                os.replace(generated, target)
            except asyncio.TimeoutError as exc:
                raise PresentationPDFError('conversion_timeout') from exc
            except OSError as exc:
                raise PresentationPDFError('conversion_io_error') from exc
            finally:
                if process is not None:
                    # Cancellation must not leave a renderer using memory after the job ends.
                    await asyncio.shield(_stop_process(process))
    finally:
        slot.release()
    return str(target)


def pdf_failure_notice(lang):
    notices = {
        'ru': 'Не удалось подготовить PDF со слайдами. Файлы не отправлены, баллы не списаны. Тема и настройки сохранены — можно повторить сборку.',
        'en': 'The slide PDF could not be prepared. No files were sent or credits charged. Your topic and settings are saved; you can retry.',
        'de': 'Das Folien-PDF konnte nicht erstellt werden. Keine Dateien versendet und keine Punkte abgebucht. Thema und Einstellungen bleiben gespeichert.',
        'fr': 'Le PDF des diapositives n’a pas pu être créé. Aucun fichier envoyé ni crédit débité. Le sujet et les réglages sont conservés.',
        'es': 'No se pudo preparar el PDF de las diapositivas. No se enviaron archivos ni se cobraron créditos. El tema y los ajustes están guardados.',
        'zh': '未能生成幻灯片 PDF。未发送文件，也未扣除积分。主题和设置已保存，可以重试。',
        'ar': 'تعذر إعداد ملف PDF للشرائح. لم تُرسل ملفات ولم تُخصم نقاط. تم حفظ الموضوع والإعدادات ويمكن إعادة المحاولة.',
    }
    return notices.get(lang, notices['en'])
