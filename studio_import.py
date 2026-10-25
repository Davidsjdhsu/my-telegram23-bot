"""Bounded PPTX -> editable constructor copy. No fetching, macros or AI calls."""
import base64
import io
import zipfile
from collections import Counter
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
from constructor_pricing import quote
from studio_export import FONTS

MAX_UPLOAD = 10_000_000
NS = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}


def preflight(raw):
    if len(raw) > MAX_UPLOAD:
        raise ValueError('Выберите PPTX до 10 МБ.')
    try:
        z = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise ValueError('Это не файл PPTX. Выберите презентацию PowerPoint.') from None
    with z:
        entries=z.infolist()
        names=[x.filename for x in entries]
        if len(entries)>2000 or len(names)!=len(set(names)) or sum(x.file_size for x in entries)>64_000_000:
            raise ValueError('В презентации слишком много данных для конструктора.')
        if 'ppt/presentation.xml' not in names or '[Content_Types].xml' not in names:
            raise ValueError('В файле нет презентации PPTX.')
        external=False
        for info in entries:
            name=info.filename
            if '..' in PurePosixPath(name).parts or name.startswith('/') or '\\' in name or info.flag_bits&1:
                raise ValueError('Файл содержит неподдерживаемые данные.')
            if any(x in name.lower() for x in ('vbaproject','/embeddings/','activex')):
                raise ValueError('Файлы с макросами или встроенными программами не поддерживаются.')
            if info.file_size>12_000_000:
                raise ValueError('Один из элементов презентации слишком большой.')
            if name.endswith(('.xml','.rels')):
                data=z.read(info)
                if len(data)>4_000_000 or b'<!DOCTYPE' in data.replace(b'\x00',b'').upper() or b'<!ENTITY' in data.replace(b'\x00',b'').upper():
                    raise ValueError('Файл содержит неподдерживаемую структуру XML.')
                if name.endswith('.rels'):
                    root=ET.fromstring(data)
                    for rel in root:
                        if rel.get('TargetMode')=='External':
                            if not rel.get('Type','').endswith('/hyperlink'):
                                raise ValueError('Связанные внешние файлы не поддерживаются. Вставьте изображения внутрь PPTX.')
                            external=True
        return external


def theme_colors(slide):
    try:
        master=slide.slide_layout.slide_master
        part=next(r.target_part for r in master.part.rels.values() if r.reltype.endswith('/theme'))
        scheme=ET.fromstring(part.blob).find('.//a:clrScheme',NS)
        result={}
        for entry in scheme:
            node=next(iter(entry));result[entry.tag.split('}')[-1]]=node.get('val') if node.tag.endswith('srgbClr') else node.get('lastClr')
        # Standard aliases; individual color maps can remap these.
        result.update(tx1=result.get('dk1','000000'),tx2=result.get('dk2','000000'),bg1=result.get('lt1','FFFFFF'),bg2=result.get('lt2','FFFFFF'))
        return result
    except (AttributeError, KeyError, StopIteration, TypeError):
        return {}


def color_node(node, colors, default):
    if node is None:return default
    for child in node:
        kind=child.tag.split('}')[-1]
        if kind=='srgbClr':value=child.get('val')
        elif kind=='schemeClr':value=colors.get(child.get('val'))
        elif kind=='sysClr':value=child.get('lastClr')
        else:continue
        if not value or len(value)!=6:return default
        rgb=[int(value[i:i+2],16)/255 for i in (0,2,4)]
        for transform in child:
            v=int(transform.get('val','100000'))/100000
            k=transform.tag.split('}')[-1]
            if k=='lumMod':rgb=[x*v for x in rgb]
            elif k=='lumOff':rgb=[x+v for x in rgb]
            elif k=='tint':rgb=[x+(1-x)*v for x in rgb]
            elif k=='shade':rgb=[x*v for x in rgb]
        return '#'+''.join(f'{round(min(1,max(0,x))*255):02X}' for x in rgb)
    return default


def background(slide, colors, warnings):
    fill=slide.background.fill
    b=dict(type='solid',color='#FFFFFF',second='#FFFFFF',angle=0)
    if fill.type==MSO_FILL_TYPE.SOLID:
        b['color']=color_node(fill._xPr.find('{'+NS['a']+'}solidFill'),colors,'#FFFFFF')
    elif fill.type==MSO_FILL_TYPE.GRADIENT:
        stops=list(fill.gradient_stops)
        b.update(type='gradient',color=color_node(stops[0]._gs,colors,'#FFFFFF'),second=color_node(stops[-1]._gs,colors,'#FFFFFF'),angle=fill.gradient_angle or 0)
        if len(stops)>2:warnings.append('Градиент упрощён до двух крайних цветов.')
    elif fill.type not in (None,MSO_FILL_TYPE.BACKGROUND):
        warnings.append('Сложный фон заменён белым. Выберите фон в редакторе.')
    else:
        # Inherited master artwork is not part of the slide's shapes.
        for owner in (slide.slide_layout,slide.slide_layout.slide_master):
            bg=owner._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}bgPr')
            if bg is not None:
                b['color']=color_node(bg.find('{'+NS['a']+'}solidFill'),colors,'#FFFFFF');break
    if b['type']=='solid':b['second']=b['color']
    return b


def import_pptx(raw, name='Презентация'):
    try:
        external=preflight(raw)
        p=Presentation(io.BytesIO(raw))
    except (zipfile.BadZipFile, KeyError, ET.ParseError):
        raise ValueError('Не удалось прочитать PPTX. Откройте и сохраните его в PowerPoint заново.') from None
    if not 1<=len(p.slides)<=20:raise ValueError('В конструкторе можно открыть от 1 до 20 слайдов.')
    if not p.slide_width or not p.slide_height:raise ValueError('Не удалось определить размер слайдов.')
    scale=min(1200/p.slide_width,675/p.slide_height)
    ox=(1200-p.slide_width*scale)/2;oy=(675-p.slide_height*scale)/2
    deck=dict(version=1,name=str(name).rsplit('.',1)[0][:80] or 'Презентация',slides=[])
    reports=[];warnings=[];total_pixels=0
    if external:warnings.append('Гиперссылки не переносятся в редактируемую копию.')
    if ox>1 or oy>1:warnings.append('Слайды вписаны в формат 16:9 с полями, без растягивания.')
    for index,slide in enumerate(p.slides,1):
        notes=[];unsupported=Counter();colors=theme_colors(slide)
        target=dict(id=f's{index}',background=background(slide,colors,notes),elements=[])
        if slide._element.xpath('.//p:timing|.//p:transition'):
            notes.append('Анимация и переходы не переносятся.')
        for owner in (slide.slide_layout,slide.slide_layout.slide_master):
            if any(not s.is_placeholder for s in owner.shapes):
                notes.append('Элементы образца слайдов не переносятся.');break
        for j,shape in enumerate(slide.shapes):
            box=dict(x=ox+shape.left*scale,y=oy+shape.top*scale,w=shape.width*scale,h=shape.height*scale,
                     id=f's{index}e{j}',rotation=((shape.rotation+180)%360)-180,opacity=1)
            if box['w']<=0 or box['h']<=0:unsupported['нулевые элементы']+=1;continue
            if box['x']<0 or box['y']<0 or box['x']+box['w']>1201 or box['y']+box['h']>676:
                unsupported['элементы за краями слайда']+=1;continue
            if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                try:
                    with Image.open(io.BytesIO(shape.image.blob)) as image:
                        pixels=image.width*image.height;total_pixels+=pixels
                        if pixels>16_000_000 or total_pixels>40_000_000:raise ValueError('В фотографиях слишком много пикселей для конструктора.')
                        im=ImageOps.exif_transpose(image).convert('RGBA')
                        crops=(shape.crop_left,shape.crop_top,shape.crop_right,shape.crop_bottom)
                        if any(v<0 for v in crops):notes.append('Отрицательная обрезка фото упрощена.')
                        left,top,right,bottom=[max(0,v) for v in crops]
                        if left+right>=1 or top+bottom>=1:raise ValueError('Некорректная обрезка фотографии.')
                        im=im.crop((round(left*im.width),round(top*im.height),round((1-right)*im.width),round((1-bottom)*im.height)))
                        # Match the original displayed aspect, including stretched pictures.
                        ratio=min(1,1600/max(box['w'],box['h']))
                        im=im.resize((max(1,round(box['w']*ratio)),max(1,round(box['h']*ratio))),Image.Resampling.LANCZOS)
                        image_data=io.BytesIO();im.save(image_data,format='PNG')
                        data=image_data.getvalue()
                        if len(data)>2_800_000:raise ValueError('Одна из фотографий слишком большая.')
                    geom=shape._element.xpath('.//a:prstGeom')
                    mask=geom[0].get('prst') if geom else 'rect'
                    if mask not in ('rect','roundRect','ellipse'):notes.append('Нестандартная маска фото заменена прямоугольной.')
                    target['elements'].append(dict(**box,type='photo',src='data:image/png;base64,'+base64.b64encode(data).decode(),iw=im.width,ih=im.height,shape={'ellipse':'circle','roundRect':'round'}.get(mask,'rect'),radius=24,zoom=1,fx=50,fy=50,border='#000000',borderWidth=0))
                    if shape._element.xpath('.//a:effectLst/*|.//a:blip/a:alphaModFix|.//a:ln'):
                        notes.append('Эффекты, рамки или прозрачность фотографии могут отличаться.')
                except (OSError,Image.DecompressionBombError):
                    unsupported['неподдерживаемые изображения']+=1
            elif shape.has_text_frame and shape.text.strip():
                tf=shape.text_frame
                text=tf.text.replace('\v','\n')
                if len(text)>3000:raise ValueError(f'Слайд {index}: текстовый блок длиннее 3000 символов.')
                paras=tf.paragraphs;runs=[r for para in paras for r in para.runs]
                font=runs[0].font if runs else paras[0].font
                fallback=paras[0].font
                size=(font.size or fallback.size)
                size=float(size)*scale if size else 24
                fname=font.name or fallback.name or 'Arial'
                if fname not in FONTS:notes.append(f'Шрифт {fname} заменён на Arial.');fname='Arial'
                solid=font._rPr.find('{'+NS['a']+'}solidFill')
                if solid is None:solid=fallback._rPr.find('{'+NS['a']+'}solidFill')
                value=color_node(solid,colors,'#000000')
                if solid is None:notes.append('Проверьте цвет текста: наследуемый цвет заменён чёрным.')
                styles={(r.font.name,r.font.size,r.font.bold,r.font.italic,ET.tostring(ET.fromstring(r.font._rPr.xml))) for r in runs}
                if len(styles)>1 or any(x.level for x in paras):notes.append('Смешанное оформление и списки приведены к стилю первого фрагмента.')
                if tf.vertical_anchor and int(tf.vertical_anchor)!=1:notes.append('Вертикальное выравнивание текста заменено верхним.')
                box.update(x=box['x']+tf.margin_left*scale,y=box['y']+tf.margin_top*scale,w=box['w']-(tf.margin_left+tf.margin_right)*scale,h=box['h']-(tf.margin_top+tf.margin_bottom)*scale)
                if box['w']<1 or box['h']<1:raise ValueError(f'Слайд {index}: слишком тесная рамка текста.')
                if not 1<=size<=200:raise ValueError(f'Слайд {index}: неподдерживаемый размер шрифта.')
                target['elements'].append(dict(**box,type='text',text=text,size=size,font=fname,bold=(font.bold if font.bold is not None else fallback.bold) is True,italic=(font.italic if font.italic is not None else fallback.italic) is True,color=value,align={2:'center',3:'right'}.get(int(paras[0].alignment or 1),'left')))
                if shape.shape_type not in (MSO_SHAPE_TYPE.TEXT_BOX,MSO_SHAPE_TYPE.PLACEHOLDER):
                    unsupported['оформление фигур с текстом']+=1
            else:
                label='таблицы' if shape.has_table else 'диаграммы' if shape.has_chart else 'группы' if shape.shape_type==MSO_SHAPE_TYPE.GROUP else 'фигуры и соединения'
                unsupported[label]+=1
        if unsupported:notes.append('Не перенесены: '+', '.join(f'{k} ({v})' for k,v in unsupported.items())+'.')
        if not target['elements'] and unsupported:raise ValueError(f'Слайд {index} не содержит поддерживаемого текста или фотографий. Исходный файл не изменён.')
        # Do not import a project which cannot pass the existing paid export limits.
        try:quote(dict(version=1,slides=[target]))
        except ValueError as e:raise ValueError(f'Слайд {index}: '+str(e).split(': ',1)[-1]) from None
        deck['slides'].append(target)
        reports.extend(f'Слайд {index}: {note}' for note in dict.fromkeys(notes))
    import json
    if len(json.dumps(deck))>12_000_000:raise ValueError('После переноса проект слишком большой. Уменьшите фотографии в PPTX.')
    warnings += reports
    deck['importWarnings']=warnings
    return dict(ok=True,deck=deck,warnings=warnings,price=quote(deck),notice='Создаётся редактируемая копия. Перенос текста и фото может изменить переносы строк. Исходный PPTX остаётся без изменений. Проверьте каждый слайд перед сборкой.')
