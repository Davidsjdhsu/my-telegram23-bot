"""Opt-in technology layouts. Pure scene planning plus an adapter for the existing bot.

No network, credentials, billing or Telegram operations. Dimensions are CSS pixels.
The same scene can be rendered by the bot or the review/export tool.
"""
import copy
import math
import re
from functools import lru_cache
from pathlib import Path

from PIL import ImageFont
from presentation_profiles import PROFILES, resolved_profile, readable_color
from presentation_catalog import TEMPLATES, composition_for
from presentation_compositions import build_catalog_scene
from presentation_visuals import fit_rect, set_slide_notes

STYLES = {'designer_soft', 'designer_strict'}
FONT = 'Arial'
LAYOUTS = {'auto', 'photo', 'statement', 'process', 'comparison', 'sections', 'conclusion'}


class DesignerValidationError(ValueError):
    pass


def clean(value, limit, field, optional=False):
    if optional and value in (None, ''):
        return ''
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise DesignerValidationError(f'Invalid {field}')
    return value.strip()


def content_instruction(photo_story=False):
    rule = '''Дизайнерское оформление: крупные заголовки, короткие смысловые блоки.
Для каждого содержательного слайда добавь designer:
{"layout":"auto|photo|statement|process|comparison|sections|conclusion",
"blocks":[{"heading":"короткий подзаголовок","text":"текст"}]}.
Для process нужно 3–4 последовательных шага, comparison ровно 2 стороны,
sections 2–4 смысловых блока, conclusion 2–3 итоговых вывода.
Для comparison добавь rows: [{"criterion":"критерий до 26 символов","left":"первая сторона до 75 символов","right":"вторая сторона до 75 символов"}].
Ровно 3 строки, одинаковый критерий для обеих сторон в каждой строке. Два blocks задают названия сторон.
Последний содержательный слайд всегда conclusion: 2–3 блока с полезным выводом,
условием применения или ограничением. Текст каждого итогового блока — полное предложение с конкретной деталью, а не два-три общих слова.
Не пересказывай предыдущий слайд и не добавляй новый факт.
content должен передавать тот же смысл, что таблица или выводы, для текстовой копии. Заголовок блока до 36 символов, текст до 120.
Вся сумма текстов и заголовков блоков до 440 символов. Для auto/photo/statement
blocks=[]; тогда используется content. content обязателен и содержит тот же смысл,
что блоки, для текстовой копии. Не добавляй фактов, процентов или чисел ради дизайна.
Для statement content до 220 символов. Для photo content до 320 символов.
Не менее половины слайдов планируй с предметными фото (photo/auto), включая обложку.
statement используй не более одного раза. process — только реальная последовательность
действий или этапов; для перечня болезней, свойств или мер профилактики используй sections.
Чередуй фото, схемы и сравнения по смыслу; не делай все слайды одинаковыми.
Если план содержит сопоставление, используй comparison, для последовательности действий — process.
Структурный макет нужен только для настоящего сравнения, последовательности или объяснения
нескольких самостоятельных разделов. Для кухни, интерьеров, отдыха и других визуальных тем
предпочитай photo с подходящей сценой. Не превращай два коротких пункта в пустой слайд.
Не выдумывай этапы или противопоставления ради оформления. Данные пользователя сохраняй. При длинном материале выбирай auto. Все поля и
количество слайдов из основного задания остаются обязательными.'''
    if photo_story:
        rule += ('\nДля личного фотоальбома исключение: последний слайд НЕ обязан иметь conclusion. '
                 'Выбирай photo/auto и одну короткую подпись, если нет настоящих итогов. '
                 'Не выдумывай 2–3 вывода, сравнение или этапы ради макета. В первую очередь фото и рассказ владельца.')
    return rule



def designer_copy(item):
    """All visible evidence for the text companion, without summarizing tables away."""
    design=validate_designer(item)
    if design['layout']=='comparison' and design.get('rows'):
        left,right=(b['heading'] for b in design['blocks'])
        return '\n\n'.join(f"{r['criterion']}. {left}: {r['left']}. {right}: {r['right']}." for r in design['rows'])
    if design['blocks']:
        return '\n\n'.join(b['heading']+': '+b['text'] for b in design['blocks'])
    return item.get('content') or ''


def plan_designer_images(slides, style):
    """Keep true diagrams; promote supporting-photo layouts in a text-heavy plan."""
    for i, item in enumerate(slides):
        item['_designer_style'] = style
        item['_designer_cover'] = i == 0
        family = item.get('_designer_design', {}).get('family')
        if family in PROFILES:
            portrait = (i == 0 and family != 'technology' and len(str(item.get('title') or '')) <= 44) or (i > 0 and (
                family == 'editorial' and (i+1)%6 in (1,2,4) or
                family in ('technology','business') and (i+1)%2 == 0))
            item['_designer_aspect_ratio'] = '3:4' if portrait else '16:9'
        if item.get('_designer_design', {}).get('template_id') in TEMPLATES:
            comp=composition_for(item,i+1,len(slides))
            template=TEMPLATES[item['_designer_design']['template_id']]
            portrait=comp=='cover_split' and template['system']!='cinema' and len(str(item.get('title') or ''))<=44
            item['_designer_aspect_ratio']='3:4' if portrait else '16:9'
        item['_skip_ai_image'] = bool(item.get('chart')) or (i > 0 and item.get('designer', {}).get('layout') in (
            'process', 'comparison', 'sections', 'statement', 'conclusion'))
        # Sparse section lists gain a meaningful photo; comparisons and processes stay native.
        blocks = item.get('designer', {}).get('blocks', [])
        if (i > 0 and family == 'editorial' and not item.get('chart')
                and item.get('designer', {}).get('layout') == 'sections' and len(blocks) == 2
                and sum(len(str(b.get('heading',''))) + len(str(b.get('text',''))) for b in blocks) < 180):
            item['_skip_ai_image'] = False
    target = min(6, math.ceil(len(slides)/2))
    current = sum(not s['_skip_ai_image'] for s in slides)
    # Choose alternate positions first; preserve all structured text when rendered with a photo.
    candidates = list(range(2, len(slides), 2)) + list(range(1, len(slides), 2))
    for i in candidates:
        s = slides[i]
        if current >= target:
            break
        if s['_skip_ai_image'] and not s.get('chart') and s.get('designer', {}).get('layout') in ('statement', 'sections'):
            s['_skip_ai_image'] = False
            current += 1


def validate_designer(item):
    """Unknown/malformed layouts fail early, rather than silently losing text."""
    raw = item.get('designer')
    if raw is None:
        return {'layout': 'auto', 'blocks': []}
    if not isinstance(raw, dict) or raw.get('layout') not in LAYOUTS:
        raise DesignerValidationError('Unknown designer layout')
    layout = raw['layout']
    blocks = raw.get('blocks', [])
    if not isinstance(blocks, list):
        raise DesignerValidationError('Blocks must be a list')
    count = {'process': (3, 4), 'comparison': (2, 2), 'sections': (2, 4), 'conclusion': (2, 3)}.get(layout, (0, 0))
    if not count[0] <= len(blocks) <= count[1]:
        raise DesignerValidationError('Layout and block count disagree')
    result = []
    for block in blocks:
        if not isinstance(block, dict):
            raise DesignerValidationError('Invalid block')
        result.append({'heading': clean(block.get('heading'), 36, 'block heading'),
                       'text': clean(block.get('text'), 120, 'block text', optional=layout=='comparison' and bool(raw.get('rows')))})
    if sum(len(b['heading']) + len(b['text']) for b in result) > 440:
        raise DesignerValidationError('Shorten blocks before rendering')
    validated = {'layout': layout, 'blocks': result}
    if layout == 'comparison' and 'rows' in raw:
        rows=raw['rows']
        if not isinstance(rows,list) or len(rows)!=3:raise DesignerValidationError('Comparison needs three shared criteria')
        validated['rows']=[{key:clean(row.get(key),26 if key=='criterion' else 75,'comparison '+key)
                            for key in ('criterion','left','right')} for row in rows if isinstance(row,dict)]
        if len(validated['rows'])!=3:raise DesignerValidationError('Invalid comparison row')
    return validated


@lru_cache(maxsize=16)
def metric_font(size, bold=False, family=FONT):
    # DejaVu is already installed on Render; its wider metrics reserve breathing
    # room for the widely available Arial used in the editable PowerPoint file.
    base = 'DejaVuSerif' if family in ('Georgia','Times New Roman') else 'DejaVuSans'
    name = base + ('-Bold.ttf' if bold else '.ttf')
    locations = [Path('/usr/share/fonts/truetype/dejavu') / name,
                 Path('/Library/Fonts') / name,
                 Path(__file__).parent / 'fonts' / name]
    # Bundled review runtime. On Render, the first system path is used.
    runtime = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
    locations.append(runtime / name)
    for path in locations:
        if path.is_file():
            return ImageFont.truetype(str(path), size)
    try:
        return ImageFont.truetype(name, size)
    except OSError:
        return ConservativeMetrics(size)


class ConservativeMetrics:
    """Reserve space when the host lacks DejaVu, without downloading fonts at runtime."""
    def __init__(self, size): self.size = size
    def getlength(self, text):
        return self.size * sum(1.1 if ord(c) > 0x2e7f or c in 'WMЖШЩЮ' else
                               .38 if c.isspace() else .76 for c in text)


def measured(font, text, size):
    if any(ord(c) > 0x2e7f for c in text):
        return sum(size * 1.05 if ord(c) > 0x2e7f else font.getlength(c) for c in text)
    return font.getlength(text)


def wrap(text, width, size, bold=False, family=FONT):
    font = metric_font(size, bold, family)
    lines = []
    for paragraph in text.split('\n'):
        current = ''
        for word in paragraph.split():
            candidate = f'{current} {word}'.strip()
            if measured(font, candidate, size) <= width:
                current = candidate
                continue
            if current:
                lines.append(current)
            current = ''
            # Long URLs / unspaced scripts still must fit without truncation.
            for letter in word:
                if current and measured(font, current + letter, size) > width:
                    lines.append(current)
                    current = ''
                current += letter
        lines.append(current)
    return '\n'.join(lines)


def text_node(text, box, color, size=30, minimum=28, bold=False, role='body', family=FONT):
    x, y, width, height = box
    for chosen in range(size, minimum - 1, -1):
        # Reduce the font before breaking ordinary words across lines.
        font = metric_font(chosen, bold, family)
        words = str(text).split()
        if any(1 < len(w) <= 40 and not any(ord(c) > 0x2e7f for c in w)
               and measured(font, w, chosen) > width - 4 for w in words):
            continue
        wrapped = wrap(str(text), width - 4, chosen, bold, family)
        lines = len(wrapped.split('\n'))
        if lines * chosen * 1.22 + 5 <= height:
            return {'type': 'text', 'text': wrapped, 'box': list(box), 'color': color,
                    'size': chosen, 'font': family, 'bold': bold, 'role': role}
    raise DesignerValidationError(f'Shorten {role}: text does not fit at a readable size')


def is_designer(data):
    return data.get('theme_name') in STYLES and not data.get('custom_colors')


def palette(dark):
    return {'bg': '#080B10' if dark else '#FFFFFF',
            'ink': '#FFFFFF' if dark else '#080B10',
            'body': '#DCE5F6' if dark else '#28323F',
            'accent': '#80F0E1' if dark else '#244DFF',
            'muted': '#B8CBCF' if dark else '#58616F'}


GRADIENT = {'angle': 0, 'stops': [(0, '#080B10', 1), (.45, '#14225B', 1),
                                 (.8, '#234DEC', 1), (1, '#32B8CC', 1)]}


def build_scene(item, asset, number, total, lang='ru', style='designer_soft'):
    """Plan the complete slide before mutating a PPTX; never truncate content."""
    if (item.get('_designer_design', {}).get('template_id') in TEMPLATES
            and TEMPLATES[item['_designer_design']['template_id']]['family']==item['_designer_design'].get('family')):
        return build_catalog_scene(item, asset, number, total, lang, style)
    if item.get('_designer_design', {}).get('family') in PROFILES:
        return build_profile_scene(item, asset, number, total, lang, style)
    title = clean(item.get('title'), 160, 'title')
    body = clean(item.get('content'), 480, 'body', optional=number == 1)
    design = validate_designer(item)
    layout, blocks = design['layout'], design['blocks']
    if not 1 <= number <= total: raise DesignerValidationError('Invalid slide number')
    if asset and (not Path(asset['path']).is_file() or asset['width'] <= 0 or asset['height'] <= 0):
        raise DesignerValidationError('Invalid image asset')
    if asset and blocks:
        # User photos take priority; preserve all structured text when using a photo layout.
        body = '\n\n'.join(b['heading'] + '\n' + b['text'] for b in blocks)
        blocks, layout = [], 'photo'
    layout = 'cover' if number == 1 else 'chart' if item.get('chart') else 'photo' if asset else layout
    if layout == 'photo' and not asset: layout = 'auto'
    radius = 36 if style == 'designer_soft' else 0
    dark = layout in ('cover', 'statement') or item.get('_designer_background') == 'dark'
    c = palette(dark)
    accent = item.get('_designer_accent')
    if isinstance(accent, str) and re.fullmatch(r'#[0-9a-fA-F]{6}', accent):
        # Keep semantic accent text readable on the chosen background.
        rgb = [int(accent[i:i+2],16)/255 for i in (1,3,5)]
        luminance = sum(v*w for v,w in zip(rgb,(.2126,.7152,.0722)))
        if (dark and luminance >= .45) or (not dark and luminance <= .65):
            c['accent'] = accent
    
    scene = {'background': '#F4F6FA' if layout in ('cover','statement') else c['bg'], 'layout': layout,
             'notes': 'AlixDocBot. ' + style, 'nodes': []}
    nodes = scene['nodes']

    def txt(value, box, role='body', size=30, minimum=26, color=None, bold=False):
        if value: nodes.append(text_node(value, box, color or c['body'], size, minimum, bold, role))

    def panel(box, gradient=None, fill=None):
        nodes.append({'type':'panel','box':list(box),'radius':radius,
                      'gradient':gradient,'fill':fill or c['bg']})

    def photo(box, cover=False):
        # Preserve complete source photographs, including evidence at the edges.
        if asset.get('origin') == 'user':
            panel(box, fill='#EAF0F7')
            x,y,w,h = box
            fitted=fit_rect(asset['width'],asset['height'],(x+16,y+16,w-32,h-32))
            nodes.append({'type':'image','path':str(asset['path']),'box':list(fitted),
                          'origin':'user','radius':0,'crop':False})
        else:
            crop = False
            image_box = fit_rect(asset['width'], asset['height'], box)
            nodes.append({'type':'image','path':str(asset['path']),'box':list(image_box),
                          'origin':'ai','radius':radius,'crop':crop})

    def heading(box=(64,72,1152,214), size=76, minimum=42):
        txt(title,box,'title',size,minimum,c['ink'])

    if layout == 'cover':
        panel((20,20,1240,680),GRADIENT)
        if asset and asset.get('origin') != 'user':
            photo((20,20,1240,680),True)
            # Native overlay guarantees text contrast even if an AI image is too bright.
            panel((20,20,1240,680),{'angle':0,'stops':[(0,'#080B10',.97),
                  (.58,'#080B10',.84),(1,'#080B10',.15)]})
            heading((64,124,1056,310),112,40)
        elif asset:
            try:
                node = text_node(title,(64,112,600,300),c['ink'],80,42,False,'title')
            except DesignerValidationError:
                # Preserve long confirmed titles alongside a complete user photograph.
                node = text_node(title,(64,64,860,310),c['ink'],64,40,False,'title')
                photo((992,64,224,166))
            else:
                photo((710,100,500,510))
            nodes.append(node)
        else: heading((64,124,1080,310),112,40)
        subtitle = item.get('subtitle') or item.get('point') or body
        if subtitle:
            subtitle = clean(subtitle,180,'subtitle')
            # Long descriptions get a full-width band below the photo/title area.
            try:
                subtitle_node = text_node(subtitle,(68,490,600 if asset else 1050,145),c['body'],32,26)
            except DesignerValidationError:
                if asset and asset.get('origin') == 'user':
                    # Rebuild the user cover so the band cannot overlap its photograph.
                    nodes[:] = nodes[:1]
                    heading((64,64,860,310),64,40)
                    photo((992,64,224,166))
                subtitle_node = text_node(subtitle,(68,444,1136,204),c['body'],32,26)
            nodes.append(subtitle_node)
    elif layout == 'photo':
        portrait = asset['width'] / asset['height'] < .9
        try:
            if portrait and len(body) <= 320 and len(title) <= 85:
                title_node = text_node(title,(608,84,608,202),c['ink'],62,40,False,'title')
                body_node = text_node(body,(608,322,608,318),c['body'],30,26)
                nodes.extend([title_node,body_node])
                photo((64,64,480,584))
            else:
                title_node = text_node(title,(64,54,1152,164),c['ink'],68,40,False,'title')
                body_node = text_node(body,(64,268,520,370),c['body'],30,26)
                nodes.extend([title_node,body_node])
                photo((640,242,576,406))
        except DesignerValidationError:
            # A wide copy column still reserves a substantial photograph, never a tiny thumbnail.
            heading((64,54,1152,164),60,40)
            txt(body,(64,250,680,398),size=30,minimum=24)
            photo((800,250,416,398))
    elif layout == 'process':
        heading((64,60,1152,210),72,42)
        panel((32,306,1216,350),GRADIENT)
        count=len(blocks); step_w=(1152-(count-1)*32)/count
        for i,b in enumerate(blocks):
            x=64+i*(step_w+32)
            txt(f'{i+1:02}',(x,332,step_w,65),'step',48,48,'#80F0E1')
            txt(b['heading'],(x,407,step_w,87),'subheading',32,26,'#FFFFFF')
            txt(b['text'],(x,505,step_w,132),size=27,minimum=24,color='#FFFFFF')
            if i < count-1:
                nodes.append({'type':'connector','box':[x+82,370,step_w-67,2],'color':'#80F0E1'})
    elif layout in ('comparison','sections'):
        heading((64,60,1152,190),72,42)
        # Three equal cards, rather than a 2x2 grid with an empty fourth cell.
        cols=3 if len(blocks)==3 else 2
        rows=math.ceil(len(blocks)/cols); row_h=376/rows
        col_w=1152/cols; text_w=col_w-48
        panel((32,250,1216,408),fill='#15202E' if dark else '#F0F4FC')
        for i,b in enumerate(blocks):
            x=64+(i%cols)*col_w; y=270+(i//cols)*row_h
            head_h=110 if cols==3 else 70
            txt(b['heading'],(x,y,text_w,head_h),'subheading',36,26,c['accent'])
            txt(b['text'],(x,y+head_h+4,text_w,row_h-head_h-22),size=29,minimum=24)
    elif layout == 'statement':
        panel((20,20,1240,680),GRADIENT)
        heading((64,72,1110,212),80,44)
        txt(body,(180,330,1000,300),size=44,minimum=28)
    elif layout == 'chart':
        heading((64,58,870 if asset else 1152,158),60,42)
        txt(body,(64,260,400,370),size=30,minimum=24)
        nodes.append({'type':'chart','box':[510,250,706,380],'data':copy.deepcopy(item['chart']),'dark':False})
        if asset: photo((980,56,236,152))
    else:
        heading((64,72,1120,208),76,42)
        txt(body,(180,330,1020,310),size=38,minimum=26)

    label=f'{number:02} / {total:02}'
    if asset and asset.get('origin') == 'ai':
        scene['notes'] += '\nImage: AI-generated illustration, not documentary evidence.'
    elif asset: scene['notes'] += '\nImage: supplied by user. Full frame preserved.'
    # Keep the original body available verbatim as speaker notes, too.
    scene['notes'] += '\n' + (item.get('content') or item.get('point') or '')
    txt(label,(64,672,1120,25),'footer',14,14,c['muted'])
    for node in nodes:
        if node.get('role') == 'title':
            lines=node['text'].split('\n')
            if layout == 'cover' and len(lines) == 2 and not (asset and asset.get('origin') == 'user'):
                if measured(metric_font(node['size']),lines[1],node['size']) + 116 < node['box'][2]:
                    node['stagger'] = 116
            elif not dark and len(lines) > 1:
                node['last_line_color'] = c['accent']
        if node['type'] == 'text' and lang == 'ar': node['rtl'] = True
    validate_scene(scene)
    return scene


def build_profile_scene(item, asset, number, total, lang, style):
    """Four art directions with shared safe, editable output primitives."""
    design = item['_designer_design']; family = design['family']; c = resolved_profile(design)
    title = clean(item.get('title'),160,'title')
    body = clean(item.get('content'),480,'body',optional=number==1)
    structured = validate_designer(item); layout = structured['layout']; blocks = structured['blocks']
    if not 1 <= number <= total: raise DesignerValidationError('Invalid slide number')
    if asset and (not Path(asset['path']).is_file() or min(asset['width'],asset['height'])<=0):
        raise DesignerValidationError('Invalid image asset')
    evidence = layout == 'conclusion' or (layout == 'comparison' and structured.get('rows'))
    if asset and blocks and not evidence:
        body='\n\n'.join(b['heading']+'\n'+b['text'] for b in blocks)
        layout='photo'; blocks=[]
    layout='cover' if number==1 else 'chart' if item.get('chart') else layout if evidence else 'photo' if asset else layout
    if layout=='photo' and not asset:layout='auto'
    scene={'background':c['bg'],'layout':layout,'family':family,'nodes':[],
           'notes':f'AlixDocBot. {style}. {family}\n'+(item.get('content') or item.get('point') or '')}
    nodes=scene['nodes'];radius=28 if style=='designer_soft' else 0
    # Keep supported script coverage when a decorative Latin/Cyrillic face is unsuitable.
    heading_font=c['font'] if lang not in ('zh','ar') else 'Arial'
    def txt(value,box,role='body',size=32,minimum=26,color=None):
        node=text_node(value,box,color or c['body'],size,minimum,
            c['bold'] if role=='title' else role=='subheading',role,
            heading_font if role=='title' else c['body_font'])
        if lang=='ar':node['rtl']=True
        nodes.append(node)
    def heading(box=(64,50,1152,190),size=72,minimum=40):
        color = (c['accent'] if number == 1 else c['secondary_ink']) if family == 'editorial' else c['ink']
        txt(title,box,'title',size,minimum,color)
    def photo(box, full=False):
        crop=False
        rect=fit_rect(asset['width'],asset['height'],box)
        nodes.append({'type':'image','path':str(asset['path']),'box':list(rect),
            'origin':asset.get('origin','ai'),'radius':radius if asset.get('origin')!='user' else 0,'crop':crop})
    def fallback_photo():
        nodes.clear();heading((64,48,1152,176),64,38)
        txt(body,(64,250,600,405),size=32,minimum=24)
        photo((720,250,496,390))
    if layout=='cover':
        try:
            subtitle=clean(item.get('subtitle') or item.get('point') or body,180,'subtitle',optional=True)
            if family=='technology' and asset and asset.get('origin')!='user':
                photo((0,0,1280,720),True)
                nodes.append({'type':'panel','box':[0,0,1280,720],'radius':0,
                    'gradient':{'angle':0,'stops':[(0,c['bg'],.98),(.58,c['bg'],.88),(1,c['bg'],.3)]}})
                heading((64,95,990,350),100,42)
                if subtitle:txt(subtitle,(68,508,760,128),size=32,minimum=26)
            elif asset and len(title)>44:
                # Give a long title the full width instead of stacking four narrow lines.
                heading((64,50,1152,240),76,38)
                if subtitle:txt(subtitle,(64,362,540,250),size=34,minimum=24)
                photo((674,302,542,330))
            elif asset:
                try:
                    heading((64,76,640,360),88 if family=='editorial' else 80,40)
                    if subtitle:txt(subtitle,(68,485,612,152),size=31,minimum=24)
                    photo((758,48,458,584))
                except DesignerValidationError:
                    nodes.clear();heading((64,54,1152,240),76,38)
                    if subtitle:txt(subtitle,(64,370,580,240),size=32,minimum=24)
                    photo((710,310,506,320))
            else:
                heading((64,100,1152,330),100,40)
                if subtitle:txt(subtitle,(68,485,960,150),size=36,minimum=26)
        except DesignerValidationError:
            nodes.clear()
            heading((64,50,1152,250),76,38)
            if subtitle:txt(subtitle,(64,350,680 if asset else 1152,300),size=32,minimum=24)
            if asset:photo((800,350,416,300))
    elif layout=='photo':
        try:
            # Narrow text is only used when it fits; complete source copy is retained.
            if len(body)>320 or len(title)>90:
                fallback_photo()
            elif family=='editorial' and number%6==5:
                heading((64,54,410,350),66,38)
                photo((520,54,696,392))
                txt(body,(64,500,1152,145),size=34,minimum=26)
            elif family=='editorial' and number%6==0:
                photo((64,48,720,405))
                heading((832,54,384,390),62,38)
                txt(body,(64,510,1152,140),size=32,minimum=24)
            elif family=='editorial' and number%6==2:
                heading((64,62,600,194),68,40)
                txt(body,(64,316,590,304),size=32,minimum=26)
                photo((730,54,486,580))
            elif (family=='editorial' and number%6 in (1,4)) or (family=='technology' and number%2==0):
                photo((64,54,536,580))
                heading((662,66,554,196),62,40)
                txt(body,(662,316,554,306),size=31,minimum=26)
            elif family=='business' and number%2==0:
                heading((64,64,580,200),68,40)
                txt(body,(64,320,560,306),size=32,minimum=26)
                photo((700,54,516,580))
            elif family=='education' and number%2==0:
                heading((64,50,1152,180),64,40)
                txt(body,(64,290,480,340),size=32,minimum=26)
                photo((620,252,596,380))
            else:
                heading((64,50,1152,180),72 if family=='editorial' else 64,40)
                photo((64,270,650,360))
                txt(body,(780,278,436,352),size=30,minimum=26)
        except DesignerValidationError:
            fallback_photo()
    elif layout=='process':
        heading((64,50,1152,190),72,40)
        width=(1152-(len(blocks)-1)*32)/len(blocks)
        for i,b in enumerate(blocks):
            x=64+i*(width+32)
            txt(f'{i+1:02}',(x,284,width,66),'subheading',48,40,c['accent'])
            txt(b['heading'],(x,372,width,110),'subheading',33,24,c['ink'])
            txt(b['text'],(x,500,width,165),size=28,minimum=22)
            if i<len(blocks)-1:
                nodes.append({'type':'connector','box':[x+75,319,width-57,2],'color':c['accent']})
    elif layout=='comparison' and structured.get('rows'):
        heading((64,48,870 if asset else 1152,154),64,38)
        if asset: photo((980,48,236,150))
        headers=['Критерий' if lang=='ru' else 'Criterion']+[b['heading'] for b in blocks]
        values=[headers]+[[r[k] for k in ('criterion','left','right')] for r in structured['rows']]
        widths=[250,451,451]; heights=[74,118,118,118]
        cells=[]
        for row,values_row in enumerate(values):
            cells.append([text_node(value,(0,0,widths[col]-32,heights[row]-20),
                c['accent'] if row==0 else c['ink'] if col==0 else c['body'],
                31 if row==0 else 29,23,row==0 or col==0,'table',c['body_font'])
                for col,value in enumerate(values_row)])
        nodes.append({'type':'table','box':[64,226,1152,sum(heights)],'widths':widths,
                      'heights':heights,'cells':cells,'background':c['bg'],'accent':c['accent']})
    elif layout=='conclusion':
        heading((64,52,870 if asset else 1152,155),72,38)
        if asset: photo((980,48,236,150))
        height=420/len(blocks)
        for i,b in enumerate(blocks):
            y=232+i*height
            txt(b['heading'],(64,y,352,height-24),'subheading',40,26,c['accent'])
            txt(b['text'],(470,y,746,height-24),size=35,minimum=26)
    elif layout in ('comparison','sections'):
        heading((64,50,1152,185),72,40)
        cols=3 if len(blocks)==3 else 2; rows=math.ceil(len(blocks)/cols)
        width=(1152-48*(cols-1))/cols; height=370/rows
        for i,b in enumerate(blocks):
            x=64+(i%cols)*(width+48);y=270+(i//cols)*height
            txt(b['heading'],(x,y,width,90 if rows==1 else 65),'subheading',38,24,c['accent'])
            txt(b['text'],(x,y+(105 if rows==1 else 70),width,height-(110 if rows==1 else 75)),size=32,minimum=22)
    elif layout=='chart':
        heading((64,50,870 if asset else 1152,180),68,40)
        txt(body,(64,280,420,340),size=30,minimum=24)
        nodes.append({'type':'chart','box':[540,260,676,370],'data':copy.deepcopy(item['chart']),
                      'dark':design['background']=='dark','palette':c})
        if asset:photo((980,54,236,176))
    elif number==total:
        # A compact conclusion balances two readable columns instead of a small low paragraph.
        try:
            heading((64,138,520,390),80 if family=='editorial' else 72,38)
            txt(body,(680,150,536,430),size=40,minimum=26,color=c['accent'] if layout=='statement' else c['body'])
        except DesignerValidationError:
            nodes.clear();heading((64,70,1152,200),72,38)
            txt(body,(64,320,1152,320),size=38,minimum=26)
    else:
        heading((64,80,1152,236),88 if family=='editorial' else 76,40)
        txt(body,(190,368,1000,272),size=40,minimum=26,color=c['accent'] if layout=='statement' else c['body'])
    label=f'{number:02} / {total:02}'
    if asset:
        if asset.get('origin')=='user':scene['notes']+='\nImage: supplied by user. Full frame preserved.'
        else:
            scene['notes']+='\nImage: AI-generated illustration, not documentary evidence.'
    txt(label,(64,675,1152,22),'footer',13,13,c['muted'])
    validate_scene(scene)
    return scene


def validate_scene(scene):
    for node in scene['nodes']:
        x,y,w,h=node['box']
        if not all(math.isfinite(v) for v in (x,y,w,h)) or min(x,y)<0 or min(w,h)<=0 or x+w>1280.1 or y+h>720.1:
            raise DesignerValidationError('Element outside slide')
    return scene


def render_scene(slide, scene, chart_renderer=None):
    """Existing bot runtime adapter; no new runtime dependency."""
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    validate_scene(scene)
    bg=slide.background.fill;bg.solid();bg.fore_color.rgb=RGBColor.from_string(scene['background'][1:])

    def rounded(shape,radius,w,h):
        geom=shape._element.spPr.find(qn('a:prstGeom'))
        if geom is None:
            geom=OxmlElement('a:prstGeom');shape._element.spPr.insert(1,geom)
        geom.set('prst','roundRect' if radius else 'rect')
        for child in list(geom): geom.remove(child)
        av=OxmlElement('a:avLst');geom.append(av)
        if radius:
            gd=OxmlElement('a:gd');gd.set('name','adj');gd.set('fmla','val '+str(round(min(.5,radius/min(w,h))*100000)));av.append(gd)

    def fill(shape,node):
        f=shape.fill
        if not node.get('gradient'):
            f.solid();f.fore_color.rgb=RGBColor.from_string(node['fill'][1:]);return
        g=node['gradient'];f.gradient()
        gs_list=f._xPr.find(qn('a:gradFill')).find(qn('a:gsLst'))
        for child in list(gs_list):gs_list.remove(child)
        for offset,color,opacity in g['stops']:
            gs=OxmlElement('a:gs');gs.set('pos',str(round(offset*100000)))
            rgb=OxmlElement('a:srgbClr');rgb.set('val',color[1:])
            alpha=OxmlElement('a:alpha');alpha.set('val',str(round(opacity*100000)));rgb.append(alpha)
            gs.append(rgb);gs_list.append(gs)
        f.gradient_angle=g['angle']

    def photo_geometry(shape,node,pw,ph):
        rounded(shape,node.get('radius',0),pw,ph)
        mask=node.get('mask','rect')
        presets={'circle':'ellipse','triangle':'triangle','hexagon':'hexagon','chamfer':'snip2DiagRect'}
        if mask in presets:
            geom=shape._element.spPr.find(qn('a:prstGeom'));geom.set('prst',presets[mask])
            for child in list(geom):geom.remove(child)
            av=OxmlElement('a:avLst');geom.append(av)
            if mask=='chamfer':
                for name in ('adj1','adj2'):
                    gd=OxmlElement('a:gd');gd.set('name',name);gd.set('fmla','val 8000');av.append(gd)

    for node in scene['nodes']:
        px,py,pw,ph=node['box'];x,y,w,h=[v/96 for v in node['box']]
        if node['type']=='panel':
            shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if node['radius'] else MSO_SHAPE.RECTANGLE,
                                        Inches(x),Inches(y),Inches(w),Inches(h))
            shape.name='Designer background';shape.line.fill.background()
            # Native theme effects otherwise add an unwanted shadow to flat rules.
            style_node=shape._element.find(qn('p:style'))
            if style_node is not None:
                effect_ref=style_node.find(qn('a:effectRef'))
                if effect_ref is not None:effect_ref.set('idx','0')
            rounded(shape,node['radius'],pw,ph);fill(shape,node)
            shape._element.spPr.append(OxmlElement('a:effectLst'))
        elif node['type']=='photo_frame':
            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
            shape.name='Designer photo frame';shape.fill.background()
            photo_geometry(shape,node,pw,ph)
            shape.line.color.rgb=RGBColor.from_string(node['color'][1:]);shape.line.width=Pt(node['width'])
            style_node=shape._element.find(qn('p:style'))
            if style_node is not None:
                effect_ref=style_node.find(qn('a:effectRef'))
                if effect_ref is not None:effect_ref.set('idx','0')
            shape._element.spPr.append(OxmlElement('a:effectLst'))
        elif node['type']=='text':
            shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));shape.name='Designer '+node['role']
            tf=shape.text_frame;tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
            for i,line in enumerate(node['text'].split('\n')):
                p=tf.paragraphs[0] if not i else tf.add_paragraph();p.text=line
                p.font.name=node['font'];p.font.size=Pt(node['size']*.75);p.font.bold=node['bold']
                p.font.color.rgb=RGBColor.from_string(node['color'][1:]);p.line_spacing=Pt(node['size']*.75*1.22)
                p.space_after=p.space_before=Pt(0)
                if i == len(node['text'].split('\n'))-1 and node.get('last_line_color'):
                    p.font.color.rgb=RGBColor.from_string(node['last_line_color'][1:])
                if i == 1 and node.get('stagger'):
                    p._p.get_or_add_pPr().set('marL',str(Inches(node['stagger']/96)))
                if node.get('rtl'):
                    from pptx.enum.text import PP_ALIGN
                    p.alignment=PP_ALIGN.RIGHT;p._p.get_or_add_pPr().set('rtl','1')
        elif node['type'] in ('image','background_image'):
            shape=slide.shapes.add_picture(node['path'],Inches(x),Inches(y),width=Inches(w),height=Inches(h))
            shape.name=('Designer texture' if node['type']=='background_image' else
                        'User photograph' if node['origin']=='user' else 'AI illustration')
            if node.get('flip_h'):shape._element.spPr.xfrm.set('flipH','1')
            if node.get('flip_v'):shape._element.spPr.xfrm.set('flipV','1')
            if node['type']=='background_image' and 'opacity' in node:
                blip=shape._element.blipFill.find(qn('a:blip'))
                alpha=OxmlElement('a:alphaModFix')
                alpha.set('amt',str(round(max(0,min(1,node['opacity']))*100000)))
                blip.append(alpha)
            if node.get('crop'):
                iw,ih=shape.image.size;source_ratio=iw/ih;target_ratio=w/h
                if source_ratio>target_ratio:
                    shape.crop_left=shape.crop_right=(1-target_ratio/source_ratio)/2
                else:shape.crop_top=shape.crop_bottom=(1-source_ratio/target_ratio)/2
            photo_geometry(shape,node,pw,ph)
            if node.get('border_color'):
                shape.line.color.rgb=RGBColor.from_string(node['border_color'][1:]);shape.line.width=Pt(node.get('border_width',1.5))
        elif node['type']=='connector':
            shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x if node.get('vertical') else x+w),Inches(y+h if node.get('vertical') else y))
            shape.line.color.rgb=RGBColor.from_string(node['color'][1:]);shape.line.width=Pt(1.5)
            tail=OxmlElement('a:tailEnd');tail.set('type','triangle');shape._element.spPr.get_or_add_ln().append(tail)
        elif node['type']=='table':
            shape=slide.shapes.add_table(len(node['cells']),len(node['widths']),Inches(x),Inches(y),Inches(w),Inches(h))
            shape.name='Designer comparison';table=shape.table
            for j,width in enumerate(node['widths']):table.columns[j].width=Inches(width/96)
            for i,height in enumerate(node['heights']):table.rows[i].height=Inches(height/96)
            for i,row in enumerate(node['cells']):
                for j,cell_node in enumerate(row):
                    cell=table.cell(i,j);cell.margin_left=cell.margin_right=Inches(16/96)
                    cell.margin_top=cell.margin_bottom=Inches(10/96);cell.vertical_anchor=MSO_ANCHOR.TOP
                    cell_bg=node.get('row_backgrounds',[node['background']]*len(node['cells']))[i]
                    cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(cell_bg[1:])
                    cell.text=cell_node['text'];cell.text_frame.word_wrap=False
                    for para in cell.text_frame.paragraphs:
                        para.font.name=cell_node['font'];para.font.size=Pt(cell_node['size']*.75)
                        para.font.bold=cell_node['bold'];para.font.color.rgb=RGBColor.from_string(cell_node['color'][1:])
                        para.space_before=para.space_after=Pt(0);para.line_spacing=Pt(cell_node['size']*.75*1.22)
                    props=cell._tc.get_or_add_tcPr()
                    for side in ('lnL','lnR','lnT','lnB'):
                        line=OxmlElement('a:'+side);line.set('w','6350')
                        fill_color=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr')
                        rgb.set('val',node['accent'][1:] if side=='lnB' else cell_bg[1:]);fill_color.append(rgb);line.append(fill_color);props.append(line)
        elif node['type']=='chart':
            if chart_renderer is None:raise DesignerValidationError('A native chart renderer is required')
            c=node.get('palette') or palette(node['dark']);conv=lambda color:tuple(bytes.fromhex(color[1:]))
            chart_renderer(slide,x,y,w,h,node['data'],{'ink':conv(c['ink']),'mid':conv(c['body']),
              'mute':conv(c['muted']),'line':conv(c['accent']),'bg':conv(c['bg'])})
    set_slide_notes(slide, scene['notes'])


def render_designer_slide(slide,item,asset,colors,number,total,lang,background,chart_renderer):
    scene=build_scene(item,asset,number,total,lang,item.get('_designer_style','designer_soft'))
    render_scene(slide,scene,chart_renderer)
    return scene
