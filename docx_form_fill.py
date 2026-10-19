"""Lossless field substitution in approved Word forms, including split text runs."""
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _set_text(node, value):
    # Preserve the surrounding run, its formatting, and unrelated page breaks/fields.
    parts = re.split(r'([\n\t])', value.replace('\r\n', '\n').replace('\r', '\n'))
    node.text = parts[0]
    node.set(qn('xml:space'), 'preserve')
    anchor = node
    for part in parts[1:]:
        new = OxmlElement('w:br' if part == '\n' else 'w:tab' if part == '\t' else 'w:t')
        if part not in ('\n', '\t'):
            new.text = part
            new.set(qn('xml:space'), 'preserve')
        anchor.addnext(new)
        anchor = new


def replace_fields(document, replacements):
    """Replace once against original text; never interpret inserted user data as tokens."""
    if not replacements:
        return
    pattern = re.compile('(?:' + '|'.join(re.escape(k) for k in sorted(replacements, key=len, reverse=True)) + r')\.?')
    roots = [document.element]
    for section in document.sections:
        roots += [section.header._element, section.footer._element]
    seen = set()
    for root in roots:
        for p in root.iter(qn('w:p')):
            if p in seen:
                continue
            seen.add(p)
            nodes = list(p.iter(qn('w:t')))
            text = ''.join(n.text or '' for n in nodes)
            spans = []
            pos = 0
            for n in nodes:
                end = pos + len(n.text or '')
                spans.append((pos, end, n))
                pos = end
            updates = {n: n.text or '' for n in nodes}
            for match in reversed(list(pattern.finditer(text))):
                token = match.group()
                dotted = token.endswith('.')
                value = replacements[token[:-1] if dotted else token]
                if dotted:
                    value = value.rstrip('.') + '.'
                covered = [(a, b, n) for a, b, n in spans if a < match.end() and b > match.start()]
                for i, (a, b, n) in enumerate(covered):
                    current = updates[n]
                    start, end = max(0, match.start()-a), min(b-a, match.end()-a)
                    updates[n] = current[:start] + (value if i == 0 else '') + current[end:]
            for n, value in updates.items():
                if value != (n.text or ''):
                    _set_text(n, value)


def adapt_filled_layout(document):
    """Let long values flow; keep short signature blocks on one page."""
    for table in document.tables:
        full = ' '.join(c.text for r in table.rows for c in r.cells)
        signature = ('Подпись' in full or 'подпись' in full or '____________ /' in full)
        keep_block = signature and len(full) < 900
        following = table._tbl.getnext()
        receipt = following is not None and 'Экземпляр договора получил' in ''.join(following.itertext())
        for ri, row in enumerate(table.rows):
            text = ' '.join(c.text for c in row.cells)
            props = row._tr.get_or_add_trPr()
            for height in list(props.findall(qn('w:trHeight'))):
                if height.get(qn('w:hRule')) == 'exact':height.set(qn('w:hRule'),'atLeast')
            if len(text) > 900:
                for x in list(props.findall(qn('w:cantSplit'))):props.remove(x)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.widow_control = True
                    if keep_block and (ri < len(table.rows)-1 or receipt):
                        p.paragraph_format.keep_with_next = True
                    elif len(text) > 900:
                        p.paragraph_format.keep_with_next = False
                        p.paragraph_format.keep_together = False
    for p in document.paragraphs:
        p.paragraph_format.widow_control = True
        if len(p.text) > 900 and not p.style.name.startswith('Heading'):
            p.paragraph_format.keep_together = False
            p.paragraph_format.keep_with_next = False
