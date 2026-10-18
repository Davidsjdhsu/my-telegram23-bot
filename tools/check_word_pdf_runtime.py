"""Build-time Word/PDF smoke check: Cyrillic, tables and source pagination."""
import asyncio
from pathlib import Path
import sys
import tempfile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Cm
from pypdf import PdfReader
from word_quality import convert_word_pdf

async def main():
    with tempfile.TemporaryDirectory(prefix='alix-word-smoke-') as td:
        root=Path(td);source=root/'check.docx';target=root/'check.pdf'
        d=Document();d.sections[0].page_width=Cm(21);d.sections[0].page_height=Cm(29.7)
        d.add_heading('Проверка документа',0)
        t=d.add_table(rows=2,cols=2);t.style='Table Grid'
        for row,values in zip(t.rows,[['Показатель','Значение'],['Количество','42']]):
            for cell,value in zip(row.cells,values):cell.text=value
        d.add_page_break();d.add_paragraph('Вторая страница');d.save(source)
        await convert_word_pdf(source,target)
        result=PdfReader(target);assert len(result.pages)==2
        assert '42' in result.pages[0].extract_text();assert 'Вторая страница' in result.pages[1].extract_text()
        print('word_pdf runtime_check=passed pages=2 table=yes')
if __name__=='__main__':asyncio.run(main())
