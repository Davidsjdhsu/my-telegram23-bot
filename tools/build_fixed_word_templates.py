"""Build independently typeset, unbranded Russian templates from versioned content."""
from pathlib import Path
import json,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]

def build(spec,path):
 d=Document();sec=d.sections[0]
 for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
 for name in ('Normal','Title','Heading 1'):
  s=d.styles[name];s.font.name='Times New Roman';s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
  s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.04
  rf=s.element.get_or_add_rPr().get_or_add_rFonts()
  for attr in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+attr),'Times New Roman')
  for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+attr),None)
 d.styles['Title'].font.size=Pt(16);d.styles['Title'].font.bold=True
 d.styles['Heading 1'].font.size=Pt(12);d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(7)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 d.add_paragraph(spec['title'],'Title')
 def text(value):
  for k,label in spec['fields'].items():value=value.replace('{'+k+'}','['+label+']')
  return value
 for block in spec['blocks']:
  kind,value=block
  if kind=='break':d.add_page_break()
  elif kind in ('p','h'):d.add_paragraph(text(value),'Heading 1' if kind=='h' else None)
  elif kind=='table':
   t=d.add_table(rows=0,cols=len(value[0]));t.autofit=False
   for rownum,row in enumerate(value):
    cells=t.add_row().cells
    for c,s in zip(cells,row):
     c.text=text(s);c.vertical_alignment=1
     pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
     for edge in ('top','left','bottom','right'):
      e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
     pr.append(bd)
     for p in c.paragraphs:
      p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
     if rownum==0:
      shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED');pr.append(shade)
      for p in c.paragraphs:
       for r in p.runs:r.bold=True
    t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.comments=''
 d.core_properties.title=spec['title'];d.core_properties.subject=spec['scope']
 d.save(path)

if __name__=='__main__':
 catalog=json.loads((ROOT/'assets/word/fixed_templates_ru.json').read_text())
 for key,spec in catalog.items():
  build(spec,ROOT/'assets/word'/spec['file']);print(key)
