"""Explicit live smoke check. Run manually on a host with existing service credentials.

Uses public fixture content only. Does not import/start the bot, read user storage,
send Telegram messages, or change balances. Calls the real image providers.
Outputs stay in a temporary directory. No automatic invocation at startup.
"""
import ast
import asyncio
import base64
import json
import logging
import os
from pathlib import Path
import sys
import tempfile
import time
import httpx
from PIL import Image
from openai import AsyncOpenAI
from pptx import Presentation
from pptx.util import Inches

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import presentation_visuals as pv
import presentation_designer as pd


def adapters(client):
    names={'generate_image','ask_presentation_vision','_prepare_vision_image','extract_json','_balanced_json_spans'}
    nodes=[n for n in ast.parse((ROOT/'main.py').read_text()).body
           if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef)) and n.name in names]
    assert {n.name for n in nodes}==names
    async def run_heavy(fn,*args):return await asyncio.to_thread(fn,*args)
    ns={'asyncio':asyncio,'base64':base64,'json':json,'os':os,'Image':Image,'httpx':httpx,
        'REPLICATE_API_TOKEN':os.environ['REPLICATE_API_TOKEN'],
        'VISION_MODEL':os.getenv('XAI_VISION_MODEL','grok-4.6'),'client':client,'run_heavy':run_heavy}
    module=ast.Module(body=[ast.ImportFrom(module='__future__',names=[ast.alias(name='annotations')],level=0)]+nodes,type_ignores=[])
    exec(compile(ast.fix_missing_locations(module),'main-adapters.py','exec'),ns)
    return ns


def fixtures():
    rows=[
      ('Океан и его удивительные обитатели','Мир морских животных и его разнообразие.','A sea turtle swimming above a colorful coral reef in clear tropical ocean water, wide underwater photograph'),
      ('Киты','Киты — млекопитающие. Они дышат воздухом и поднимаются к поверхности.','A humpback whale seen underwater, its entire body and long pectoral fins visible in blue ocean water'),
      ('Акулы','Акулы — рыбы. В отличие от китов, они получают кислород из воды с помощью жабр.','A reef shark swimming sideways over the sandy seafloor, full body visible, natural underwater photograph'),
      ('Осьминоги','Осьминоги исследуют окружающий мир щупальцами. Многие виды меняют окраску и прячутся среди камней.','An octopus resting on a rocky seabed, its arms spread naturally over the rocks, detailed underwater photograph'),
      ('Коралловые рифы','Рифы дают укрытие множеству морских животных. Кораллы — это животные, а не растения.','A thriving tropical coral reef with branching corals and small colorful reef fish, clear underwater photograph'),
      ('Глубоководные животные','На больших глубинах солнечного света почти нет. Некоторые животные создают собственное свечение.','A single deep sea anglerfish seen from the side in dark blue water, a small glowing lure above its mouth, realistic natural history illustration'),
      ('Защита океана','Меньше одноразового пластика и ответственное обращение с отходами помогают снижать загрязнение.','Volunteers collecting discarded plastic bottles on a beach, plain unmarked clothing and bags, documentary photograph'),
      ('Океан нуждается в защите','Разные животные приспособились к разным условиям. Сохранение их мест обитания помогает сберечь разнообразие жизни.','A pod of dolphins swimming together beneath the ocean surface, sunbeams in clear blue water, wide underwater photograph')]
    return [{'title':t,'content':c,'image_prompt':p,'designer':{'layout':'photo','blocks':[]}} for t,c,p in rows]


async def main():
    if not all(os.environ.get(k) for k in ('XAI_API_KEY','REPLICATE_API_TOKEN')):
        print('SMOKE_BLOCKED missing_provider_credentials');return 2
    out=Path(tempfile.mkdtemp(prefix='alix-image-smoke-'))
    logging.basicConfig(level=logging.WARNING,format='%(message)s')
    start=time.monotonic()
    async with AsyncOpenAI(api_key=os.environ['XAI_API_KEY'],base_url='https://api.x.ai/v1',max_retries=0) as client:
        ns=adapters(client);slides=fixtures();pd.plan_designer_images(slides,'designer_soft')
        try:
            assets,warnings=await pv.prepare_images('Океан и его удивительные обитатели',slides,[],{},out,
                ns['generate_image'],ns['ask_presentation_vision'],ns['extract_json'])
            status={'slides':len(slides),'images':sum(bool(a) for a in assets),'missing':[i+1 for i,a in enumerate(assets) if not a],
                    'seconds':round(time.monotonic()-start),'warnings':warnings}
            print('SMOKE_IMAGES '+json.dumps(status),flush=True)
            pv.require_images(slides,assets,warnings,True)
            prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
            for i,(item,asset) in enumerate(zip(slides,assets),1):
                pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),pd.build_scene(item,asset,i,len(slides),'ru','designer_soft'))
            path=out/'ocean.pptx';prs.save(path)
            reopened=Presentation(path)
            pictures=[sum(shape.shape_type==13 for shape in slide.shapes) for slide in reopened.slides]
            assert pictures==[1]*8
            print('SMOKE_PASS '+json.dumps({'pictures_per_slide':pictures,'pptx_bytes':path.stat().st_size,'output':str(path)}),flush=True)
            return 0
        except Exception as error:
            print('SMOKE_FAIL '+type(error).__name__,flush=True)
            return 1

if __name__=='__main__':
    raise SystemExit(asyncio.run(main()))
