"""Manual public-fixture check with existing server credentials; no bot/user storage."""
import ast
import asyncio
import json
import os
from pathlib import Path
import sys
import tempfile
import time
from openai import AsyncOpenAI

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import presentation_support as ps
import presentation_designer as pd
import presentation_visuals as pv


def adapters(client):
    names = {'ask_grok', '_create_presentation_plan', 'extract_json', '_balanced_json_spans', 'grok_failed'}
    nodes = [n for n in ast.parse((ROOT/'main.py').read_text()).body
             if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name in names]
    assert {n.name for n in nodes} == names
    ns = dict(vars(ps), asyncio=asyncio, json=json, client=client, GROK_ERROR_PREFIX='MODEL_ERROR:',
              content_gen_lang=lambda *a: 'ru', grok_json_lang_instruction=lambda *a: '\nВсе тексты на русском.')
    module = ast.Module(body=[ast.ImportFrom(module='__future__', names=[ast.alias(name='annotations')], level=0)] + nodes, type_ignores=[])
    exec(compile(ast.fix_missing_locations(module), 'quality-adapters.py', 'exec'), ns)
    return ns


async def main():
    if not os.environ.get('XAI_API_KEY'):
        print('QUALITY_BLOCKED missing_provider_credentials'); return 2
    started = time.monotonic()
    data = {'topic': 'Отели будущего: технологии и комфорт', 'slides': 8,
            'theme_name': 'designer_soft', 'mode': 'ai',
            'extra': 'Умные номера, роботы-помощники, персональный сервис, экологичные решения, отдых гостей. Тёмный фон, бирюзовый акцент, мягкие скругления. Коротко, без вымышленных цифр. Возможные решения обозначай как концепции.'}
    async with AsyncOpenAI(api_key=os.environ['XAI_API_KEY'], base_url='https://api.x.ai/v1', max_retries=0) as client:
        ns = adapters(client)
        plan = await ns['_create_presentation_plan'](data, 'ru')
        print('QUALITY_PLAN ' + json.dumps({'count': len(plan['slides']), 'design': plan['design'], 'titles': [s['title'] for s in plan['slides']]}, ensure_ascii=False), flush=True)
        prompt = ps.content_prompt(data, plan, '\nВсе тексты на русском.') + '\n' + pd.content_instruction()
        for attempt in range(2):
            raw = await ns['ask_grok'](prompt, max_tokens=10000, system_prompt=ps.PRESENTATION_SYSTEM, temperature=0.35)
            if ns['grok_failed'](raw):
                raise RuntimeError('Content provider failed')
            try:
                content = ps.validate_content(ns['extract_json'](raw), plan)
                for i, item in enumerate(content['slides'], 2):
                    if 'designer' not in item:
                        raise ps.PresentationValidationError('Missing designer layout')
                    item['designer'] = pd.validate_designer(item)
                    pd.build_scene(item, None, i, 8, 'ru', 'designer_soft')
                break
            except (ValueError, TypeError, KeyError, AttributeError):
                if attempt: raise
                prompt += '\nВерни исправленный полный JSON: заголовки из плана, короткие тексты, обязательный designer, блоки должны помещаться в макеты.'
        for item in content['slides']:
            print('QUALITY_SLIDE ' + json.dumps({'title': item['title'], 'content': item['content'], 'layout': item['designer']['layout'], 'image': item['image_prompt']}, ensure_ascii=False), flush=True)
        output = Path(tempfile.mkdtemp(prefix='alix-quality-'))/'result.json'
        output.write_text(json.dumps({'plan': plan, 'content': content}, ensure_ascii=False), encoding='utf-8')
        if '--images' in sys.argv:
            from check_presentation_images import adapters as image_adapters
            from pptx import Presentation
            from pptx.util import Inches
            from pptx.dml.color import RGBColor
            image_api = image_adapters(client)
            slides = [{'title': content['title'], 'point': plan['slides'][0]['point'], 'image_prompt': content['cover_image_prompt']}] + content['slides']
            for item in slides:
                item['_designer_background'] = plan['design']['background']
                item['_designer_accent'] = plan['design']['accent']
            pd.plan_designer_images(slides, 'designer_soft')
            assets, warnings = await pv.prepare_images(content['title'], slides, [], {}, output.parent,
                image_api['generate_image'], image_api['ask_presentation_vision'], image_api['extract_json'])
            pv.require_images(slides, assets, warnings, True)
            colors = {'bg': (12,24,36), 'ink': (246,248,250), 'mid': (221,230,237), 'mute': (158,180,195), 'line': (112,213,207)}
            def background(slide, palette, *args, **kwargs):
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(*palette['bg'])
            for variant in ('ordinary', 'designer'):
                prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
                for i, (item, asset) in enumerate(zip(slides, assets), 1):
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    if variant == 'designer':
                        pd.render_scene(slide, pd.build_scene(item, asset, i, 8, 'ru', 'designer_soft'))
                    else:
                        pv.render_slide(slide, item, asset, colors, i, 8, 'ru', background, lambda *a: None)
                path = output.parent/(variant+'.pptx'); prs.save(path)
                reopened = Presentation(path)
                assert len(reopened.slides) == 8
                print('QUALITY_FILE '+json.dumps({'variant': variant, 'bytes': path.stat().st_size, 'images': sum(bool(a) for a in assets)}), flush=True)
        print('QUALITY_PASS ' + json.dumps({'slides': 8, 'seconds': round(time.monotonic()-started), 'output': str(output)}), flush=True)
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(asyncio.run(main()))
    except Exception as error:
        print('QUALITY_FAIL ' + type(error).__name__, flush=True)
        raise SystemExit(1)
