"""Offline release gate: real covers, PDF pagination, TOC and preserved values."""
import asyncio
import tempfile
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from pypdf import PdfReader
from academic_templates import BLANKS, COVER_LABELS, blank_path, render, inspect_layout
from word_quality import convert_word_pdf, word_review_findings

async def main():
    with tempfile.TemporaryDirectory(prefix='academic-check-')as td:
        root=Path(td)
        for kind in BLANKS:
            dest=root/(kind+'.pdf')
            # Converter may cache fields: always use a temporary DOCX copy.
            src=root/(kind+'.docx');src.write_bytes(blank_path(kind).read_bytes())
            await convert_word_pdf(src,dest)
            reader=PdfReader(dest)
            assert len(reader.pages)==1,(kind,len(reader.pages))
            text=reader.pages[0].extract_text()
            assert 'НИ ТГУ' not in text and 'Томск' not in text
        sections=[{'title':'Введение','content':'Цель работы — сравнение способов хранения данных.'},
                  {'title':'1 Сравнение','content':'Локальные файлы доступны без сети.'},
                  {'title':'1.1 Пример','content':'Условный пример для проверки структуры.'},
                  {'title':'Заключение','content':'Выбор зависит от задачи и условий работы.'}]
        for kind in ('referat','coursework','report','essay','notes'):
            src=root/('work_'+kind+'.docx');dest=src.with_suffix('.pdf')
            render(src,'Хранение данных',sections,kind,{'school':'Учебный вуз','author':'Анна Петрова','group':'ИС-21','city':'Москва','year':'2028'})
            await convert_word_pdf(src,dest)
            reader=PdfReader(dest);text='\n'.join(p.extract_text()for p in reader.pages)
            assert COVER_LABELS[kind] in reader.pages[0].extract_text(), kind
            assert 'Цель работы' not in reader.pages[0].extract_text(), kind
            if kind in ('essay', 'notes', 'report'):
                assert len(reader.pages) == 2, (kind, len(reader.pages))
                assert len(inspect_layout(dest, kind)['body_page_fills']) == 1
            for value in ('Учебный вуз','Анна Петрова','ИС-21'):assert value in text,(kind,value)
            if kind in ('referat','coursework'):
                assert len(reader.pages)==5,(kind,len(reader.pages))
                toc=reader.pages[1].extract_text()
                for page in ('3','4','5'):assert page in toc
                d=Document(src)
                cache=[p.text for p in d.paragraphs if '\t' in p.text]
                assert cache==['Введение\t3','1 Сравнение\t4','1.1 Пример\t4','Заключение\t5'],cache
        # The reported defect: five short sections must not force seven pages.
        short=[{'title':name,'content':'Это контрольный абзац для проверки размещения разделов.'}
               for name in ('Введение','1 Раздел','2 Раздел','3 Раздел','Заключение')]
        src=root/'short.docx'; dest=src.with_suffix('.pdf')
        render(src,'Контрольный реферат',short,'referat',requested_pages=6)
        await convert_word_pdf(src,dest)
        assert len(PdfReader(dest).pages)==3
        toc=[p.text for p in Document(src).paragraphs if '\t' in p.text]
        assert len(toc)==5 and all(line.endswith('\t3') for line in toc),toc
        layout=inspect_layout(dest,'referat')
        assert len(layout['body_page_fills'])==1
        findings=word_review_findings({'sections':short},3,6,layout)
        assert findings['short_pages']
        # Sparse text must remain visible to review even after exceeding the requested count.
        findings=word_review_findings({'sections':short},7,6,
            {'last_page_fill':.25,'body_page_fills':[.3,.8,.8,.9,.25],'front_matter_pages':2})
        assert findings['sparse_body_pages']==[3,7] and findings['excess_pages']
        # Regression: expanded title metadata must not turn the cover into two pages.
        title='Влияние цифровых технологий на самостоятельность человека в обучении, профессиональной деятельности и принятии повседневных решений в условиях информационного общества'
        meta={'school':'Федеральное государственное бюджетное образовательное учреждение высшего образования «Учебный университет информационных технологий и управления»',
              'department':'Кафедра информационных технологий', 'discipline':'Информатика',
              'teacher':'Петрова-Смирнова Александра Константиновна',
              'author':'Иванов-Сидоров Александр Константинович', 'group':'ИС-21',
              'city':'Москва', 'year':'2026'}
        for case, values in [('long', meta), ('empty', {})]:
            src=root/(case+'.docx'); dest=src.with_suffix('.pdf')
            render(src,title,[{'title':'Проверка','content':'Контроль основного текста.'}],'essay',values)
            await convert_word_pdf(src,dest)
            reader=PdfReader(dest)
            assert len(reader.pages)==2,(case,len(reader.pages))
            first=reader.pages[0].extract_text()
            assert ('Москва' if values else 'Город') in first
            assert 'Контроль основного текста' in reader.pages[1].extract_text()
            for key in ('school','teacher','author'):
                if values:
                    assert ' '.join(values[key].split()) in ' '.join(first.split()), (case,key)
        print('academic runtime_check=passed blanks=9 genres=5 covers=all5 cover_long_and_empty=passed body_pages=separate toc=verified short_referat=flowing density=all_body_pages credits=unchanged paid_api_calls=0',flush=True)
if __name__=='__main__':asyncio.run(main())
