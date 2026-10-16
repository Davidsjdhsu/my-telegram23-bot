"""Offline build gate: real PPTX -> PDF, no bot startup, users or API access."""
import asyncio
from pathlib import Path
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pypdf import PdfReader
from presentation_pdf import convert_presentation_pdf


async def main():
    with tempfile.TemporaryDirectory(prefix='alix-pdf-smoke-') as td:
        root = Path(td)
        image = root / 'fixture.png'
        Image.new('RGB', (320, 180), (22, 119, 151)).save(image)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(11), Inches(1))
        box.text = 'Проверка PDF: фотографии и оформление'
        box.text_frame.paragraphs[0].font.name = 'Calibri'
        box.text_frame.paragraphs[0].font.size = Pt(30)
        slide.shapes.add_picture(str(image), Inches(1), Inches(2), width=Inches(5))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData(); data.categories = ['A', 'B']; data.add_series('Count', [2, 4])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(10), Inches(5), data)
        source, target = root / 'fixture.pptx', root / 'fixture.pdf'
        prs.save(source)
        await convert_presentation_pdf(source, target, 2)
        pdf = PdfReader(target)
        assert 'Проверка PDF' in pdf.pages[0].extract_text()
        assert 'Count' in pdf.pages[1].extract_text()
        print('presentation_pdf runtime_check=passed pages=2 pictures=yes chart=yes')


if __name__ == '__main__':
    asyncio.run(main())
