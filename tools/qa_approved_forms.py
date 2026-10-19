"""Real word_build + render/PDF path; fake Telegram and balances, no AI/network."""
import sys,asyncio,json,shutil
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT),str(ROOT/'tests')]
from docx import Document
from pypdf import PdfReader
import fixed_word_templates as fixed, car_sale_template as car,word_quality
from test_car_sale_template import CarWorkflowTests
from test_fixed_word_templates import WorkflowTests
OUT=ROOT/'qa'/'filled';OUT.mkdir(parents=True,exist_ok=True)
def text(d):return '\n'.join(n.text or '' for n in d.element.iter() if n.tag.endswith('}t'))
async def main():
    report=[]
    cases=[(k,s['sample'].copy()) for k,s in fixed.CATALOG.items() if k!='services']
    for k in ['services_ru','services_act_ru','rent_flat_ru','claim','act','work_act']:
        if k not in fixed.CATALOG:continue
        spec=fixed.CATALOG[k];v=spec['sample'].copy()
        for key in list(v):
            if key.endswith('_details') or key in ('extra','services','remarks','property','works','work','items','facts'):
                v[key]=(v[key]+'; Дополнительное описание предоставленных данных. '*35)[:1350]+' КОНЕЦ'+key.upper()
        cases.append((k+'__long',v))
    for name,values in cases:
        kind=name.split('__')[0];spec=fixed.CATALOG[kind]
        source='\n'.join(spec['fields'][k]+': '+v for k,v in values.items())
        e,m,state=WorkflowTests().env(kind,text=source)
        if len(source)>4000:
            lines=source.splitlines();cut=1
            while cut<len(lines) and len('\n'.join(lines[:cut+1]))<=4000:cut+=1
            state.data.update(topic='\n'.join(lines[:cut]),extra='\n'.join(lines[cut:]))
            assert len(state.data['extra'])<=12000
        e['prepare_word_files']=word_quality.prepare_word_files;e['convert_word_pdf']=word_quality.convert_word_pdf
        events=[]
        async def capture(path,**kw):
            shutil.copyfile(path,OUT/(name+Path(path).suffix));events.append('send')
        m.answer_document.side_effect=capture;e['spend_credits'].side_effect=lambda *a,**kw:events.append('charge')
        await e['word_build'](m,state);await e['word_build'](m,state)
        assert events==['send','send','charge'],(name,events,str(m.answer.call_args_list)[-500:])
        e['ask_grok'].assert_not_awaited()
        filled=Document(OUT/(name+'.docx'));blank=Document(fixed.ASSETS/spec['file'])
        assert len(filled.tables)==len(blank.tables)
        assert [len(t.rows) for t in filled.tables]==[len(t.rows) for t in blank.tables]
        for p,q in zip(blank.paragraphs,filled.paragraphs):
            if '[' not in p.text:assert p.text==q.text,(kind,p.text)
        pdf=PdfReader(OUT/(name+'.pdf'));pdftext=' '.join(' '.join(p.extract_text().split()) for p in pdf.pages)
        for v in values.values():
            assert v.rstrip('.') in text(filled),(name,'DOCX value',v[:70])
            if '__long' in name and len(v)>900:
                assert v.split()[-1] in ''.join(pdftext.split()),(name,'end marker',v[-60:])
            else:assert ' '.join(v.rstrip('.').split()) in pdftext,(name,v[:70])
        assert '[' not in text(filled),(name,'unfilled placeholder')
        report.append({'kind':name,'pages':len(pdf.pages),'input_chars':len(source),'structure_preserved':True,'charge_after_pair':True})
        print(report[-1],flush=True)
    carvalues={k:'Проверочное значение' for k in car.LABELS}
    carvalues.update(seller_name='Петров Александр Константинович',buyer_name='Иванов Иван Иванович',price='1200000',price_words='один миллион двести тысяч',received_amount='1200000',balance='0')
    source='\n'.join(car.LABELS[k]+': '+v for k,v in carvalues.items())
    e,m,state=CarWorkflowTests().env({'word_kind':'dkp_car','topic':source})
    e['prepare_word_files']=word_quality.prepare_word_files;e['convert_word_pdf']=word_quality.convert_word_pdf
    async def capturecar(path,**kw):shutil.copyfile(path,OUT/('dkp_car'+Path(path).suffix))
    m.answer_document.side_effect=capturecar
    await e['word_build'](m,state);await e['word_build'](m,state)
    assert m.answer_document.await_count==2; e['ask_grok'].assert_not_awaited();e['spend_credits'].assert_called_once()
    report.append({'kind':'dkp_car','pages':len(PdfReader(OUT/'dkp_car.pdf').pages),'input_chars':len(source)})
    (OUT.parent/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
asyncio.run(main())
