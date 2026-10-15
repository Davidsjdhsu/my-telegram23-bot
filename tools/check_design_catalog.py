"""Offline QA: real renderer, existing images, all 600 compositions. No paid calls."""
import sys, json, os, subprocess, re, argparse, shutil
from pathlib import Path
sys.path[:0]=[str(Path(__file__).resolve().parents[1]),str(Path(__file__).resolve().parents[1]/'tests')]
from test_presentation_catalog import specimen
from test_presentation_support import load_functions
from presentation_catalog import TEMPLATES, COMPOSITIONS
from presentation_designer import build_scene, render_scene
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from PIL import Image, ImageDraw
from pypdf import PdfReader
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source-pptx',required=True,help='Existing deck containing portrait and landscape images; used unchanged as offline fixtures')
parser.add_argument('--output',default='qa/catalog')
parser.add_argument('--libreoffice',default=shutil.which('soffice'))
parser.add_argument('--pdftoppm',default=shutil.which('pdftoppm'))
parser.add_argument('--fontconfig')
args=parser.parse_args()
if not args.libreoffice or not args.pdftoppm:parser.error('LibreOffice and pdftoppm are required for visual QA')
root=Path(__file__).resolve().parents[1];qa=Path(args.output).resolve();qa.mkdir(parents=True,exist_ok=True)
# Reuse a user-supplied bot image; source deck is unchanged.
source=Path(args.source_pptx)
prs=Presentation(source)
photo=next(s for slide in prs.slides for s in slide.shapes if s.shape_type==13)
path=qa/('existing-bot-image.'+photo.image.ext);path.write_bytes(photo.image.blob)
iw,ih=photo.image.size;asset={'path':str(path),'width':iw,'height':ih,'origin':'ai'}
wide_photo=next(s for slide in prs.slides for s in slide.shapes if s.shape_type==13 and s.image.size[0]>s.image.size[1])
wide_path=qa/('existing-wide-image.'+wide_photo.image.ext);wide_path.write_bytes(wide_photo.image.blob)
wide_asset={'path':str(wide_path),'width':wide_photo.image.size[0],'height':wide_photo.image.size[1],'origin':'ai'}
env=load_functions(['add_chart','safe_line','_relative_luminance'],{'CategoryChartData':CategoryChartData,'XL_CHART_TYPE':XL_CHART_TYPE,
                  'XL_LEGEND_POSITION':XL_LEGEND_POSITION,'Inches':Inches,'Pt':Pt,'RGBColor':RGBColor})
for key,t in TEMPLATES.items():
    prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
    for comp in COMPOSITIONS:
        item,n=specimen(key,comp);a=(asset if comp=='cover_split' and t['system']!='cinema' else wide_asset) if comp.startswith(('cover','photo')) else None
        scene=build_scene(item,a,n,12)
        render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene,env['add_chart'])
    prs.save(qa/(key+'.pptx'))
print('Created 50 decks / 600 slides using the actual renderer; no provider calls.',flush=True)

pdfs=qa/'pdf';pdfs.mkdir(exist_ok=True)
env=dict(os.environ)
if args.fontconfig:env['FONTCONFIG_FILE']=str(Path(args.fontconfig).resolve())
subprocess.run([args.libreoffice,'-env:UserInstallation='+(qa/'office-profile').as_uri(),
                '--headless','--convert-to','pdf','--outdir',str(pdfs)]+[str(qa/(key+'.pptx')) for key in TEMPLATES],
                env=env,check=True,stdout=subprocess.DEVNULL)
missing=[];pages=0
norm=lambda text:re.sub(r'\W+','',text.casefold())
for key in TEMPLATES:
    prs=Presentation(qa/(key+'.pptx'));doc=PdfReader(pdfs/(key+'.pdf')).pages
    if len(doc)!=12:raise RuntimeError((key,'Wrong PDF page count'))
    pages+=len(doc)
    for i,slide in enumerate(prs.slides):
        printed=norm(doc[i].extract_text())
        for sh in slide.shapes:
            texts=[sh.text] if sh.has_text_frame else [c.text for row in sh.table.rows for c in row.cells] if sh.has_table else []
            for text in texts:
                if text.strip() and norm(text) not in printed:missing.append((key,i+1,text))
    # Inspect every cover and a photo composition, plus all layouts for six different systems.
    selected=list(range(1,13)) if key in ('d01','d06','d11','d12','d21','d44') else [1,3]
    for page in selected:
        subprocess.run([args.pdftoppm,'-f',str(page),'-l',str(page),'-singlefile','-scale-to','1000','-png',
                        str(pdfs/(key+'.pdf')),str(qa/f'{key}-{page:02}')],check=True,stdout=subprocess.DEVNULL)
report={'templates':50,'compositions_per_template':12,'rendered_pages':pages,'missing_text':missing,
        'paid_calls':0,'source':'Production rendering functions with offline fixtures and existing image. Not live Telegram generation.'}
(qa/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
# Contact sheets contain rendered slides, never substitute mockups.
for group in range(5):
    canvas=Image.new('RGB',(1200,5*368),'#E4E4E4');draw=ImageDraw.Draw(canvas)
    for j,key in enumerate(list(TEMPLATES)[group*10:(group+1)*10]):
        im=Image.open(qa/f'{key}-01.png');im.thumbnail((590,332));x=(j%2)*600;y=(j//2)*368
        canvas.paste(im,(x,y));draw.text((x+8,y+338),f'{key} / {TEMPLATES[key]["system"]}',fill='black')
    canvas.save(qa/f'covers-{group+1}.jpg')
for key in ('d01','d06','d11','d12','d21','d44'):
    canvas=Image.new('RGB',(1200,6*352),'#E4E4E4');draw=ImageDraw.Draw(canvas)
    for i in range(12):
        im=Image.open(qa/f'{key}-{i+1:02}.png');im.thumbnail((590,332));x=(i%2)*600;y=(i//2)*352
        canvas.paste(im,(x,y));draw.text((x+8,y+334),COMPOSITIONS[i],fill='black')
    canvas.save(qa/f'{key}-contact.jpg')
print(json.dumps(report,ensure_ascii=False),flush=True)
if missing:raise SystemExit(1)
