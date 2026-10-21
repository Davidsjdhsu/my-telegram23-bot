"""Editable price catalogues: optional discounts, no order quantities or totals."""
import math
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table
from openpyxl.workbook.properties import CalcProperties
from excel_styles import apply_formula_protection, PALETTES, normalize_excel_appearance


def build_price_list(path, title, rows, *, style='standard', palette='green', subtitle=None, number=None, protect=False):
    if not rows:
        raise ValueError('В прайс-листе нет позиций')
    number = number or (lambda value, default=None: default if value in (None, '') else float(value))
    style, palette = normalize_excel_appearance(style, palette)
    designer = style == 'designer'
    dark, accent, soft, stripe = PALETTES[palette]
    has = lambda key: any(row.get(key) not in (None, '') for row in rows)
    discounts = has('discount')
    columns = []
    if has('sku'): columns.append(('sku', 'Артикул', 15))
    columns.append(('name', 'Наименование', 42))
    if has('unit'): columns.append(('unit', 'Ед.', 9))
    columns.append(('price', 'Цена, ₽', 22))
    if discounts:
        columns.extend([('discount', 'Скидка, %', 13), ('final', 'Цена со скидкой, ₽', 23)])
    if has('category'): columns.append(('category', 'Категория', 19))
    if has('note'): columns.append(('note', 'Примечание', 34))
    wb = Workbook()
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    ws = wb.active
    ws.title = 'Прайс-лист'
    ws.sheet_view.showGridLines = not designer
    ws.sheet_properties.tabColor = dark
    start, title_row, header = (2, 2, 6) if designer else (1, 1, 4)
    end = start + len(columns) - 1
    mapping = {key: get_column_letter(start+i) for i, (key, _, _) in enumerate(columns)}
    cache = {}
    if designer: ws.column_dimensions['A'].width = 3
    edge = Side(style='thin', color='A2B3A9')
    border = Border(left=edge, right=edge, top=edge, bottom=edge)

    def text(r, c, value):
        cell = ws.cell(r, c, '' if value is None else str(value))
        cell.data_type = 's'
        return cell

    for r, value, size, bold in [(title_row, title, 17 if designer else 12, True),
                               (title_row+1, subtitle or 'Цены в рублях', 10, False)]:
        ws.merge_cells(start_row=r, start_column=start, end_row=r, end_column=end)
        c = text(r, start, value)
        c.font = Font(name='Arial', size=size, bold=bold, color=dark if bold and designer else '333333')
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = max(32 if designer else 24, math.ceil(len(str(value))/70)*19)
    for col, (key, label, width) in enumerate(columns, start):
        ws.column_dimensions[get_column_letter(col)].width = width
        c = text(header, col, label)
        c.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=dark)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border if not designer else Border(right=Side(style='thin', color='FFFFFF'))
    ws.row_dimensions[header].height = 32
    first, last = header+1, header+len(rows)
    for r, row in enumerate(rows, first):
        price = number(row.get('price'), None)
        discount = number(row.get('discount'), None)
        if price is not None and (not math.isfinite(price) or price < 0):
            raise ValueError('Цена должна быть конечным числом от 0')
        if row.get('discount') not in (None, '') and discount is None:
            raise ValueError('Непонятная скидка')
        if discount is not None and (not math.isfinite(discount) or not 0 <= discount <= 100):
            raise ValueError('Скидка должна быть от 0 до 100 процентов')
        height = 30 if designer else 24
        for col, (key, _, width) in enumerate(columns, start):
            c = ws.cell(r, col)
            c.font = Font(name='Arial', size=11, color='222222')
            c.fill = PatternFill('solid', fgColor=stripe if designer and (r-first)%2 else 'FFFFFF')
            c.alignment = Alignment(vertical='center', wrap_text=True)
            if not designer: c.border = border
            if key in ('price', 'discount', 'final'):
                c.alignment = Alignment(horizontal='right', vertical='center')
                c.number_format = '0.##%' if key == 'discount' else '#,##0.00'
                if key == 'final':
                    p, d = f'{mapping["price"]}{r}', f'{mapping["discount"]}{r}'
                    c.value = (f'=IF(NOT(ISNUMBER({p})),"Укажите цену",IF({p}<0,"Проверьте цену",'
                               f'IF({d}="",ROUND({p},2),IF(NOT(ISNUMBER({d})),"Проверьте скидку",'
                               f'IF(OR({d}<0,{d}>1),"Проверьте скидку",ROUND({p}*(1-{d}),2))))))')
                    cache[c.coordinate] = ('Укажите цену' if price is None else
                        float((Decimal(str(price))*(1-Decimal(str(discount or 0))/100)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)))
                    c.font = Font(name='Arial', size=11, bold=True, color=dark)
                else:
                    c.value = price if key == 'price' else (None if discount is None else discount/100)
                    c.fill = PatternFill('solid', fgColor=soft)
                    c.font = Font(name='Arial', size=11, color='1767B0' if designer else '000000')
            else:
                text(r, col, row.get(key))
                c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
                if key == 'sku': c.number_format = '@'
                value = str(row.get(key) or '')
                lines = sum(max(1, math.ceil(len(line)/max(1, width-3))) for line in value.split('\n'))
                height = max(height, lines*16+8)
        ws.row_dimensions[r].height = height
    # A native Excel Table supports filtering. Sorting locked formulas requires unprotecting.
    ref = f'{get_column_letter(start)}{header}:{get_column_letter(end)}{last}'
    ws.add_table(Table(displayName='PriceCatalogue', ref=ref))
    ws.freeze_panes = f'{mapping["price"]}{first}'
    p = f'{mapping["price"]}{first}'
    ws.conditional_formatting.add(f'{p}:{mapping["price"]}{last}', FormulaRule(
        formula=[f'OR(NOT(ISNUMBER({p})),{p}<0)'],
        fill=PatternFill('solid',fgColor='FFF2CC')))
    if discounts:
        result = f'{mapping["final"]}{first}'
        ws.conditional_formatting.add(f'{result}:{mapping["final"]}{last}', FormulaRule(
            formula=[f'NOT(ISNUMBER({result}))'],
            font=Font(color='9C0006',bold=True),fill=PatternFill('solid',fgColor='FCE4D6')))
    for key, maximum in [('price', 1e15), ('discount', 1)]:
        if key not in mapping: continue
        rule = DataValidation(type='decimal', operator='between', formula1=0, formula2=maximum,
                              allow_blank=True, showErrorMessage=True, errorStyle='stop',
                              error='Введите цену от 0.' if key == 'price' else 'Введите процент от 0% до 100%, например 10%.')
        ws.add_data_validation(rule)
        rule.add(f'{mapping[key]}{first}:{mapping[key]}{last}')
    note_row = last+2
    ws.merge_cells(start_row=note_row, start_column=start, end_row=note_row, end_column=end)
    note = 'Цены и описания можно менять. Пустая цена — не указана. Фильтры в заголовках позволяют отбирать и сортировать позиции.'
    if discounts: note += ' Скидку вводите с %, например 10%. Цена со скидкой пересчитывается автоматически.'
    if protect: note += ' Защита формул включена. Для сортировки строк с формулами снимите защиту листа без пароля.'
    note += ' Эту подсказку можно удалить.'
    c = text(note_row, start, note)
    c.font = Font(name='Arial', size=10, color='66726B')
    c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[note_row].height = math.ceil(len(note)/max(50, sum(x[2] for x in columns)*.9))*16+8
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if sum(x[2] for x in columns)>100 else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=.3,right=.3,top=.3,bottom=.3,header=.1,footer=.1)
    ws.print_title_rows = f'{header}:{header}'
    ws.print_area = f'A1:{get_column_letter(end)}{note_row}'
    apply_formula_protection(wb, protect)
    wb.save(path)
    return cache
