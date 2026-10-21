"""Editable monthly cash model; all calculations are native Excel formulas.

This is a single-currency, single-product cash scenario. Receipts and payments
occur in the month shown. It does not infer tax regimes, FX rates or financing.
"""
import math
import re
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.workbook.properties import CalcProperties
from excel_styles import PALETTES, normalize_excel_appearance

MISSING = 'Заполните данные'
INVALID = 'Проверьте данные'
NO_PAYBACK = 'Не достигнут за период'


class StartupValidationError(ValueError):
    pass


STARTUP_PROMPT = """
Извлеки данные для помесячной денежной модели одного продукта/услуги.
Используй ТОЛЬКО числа пользователя, ничего не придумывай. Неизвестное поле — null,
а явно указанный ноль — 0. Не добавляй расходы, которых пользователь не указал.
Неизвестные суммы расходов — null. Если явно сказано, что постоянных расходов нет,
expenses:[], expenses_confirmed_zero:true; иначе не ставь этот флаг.
Инвестиции investment — разовые затраты до первого месяца. starting_cash — деньги,
доступные ДО этих затрат; не приравнивай их автоматически к инвестициям.
price — цена единицы, monthly_sales — число продаж, variable_cost — затраты на одну
продажу, expenses — только постоянные расходы за месяц, tax_rate — прямо заданный
процент от ВЫРУЧКИ (6 означает 6%). Если налоги не заданы — null; если явно исключены — 0.
Не определяй налоговый режим и не подставляй типовую ставку. Не дублируй затраты:
переменные расходы, налоги с выручки и разовые вложения не входят в expenses.
Для изменений по месяцам используй monthly: список объектов с month (1..36) и только
указанными переопределениями price, monthly_sales, variable_cost, fixed_expenses,
tax_rate. fixed_expenses — общая сумма постоянных расходов этого месяца, не добавка.
horizon_months: указанный срок 1..36, при отсутствии срока — 12. Не сокращай явно
заданный срок. Сохраняй валюту, все числа — в исходной валюте, без конвертации.
currency — код ISO, currencies — все реально встреченные валюты (если нет указания,
RUB). Разные валюты не смешивай: перечисли их в currencies. Никаких курсов от себя.
Если нужны кредиты, погашение долга, НДС, налог с прибыли, несколько продуктов с
разной экономикой или иная неподдержанная схема, верни unsupported:true. Не заменяй
такую задачу упрощённой, не превращай налог с прибыли в налог с выручки.
Только JSON следующей структуры (null означает неизвестное):
{"title":"Название проекта","currency":"RUB","currencies":["RUB"],
 "investment":null,"starting_cash":null,"price":null,"monthly_sales":null,
 "variable_cost":null,"tax_rate":null,"horizon_months":12,
 "expenses":[{"name":"Указанная статья","amount":null}],
 "expenses_confirmed_zero":false,"monthly":[]}
"""

STARTUP_HINT_RU = ('Пришли данные одним сообщением: вложения; деньги на старте до вложений; '
    'цену и количество продаж в месяц; затраты на одну продажу; постоянные расходы списком; '
    'заданный налог с выручки в % (или явно 0, если не учитывать); срок до 36 месяцев и валюту. '
    'Можно указать разные продажи и расходы по месяцам. Неизвестные суммы останутся пустыми для заполнения в Excel.')
STARTUP_HINT_EN = ('Send investment, cash available before investment, unit price, monthly sales, '
    'cost per unit, fixed monthly expenses, your revenue tax rate (or explicitly 0 to exclude it), '
    'period up to 36 months and currency. You can specify different monthly values. Unknown inputs stay blank.')


def normalize_model(data, number=None):
    if not isinstance(data, dict) or data.get('unsupported'):
        raise StartupValidationError('Нужна другая схема расчёта')

    def numeric(value, maximum=1e12):
        if value in (None, ''): return None
        try:
            n = number(value, None) if number else float(str(value).replace(',', '.'))
            if isinstance(value, bool) or n is None or not math.isfinite(n) or not 0 <= n <= maximum:
                raise ValueError()
            return float(n)
        except (ValueError, TypeError, OverflowError):
            raise StartupValidationError('Некорректное исходное число') from None

    currency = str(data.get('currency') or 'RUB').upper().strip()
    aliases = {'RUR':'RUB','₽':'RUB','$':'USD','€':'EUR','РУБ':'RUB'}
    currency = aliases.get(currency, currency)
    currencies = data.get('currencies') or [currency]
    if not isinstance(currencies, list): raise StartupValidationError('Некорректная валюта')
    source_currencies = currencies + [data.get(k) for k in ('investment_currency','price_currency') if data.get(k)]
    expenses = data.get('expenses')
    if expenses is None: expenses = []
    if not isinstance(expenses, list) or len(expenses) > 100:
        raise StartupValidationError('Некорректный список расходов')
    if any(not isinstance(e, dict) for e in expenses): raise StartupValidationError('Некорректная статья расходов')
    source_currencies += [e['currency'] for e in expenses if e.get('currency')]
    seen = {aliases.get(str(c).upper().strip(),str(c).upper().strip()) for c in source_currencies}
    if not re.fullmatch('[A-Z]{3}',currency) or seen != {currency}:
        raise StartupValidationError('Нужна одна валюта без автоматической конвертации')
    horizon = numeric(data.get('horizon_months',12),36)
    if horizon is None: horizon = 12
    if horizon < 1 or not horizon.is_integer(): raise StartupValidationError('Срок должен быть от 1 до 36 месяцев')
    result = {key:numeric(data.get(key),100 if key=='tax_rate' else 1e12) for key in
              ('investment','starting_cash','price','monthly_sales','variable_cost','tax_rate')}
    result.update(currency=currency,horizon=int(horizon))
    result['expenses'] = [dict(name=str(e.get('name') or 'Постоянные расходы'), amount=numeric(e.get('amount'))) for e in expenses]
    if not result['expenses']:
        result['expenses'] = [dict(name='Постоянные расходы',amount=0 if data.get('expenses_confirmed_zero') is True else None)]
    overrides = data.get('monthly') or []
    if not isinstance(overrides,list) or len(overrides)>36: raise StartupValidationError('Некорректный план по месяцам')
    result['monthly']={}
    for row in overrides:
        if not isinstance(row,dict): raise StartupValidationError('Некорректный месяц')
        month=numeric(row.get('month'),36)
        if month is None or not month.is_integer() or not 1<=month<=horizon or int(month) in result['monthly']:
            raise StartupValidationError('Повторный месяц или месяц вне срока расчёта')
        result['monthly'][int(month)]={k:numeric(row[k],100 if k=='tax_rate' else 1e12) for k in
            ('price','monthly_sales','variable_cost','fixed_expenses','tax_rate') if k in row}
    return result


def _money(value):
    return float(Decimal(str(value)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))


def build_startup(path, title, data, *, style='standard', palette='green', number=None):
    model = normalize_model(data,number)
    style,palette = normalize_excel_appearance(style,palette)
    designer=style=='designer'
    dark,accent,soft,stripe=PALETTES[palette]
    unit={'RUB':'₽','USD':'$','EUR':'€','GBP':'£'}.get(model['currency'],model['currency'])
    wb=Workbook();summary=wb.active;summary.title='Итоги'
    inputs=wb.create_sheet('Исходные данные');plan=wb.create_sheet('План')
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    caches={1:{},2:{},3:{}}
    border=Border(**{k:Side(style='thin',color='A2B3A9') for k in ('left','right','top','bottom')})
    header=6 if designer else 4
    title_row=2 if designer else 1

    def text(ws,r,c,value):
        cell=ws.cell(r,c,str(value));cell.data_type='s';return cell

    def row(ws,r,values,*,heading=False,total=False,editable=()):
        for col,value in enumerate(values,1):
            cell=ws.cell(r,col)
            if value is not None:
                if isinstance(value,str):text(ws,r,col,value)
                else:cell.value=value
            cell.font=Font(name='Arial',size=11,bold=heading or total,color='FFFFFF' if heading else dark if total else '222222')
            cell.fill=PatternFill('solid',fgColor=dark if heading else soft if total or col in editable else stripe if designer and r%2 else 'FFFFFF')
            cell.border=border if not designer or heading else Border(bottom=Side(style='hair',color='D9E2DC'))
            cell.alignment=Alignment(horizontal='left' if col==1 else 'right',vertical='center',wrap_text=True)
            cell.protection=Protection(locked=col not in editable)
            cell.number_format='#,##0.00####' if col in editable else '#,##0.00'
            if col in editable and designer:cell.font=Font(name='Arial',size=11,color='1767B0')
        ws.row_dimensions[r].height=34 if heading else 28 if designer else 25

    def merged_label(ws,r,label):
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
        text(ws,r,1,label)
        ws.cell(r,1).alignment=Alignment(vertical='center',wrap_text=True,indent=1)
        ws.row_dimensions[r].height=max(ws.row_dimensions[r].height or 25, math.ceil(len(label)/60)*16+8)

    def formula(ws,r,col,expr,value,index):
        cell=ws.cell(r,col);cell.value=expr;caches[index][cell.coordinate]=value
        cell.font=Font(name='Arial',size=11,bold=True,color=dark)
        cell.alignment=Alignment(horizontal='right',vertical='center',shrink_to_fit=True)

    def guarded(expr, refs, *, nonnegative=False, percent=None, constraints=()):
        nums=','.join(f'ISNUMBER({ref})' for ref in refs)
        valid=','.join(f'{ref}>=0' for ref in refs) if nonnegative else ''
        if percent:valid+=(',' if valid else '')+f'{percent}<=1'
        if constraints:valid+=(',' if valid else '')+','.join(constraints)
        inside=f'IFERROR({expr},"{INVALID}")'
        if valid:inside=f'IF(AND({valid}),{inside},"{INVALID}")'
        return f'=IF(AND({nums}),{inside},"{MISSING}")'

    def validation(ws,cell,maximum=1e12):
        dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=maximum,allow_blank=True,
                          showErrorMessage=True,errorStyle='stop',error='Введите число от 0.' if maximum!=1 else 'Введите процент от 0% до 100%.')
        ws.add_data_validation(dv);dv.add(cell)
        ws.conditional_formatting.add(cell,FormulaRule(formula=[f'OR(NOT(ISNUMBER({cell})),{cell}<0,{cell}>{maximum})'],fill=PatternFill('solid',fgColor='FFF2CC')))

    def setup(ws,widths,subtitle):
        for c,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(c)].width=width
        ws.sheet_view.showGridLines=not designer;ws.sheet_properties.tabColor=dark
        for r,value,size in [(title_row,title,17 if designer else 13),(title_row+1,subtitle,10)]:
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=len(widths))
            cell=text(ws,r,1,value);cell.font=Font(name='Arial',size=size,bold=r==title_row,color=dark if r==title_row else '66726B')
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            ws.row_dimensions[r].height=max(30,math.ceil(len(str(value))/100)*17)
        ws.protection.sheet=True;ws.protection.selectLockedCells=False;ws.protection.selectUnlockedCells=False
        ws.protection.formatColumns=False;ws.protection.formatRows=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.orientation='landscape'
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.page_margins=PageMargins(left=.3,right=.3,top=.3,bottom=.3,header=.1,footer=.1)
        ws.freeze_panes=f'B{header+1}'

    def note(ws,r,message,end):
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=end)
        c=text(ws,r,1,message+' Эту подсказку можно удалить.');c.font=Font(name='Arial',size=10,color='66726B')
        c.alignment=Alignment(wrap_text=True,vertical='center');c.protection=Protection(locked=False)
        ws.row_dimensions[r].height=max(45,math.ceil(len(message)/110)*16+10)
        ws.print_area=f'A1:{get_column_letter(end)}{r}'

    setup(summary,[18,20,20,20,29,12],f'Денежная модель · {model["horizon"]} мес. · Валюта: {model["currency"]}')
    setup(inputs,[14,23,23,23,23,18],f'Заполняйте цветные ячейки · Валюта: {model["currency"]}')
    setup(plan,[12,21,21,21,21,22,24,24],f'Поступления и выплаты в том же месяце · Валюта: {model["currency"]}')
    row(inputs,header,['Базовые параметры',None,None,None,'Значение','Ед.'],heading=True)
    merged_label(inputs,header,'Базовые параметры')
    fields=[('investment','Вложения до первого месяца',unit),('starting_cash','Деньги на старте, до вложений',unit),
            ('price','Цена за единицу',unit),('monthly_sales','Продажи за месяц','ед.'),
            ('variable_cost','Переменные затраты на единицу',unit),('tax_rate','Заданный налог с выручки','')]
    addr={}
    for r,(key,label,u) in enumerate(fields,header+1):
        value=model[key];value=value/100 if key=='tax_rate' and value is not None else value
        row(inputs,r,[label,None,None,None,value,u],editable=(5,));merged_label(inputs,r,label)
        addr[key]=f'$E${r}'
        if key=='tax_rate':inputs.cell(r,5).number_format='0.00####%'
        if key=='monthly_sales':inputs.cell(r,5).number_format='#,##0.######'
        validation(inputs,f'E{r}',1 if key=='tax_rate' else 1e12)
    expense_header=header+8
    row(inputs,expense_header,['Постоянные расходы в месяц',None,None,None,'Сумма',unit],heading=True)
    merged_label(inputs,expense_header,'Постоянные расходы в месяц')
    exp_first=expense_header+1
    for r,e in enumerate(model['expenses'],exp_first):
        row(inputs,r,[e['name'],None,None,None,e['amount'],unit],editable=(1,5));merged_label(inputs,r,e['name']);validation(inputs,f'E{r}')
    exp_last=exp_first+len(model['expenses'])-1;total=exp_last+1
    fixed_values=[e['amount'] for e in model['expenses']]
    fixed=None if None in fixed_values else _money(sum(Decimal(str(v)) for v in fixed_values))
    row(inputs,total,['Итого постоянных расходов',None,None,None,None,unit],total=True);merged_label(inputs,total,'Итого постоянных расходов')
    formula(inputs,total,5,f'=IF(COUNT(E{exp_first}:E{exp_last})={len(fixed_values)},IF(MIN(E{exp_first}:E{exp_last})<0,"{INVALID}",ROUND(SUM(E{exp_first}:E{exp_last}),2)),"{MISSING}")',fixed if fixed is not None else MISSING,2)
    addr['fixed_expenses']=f'$E${total}'
    monthly_header=total+3;first_input=monthly_header+1
    row(inputs,monthly_header,['Месяц',f'Цена, {unit}','Продажи, ед.',f'Затраты/ед., {unit}',f'Постоянные, {unit}','Налог с выручки'],heading=True)
    row(plan,header,['Месяц',f'Выручка, {unit}',f'Постоянные, {unit}',f'Переменные, {unit}',f'Налог, {unit}',f'Результат, {unit}',f'Возврат вложений, {unit}',f'Остаток денег, {unit}'],heading=True)
    month_keys=['price','monthly_sales','variable_cost','fixed_expenses','tax_rate']
    monthly_results=[];cumulative=-model['investment'] if model['investment'] is not None else None
    cash_values=[]
    for month in range(1,model['horizon']+1):
        ir=first_input+month-1;pr=header+month
        row(inputs,ir,[month,None,None,None,None,None],editable=(2,3,4,5,6))
        inputs.cell(ir,1).number_format='0'
        overrides=model['monthly'].get(month,{})
        vals={}
        for col,key in enumerate(month_keys,2):
            val=overrides.get(key, fixed if key=='fixed_expenses' else model.get(key))
            if key=='tax_rate' and val is not None:val/=100
            vals[key]=val
            cell=inputs.cell(ir,col)
            if key in overrides:cell.value=val
            else:
                formula(inputs,ir,col,f'=IF(ISNUMBER({addr[key]}),{addr[key]},"")',val if val is not None else '',2)
                cell.font=Font(name='Arial',size=11,color='1767B0' if designer else '000000')
            if key=='tax_rate':cell.number_format='0.00####%'
            if key=='monthly_sales':cell.number_format='#,##0.######'
            validation(inputs,cell.coordinate,1 if key=='tax_rate' else 1e12)
        row(plan,pr,[month,None,None,None,None,None,None,None]);plan.cell(pr,1).number_format='0'
        refs={k:f"'Исходные данные'!{get_column_letter(c)}{ir}" for c,k in enumerate(month_keys,2)}
        price,qty,var,fx,tax=[vals[k] for k in month_keys]
        fx = None if fx is None else _money(fx)
        revenue=None if price is None or qty is None else _money(Decimal(str(price))*Decimal(str(qty)))
        variable=None if var is None or qty is None else _money(Decimal(str(var))*Decimal(str(qty)))
        taxes=None if revenue is None or tax is None else _money(Decimal(str(revenue))*Decimal(str(tax)))
        result=None if any(v is None for v in (revenue,variable,fx,taxes)) else _money(Decimal(str(revenue))-Decimal(str(variable))-Decimal(str(fx))-Decimal(str(taxes)))
        cumulative=None if cumulative is None or result is None else _money(Decimal(str(cumulative))+Decimal(str(result)))
        cash=None if cumulative is None or model['starting_cash'] is None else _money(Decimal(str(cumulative))+Decimal(str(model['starting_cash'])))
        monthly_results.append(result);cash_values.append(cash)
        formulas=[
            guarded(f'ROUND({refs["price"]}*{refs["monthly_sales"]},2)',[refs['price'],refs['monthly_sales']],nonnegative=True),
            guarded(f'ROUND({refs["fixed_expenses"]},2)',[refs['fixed_expenses']],nonnegative=True),
            guarded(f'ROUND({refs["variable_cost"]}*{refs["monthly_sales"]},2)',[refs['variable_cost'],refs['monthly_sales']],nonnegative=True),
            guarded(f'ROUND(B{pr}*{refs["tax_rate"]},2)',[f'B{pr}',refs['tax_rate']],nonnegative=True,percent=refs['tax_rate']),
            guarded(f'ROUND(B{pr}-C{pr}-D{pr}-E{pr},2)',[f'{c}{pr}' for c in 'BCDE']),
            guarded(f'ROUND(F{pr}-\'Исходные данные\'!{addr["investment"]},2)',[f'F{pr}',f"'Исходные данные'!{addr['investment']}"],constraints=[f"'Исходные данные'!{addr['investment']}>=0"]) if month==1 else guarded(f'ROUND(G{pr-1}+F{pr},2)',[f'G{pr-1}',f'F{pr}']),
            guarded(f'ROUND(G{pr}+\'Исходные данные\'!{addr["starting_cash"]},2)',[f'G{pr}',f"'Исходные данные'!{addr['starting_cash']}"],constraints=[f"'Исходные данные'!{addr['starting_cash']}>=0"]),
        ]
        for col,expr,val in zip(range(2,9),formulas,[revenue,fx,variable,taxes,result,cumulative,cash]):
            formula(plan,pr,col,expr,MISSING if val is None else val,3)
    pf,pl=header+1,header+model['horizon'];last_input=first_input+model['horizon']-1
    row(plan,pl+1,['За период',None,None,None,None,None,None,None],total=True)
    for col in range(2,7):
        letter=get_column_letter(col);v=[caches[3][f'{letter}{r}'] for r in range(pf,pl+1)]
        value=MISSING if any(isinstance(x,str) for x in v) else _money(sum(Decimal(str(x)) for x in v))
        formula(plan,pl+1,col,f'=IF(COUNT({letter}{pf}:{letter}{pl})={model["horizon"]},ROUND(SUM({letter}{pf}:{letter}{pl}),2),"{MISSING}")',value,3)
    for col in (7,8):
        letter=get_column_letter(col);formula(plan,pl+1,col,f'={letter}{pl}',caches[3][f'{letter}{pl}'],3)
    # Summary: references keep edits to month-specific inputs live.
    row(summary,header,['Показатель',None,None,None,'Значение','Ед.'],heading=True)
    merged_label(summary,header,'Показатель')
    sr=header+1;inv=f"'Исходные данные'!{addr['investment']}";fund=f"'Исходные данные'!{addr['starting_cash']}"
    initial=None if model['investment'] is None or model['starting_cash'] is None else _money(model['starting_cash']-model['investment'])
    balance_values=[initial]+cash_values
    payback=MISSING if any(v is None for v in monthly_results) or model['investment'] is None else (0 if model['investment']==0 else next((m for m in range(1,model['horizon']+1) if caches[3][f'G{header+m}']>=0),NO_PAYBACK))
    p=f"'Исходные данные'!{addr['price']}";vc=f"'Исходные данные'!{addr['variable_cost']}";tx=f"'Исходные данные'!{addr['tax_rate']}";fxref=f"'Исходные данные'!{addr['fixed_expenses']}"
    margin=None if any(model[k] is None for k in ('price','variable_cost','tax_rate')) else model['price']*(1-model['tax_rate']/100)-model['variable_cost']
    breakeven=MISSING if fixed is None or margin is None else (math.ceil(fixed/margin) if margin>0 else 0 if margin==0 and fixed==0 else 'Нет положит. маржи')
    metrics=[
        ('Выручка за период',f"='План'!B{pl+1}",caches[3][f'B{pl+1}'],unit),
        ('Результат за период',f"='План'!F{pl+1}",caches[3][f'F{pl+1}'],unit),
        ('Результат после возврата вложений',f"='План'!G{pl}",caches[3][f'G{pl}'],unit),
        ('Остаток денег в конце периода',f"='План'!H{pl}",caches[3][f'H{pl}'],unit),
        ('Деньги после стартовых вложений',guarded(f'ROUND({fund}-{inv},2)',[fund,inv],nonnegative=True),MISSING if initial is None else initial,unit),
        ('Первый возврат вложений, месяц',f'=IF(AND(ISNUMBER({inv}),COUNT(\'План\'!G{pf}:G{pl})={model["horizon"]}),IF({inv}=0,0,IF(COUNTIF(\'План\'!G{pf}:G{pl},">=0")=0,"{NO_PAYBACK}",MATCH(TRUE,INDEX(\'План\'!G{pf}:G{pl}>=0,0),0))),"{MISSING}")',payback,'мес.'),
        ('Безубыточность: продажи при базовых параметрах',guarded(f'IF(({p}*(1-{tx})-{vc})>0,ROUNDUP({fxref}/({p}*(1-{tx})-{vc}),0),IF(AND(({p}*(1-{tx})-{vc})=0,{fxref}=0),0,"Нет положит. маржи"))',[p,vc,tx,fxref],nonnegative=True,percent=tx),breakeven,'ед./мес.'),
        ('Дополнительно нужно денег (минимальный запас)',f'=IF(AND(ISNUMBER(E{sr+4}),COUNT(\'План\'!H{pf}:H{pl})={model["horizon"]}),MAX(0,-MIN(E{sr+4},\'План\'!H{pf}:H{pl})),"{MISSING}")',MISSING if None in balance_values else max(0,-min(balance_values)),unit),
    ]
    for r,(label,expr,value,u) in enumerate(metrics,sr):
        row(summary,r,[label,None,None,None,None,u],total=True);merged_label(summary,r,label);formula(summary,r,5,expr,value,1)
        if u in ('мес.','ед./мес.'):summary.cell(r,5).number_format='0'
    # Negative cash is a financing shortfall, not automatically a failed payback.
    for ws,cols,first,last in [(plan,'FGH',pf,pl),(summary,'E',sr,sr+7)]:
        for col in cols:
            cell=f'{col}{first}';rng=f'{cell}:{col}{last}'
            ws.conditional_formatting.add(rng,FormulaRule(formula=[f'AND(ISNUMBER({cell}),{cell}<0)'],font=Font(color='9C0006',bold=True),fill=PatternFill('solid',fgColor='FCE4D6')))
            ws.conditional_formatting.add(rng,FormulaRule(formula=[f'ISTEXT({cell})'],font=Font(color='9C0006'),fill=PatternFill('solid',fgColor='FFF2CC')))
    note(inputs,last_input+2,'Цветные ячейки доступны для ввода. Месяцы сначала связаны с базовыми параметрами; можно заменить число для отдельного месяца. Пустое поле не равно нулю. Если затраты или налог исключены, явно впишите 0. Налог вводится с %, например 6%. Защита снимается без пароля.',6)
    note(plan,pl+3,'Результат = выручка − постоянные расходы − переменные расходы − заданный налог с выручки. Возврат вложений = накопленный результат минус стартовые вложения. Остаток денег учитывает деньги на старте. Отрицательный остаток означает нехватку средств.',8)
    explanation=('Упрощённый денежный план одного продукта: оплата и затраты приходятся на месяц продажи. '
        'Кредиты, отсрочки, НДС и амортизация не моделируются. Налог — только заданный процент от выручки. '
        'Первый возврат вложений означает первое пересечение нуля; позднее результат может снова стать отрицательным. '
        'Неизвестные данные нужно заполнить на листе «Исходные данные».')
    note(summary,sr+9,explanation,6)
    # Hidden chart-only series use #N/A for unknowns: never plot missing cash as zero.
    for source,target in [('G','I'),('H','J')]:
        plan[f'{target}{header}']=plan[f'{source}{header}'].value
        for r in range(pf,pl+1):
            value=caches[3][f'{source}{r}']
            formula(plan,r,9 if target=='I' else 10,
                    f'=IF(ISNUMBER({source}{r}),{source}{r},NA())',
                    value if isinstance(value,(int,float)) else '#N/A',3)
        plan.column_dimensions[target].hidden=True
    chart=LineChart();chart.title='Возврат вложений и остаток денег';chart.style=12
    chart.visible_cells_only=False
    chart.display_blanks='gap'
    chart.y_axis.title=unit;chart.x_axis.title='Месяц'
    chart.add_data(Reference(plan,min_col=9,max_col=10,min_row=header,max_row=pl),titles_from_data=True)
    chart.set_categories(Reference(plan,min_col=1,min_row=pf,max_row=pl))
    chart.width=23;chart.height=7
    for series,color in zip(chart.series,[dark,accent]):series.graphicalProperties.line.solidFill=color
    chart_row=sr+10
    if any(isinstance(caches[3][f'G{r}'],(int,float)) for r in range(pf,pl+1)):
        summary.add_chart(chart,f'A{chart_row}')
        summary.print_area=f'A1:F{chart_row+13}'
    summary.page_setup.fitToHeight=1
    # Repeated headings on long parameter and monthly plan sheets.
    inputs.print_title_rows=f'{monthly_header}:{monthly_header}'
    plan.print_title_rows=f'{header}:{header}'
    wb.active=0
    wb.save(path)
    return caches
