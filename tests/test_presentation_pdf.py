import asyncio
import io
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, Mock, patch

from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from PIL import Image
import presentation_pdf as pdf
import test_presentation_support as fixtures


class PDFValidationTests(unittest.TestCase):
    def fixture(self, root, *, picture=False):
        prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        if picture:
            path=root/'image.png';Image.new('RGB',(80,40),'blue').save(path)
            slide.shapes.add_picture(str(path),Inches(1),Inches(1))
        source=root/'source.pptx';prs.save(source)
        return source

    def test_rejects_wrong_count_geometry_blank_page_and_lost_photo(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=self.fixture(root);target=root/'slides.pdf'
            c=canvas.Canvas(str(target),pagesize=(960,540));c.drawString(30,30,'Slide');c.showPage();c.save()
            pdf.validate_pdf(target,source,1)
            with self.assertRaisesRegex(pdf.PresentationPDFError,'page_count'):pdf.validate_pdf(target,source,2)
            c=canvas.Canvas(str(target),pagesize=(595,842));c.drawString(30,30,'Text copy');c.showPage();c.save()
            with self.assertRaisesRegex(pdf.PresentationPDFError,'geometry'):pdf.validate_pdf(target,source,1)
            c=canvas.Canvas(str(target),pagesize=(960,540));c.showPage();c.save()
            # ReportLab's blank page still has graphics setup, so require paint operations.
            with self.assertRaises(pdf.PresentationPDFError):pdf.validate_pdf(target,source,1)
            source=self.fixture(root,picture=True)
            c=canvas.Canvas(str(target),pagesize=(960,540));c.drawString(30,30,'Image missing');c.showPage();c.save()
            with self.assertRaisesRegex(pdf.PresentationPDFError,'missing_pdf_images'):pdf.validate_pdf(target,source,1)


class PDFWorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def test_converter_unavailable_stops_before_paid_generation(self):
        env,state,sent=fixtures.BuildTests().build_env(3)
        env['ensure_pdf_converter']=Mock(side_effect=pdf.PresentationPDFError('converter_unavailable'))
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertFalse(sent);env['ask_grok'].assert_not_awaited();env['generate_image'].assert_not_awaited()
        env['spend_credits'].assert_not_called();env['finish_job'].assert_called_once()

    async def test_conversion_failure_sends_no_partial_pair_and_no_charge(self):
        env,state,sent=fixtures.BuildTests().build_env(3)
        env['convert_presentation_pdf']=AsyncMock(side_effect=pdf.PresentationPDFError('conversion_timeout'))
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertFalse(sent);env['spend_credits'].assert_not_called();self.assertTrue(state.data)
        source=env['convert_presentation_pdf'].await_args.args[0]
        self.assertFalse(Path(source).exists())

    async def test_actual_saved_pptx_is_converted_before_delivery(self):
        env,state,sent=fixtures.BuildTests().build_env(3)
        wrapped=env['convert_presentation_pdf'].side_effect
        captured=[]
        async def inspect(source,target,count):
            self.assertFalse(sent)
            captured.append(Path(source).read_bytes())
            await wrapped(source,target,count)
        env['convert_presentation_pdf'].side_effect=inspect
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertEqual(captured[0],sent[0][1]);self.assertEqual(len(sent),2)
        env['spend_credits'].assert_called_once()

    async def test_timeout_kills_child_and_releases_slot(self):
        async def wait():await asyncio.sleep(30)
        process=SimpleNamespace(pid=123456789,returncode=None,wait=AsyncMock(side_effect=wait))
        def killed(*args):process.returncode=-9;process.wait.side_effect=None
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'source.pptx';source.write_bytes(b'fixture')
            with patch.object(pdf,'ensure_pdf_converter',return_value='soffice'), \
                 patch.object(pdf.asyncio,'create_subprocess_exec',AsyncMock(return_value=process)), \
                 patch.object(pdf.os,'killpg',side_effect=killed) as kill:
                with self.assertRaisesRegex(pdf.PresentationPDFError,'conversion_timeout'):
                    await pdf.convert_presentation_pdf(source,Path(td)/'out.pdf',1,timeout=.01)
                kill.assert_called_once()
                self.assertFalse(pdf._slots[asyncio.get_running_loop()].locked())
                self.assertFalse(list(Path(td).glob('pdf-convert-*')))

    async def test_cancellation_kills_child_and_releases_slot(self):
        started=asyncio.Event()
        async def wait():started.set();await asyncio.sleep(30)
        process=SimpleNamespace(pid=123456789,returncode=None,wait=AsyncMock(side_effect=wait))
        def killed(*args):process.returncode=-9;process.wait.side_effect=None
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'source.pptx';source.write_bytes(b'fixture')
            with patch.object(pdf,'ensure_pdf_converter',return_value='soffice'), \
                 patch.object(pdf.asyncio,'create_subprocess_exec',AsyncMock(return_value=process)), \
                 patch.object(pdf.os,'killpg',side_effect=killed) as kill:
                task=asyncio.create_task(pdf.convert_presentation_pdf(source,Path(td)/'out.pdf',1))
                await started.wait();task.cancel()
                with self.assertRaises(asyncio.CancelledError):await task
                kill.assert_called_once();self.assertFalse(pdf._slots[asyncio.get_running_loop()].locked())

    async def test_busy_queue_never_launches_another_converter(self):
        loop=asyncio.get_running_loop();pdf._slots[loop]=asyncio.Semaphore(0)
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'source.pptx';source.write_bytes(b'fixture')
            with patch.object(pdf,'ensure_pdf_converter',return_value='soffice'), \
                 patch.object(pdf.asyncio,'create_subprocess_exec',AsyncMock()) as spawn:
                with self.assertRaisesRegex(pdf.PresentationPDFError,'converter_busy'):
                    await pdf.convert_presentation_pdf(source,Path(td)/'out.pdf',1,queue_timeout=.01)
                spawn.assert_not_awaited()
        del pdf._slots[loop]
