import asyncio
import copy
import json
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from aiohttp import web
from openpyxl import load_workbook
from openpyxl.utils.protection import hash_password
from presentation_orders import PresentationOrders
from order_billing import charge_callback
from excel_password import protection
from studio_excel_catalog import normalize
from studio_private import Studio, COOKIE, PILOT_UID
from studio_outline import OutlinePreview
from studio_word import atomic
from test_office import office_env, COLORS
from test_excel_startup import MODEL
from test_excel_study import MOTION


class Recovery(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.user={'credits':1000,'history':[],'generations':0};self.disk=copy.deepcopy(self.user)
        async def persist():self.disk=copy.deepcopy(self.user);return True
        self.api={'get_user':lambda uid:self.user,'save_users_now':persist,'can_afford':lambda uid,c:self.user['credits']>=c,'ADMIN_IDS':set(),'note_success':Mock()}
        self.send=AsyncMock(return_value=SimpleNamespace(message_id=4));self.notify=AsyncMock()
        self.service=self.restore()
        self.data={'topic':'Тест','slides':6,'pres_request_id':'test-request'}
    def restore(self):
        return PresentationOrders(self.root,self.send,charge_callback(self.api,'presentation_payments','Презентация'),self.api['can_afford'],self.notify)
    async def asyncTearDown(self):self.tmp.cleanup()
    def files(self,job):
        from pptx import Presentation
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[0]);slide.shapes.title.text='Тест'
        deck.save(job.directory/'document.pptx')
        (job.directory/'document.pdf').write_bytes(b'fixture')
        job.row.update(files_ready=True,title='Тест');job.save()
    async def test_restart_retains_request_and_completed_stage_no_regeneration(self):
        job=self.service.get(1,self.data,96);self.files(job);job.row.update(status='building',sent=['pptx']);job.save()
        new=self.restore();self.assertEqual(new.latest(1)['status'],'interrupted')
        await new.notify_restarts();self.notify.assert_awaited_once()
        build=AsyncMock();result=await new.run(1,self.data,96,build)
        self.assertTrue(result['delivered']);build.assert_not_awaited()
        self.assertEqual(self.send.await_count,1);self.assertEqual(self.send.call_args.args[-1],'pdf')
        self.assertEqual(self.user['credits'],904)
        again=self.restore();await again.run(1,self.data,96,build)
        self.assertEqual(self.send.await_count,1);self.assertEqual(self.user['credits'],904)
    async def test_uncertain_send_is_not_repeated_or_charged(self):
        job=self.service.get(1,self.data,96);self.files(job);self.send.side_effect=ConnectionError()
        await self.service.run(1,self.data,96,AsyncMock())
        self.assertEqual(self.user['credits'],1000)
        new=self.restore();result=await new.run(1,self.data,96,AsyncMock())
        self.assertEqual(result['status'],'delivery_unknown');self.assertEqual(self.send.await_count,1)
    async def test_repeated_click_one_build_one_charge(self):
        entered=asyncio.Event();release=asyncio.Event()
        async def build(job):
            entered.set();await release.wait();self.files(job);return await self.service.deliver(job)
        first=asyncio.create_task(self.service.run(1,self.data,96,build));await entered.wait()
        second=await self.service.run(1,self.data,96,build);self.assertEqual(second['status'],'building')
        release.set();await first;self.assertEqual(self.user['credits'],904);self.assertEqual(self.send.await_count,2)
    async def test_pending_billing_resume_and_concurrent_charge_are_idempotent(self):
        self.user['presentation_payments']={'one':{'status':'pending','charged':96,'cost':96,'version':1}}
        charge=charge_callback(self.api,'presentation_payments','Презентация')
        self.assertEqual(await asyncio.gather(charge(1,'one',96,'Тест'),charge(1,'one',96,'Тест')),[96,96])
        self.assertEqual(self.user['credits'],904);self.assertEqual(len(self.user['history']),1)
    async def test_save_failure_then_restart_charges_once(self):
        counter=0
        async def persist():
            nonlocal counter
            counter+=1
            if counter==2:return False
            self.disk=copy.deepcopy(self.user);return True
        self.api['save_users_now']=persist
        charge=charge_callback(self.api,'presentation_payments','Презентация')
        with self.assertRaises(RuntimeError):await charge(1,'one',96,'Тест')
        self.assertEqual(self.disk['credits'],1000)
        self.user.clear();self.user.update(copy.deepcopy(self.disk))
        await charge(1,'one',96,'Тест');await charge(1,'one',96,'Тест')
        self.assertEqual(self.disk['credits'],904)
    async def test_completed_receipt_after_failed_disk_ack_never_double_debits(self):
        count=0
        async def persist():
            nonlocal count
            count+=1;self.disk=copy.deepcopy(self.user);return count!=2
        self.api['save_users_now']=persist
        charge=charge_callback(self.api,'presentation_payments','Презентация')
        with self.assertRaises(RuntimeError):await charge(1,'one',96,'Тест')
        self.user.clear();self.user.update(copy.deepcopy(self.disk))
        await charge(1,'one',96,'Тест');self.assertEqual(self.user['credits'],904)
    async def test_session_and_outline_survive_process_restart_with_owner_checks(self):
        import hashlib
        s=Studio(None,None,None,'https://test',state_root=self.root/'studio')
        atomic(s.state_root/'sessions.json',{hashlib.sha256(b'secret-cookie-token-123456').hexdigest():{'expires':time.time()+600,'user':{'id':PILOT_UID}}})
        request=SimpleNamespace(cookies={COOKIE:'secret-cookie-token-123456'})
        self.assertEqual(s.authorized(request)['user']['id'],PILOT_UID)
        o=OutlinePreview(s,AsyncMock(),AsyncMock())
        o.rows['one']={'id':'one','uid':PILOT_UID,'created':time.time(),'expires':time.time()+600,'status':'building','data':{'pres_plan':{'title':'Тест','audience':'Студенты','slides':[{'title':'Начало'}]}},'cost':96,'palette':{}}
        o.save();new=OutlinePreview(s,AsyncMock(),AsyncMock())
        self.assertEqual(new.rows['one']['status'],'interrupted');self.assertEqual(new.rows['one']['data'],o.rows['one']['data'])
        for archive in s.archives:archive.close()


class Password(unittest.TestCase):
    def test_never_persists_plaintext_and_requires_password_when_enabled(self):
        self.assertEqual(protection({'excel_protect':True,'excel_password':'Test-2026'})['excel_password_hash'],hash_password('Test-2026'))
        with self.assertRaises(ValueError):protection({'excel_protect':True})
        self.assertEqual(protection({'excel_protect':False,'excel_password':'Test-2026'}),{'excel_protect':False,'excel_password_hash':None})
        for kind in ('price_list','family_budget','project_budget','calc_table','startup_model'):
            p=normalize({'kind':kind,'title':'Тест','source':'Мои числа','excel_protect':True,'excel_password':'Test-2026'})
            self.assertNotIn('Test-2026',json.dumps(p));self.assertTrue(p['excel_protect'])
    def test_all_kinds_styles_password_input_unlock_and_cached_values(self):
        env=office_env()
        cases={'expense_estimate':[{'name':'Работа','qty':2,'price':500}],
               'family_budget':[{'name':'Доход','type':'income','amount':1000},{'name':'Расход','type':'expense','amount':200}],
               'project_budget':[{'name':'Доход','type':'income','amount':1000},{'name':'Расход','type':'expense','amount':200}],
               'price_list':[{'name':'Товар','price':100,'discount':10}],'calc_table':MOTION,'startup_model':MODEL}
        for kind,data in cases.items():
            for style in ('standard','designer'):
                with self.subTest(kind=kind,style=style),tempfile.TemporaryDirectory() as d:
                    path=Path(d)/'file.xlsx';kw={'excel_style':style,'excel_protect':True,'excel_password_hash':hash_password('Test-2026')}
                    if kind=='startup_model':env['build_excel_startup'](str(path),'Тест',COLORS,data,**kw)
                    else:env['build_excel_items'](str(path),'Тест',kind,COLORS,data,**kw)
                    wb=load_workbook(path);cache=load_workbook(path,data_only=True)
                    formulas=0
                    for sheet in wb:
                        self.assertTrue(sheet.protection.sheet);self.assertEqual(sheet.protection.password,hash_password('Test-2026'))
                        self.assertNotEqual(sheet.protection.password,hash_password('Wrong-password'))
                        for row in sheet:
                            for c in row:
                                if c.data_type=='f':
                                    formulas+=1;self.assertIsNotNone(cache[sheet.title][c.coordinate].value)
                                elif c.value is not None:self.assertFalse(c.protection.locked)
                    self.assertGreater(formulas,0);wb.close();cache.close()

if __name__=='__main__':unittest.main()

class FullPresentationRecovery(unittest.IsolatedAsyncioTestCase):
 async def test_pdf_failure_resumes_saved_pptx_without_new_text_or_images(self):
  import test_presentation_support as f
  env,state,sent=f.BuildTests().build_env(3)
  with tempfile.TemporaryDirectory() as d:
   charge=AsyncMock(return_value=96)
   send=AsyncMock(return_value=SimpleNamespace(message_id=42))
   service=PresentationOrders(d,send,charge,lambda *a:True,AsyncMock())
   env['_presentation_orders']=service
   env['convert_presentation_pdf']=AsyncMock(side_effect=RuntimeError('fixture pdf outage'))
   message=SimpleNamespace(from_user=SimpleNamespace(id=-987654321),answer=AsyncMock())
   await env['_build_presentation'](message,state)
   row=service.latest(message.from_user.id);self.assertTrue(row['pptx_ready']);self.assertEqual(row['status'],'interrupted')
   before=env['ask_grok'].await_count
   async def convert(pptx,pdf,count):Path(pdf).write_bytes(b'fixture pdf')
   env['convert_presentation_pdf']=AsyncMock(side_effect=convert)
   env['_presentation_orders']=PresentationOrders(d,send,charge,lambda *a:True,AsyncMock())
   result=await env['_build_presentation'](message,state)
   self.assertTrue(result['delivered']);self.assertEqual(env['ask_grok'].await_count,before);charge.assert_awaited_once();self.assertEqual(send.await_count,2)
