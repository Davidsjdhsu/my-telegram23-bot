import io
import unittest
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from presentation_visuals import set_slide_notes, register_notes_master, render_slide
from presentation_designer import build_scene, render_scene


class AppleCompatibilityTests(unittest.TestCase):
    def check_saved(self, prs, expected):
        buf=io.BytesIO(); prs.save(buf); buf.seek(0)
        with ZipFile(buf) as z:
            root=etree.fromstring(z.read('ppt/presentation.xml'))
            entries=root.findall('p:notesMasterIdLst/p:notesMasterId', {'p':root.nsmap['p']})
            self.assertEqual(len(entries),1)
            rels=etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'))
            target=next(r for r in rels if r.get('Id')==entries[0].get(qn('r:id')))
            self.assertTrue(target.get('Type').endswith('/notesMaster'))
            self.assertIn('ppt/'+target.get('Target'),z.namelist())
            tags=[etree.QName(c).localname for c in root]
            self.assertLess(tags.index('sldMasterIdLst'),tags.index('notesMasterIdLst'))
            self.assertLess(tags.index('notesMasterIdLst'),tags.index('sldIdLst'))
        buf.seek(0)
        reopened=Presentation(buf)
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in reopened.slides],expected)

    def test_multiple_notes_are_registered_once_and_survive_save(self):
        p=Presentation()
        for text in ['Источник: пользователь', 'AI illustration']:
            s=p.slides.add_slide(p.slide_layouts[6]);set_slide_notes(s,text);set_slide_notes(s,text)
        self.check_saved(p,['Источник: пользователь','AI illustration'])

    def test_existing_reference_is_preserved_and_missing_slide_master_supported(self):
        root=etree.fromstring('<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst/></p:presentation>')
        register_notes_master(root,'rId1');register_notes_master(root,'rId1')
        self.assertEqual(len(root),2)
        self.assertEqual(root[0].tag,qn('p:notesMasterIdLst'))
        self.assertEqual(len(root[0]),1)

    def test_ordinary_renderer_registers_notes(self):
        from unittest.mock import patch
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
        colors={'ink':(0,0,0),'mute':(90,90,90),'mid':(30,30,30)}
        with patch('presentation_visuals.add_image'):
            render_slide(s,{'title':'Test','content':''},{'width':1600,'height':900,'origin':'ai'},colors,1,1,'ru',lambda *a:None,lambda *a:None)
        self.check_saved(p,[s.notes_slide.notes_text_frame.text])

    def test_designer_renderer_registers_notes(self):
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
        scene=build_scene({'title':'Океан','content':'Защита океана','designer':{'layout':'statement','blocks':[]}},None,1,1,'ru','designer_soft')
        render_scene(s,scene)
        self.check_saved(p,[scene['notes']])

if __name__=='__main__': unittest.main()
