"""Both protection modes preserve formulas/caches and all intended input cells."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from test_office import office_env, COLORS
from test_excel_startup import MODEL
from test_excel_study import MOTION


class ProtectionTests(unittest.TestCase):
    def test_every_kind_and_style_keeps_calculations_and_editable_inputs(self):
        e = office_env()
        rows = {
            'expense_estimate': [dict(name='Работа', qty=2, price=500)],
            'family_budget': [dict(name='Доход', type='income', amount=1000), dict(name='Расход', type='expense', amount=200)],
            'project_budget': [dict(name='Доход', type='income', amount=1000), dict(name='Расход', type='expense', amount=200)],
            'price_list': [dict(name='Товар', price=100, discount=10)],
            'calc_table': MOTION,
            'startup_model': MODEL,
        }
        for kind, data in rows.items():
            for style in ['standard', 'designer']:
                with self.subTest(kind=kind, style=style), tempfile.TemporaryDirectory() as td:
                    snapshots=[]
                    for protect in [False, True]:
                        p=Path(td)/f'{protect}.xlsx'
                        kw=dict(excel_style=style, excel_protect=protect)
                        if kind=='startup_model': e['build_excel_startup'](str(p),'Тест',COLORS,data,**kw)
                        else: e['build_excel_items'](str(p),'Тест',kind,COLORS,data,**kw)
                        wb=load_workbook(p);cached=load_workbook(p,data_only=True)
                        values={};formulas=0
                        for sheet in wb:
                            self.assertEqual(sheet.protection.sheet,protect)
                            self.assertIsNone(sheet.protection.password)
                            for row in sheet:
                                for c in row:
                                    if c.data_type=='f':
                                        formulas+=1
                                        # Monthly inputs are intentionally replaceable default formulas.
                                        override=sheet.title=='Исходные данные' and c.column in range(2,7) and c.row>next(row[0].row for row in sheet if row[0].value=='Месяц')
                                        self.assertEqual(c.protection.locked,not override)
                                        values[(sheet.title,c.coordinate)]=(c.value,cached[sheet.title][c.coordinate].value)
                                    elif c.value is not None:
                                        self.assertFalse(c.protection.locked, (sheet.title,c.coordinate))
                        self.assertGreater(formulas,0)
                        snapshots.append(values)
                        if kind=='startup_model':
                            inputs=wb['Исходные данные'];row=7 if style=='designer' else 5
                            self.assertFalse(inputs.cell(row,5).protection.locked)
                            inputs.cell(row,5).value=123456
                            wb.save(Path(td)/'edited.xlsx')
                            self.assertEqual(load_workbook(Path(td)/'edited.xlsx')['Исходные данные'].cell(row,5).value,123456)
                    self.assertEqual(*snapshots)

    def test_default_is_editable(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'default.xlsx'
            office_env()['build_excel_startup'](str(p),'Тест',COLORS,MODEL)
            self.assertTrue(all(not s.protection.sheet for s in load_workbook(p)))
