import io
import tempfile
import unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import presentation_profiles as pp
import presentation_designer as pd
import presentation_support as ps
import test_presentation_support as f


class ProfileTests(unittest.TestCase):
    def test_photo_shapes_follow_selected_family(self):
        slides=[{'_designer_design':{'family':'editorial'}} for _ in range(5)]
        pd.plan_designer_images(slides,'designer_soft')
        self.assertEqual([s['_designer_aspect_ratio'] for s in slides],['3:4','3:4','16:9','3:4','16:9'])

    def test_long_cover_keeps_subtitle_in_every_family_and_language(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'photo.png';Image.new('RGB',(800,600)).save(path)
            for family in pp.PROFILES:
                for lang,text in [('ru','Современные исследования и результаты '),('en','Research methods and applications '),('zh','现代科学技术的应用和未来发展方向'),('ar','البحث العلمي وتطبيقاته العملية ')]:
                    for origin in (None,'user','ai'):
                        d={'family':family,'background':'dark','accent':'#C48356','secondary':'#EAD5B8'}
                        asset={'path':str(path),'width':800,'height':600,'origin':origin} if origin else None
                        title=(text*8)[:120];subtitle=(text*10)[:180]
                        scene=pd.build_scene({'title':title,'point':subtitle,'_designer_design':d},asset,1,8,lang)
                        actual=''.join(x['text'].replace('\n','') for x in scene['nodes'] if x.get('role')=='body')
                        self.assertEqual(''.join(actual.split()),''.join(subtitle.split()))

    def test_topic_and_audience_selection(self):
        for topic, family in [('Кофейня будущего','editorial'),('Роботы и искусственный интеллект','technology'),
                              ('Компания и её проекты','business'),('Океан и киты','education')]:
            self.assertEqual(pp.choose_profile({'topic':topic}),family)
        self.assertEqual(pp.choose_profile({'topic':'Космос','audience':'ученики 9 класса'}),'education')
        self.assertEqual(pp.choose_profile({'topic':'Кофейня','extra':'Технологический стиль'}),'technology')

    def test_explicit_palette_and_background_override_defaults(self):
        d=pp.choose_design({'topic':'Кофейня будущего','extra':'Тёмный фон, кремовые и медные акценты.'})
        self.assertEqual({k:d[k] for k in ('family','background','accent','secondary')},{'family':'editorial','background':'dark','accent':'#C48356','secondary':'#EAD5B8'})
        self.assertEqual(pp.choose_design({'topic':'Синие киты'})['accent'],pp.TEMPLATES['d03']['accent'])
        self.assertEqual(pp.choose_design({'topic':'Технологии','extra':'Светлый фон. Акцент #991122'})['accent'],'#991122')

    def test_confirmed_design_survives_validation(self):
        d={'topic':'Кофейня','theme_name':'designer_soft','extra':'Тёмный фон, медные акценты'}
        base=f.make_plan();base['audience']='широкая аудитория'
        p=ps.finalize_plan_design(base,d)
        self.assertEqual(ps.validate_plan(p,8)['design'],p['design'])
        self.assertIn('Кофейная редакция',ps.format_plan(p))
        self.assertIn('Georgia',ps.format_plan(p))
        self.assertEqual(ps.finalize_plan_design(f.make_plan(),{'theme_name':'minimal'}),f.make_plan())

    def test_variants_render_without_text_loss_and_preserve_notes_master(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'photo.png';Image.new('RGB',(1000,750),'brown').save(path)
            asset={'path':str(path),'width':1000,'height':750,'origin':'ai'}
            families=[]
            for family in pp.PROFILES:
                d=pp.choose_design({'topic':'Кофейня'});d['family']=family;d.pop('template_id',None)
                # A consistent deck direction, and more than one photo composition.
                geometries=set();fonts=set()
                prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
                for n in range(1,6):
                    item={'title':'Кофе и работа бариста','content':'Бариста регулирует помол. Гость выбирает напиток.',
                          'point':'Как готовят кофе','_designer_design':d}
                    scene=pd.build_scene(item,asset,n,8)
                    geometries.add(tuple(next(x['box'] for x in scene['nodes'] if x['type']=='image')))
                    fonts.add(next(x['font'] for x in scene['nodes'] if x.get('role')=='title'))
                    pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
                self.assertGreaterEqual(len(geometries),3)
                families.append(next(iter(fonts)))
                stream=io.BytesIO();prs.save(stream);stream.seek(0)
                result=Presentation(stream)
                self.assertEqual(len(result.slides),5)
                self.assertIsNotNone(result._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}notesMasterIdLst'))
            self.assertEqual(len(set(families)),4)

    def test_text_contrast_and_dense_multilingual_content(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'photo.png';Image.new('RGB',(800,600)).save(path)
            asset={'path':str(path),'width':800,'height':600,'origin':'user'}
            for family in pp.PROFILES:
                for bg in ('light','dark'):
                    d={'family':family,'background':bg,'accent':'#EAD5B8','secondary':'#EAD5B8'}
                    palette=pp.resolved_profile(d)
                    a,b=pp.relative_luminance(palette['accent']),pp.relative_luminance(palette['bg'])
                    self.assertGreaterEqual((max(a,b)+.05)/(min(a,b)+.05),4.5)
                    body=('Проверенное предложение. '*19).strip()
                    scene=pd.build_scene({'title':'Особенности работы и взаимодействия с гостями','content':body,'_designer_design':d},asset,2,8)
                    actual=' '.join(x['text'] for x in scene['nodes'] if x.get('role')=='body')
                    self.assertEqual(' '.join(actual.split()),body)
                    photo=next(x for x in scene['nodes'] if x['type']=='image')
                    self.assertFalse(photo['crop'])

    def test_full_bot_builder_uses_profile_on_all_slides(self):
        import asyncio
        from types import SimpleNamespace
        env,state,sent=f.BuildTests().build_env(3)
        state.data.update(topic='Кофейня',theme_name='designer_soft',extra='Тёмный фон, медные акценты')
        env['_create_presentation_plan'].return_value['audience']='широкая аудитория'
        original=f.make_content(3)
        for s in original['slides']:s['designer']={'layout':'photo','blocks':[]}
        import json
        env['ask_grok'].return_value=json.dumps(original)
        asyncio.run(env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state))
        self.assertEqual(len(sent),2)
        prs=Presentation(io.BytesIO(sent[0][1]))
        for slide in prs.slides:
            self.assertEqual(str(slide.background.fill.fore_color.rgb),'1F2529')
            title=next(s for s in slide.shapes if s.name=='Designer title')
            self.assertEqual(title.text_frame.paragraphs[0].font.name,'Georgia')
