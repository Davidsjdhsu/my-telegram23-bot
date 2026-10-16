import hashlib
import io
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from presentation_patterns import ROOT, pattern_node
from presentation_designer import build_scene, render_scene
from presentation_style_direction import decorate_fallback
from test_presentation_catalog import specimen


class PatternTests(unittest.TestCase):
    def test_all_100_distinct_assets_survive_export_and_reopen(self):
        hashes=set()
        for style in ('designer_soft','designer_strict'):
            prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
            expected=[]
            for i in range(1,51):
                item,n=specimen(f'd{i:02}','cover_split')
                scene=build_scene(item,None,n,8,style=style)
                self.assertTrue(any(x['type']=='text' for x in scene['nodes']))
                backgrounds=[x for x in scene['nodes'] if x['type']=='background_image']
                self.assertEqual(len(backgrounds),1)
                data=Path(backgrounds[0]['path']).read_bytes()
                hashes.add(hashlib.sha256(data).hexdigest());expected.append(hashlib.sha256(data).hexdigest())
                with Image.open(io.BytesIO(data)) as im:
                    self.assertGreaterEqual(im.width,1280)
                    self.assertAlmostEqual(im.width/im.height,16/9,delta=.02)
                render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
            out=io.BytesIO();prs.save(out);out.seek(0)
            reopened=Presentation(out)
            for slide,digest in zip(reopened.slides,expected,strict=True):
                texture=next(s for s in slide.shapes if s.name=='Designer texture')
                self.assertEqual(hashlib.sha256(texture.image.blob).hexdigest(),digest)
                self.assertTrue(any(s.has_text_frame and s.text for s in slide.shapes))
        self.assertEqual(len(hashes),100)

    def test_retry_keeps_topic_artwork_and_upright_motifs(self):
        item,_=specimen('d18','photo_left');design=item['_designer_design']
        for style in ('designer_soft','designer_strict'):
            nodes=[pattern_node(design,style,n) for n in range(1,11)]
            self.assertEqual(len({x['path'] for x in nodes}),1)
            self.assertTrue(all(not x.get('flip_v') for x in nodes))
            self.assertEqual(nodes,[pattern_node(design,style,n) for n in range(1,11)])

    def test_fallback_keeps_content_and_custom_background(self):
        item,_=specimen('d02','photo_left')
        for custom in (False,True):
            original={'type':'text','role':'body','text':'Текст остаётся целиком','color':'#FFFFFF','box':[64,64,500,200]}
            scene={'background':'#102030','nodes':[original]}
            decorate_fallback(scene,item,None,3,'designer_soft',custom=custom)
            self.assertEqual(scene['nodes'][1]['text'],'Текст остаётся целиком')
            self.assertEqual(scene['pattern_id'],'soft-d02')
            if custom:
                self.assertEqual(scene['background'],'#102030')
                self.assertLessEqual(scene['nodes'][0]['opacity'],.2)

    def test_non_designer_and_missing_artwork_fail_safely(self):
        item,_=specimen('d02','photo_left');design=item['_designer_design']
        self.assertIsNone(pattern_node(design,'minimal',1))
        with patch('pathlib.Path.is_file',return_value=False):
            self.assertIsNone(pattern_node(design,'designer_soft',1))

if __name__=='__main__':unittest.main()
