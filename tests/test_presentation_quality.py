import ast
import json
import re
import unittest
from pathlib import Path
from unittest.mock import AsyncMock
import presentation_visuals as pv
import presentation_designer as pd
import test_presentation_support as fixtures


class QualityTests(unittest.TestCase):
    def style(self,text):
        root=ast.parse(Path('main.py').read_text());ns={'re':re}
        nodes=[n for n in root.body if (isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='STYLE_KEYWORDS' for t in n.targets)) or (isinstance(n,ast.FunctionDef) and n.name=='extract_style')]
        exec(compile(ast.Module(body=nodes,type_ignores=[]),'style','exec'),ns)
        return ns['extract_style'](text)

    def test_designer_inflections_override_topic(self):
        for word in ['дизайнерскую','дизайнерской','дизайнерский','дизайнерские']:
            self.assertEqual(self.style(f'{word} презентация про технологии'), 'designer_soft')
        self.assertEqual(self.style('Мягкие скругления, технологии отелей'), 'designer_soft')
        self.assertEqual(self.style('Строгую дизайнерскую презентацию про технологии'),'designer_strict')
        self.assertEqual(self.style('Без дизайнерского стиля, технологии'),'tech')

    def test_image_ratio_matches_layout_instead_of_slide_number(self):
        self.assertEqual(pv.ordinary_image_ratio({'visual_layout':'wide'},4),'16:9')
        self.assertEqual(pv.ordinary_image_ratio({'visual_layout':'focus'},2),'3:4')
        self.assertEqual(pv.ordinary_image_ratio({},2),'4:3')
        self.assertEqual(pv.ordinary_image_ratio({},5),'4:3')

    def test_headline_shrinks_before_breaking_words(self):
        node=pd.text_node('Коралловые рифы',(0,0,440,244),'#000000',86,42,role='title')
        self.assertIn('Коралловые',node['text'].splitlines())
        self.assertNotIn('Кораллов\n',node['text'])

    def test_requested_dark_designer_background(self):
        scene=pd.build_scene({'title':'Умный номер','content':'Гость выбирает освещение.',
          '_designer_background':'dark','_designer_accent':'#70D5CF'},None,2,8)
        self.assertEqual(scene['background'],'#080B10')
        title=next(n for n in scene['nodes'] if n.get('role')=='title')
        self.assertEqual(title['color'],'#FFFFFF')


class RetryTests(unittest.IsolatedAsyncioTestCase):
    async def test_transient_plan_failure_retries_and_validates(self):
        model=AsyncMock(side_effect=['ERROR',json.dumps(fixtures.make_plan())])
        env=fixtures.load_functions(['_create_presentation_plan'],{
          'ask_grok':model,'grok_failed':lambda x:x=='ERROR','extract_json':json.loads,
          'content_gen_lang':lambda *a:'ru','grok_json_lang_instruction':lambda *a:''})
        result=await env['_create_presentation_plan']({'slides':8},'ru')
        self.assertEqual(len(result['slides']),8)
        self.assertEqual(model.await_count,2)


import asyncio
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import presentation_quality as pq
import presentation_designer as pd
import presentation_profiles as pp
import test_presentation_support as f


class EditorialTests(unittest.IsolatedAsyncioTestCase):
    async def test_invalid_review_retries_and_never_approves(self):
        model=AsyncMock(side_effect=['oops','{"issues":[]}'])
        self.assertEqual(await pq.review_content({}, {},f.make_content(),model,json.loads),[])
        self.assertEqual(model.await_count,2)
        model=AsyncMock(return_value='{"issues":null}')
        with self.assertRaises(pq.DeckQualityError):
            await pq.review_content({}, {},f.make_content(),model,json.loads)
        self.assertEqual(model.await_count,2)

    def test_rejects_invalid_slide_and_unbounded_or_unknown_issue(self):
        issue={'slide':2,'problem':'unsupported','fix':'Make the claim conditional'}
        self.assertEqual(pq.parse_review({'issues':[issue]},{2}),[issue])
        for change in ({'slide':True},{'slide':9},{'problem':'taste'},{'fix':'x'*451}):
            with self.assertRaises(pq.DeckQualityError):pq.parse_review({'issues':[{**issue,**change}]},{2})

    async def test_review_repair_occurs_before_images(self):
        env,state,sent=f.BuildTests().build_env(3)
        review=AsyncMock(side_effect=[[{'slide':2,'problem':'unsupported','fix':'Use a conditional statement'}],[]])
        env['review_content']=review
        original=env['generate_image']
        async def image(*a,**k):
            self.assertEqual(review.await_count,2)
            return await original(*a,**k)
        env['generate_image']=image
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertEqual(len(sent),2);self.assertEqual(env['ask_grok'].await_count,2)
        self.assertIn('Use a conditional statement',env['ask_grok'].await_args.args[0])

    async def test_bad_review_stops_before_images_delivery_billing(self):
        env,state,sent=f.BuildTests().build_env(3)
        env['review_content']=AsyncMock(side_effect=pq.DeckQualityError('unavailable'))
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertEqual(sent,[]);env['generate_image'].assert_not_awaited();env['spend_credits'].assert_not_called()
        env['finish_job'].assert_called_once();self.assertEqual(state.current,'confirm')

    def test_local_final_comparison_and_repetition_rules(self):
        content=f.make_content(4)
        content['slides'][0]['designer']={'layout':'comparison','blocks':[]}
        problems={x['problem'] for x in pq.local_issues(content,True)}
        self.assertEqual(problems,{'comparison','conclusion','repetition'})


class EvidenceTests(unittest.TestCase):
    def test_native_comparison_and_conclusion_preserve_uploaded_photos(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'photo.png';Image.new('RGB',(800,600),'navy').save(path)
            asset={'path':str(path),'width':800,'height':600,'origin':'user'}
            for family in pp.PROFILES:
                for with_photo in (None,asset):
                    item={'title':'Человек и робот','content':'Разделение задач зависит от оборудования.',
                          '_designer_design':{**pp.choose_design({'topic':'Кофе'}),'family':family},
                          'designer':{'layout':'comparison','blocks':[{'heading':'Робот','text':'Выполняет заданную работу.'},{'heading':'Человек','text':'Оценивает ситуацию.'}],
                          'rows':[{'criterion':'Повторяемость','left':'Повторяет заданный маршрут','right':'Меняет действия по ситуации'},
                                  {'criterion':'Исключения','left':'Может требовать помощи оператора','right':'Проверяет причину сбоя'},
                                  {'criterion':'Подготовка','left':'Нуждается в настройке','right':'Нуждается в обучении'}]}}
                    scene=pd.build_scene(item,with_photo,7,8)
                    self.assertEqual(scene['layout'],'comparison')
                    prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
                    pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
                    item['designer']={'layout':'conclusion','blocks':[{'heading':'Где применять','text':'Начните с повторяемых маршрутов.'},{'heading':'Что учесть','text':'Проверьте условия безопасности и исключения.'}]}
                    pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),pd.build_scene(item,with_photo,8,8))
                    stream=io.BytesIO();prs.save(stream);stream.seek(0);saved=Presentation(stream)
                    table=next(s.table for s in saved.slides[0].shapes if s.has_table)
                    self.assertEqual((len(table.rows),len(table.columns)),(4,3))
                    self.assertIn('Помощи'.lower(),table.cell(2,1).text.lower())
                    for slide in saved.slides:
                        pics=[s for s in slide.shapes if s.shape_type==13]
                        self.assertEqual(len(pics),int(with_photo is not None))
                        for pic in pics:self.assertEqual((pic.crop_left,pic.crop_right,pic.crop_top,pic.crop_bottom),(0,0,0,0))
                    self.assertIsNotNone(saved._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}notesMasterIdLst'))

class StructuredCopyTests(unittest.TestCase):
    def test_missing_duplicate_prose_recovers_exact_blocks(self):
        c=f.make_content(3);c['slides'][-1]['content']=''
        c['slides'][-1]['designer']={'layout':'conclusion','blocks':[
            {'heading':'Применение','text':'Выберите повторяемую задачу.'},
            {'heading':'Ограничение','text':'Предусмотрите участие человека при сбое.'}]}
        result=f.ps.validate_content(c,f.ps.validate_plan(f.make_plan(3),3))
        self.assertIn('Предусмотрите участие человека при сбое.',result['slides'][-1]['content'])
        self.assertIn('Применение',pd.designer_copy(result['slides'][-1]))
        c['slides'][-1]['designer']['blocks']=[]
        with self.assertRaises(f.ps.PresentationValidationError):f.ps.validate_content(c,f.ps.validate_plan(f.make_plan(3),3))

    def test_text_companion_keeps_all_comparison_cells(self):
        item={'designer':{'layout':'comparison','blocks':[{'heading':'Первый','text':'Пояснение'},{'heading':'Второй','text':'Пояснение'}],
             'rows':[{'criterion':str(i),'left':'Левый '+str(i),'right':'Правый '+str(i)} for i in range(3)]}}
        copy_text=pd.designer_copy(item)
        for i in range(3):
            self.assertIn('Левый '+str(i),copy_text);self.assertIn('Правый '+str(i),copy_text)


class ComparisonHeaderTests(unittest.TestCase):
    def test_comparison_headers_do_not_require_hidden_duplicate_paragraphs(self):
        item={'designer':{'layout':'comparison','blocks':[{'heading':'Человек','text':''},{'heading':'Робот'}],
             'rows':[{'criterion':str(i),'left':'Слева','right':'Справа'} for i in range(3)]}}
        result=pd.validate_designer(item)
        self.assertEqual([b['text'] for b in result['blocks']],['',''])
        del item['designer']['rows']
        with self.assertRaises(pd.DesignerValidationError):pd.validate_designer(item)

class StructuredModelTests(unittest.IsolatedAsyncioTestCase):
    async def test_presentation_model_isolated_from_general_chat(self):
        from unittest.mock import Mock
        result=SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"issues":[]}'),finish_reason='stop')])
        client=Mock();client.with_options.return_value=client
        client.chat.completions.create=AsyncMock(return_value=result)
        env=f.load_functions(['ask_grok'],{'client':client,'asyncio':asyncio,'REVIEW_SYSTEM':pq.REVIEW_SYSTEM,'GROK_ERROR_PREFIX':'ERROR'})
        await env['ask_grok']('review',system_prompt=pq.REVIEW_SYSTEM)
        args=client.chat.completions.create.await_args.kwargs
        self.assertEqual(args['model'],'grok-4.6');self.assertEqual(args['extra_body'],{'reasoning_effort':'low'})
        self.assertEqual(args['response_format'],{'type':'json_object'})
        await env['ask_grok']('other request')
        args=client.chat.completions.create.await_args.kwargs
        self.assertEqual(args['model'],'grok-4.3');self.assertNotIn('response_format',args)


class RepairContextTests(unittest.IsolatedAsyncioTestCase):
    async def test_repair_receives_its_previous_content_and_targeted_note(self):
        env,state,sent=f.BuildTests().build_env(3)
        env['review_content']=AsyncMock(side_effect=[[{'slide':2,'problem':'wording','fix':'Уточни действие панели'}],[]])
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        prompt=env['ask_grok'].await_args.args[0]
        self.assertIn('Свет передаёт энергию',prompt)
        self.assertIn('Уточни действие панели',prompt)
        self.assertEqual(len(sent),2)

    def test_editor_sees_only_rendered_comparison_cells(self):
        item={'title':'Сравнение','content':'Unused prose','designer':{'layout':'comparison',
              'blocks':[{'heading':'A','text':'Hidden slogan'},{'heading':'B','text':'Hidden slogan'}],
              'rows':[{'criterion':'Условие','left':'Точный факт','right':'Другой факт'}]}}
        text=pq.visible_copy(item)
        self.assertIn('Точный факт',text);self.assertIn('A:',text)
        self.assertNotIn('Hidden slogan',text);self.assertNotIn('Unused prose',text)


class ImageContextTests(unittest.TestCase):
    def test_image_review_includes_original_request_constraints(self):
        import presentation_visuals as pv
        prompt=pv.review_prompt('Роботы',{'title':'Доставка','content':'Робот везёт груз.',
          'image_prompt':'Warehouse robot','_request_context':'Реальные мобильные платформы, без человекоподобных машин.'})
        self.assertIn('Реальные мобильные платформы',prompt)
        self.assertIn('request_context',prompt)

    def test_long_photo_slide_keeps_a_useful_image_size(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'photo.png';Image.new('RGB',(1344,768),'navy').save(path)
            item={'title':'Техника на складе','content':('Повторяемая работа требует настройки. '*12).strip(),
              '_designer_design':pp.choose_design({'topic':'Роботы'}),'designer':{'layout':'photo','blocks':[]}}
            scene=pd.build_scene(item,{'path':str(path),'width':1344,'height':768,'origin':'ai'},3,8)
            image=next(n for n in scene['nodes'] if n['type']=='image')
            self.assertGreaterEqual(image['box'][2],490)


class BoundedEditorialTests(unittest.IsolatedAsyncioTestCase):
    async def test_wording_preferences_do_not_block_after_two_repairs(self):
        env,state,sent=f.BuildTests().build_env(3)
        env['review_content']=AsyncMock(return_value=[{'slide':3,'problem':'wording','fix':'Сделай фразу естественнее'}])
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertEqual(env['review_content'].await_count,3)
        self.assertEqual(len(sent),2);env['spend_credits'].assert_called_once()

    async def test_unsupported_claims_still_block_after_repairs(self):
        env,state,sent=f.BuildTests().build_env(3)
        env['review_content']=AsyncMock(return_value=[{'slide':2,'problem':'unsupported','fix':'Не обещай гарантированный результат'}])
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertEqual(sent,[]);env['spend_credits'].assert_not_called();env['generate_image'].assert_not_awaited()
        self.assertEqual(state.current,'confirm')
