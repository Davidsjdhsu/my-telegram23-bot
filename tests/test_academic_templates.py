import ast
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
import academic_templates as academic
from test_office import office_env
import test_office
from test_presentation_support import load_functions

SECTIONS=[{'title':'Содержание','content':'НЕВЕРНОЕ ОГЛАВЛЕНИЕ 99'},
          {'title':'Введение','content':'Цель исследования.'},
          {'title':'1 Хранение','content':'Сохранённые данные доступны локально.','table':{'headers':['Тип','Пример'],'rows':[['Локальный','Файл']]}},
          {'title':'1.1 Пример','content':'Условный пример.'},
          {'title':'Заключение','content':'Выводы по поставленной задаче.'}]

class LayoutTests(unittest.TestCase):
    def test_all_blanks_are_neutral_editable_and_keep_source_geometry(self):
        self.assertEqual(len(academic.BLANKS),12)
        for key in academic.BLANKS:
            p=academic.blank_path(key);d=Document(p)
            text=' '.join(n.text or '' for n in d.element.iter(qn('w:t')))
            self.assertNotIn('НИ ТГУ',text);self.assertNotIn('ТОМСКИЙ',text);self.assertNotIn('Томск',text)
            self.assertIn('[Название учебного заведения]',text)
            self.assertAlmostEqual(d.sections[0].left_margin.cm,3,places=2)
            with ZipFile(p)as z:self.assertFalse(any(n.startswith('word/media/')for n in z.namelist()))
        self.assertEqual(len(Document(academic.blank_path('assignment')).tables),10)
        with self.assertRaises(ValueError):academic.blank_path('../../main.py')

    def test_real_builder_replaces_model_toc_and_preserves_values_and_tables(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'doc.docx'
            office_env()['build_word'](path,'Тема',SECTIONS,'coursework',{'school':'Мой вуз','author':'Анна Петрова','group':'ИС-21','year':'2028'})
            d=Document(path);text='\n'.join(p.text for p in d.paragraphs)
            self.assertNotIn('НЕВЕРНОЕ',text);self.assertIn('ОГЛАВЛЕНИЕ',text)
            for value in ('Мой вуз','Анна Петрова','ИС-21','2028'):self.assertIn(value,text)
            self.assertNotIn('[',text);self.assertNotIn('Министерство',text)
            self.assertEqual(len(d.tables),1)
            self.assertIn('PAGEREF study_',d.element.xml)
            self.assertTrue(d.sections[0].different_first_page_header_footer)
            self.assertEqual(d.styles['Normal'].font.size.pt,14)
            self.assertEqual(d.styles['Normal'].paragraph_format.line_spacing,1.5)
            self.assertEqual(d.styles['Heading 2'].paragraph_format.left_indent.cm,0)
            self.assertTrue(any(p.paragraph_format.page_break_before for p in d.paragraphs if p.text=='Заключение'))

    def test_essay_and_notes_have_correct_cover_without_toc(self):
        for kind in ('essay','notes'):
            with tempfile.TemporaryDirectory()as td:
                p=Path(td)/'d.docx';office_env()['build_word'](p,'Тема',SECTIONS,kind,{'author':'Автор'})
                d=Document(p);text='\n'.join(p.text for p in d.paragraphs)
                self.assertNotIn('ОГЛАВЛЕНИЕ',text);self.assertNotIn('ТЕМА РАБОТЫ',text)
                self.assertIn('Автор',text)
                self.assertTrue(d.sections[0].different_first_page_header_footer)
                self.assertIn(academic.COVER_LABELS[kind],text)
                if kind == 'essay':
                    self.assertFalse(any(p.style and p.style.name.startswith('Heading') for p in d.paragraphs))

    def test_explicit_corrections_and_empty_values_override_model(self):
        fields=academic.explicit_meta({'school':'Вымышленный вуз','year':'2026'},'Вуз: Мой вуз\nГод: 2028\nФИО: Иван\nГод: ')
        self.assertEqual(fields,{'school':'Мой вуз','year':'','author':'Иван'})

    def test_short_referat_keeps_cover_toc_and_flows_sections_without_erasing_content(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'short.docx'
            office_env()['build_word'](path,'Тема',SECTIONS,'referat',{'author':'Автор'},
                                      extra_data={'academic':True,'requested_pages':6})
            d=Document(path)
            headings=[p for p in d.paragraphs if p.style and p.style.name.startswith('Heading')]
            self.assertTrue(headings[0].paragraph_format.page_break_before)
            self.assertFalse(any(p.paragraph_format.page_break_before for p in headings[1:]))
            self.assertEqual(len(d.tables),1)
            self.assertEqual(d.tables[0].cell(1,1).text,'Файл')
            self.assertIn('ОГЛАВЛЕНИЕ','\n'.join(p.text for p in d.paragraphs))
            self.assertIn('6 страниц основного текста',academic.page_guidance('referat',6))
            self.assertIn('1200 слов',academic.page_guidance('referat',6))
            office_env()['build_word'](path,'Тема',SECTIONS,'coursework',{},
                                      extra_data={'academic':True,'requested_pages':6})
            d=Document(path)
            self.assertTrue(next(p for p in d.paragraphs if p.text=='Заключение').paragraph_format.page_break_before)

class HandlerTests(unittest.IsolatedAsyncioTestCase):
    async def test_actual_word_handler_preserves_metadata_and_delivers_before_charge(self):
        for kind in academic.KINDS:
            e,m,state=test_office.WordWorkflowTests().env()
            from word_quality import prepare_word_files
            async def prepare(*args, **kwargs):
                return await prepare_word_files(*args, **kwargs, measure=lambda _: academic.front_matter_pages(kind) + 5)
            e['prepare_word_files'] = prepare
            data={'word_kind':kind,'word_pages':5,'topic':'Тема: Файлы\nВуз: Мой вуз\nФИО: Анна Петрова\nГруппа: ИС-21\nКурс: 1\nГод: 2028\nДисциплина: Информатика\nПреподаватель: Ирина Петрова\nГород: Москва'}
            state.get_data.return_value=data
            e.update(WORD_KIND_LABELS={k:{'ru':k}for k in academic.KINDS})
            e['ask_grok'].return_value=json.dumps({'title':'Файлы','meta':{'school':'Ошибочный вуз'},'sections':SECTIONS},ensure_ascii=False)
            async def heavy(fn,*args,**kwargs):return fn(*args,**kwargs)
            e['run_heavy']=heavy
            delivered=[]
            async def capture(path,**kwargs):
                if str(path).endswith('.docx'):
                    doc=Document(path);text=' '.join(p.text for p in doc.paragraphs)
                    self.assertIn('Мой вуз',text);self.assertIn('Анна Петрова',text);self.assertNotIn('Ошибочный вуз',text)
                delivered.append(path)
            m.answer_document.side_effect=capture
            e['spend_credits'].side_effect=lambda *a,**kw:self.assertEqual(len(delivered),2)
            await e['word_build'](m,state)
            self.assertEqual(len(delivered),2,kind)
            e['spend_credits'].assert_called_once()
            self.assertIn('НЕ включай титульный',e['ask_grok'].call_args_list[0].args[0])

    async def test_catalog_callback_sends_all_blanks_free_and_releases_job(self):
        env={'start_job':Mock(return_value=(True,'')),'finish_job':Mock(),'FSInputFile':lambda p,**kw:p}
        fn=load_functions(['academic_blank'],env)['academic_blank']
        for key in academic.BLANKS:
            c=SimpleNamespace(data='study_blank:'+key,from_user=SimpleNamespace(id=123),answer=AsyncMock(),message=SimpleNamespace(answer_document=AsyncMock()))
            await fn(c)
            self.assertTrue(c.message.answer_document.call_args.args[0].is_file())
        self.assertEqual(env['finish_job'].call_count,12)
        c.message.answer_document.side_effect=RuntimeError('offline')
        with self.assertRaises(RuntimeError):await fn(c)
        self.assertEqual(env['finish_job'].call_count,13)
        c.data='study_blank:../../main.py';c.message.answer_document.reset_mock()
        await fn(c);c.message.answer_document.assert_not_awaited()
