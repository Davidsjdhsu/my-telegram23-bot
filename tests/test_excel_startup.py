import copy,json,os,tempfile,unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock,AsyncMock
from openpyxl import load_workbook
from excel_startup import normalize_model,StartupValidationError,MISSING,NO_PAYBACK
from excel_styles import PALETTES,STYLE_LABELS,PALETTE_LABELS,normalize_excel_appearance
from test_office import office_env,COLORS
from test_presentation_support import load_functions

MODEL=dict(investment=500000,starting_cash=600000,price=1000,monthly_sales=300,
           variable_cost=400,tax_rate=6,horizon_months=12,currency='RUB',
           expenses=[dict(name='Аренда',amount=40000),dict(name='Зарплата',amount=60000)])

def metrics(w):
    return {r[0].value:r[4].value for r in w['Итоги'] if r[0].value}

def build(path,data,**kw):
    office_env()['build_excel_startup'](str(path),'Мастерская',COLORS,data,**kw)

class StartupTests(unittest.TestCase):
    def test_six_styles_formulas_cached_results_and_editable_months(self):
        for style,palette in [('standard','copper')]+[('designer',p) for p in PALETTES]:
            with self.subTest(style=style,palette=palette),tempfile.TemporaryDirectory() as td:
                p=Path(td)/'model.xlsx';build(p,MODEL,excel_style=style,excel_palette=palette)
                w=load_workbook(p);c=load_workbook(p,data_only=True);m=metrics(c)
                self.assertEqual(w.sheetnames,['Итоги','Исходные данные','План'])
                self.assertEqual(m['Выручка за период'],3600000)
                self.assertEqual(m['Результат за период'],744000)
                self.assertEqual(m['Результат после возврата вложений'],244000)
                self.assertEqual(m['Остаток денег в конце периода'],844000)
                self.assertEqual(m['Первый возврат вложений, месяц'],9)
                self.assertEqual(m['Безубыточность: продажи при базовых параметрах'],186)
                self.assertEqual(m['Дополнительно нужно денег (минимальный запас)'],0)
                h=6 if style=='designer' else 4
                self.assertEqual(w['Итоги'].cell(h,1).fill.fgColor.rgb[-6:],PALETTES[palette if style=='designer' else 'green'][0])
                self.assertEqual(w['План'].cell(h+1,2).data_type,'f')
                self.assertEqual(c['План'].cell(h+1,5).value,18000)
                s=w['Исходные данные'];mh=next(r[0].row for r in s if r[0].value=='Месяц')
                self.assertFalse(s.cell(mh+1,2).protection.locked)
                self.assertEqual(s.cell(mh+1,2).data_type,'f')
                self.assertFalse(s.cell(h+1,5).protection.locked)
                self.assertTrue(w['План'].cell(h+1,2).protection.locked)
                self.assertEqual(w.calculation.calcMode,'auto');self.assertEqual(len(w['Итоги']._charts),1)
                self.assertNotIn('#REF!',str(w['План'].cell(h+1,2).value))

    def test_unknown_costs_and_taxes_stay_missing_zero_is_explicit(self):
        data=copy.deepcopy(MODEL);data['variable_cost']=None;data['tax_rate']=None
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'model.xlsx';build(p,data);w=load_workbook(p,data_only=True);m=metrics(w)
            self.assertEqual(m['Результат за период'],MISSING)
            self.assertEqual(m['Выручка за период'],3600000)
            self.assertEqual(m['Первый возврат вложений, месяц'],MISSING)
            self.assertEqual(w['План']['I5'].value,'#N/A')
            self.assertEqual(w['План']['I5'].data_type,'e')
            self.assertTrue(w['План'].column_dimensions['I'].hidden)
            self.assertEqual(len(w['Итоги']._charts),0)
        self.assertIsNone(normalize_model({})['expenses'][0]['amount'])
        self.assertEqual(normalize_model(dict(expenses=[],expenses_confirmed_zero=True))['expenses'][0]['amount'],0)
        data.update(variable_cost=0,tax_rate=0,investment=0,starting_cash=0,expenses=[],expenses_confirmed_zero=True)
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'model.xlsx';build(p,data);m=metrics(load_workbook(p,data_only=True))
            self.assertEqual(m['Результат за период'],3600000)
            self.assertEqual(m['Первый возврат вложений, месяц'],0)

    def test_monthly_overrides_preserve_values_and_cash_shortfall(self):
        data=copy.deepcopy(MODEL);data.update(starting_cash=300000,monthly=[dict(month=1,monthly_sales=100),dict(month=2,monthly_sales=200)])
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'model.xlsx';build(p,data);w=load_workbook(p,data_only=True);m=metrics(w)
            self.assertEqual(w['План']['F5'].value,-46000)
            self.assertEqual(w['План']['F6'].value,8000)
            self.assertEqual(m['Результат за период'],582000)
            self.assertEqual(m['Первый возврат вложений, месяц'],11)
            self.assertEqual(m['Дополнительно нужно денег (минимальный запас)'],246000)

    def test_no_positive_margin_never_claims_feasible_breakeven(self):
        data=copy.deepcopy(MODEL);data.update(variable_cost=1000)
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'model.xlsx';build(p,data);m=metrics(load_workbook(p,data_only=True))
            self.assertEqual(m['Первый возврат вложений, месяц'],NO_PAYBACK)
            self.assertEqual(m['Безубыточность: продажи при базовых параметрах'],'Нет положит. маржи')

    def test_fractional_costs_round_consistently_with_monthly_cash(self):
        data=dict(MODEL,investment=0,starting_cash=0,price=9.99,monthly_sales=3,
                  variable_cost=2.555,tax_rate=6,horizon_months=1,monthly=[dict(month=1,fixed_expenses=12.345)])
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'model.xlsx';build(p,data);w=load_workbook(p,data_only=True)
            self.assertEqual(w['План']['C5'].value,12.35)
            self.assertEqual(w['План']['D5'].value,7.67)
            self.assertEqual(metrics(w)['Результат за период'],8.15)

    def test_invalid_currencies_numbers_and_months_rejected(self):
        cases=[dict(price=-1),dict(tax_rate=101),dict(price=float('inf')),dict(horizon_months=37),
               dict(horizon_months=2.5),dict(currencies=['USD','RUB']),dict(unsupported=True),
               dict(monthly=[dict(month=13,price=1)]),dict(monthly=[dict(month=1),dict(month=1)]),
               dict(expenses=[dict(name='Аренда',amount=-1)]),dict(starting_cash=True)]
        for extra in cases:
            with self.subTest(extra=extra),self.assertRaises(StartupValidationError):normalize_model(dict(MODEL,**extra))
        self.assertEqual(normalize_model(dict(MODEL,currency='USD',currencies=['USD']))['price'],1000)

    def test_literal_text_and_missing_cash_do_not_change_payback(self):
        data=copy.deepcopy(MODEL);data['starting_cash']=None;data['expenses'][0]['name']='=1+1'
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'model.xlsx';build(p,data);w=load_workbook(p);m=metrics(load_workbook(p,data_only=True))
            self.assertEqual(m['Первый возврат вложений, месяц'],9)
            self.assertEqual(m['Остаток денег в конце периода'],MISSING)
            cell=next(c for row in w['Исходные данные'] for c in row if c.value=='=1+1');self.assertEqual(cell.data_type,'s')

class StartupFlowTests(unittest.IsolatedAsyncioTestCase):
    async def test_miniapp_calls_real_builder_and_charges_only_after_delivery(self):
        for fail,invalid,protect in [(f,i,p) for f,i in [(False,False),(True,False),(False,True)] for p in [False,True,'false']]:
            data={}
            async def update(**kw):data.update(kw)
            state=SimpleNamespace(update_data=update,get_data=AsyncMock(side_effect=lambda:data),clear=AsyncMock(),set_state=AsyncMock())
            order=[];user={'generations':0,'history':[]}
            async def send(file,**kw):
                w=load_workbook(file);self.assertTrue(all(s.protection.sheet == (protect is True) for s in w));self.assertEqual(w['Итоги']['A6'].fill.fgColor.rgb[-6:],PALETTES['burgundy'][0]);order.append('send')
                if fail:raise RuntimeError('test send failure')
            msg=SimpleNamespace(from_user=SimpleNamespace(id=-903045),answer=AsyncMock(),answer_document=send)
            async def run(fn,*args,**kw):return fn(*args,**kw)
            ask=AsyncMock(return_value=json.dumps(dict(MODEL,unsupported=invalid,title='Мастерская')))
            env=load_functions(['excel_build','_handle_miniapp_payload'],dict(ensure_answerable=lambda m:m,user_lang=lambda _: 'ru',
                get_user=lambda _:user,_clip_text=lambda v,n:(v or '')[:n],can_afford=lambda *a:True,CREDIT_COSTS={'excel':5},EXCEL_KIND_DESC={'startup_model':'Финмодель'},
                normalize_excel_appearance=normalize_excel_appearance,start_job=lambda _: (True,''),finish_job=Mock(),
                tr=lambda *a,**kw:'notice',pick_theme=lambda _:('default',COLORS),grok_json_lang_instruction=lambda _: '',
                content_gen_lang=lambda *a,**kw:'ru',ask_grok=ask,extract_json=json.loads,run_heavy=run,
                build_excel_startup=office_env()['build_excel_startup'],safe_filename=lambda *a,**kw:'startup',
                FSInputFile=lambda path,**kw:path,datetime=datetime,os=os,spend_credits=lambda *a,**kw:order.append('charge'),
                note_success=Mock(),signature_line=AsyncMock(return_value=''),main_kb=lambda *a,**k:None,
                Form=SimpleNamespace(waiting_excel_confirm='confirm'),excel_confirm_kb=lambda _:None))
            await env['_handle_miniapp_payload'](msg,state,dict(kind='startup_model',content='Данные',mode='ai',excel_style='designer',excel_palette='burgundy',excel_protect=protect,excel_password='Test-2026' if protect else ''),'gen_excel')
            self.assertEqual(data['excel_mode'],'user')
            self.assertEqual(order,[] if invalid else ['send'] if fail else ['send','charge'])
            if invalid:state.clear.assert_not_awaited();state.set_state.assert_awaited_with('confirm')
            self.assertIn('ничего не придумывай',ask.call_args.args[0])
            if fail:os.remove('/tmp/xl_-903045.xlsx')

    async def test_chat_style_routes_to_real_startup_data(self):
        for style in ['Стандартный','Дизайнерский']:
            data={'excel_kind':'startup_model'}
            async def update(**kw):data.update(kw)
            state=SimpleNamespace(update_data=update,get_data=AsyncMock(side_effect=lambda:data),set_state=AsyncMock())
            msg=SimpleNamespace(text=style,from_user=SimpleNamespace(id=1),answer=AsyncMock())
            env=load_functions(['excel_style_choice','excel_palette_choice'],dict(user_lang=lambda _:'ru',
                EXCEL_STYLES=STYLE_LABELS,EXCEL_PALETTES=PALETTE_LABELS,
                Form=SimpleNamespace(waiting_excel_palette='palette',waiting_excel_startup_data='startup'),
                excel_appearance_kb=lambda *a:None,cancel_kb=lambda *a:None))
            await env['excel_style_choice'](msg,state)
            if style=='Дизайнерский':
                msg.text='Графит и медь';await env['excel_palette_choice'](msg,state)
            state.set_state.assert_awaited_with('startup');self.assertEqual(data['excel_mode'],'user')
            self.assertIn('деньги на старте',msg.answer.call_args.args[0])
