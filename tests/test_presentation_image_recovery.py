"""Regression cases for the real eight-slide order delivered without pictures."""
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from PIL import Image
from pptx import Presentation
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import presentation_visuals as pv
import presentation_designer as pd
import presentation_support as ps
import test_presentation_support as f

GOOD={'relevant':True,'has_text':False,'has_chart':False,'defective':False}

class ReviewRecovery(unittest.IsolatedAsyncioTestCase):
    async def test_transient_review_retries_same_image_not_generator(self):
        for first in [TimeoutError(), ConnectionError('SECRET'), 'not JSON', '{}', json.dumps({**GOOD,'has_text':'false'})]:
            with self.subTest(first=type(first).__name__), tempfile.TemporaryDirectory() as td:
                async def gen(prompt,path,**kwargs):
                    Image.new('RGB',(96,72),'blue').save(path)
                    return True
                generator=AsyncMock(side_effect=gen)
                inspector=AsyncMock(side_effect=[first,json.dumps(GOOD)])
                with self.assertLogs(pv.log,level='WARNING') as logs:
                    assets,w=await pv.prepare_images('Dogs',[{'title':'Dogs'}],[],{},td,generator,inspector,json.loads)
                self.assertNotIn('SECRET',' '.join(logs.output))
                self.assertEqual(generator.await_count,1)
                self.assertEqual(inspector.await_count,2)
                self.assertEqual(inspector.await_args_list[0].args[0],inspector.await_args_list[1].args[0])
                self.assertGreaterEqual(inspector.await_args.kwargs['max_tokens'],512)
                self.assertEqual(assets[0]['origin'],'ai')
                self.assertEqual(w['ai_omitted'],0)

    async def test_missing_required_photo_never_delivered_or_charged(self):
        for designer in [False,True]:
            env,state,sent=f.BuildTests().build_env(8)
            if designer:
                state.data['theme_name']='designer_soft'
                state.data['pres_plan_key']=ps.plan_key(state.data)
                content=f.make_content(8)
                for slide in content['slides']:slide['designer']={'layout':'auto','blocks':[]}
                env['ask_grok']=AsyncMock(return_value=json.dumps(content))
            env['generate_image']=AsyncMock(return_value=False)
            before=copy.deepcopy(state.data)
            await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)),state)
            self.assertEqual(sent,[])
            env['spend_credits'].assert_not_called()
            env['note_success'].assert_not_called()
            env['finish_job'].assert_called_once()
            self.assertEqual(state.data,before)
            self.assertEqual(state.current,'confirm')
            self.assertIn('Кредиты не списаны',env['bot'].send_message.await_args.args[1])

    async def test_partial_images_also_block_delivery(self):
        env,state,sent=f.BuildTests().build_env(3)
        original=env['generate_image']
        async def gen(prompt,path,**kw):
            if 'generated-1-' in path:return False
            return await original(prompt,path,**kw)
        env['generate_image']=gen
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)),state)
        self.assertEqual(sent,[]);env['spend_credits'].assert_not_called()

    async def test_same_draft_can_succeed_after_retry(self):
        env,state,sent=f.BuildTests().build_env(3)
        original=env['generate_image'];env['generate_image']=AsyncMock(return_value=False)
        m=SimpleNamespace(from_user=SimpleNamespace(id=-987654321))
        await env['_build_presentation'](m,state)
        env['generate_image']=original
        await env['_build_presentation'](m,state)
        self.assertEqual(len(sent),2)
        deck=Presentation(io.BytesIO(sent[0][1]))
        self.assertEqual([sum(s.shape_type==13 for s in slide.shapes) for slide in deck.slides],[1,1,1])
        env['spend_credits'].assert_called_once()
        self.assertFalse(state.data)

    async def test_ai_choice_persists_for_retry(self):
        env=f.load_functions(['photo_source_ai'],{'_build_presentation':AsyncMock()})
        state=f.State({})
        await env['photo_source_ai'](None,state)
        self.assertEqual(state.data['photo_mode'],'ai')
        env['_build_presentation'].assert_awaited_once()

    async def test_serialized_deck_missing_images_is_not_delivered(self):
        env,state,sent=f.BuildTests().build_env(3)
        env['render_slide']=lambda *args: None
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)),state)
        self.assertEqual(sent,[])
        env['spend_credits'].assert_not_called()
        self.assertEqual(state.current,'confirm')

class PhotoPlanning(unittest.TestCase):
    def test_dogs_layouts_get_four_photos_without_destroying_diagrams(self):
        slides=[{'title':'Dogs'}]
        for layout in ['auto','statement','sections','comparison','statement','process','auto']:
            slides.append({'title':layout,'designer':{'layout':layout,'blocks':[]}})
        before=copy.deepcopy(slides)
        pd.plan_designer_images(slides,'designer_soft')
        self.assertEqual(sum(not s['_skip_ai_image'] for s in slides),4)
        for i in [4,6]:self.assertTrue(slides[i]['_skip_ai_image'])
        for i in range(1,8):self.assertEqual(slides[i]['designer'],before[i]['designer'])

    def test_gate_allows_real_diagrams_without_photos(self):
        slides=[{'title':'cover'},{'title':'process','_skip_ai_image':True}]
        pv.require_images(slides,[{'origin':'ai'},None],{},True)
        for assets in [[],[None,None],[None]]:
            with self.assertRaises(pv.ImageDeliveryError):pv.require_images(slides,assets,{})

    def test_three_sections_use_three_equal_columns(self):
        blocks=[{'heading':h,'text':'Короткое объяснение.'} for h in ['Первое','Второе','Третье']]
        scene=pd.build_scene({'title':'Три категории','content':'Три объяснения.','designer':{'layout':'sections','blocks':blocks}},None,2,8)
        heads=[n for n in scene['nodes'] if n.get('role')=='subheading']
        self.assertEqual(len(heads),3)
        self.assertEqual(len({n['box'][1] for n in heads}),1)
        self.assertEqual(len({n['box'][0] for n in heads}),3)
