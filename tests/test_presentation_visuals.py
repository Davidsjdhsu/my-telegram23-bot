import asyncio
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, Mock
from types import SimpleNamespace
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import presentation_visuals as pv
import test_presentation_support as fixtures

GOOD={'relevant':True,'has_text':False,'has_chart':False,'defective':False}

class VisualTests(unittest.TestCase):
    def test_verdict_requires_typed_complete_json(self):
        self.assertTrue(pv.approved(GOOD))
        for v in [None,{}, {**GOOD,'has_text':True},{**GOOD,'has_chart':True},
                  {**GOOD,'relevant':False},{**GOOD,'defective':True},{**GOOD,'has_text':'false'}]:
            self.assertFalse(pv.approved(v))

    def test_fit_preserves_full_frame_and_aspect(self):
        for width,height in [(1600,900),(600,1600),(1000,1000),(3000,500)]:
            x,y,w,h=pv.fit_rect(width,height,(1,2,5,4))
            self.assertAlmostEqual(w/h,width/height)
            self.assertGreaterEqual(x,1);self.assertGreaterEqual(y,2)
            self.assertLessEqual(x+w,6.00001);self.assertLessEqual(y+h,6.00001)

    def test_normalization_keeps_exif_frame_and_original(self):
        with tempfile.TemporaryDirectory() as td:
            src=Path(td)/'own.jpg';dst=Path(td)/'copy.jpg'
            im=Image.new('RGB',(40,80),(90,160,230));exif=im.getexif();exif[274]=6
            im.save(src,exif=exif);before=src.read_bytes()
            asset=pv.normalize_image(src,dst)
            self.assertEqual((asset['width'],asset['height']),(80,40))
            self.assertEqual(src.read_bytes(),before)
            pixel=Image.open(dst).getpixel((20,20))
            self.assertTrue(all(abs(a-b)<4 for a,b in zip(pixel,(90,160,230))))

    def test_prompt_has_specific_meaning_and_no_theme_architecture(self):
        prompt=pv.image_prompt('Геология',{'title':'Отбор проб','content':'Геолог собирает керн','image_prompt':'rock core samples'})
        for text in ['Геология','Отбор проб','rock core samples','ABSOLUTELY NO text','charts','empty room']:
            self.assertIn(text,prompt)

    def test_photo_mapping_validates_model_numbers(self):
        self.assertEqual(pv.candidate_slides({'slides':[2,2,True,0,31,'3',3]},8),[1,2])
        self.assertEqual(pv.explicit_slide('Для слайда 4: керн',8),3)
        self.assertIsNone(pv.explicit_slide('слайд 40',8))

    def test_readable_layouts_and_image_geometry(self):
        env,_,_=fixtures.BuildTests().build_env(3)
        colors=env['THEMES']['minimal']
        with tempfile.TemporaryDirectory() as td:
            for dims in [(1600,900),(600,1200),(1000,1000)]:
                path=Path(td)/'image.jpg';Image.new('RGB',dims,'gray').save(path)
                asset={**pv.normalize_image(path,Path(td)/'norm.jpg'),'origin':'user'}
                for num,title,body in [(1,'Геологические проекты и их участники с вложениями',''),
                        (2,'Ключевые участники','Геолог изучает образцы пород. Инженер планирует бурение.'),
                        (3,'Очень длинный заголовок для проверки переносов и читаемости текста на слайде презентации',('Геологи исследуют горные породы и планируют работу. '*9).strip())]:
                    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
                    slide=prs.slides.add_slide(prs.slide_layouts[6])
                    pv.render_slide(slide,{'title':title,'content':body},asset,colors,num,3,'ru',env['slide_background'],Mock())
                    for shape in slide.shapes:
                        self.assertGreaterEqual(shape.left,0);self.assertGreaterEqual(shape.top,0)
                        self.assertLessEqual(shape.left+shape.width,prs.slide_width+10)
                        self.assertLessEqual(shape.top+shape.height,prs.slide_height+10)
                        if shape.shape_type==13:
                            self.assertAlmostEqual(shape.width/shape.height,dims[0]/dims[1],places=5)
                            self.assertEqual(shape.crop_left+shape.crop_right+shape.crop_top+shape.crop_bottom,0)
                        if shape.has_text_frame and num>1:
                            for p in shape.text_frame.paragraphs:
                                if p.text and '/' not in p.text:self.assertGreaterEqual(p.font.size.pt,20)

class ChartLayoutTests(unittest.TestCase):
    def test_real_chart_and_own_photo_both_survive(self):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        env,_,_=fixtures.BuildTests().build_env(3)
        chart_env=fixtures.load_functions(['add_chart'],{'CategoryChartData':CategoryChartData,
            'XL_CHART_TYPE':XL_CHART_TYPE,'XL_LABEL_POSITION':XL_LABEL_POSITION,
            'XL_LEGEND_POSITION':XL_LEGEND_POSITION,'Inches':Inches,'Pt':Pt,'RGBColor':RGBColor,'safe_line':env['safe_line']})
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'own.jpg';Image.new('RGB',(400,800),'blue').save(p)
            asset={**pv.normalize_image(p,Path(td)/'norm.jpg'),'origin':'user'}
            prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            chart={'chart_type':'bar','title':'Данные пользователя','categories':['A','B'],
                   'series':[{'name':'Значение','values':[10,20]}]}
            pv.render_slide(slide,{'title':'Данные проекта и фотография участка','content':'Данные предоставлены пользователем.','chart':chart},
                            asset,env['THEMES']['minimal'],2,3,'ru',env['slide_background'],chart_env['add_chart'])
            self.assertEqual(sum(s.shape_type==13 for s in slide.shapes),1)
            self.assertEqual(sum(s.has_chart for s in slide.shapes),1)
            self.assertEqual(list(next(s.chart for s in slide.shapes if s.has_chart).series[0].values),[10,20])

class ImageFlowTests(unittest.IsolatedAsyncioTestCase):
    def slides(self,n=3):
        return [{'title':f'Тема {i}','content':'Образцы пород','image_prompt':'rock samples'} for i in range(n)]

    async def test_reject_bad_image_then_retry_once_and_keep_good(self):
        with tempfile.TemporaryDirectory() as td:
            count=0
            async def gen(prompt,path,**kw):
                nonlocal count;count+=1;Image.new('RGB',(80,60),(count*30,50,60)).save(path);return True
            inspect=AsyncMock(side_effect=[json.dumps({**GOOD,'has_text':True}),json.dumps(GOOD)])
            assets,w=await pv.prepare_images('Геология',self.slides(1),[],{},td,gen,inspect,json.loads)
            self.assertEqual(count,2);self.assertEqual(w['ai_omitted'],0)
            self.assertEqual(assets[0]['origin'],'ai')

    async def test_two_bad_images_use_text_fallback(self):
        with tempfile.TemporaryDirectory() as td:
            async def gen(prompt,path,**kw):Image.new('RGB',(80,60),'red').save(path);return True
            generate=AsyncMock(side_effect=gen)
            inspect=AsyncMock(return_value=json.dumps({**GOOD,'has_chart':True}))
            assets,w=await pv.prepare_images('Геология',self.slides(1),[],{},td,generate,inspect,json.loads)
            self.assertEqual(assets,[None]);self.assertEqual(generate.await_count,2)
            self.assertEqual(w['ai_omitted'],1)

    async def test_checker_failure_never_publishes_unchecked_image(self):
        with tempfile.TemporaryDirectory() as td:
            async def gen(prompt,path,**kw):Image.new('RGB',(80,60),'red').save(path);return True
            assets,w=await pv.prepare_images('Тема',self.slides(1),[],{},td,gen,AsyncMock(side_effect=TimeoutError),json.loads)
            self.assertEqual(assets,[None]);self.assertEqual(w['ai_omitted'],1)

    async def test_explicit_caption_wins_over_semantic_photo_match(self):
        with tempfile.TemporaryDirectory() as td:
            paths=[]
            for i,color in enumerate(['red','blue']):
                p=Path(td)/f'{i}.jpg';Image.new('RGB',(60,100),color).save(p);paths.append(str(p))
            inspect=AsyncMock(return_value='{"slides":[2]}')
            generate=AsyncMock(return_value=False)
            assets,w=await pv.prepare_images('Тема',self.slides(2),paths,{paths[1]:'слайд 2'},td,generate,inspect,json.loads)
            self.assertEqual(assets[1]['path'],str(Path(td)/'own-1.jpg'))
            self.assertEqual(assets[0]['path'],str(Path(td)/'own-0.jpg'))
            generate.assert_not_awaited()

    async def test_own_photo_survives_failed_classifier(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'own.png';Image.new('RGB',(60,100),'blue').save(p)
            generate=AsyncMock(return_value=False)
            assets,w=await pv.prepare_images('Тема',self.slides(1),[str(p)],{},td,generate,AsyncMock(side_effect=TimeoutError),json.loads)
            self.assertEqual(assets[0]['origin'],'user');generate.assert_not_awaited()
            self.assertTrue(p.exists())

    async def test_corrupt_photo_does_not_break_deck(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'bad.jpg';p.write_text('bad')
            assets,w=await pv.prepare_images('Тема',self.slides(1),[str(p)],{},td,AsyncMock(return_value=False),AsyncMock(),json.loads)
            self.assertEqual(w['own_unreadable'],1);self.assertEqual(assets,[None])

    async def test_concurrency_is_bounded_and_chart_skips_ai(self):
        active=peak=0
        async def gen(*a,**kw):
            nonlocal active,peak
            active+=1;peak=max(peak,active);await asyncio.sleep(.001);active-=1;return False
        with tempfile.TemporaryDirectory() as td:
            slides=self.slides(8);slides[1]['chart']={'categories':['A']}
            generate=AsyncMock(side_effect=gen)
            await pv.prepare_images('Тема',slides,[],{},td,generate,AsyncMock(),json.loads)
            self.assertEqual(generate.await_count,7);self.assertEqual(peak,2)

    async def test_duplicate_generated_frames_not_repeated(self):
        with tempfile.TemporaryDirectory() as td:
            async def gen(prompt,path,**kw):Image.new('RGB',(80,60),'red').save(path);return True
            assets,w=await pv.prepare_images('Тема',self.slides(3),[],{},td,gen,AsyncMock(return_value=json.dumps(GOOD)),json.loads)
            self.assertEqual(sum(a is not None for a in assets),1)
            self.assertEqual(w['ai_omitted'],2)

    async def test_real_builder_inserts_reviewed_images_and_notes(self):
        env,state,sent=fixtures.BuildTests().build_env(3)
        count=0
        async def gen(prompt,path,**kw):
            nonlocal count;count+=1;Image.new('RGB',(120,160),(count*40,50,60)).save(path);return True
        env['generate_image']=gen;env['ask_grok_vision']=AsyncMock(return_value=json.dumps(GOOD))
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)),state)
        self.assertEqual(len(sent),2)
        prs=Presentation(io.BytesIO(sent[0][1]))
        self.assertEqual(len(prs.slides),3)
        for slide in prs.slides:
            self.assertEqual(sum(s.shape_type==13 for s in slide.shapes),1)
            self.assertIn('AI-generated',slide.notes_slide.notes_text_frame.text)
        env['spend_credits'].assert_called_once()

if __name__=='__main__':unittest.main()
