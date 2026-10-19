import ast,json,tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from docx import Document
import fixed_word_templates as fixed
import test_car_sale_template as car_tests

class FieldsTests(unittest.TestCase):
    def test_all_fields_replace_and_no_branding(self):
        for kind,spec in fixed.CATALOG.items():
            with self.subTest(kind=kind),tempfile.TemporaryDirectory() as td:
                p=Path(td)/'test.docx';fixed.render(p,kind,spec['sample'])
                d=Document(p);text='\n'.join(p.text for p in d.paragraphs)+str([[c.text for c in r.cells] for t in d.tables for r in t.rows])
                for key,label in spec['fields'].items():
                    self.assertNotIn('['+label+']',text);self.assertIn(spec['sample'][key].rstrip('.'),text)
                self.assertNotIn('..',text)
                with ZipFile(p) as z:
                    self.assertIsNone(z.testzip());self.assertFalse(any(n.startswith('word/media/') for n in z.namelist()))
                    xml=''.join(z.read(n).decode() for n in z.namelist() if n.endswith('.xml'))
                    for brand in ['auto.ru','авто.ру','dogovor.expert','яндекс','yandex']:self.assertNotIn(brand,xml.casefold())
    def test_empty_values_keep_placeholders(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'test.docx';fixed.render(p,'services_ru',{'number':''})
            self.assertEqual(p.read_bytes(),(fixed.ASSETS/fixed.CATALOG['services_ru']['file']).read_bytes())
    def test_grounding_rejects_invented_values(self):
        self.assertEqual(fixed.validate_fields({'provider':'ООО Тест','price':'100','unknown':'test'},'ООО Тест',fixed.CATALOG['services_ru']),{'provider':'ООО Тест'})
    def test_explicit_ru_kind_cannot_use_generic_generation(self):
        self.assertTrue(fixed.applies({'word_kind':'services_ru','content_lang':'en'}))
        self.assertTrue(fixed.applies({'word_kind':'services'}))
        self.assertFalse(fixed.applies({'word_kind':'services','content_lang':'en'}))
    def test_actual_entrypoints(self):
        source=(Path(__file__).resolve().parents[1]/'main.py').read_text();tree=ast.parse(source)
        funcs={n.name:ast.get_source_segment(source,n) for n in tree.body if isinstance(n,ast.AsyncFunctionDef)}
        for name in ['word_build','word_build_draft','word_extra']:self.assertIn('fixed_template_applies',funcs[name])
        self.assertIn('await run_heavy(render_fixed_word, tpl_path, kind)',funcs['word_mode_template'])
        self.assertIn('await word_build(m, state)',funcs['_handle_miniapp_payload'])

class WorkflowTests(unittest.IsolatedAsyncioTestCase):
    def env(self,kind='services_ru',text=None):
        spec=fixed.CATALOG[kind]
        data={'word_kind':kind,'topic':text if text is not None else '\n'.join(spec['fields'][k]+': '+v for k,v in spec['sample'].items())}
        return car_tests.CarWorkflowTests().env(data)
    async def test_each_handler_reviews_then_delivers_without_model(self):
        for kind in fixed.CATALOG:
            e,m,s=self.env(kind);await e['word_build'](m,s)
            self.assertEqual(s.current,'confirm');m.answer_document.assert_not_awaited();e['spend_credits'].assert_not_called()
            events=[];m.answer_document.side_effect=lambda *a,**kw:events.append('send');e['spend_credits'].side_effect=lambda *a,**kw:events.append('charge')
            await e['word_build'](m,s)
            self.assertEqual(events,['send','send','charge']);e['ask_grok'].assert_not_awaited();self.assertEqual(s.data,{})
    async def test_pdf_send_failure_resumes_failed_delivery_only(self):
        e,m,s=self.env();await fixed.preview(m,s,e);m.answer_document.side_effect=[None,RuntimeError('offline')]
        await fixed.build(m,s,e);e['spend_credits'].assert_not_called();self.assertTrue(s.data['fixed_word']['sent_docx'])
        m.answer_document.side_effect=None;await fixed.build(m,s,e)
        self.assertEqual(m.answer_document.await_count,3);e['prepare_word_files'].assert_awaited_once();e['spend_credits'].assert_called_once()
    async def test_final_notice_failure_no_double_charge(self):
        e,m,s=self.env();await fixed.preview(m,s,e)
        async def answer(text,**kw):
            if text.startswith('Готово.'):raise RuntimeError('offline')
        m.answer.side_effect=answer;await fixed.build(m,s,e);m.answer.side_effect=None;await fixed.build(m,s,e)
        self.assertEqual(m.answer_document.await_count,2);e['spend_credits'].assert_called_once()
    async def test_missing_fields_reviewed_before_build(self):
        e,m,s=self.env(text='Исполнитель: ООО Тест');await fixed.build(m,s,e)
        self.assertIn('Не заполнено:',str(m.answer.call_args_list));m.answer_document.assert_not_awaited()
    async def test_corrections_clear_and_preserve_other_fields(self):
        e,m,s=self.env();await fixed.preview(m,s,e);await s.update_data(extra='Стоимость услуг с валютой: 500 рублей\nГород: ');await fixed.preview(m,s,e)
        f=s.data['fixed_word']['fields'];self.assertEqual(f['price'],'500 рублей');self.assertNotIn('city',f);self.assertIn('provider',f);e['ask_grok'].assert_not_awaited()
    async def test_switch_template_drops_old_fields(self):
        e,m,s=self.env();await fixed.preview(m,s,e);await s.update_data(word_kind='rent_flat_ru',topic='ФИО наймодателя: Петров',extra='');await fixed.build(m,s,e)
        self.assertEqual(s.data['fixed_word']['fields'],{'owner':'Петров'});m.answer_document.assert_not_awaited()
    async def test_no_balance_no_delivery(self):
        e,m,s=self.env();await fixed.preview(m,s,e);e['can_afford']=lambda *a:False
        await fixed.build(m,s,e);m.answer_document.assert_not_awaited();e['start_job'].assert_not_called()
    async def test_model_failure_preserves_input_and_requests_form(self):
        e,m,s=self.env(text='Консультации для ООО Тест');e['ask_grok'].side_effect=RuntimeError('offline');await fixed.build(m,s,e)
        self.assertEqual(s.current,'extra');self.assertIn('ООО Тест',s.data['topic']);m.answer_document.assert_not_awaited();e['spend_credits'].assert_not_called()
    async def test_out_of_scope_not_built_or_charged(self):
        for kind,text in [('rent_flat_ru','Арендатор ООО Тест'),('claim_replace_ru','Требую вернуть деньги'),('services_act_ru','Сделать КС-2')]:
            e,m,s=self.env(kind,text);await fixed.build(m,s,e);e['ask_grok'].assert_not_awaited();m.answer_document.assert_not_awaited();e['spend_credits'].assert_not_called()
    async def test_export_failure_keeps_fields_for_retry(self):
        e,m,s=self.env();await fixed.preview(m,s,e);original=e['prepare_word_files'].side_effect;e['prepare_word_files'].side_effect=RuntimeError('pdf')
        await fixed.build(m,s,e);e['spend_credits'].assert_not_called();e['prepare_word_files'].side_effect=original;await fixed.build(m,s,e)
        e['ask_grok'].assert_not_awaited();e['spend_credits'].assert_called_once()
