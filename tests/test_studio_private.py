import hashlib,hmac,io,json,time,unittest
from unittest.mock import AsyncMock,patch
from urllib.parse import urlencode,parse_qsl
from aiohttp import web,CookieJar
from aiohttp.test_utils import TestClient,TestServer
from studio_private import Studio,PILOT_UID,PILOT_UIDS,COOKIE,menu_url
from studio_export import build_pptx
from constructor_billing import ConstructorBilling
from constructor_pricing import VERSION
from pptx import Presentation
from test_presentation_support import load_functions
ORIGIN='https://example.test'
HEADERS={'Origin':ORIGIN,'X-Studio-Request':'1'}
class StudioTests(unittest.IsolatedAsyncioTestCase):
 async def asyncSetUp(self):
  self.send=AsyncMock();self.action=AsyncMock(side_effect=lambda *a: web.json_response({'ok':True},status=202))
  verify=load_functions(['verify_telegram_init_data','miniapp_verified_user'],{'BOT_TOKEN':'test-token','INIT_DATA_MAX_AGE_SECONDS':604800,'parse_qsl':parse_qsl,'hmac':hmac,'hashlib':hashlib,'json':json,'time':time})['miniapp_verified_user']
  self.users={uid:{"credits":1000} for uid in PILOT_UIDS}
  billing=ConstructorBilling(lambda uid:self.users[uid],AsyncMock(return_value=True),lambda uid:(True,""),lambda uid:None,set())
  self.studio=Studio(verify,self.send,self.action,ORIGIN,billing=billing)
  app=web.Application();self.studio.register(app)
  self.client=TestClient(TestServer(app),cookie_jar=CookieJar(unsafe=True));await self.client.start_server()
 async def asyncTearDown(self):
  await self.client.close()
  for archive in self.studio.archives:archive.close()
 def signed(self,uid=PILOT_UID,age=0):
  p={'auth_date':str(int(time.time())-age),'user':json.dumps({'id':uid,'first_name':'Test'})}
  secret=hmac.new(b'WebAppData',b'test-token',hashlib.sha256).digest()
  p['hash']=hmac.new(secret,'\n'.join(f'{k}={v}' for k,v in sorted(p.items())).encode(),hashlib.sha256).hexdigest()
  return urlencode(p)
 async def login(self,uid=PILOT_UID):
  r=await self.client.post('/studio/session',json={'initData':self.signed(uid)},headers=HEADERS);self.assertEqual(r.status,200)
  c=r.cookies[COOKIE];self.assertTrue(c['httponly']);self.assertTrue(c['secure']);self.assertEqual(c['samesite'],'Strict')
  self.client.session.cookie_jar.update_cookies({COOKIE:c.value})
 async def test_auth_rejects_unsigned_other_user_expired_future_missing_date(self):
  for v in ['',self.signed(123),self.signed(age=3602),self.signed(age=-100),self.signed().replace('auth_date=','not_date=')]:
   r=await self.client.post('/studio/session',json={'initData':v,'user_id':PILOT_UID},headers=HEADERS);self.assertEqual(r.status,403)
  for path in ['/studio/constructor/','/studio/templates/catalog.json','/studio/constructor/app.js']:
   self.assertEqual((await self.client.get(path)).status,401)
  self.assertEqual((await self.client.post('/studio/export',json={})).status,401)
 async def test_entry_shell_is_public_but_never_exposes_private_data(self):
  self.studio.profile_callback=AsyncMock(return_value={'name':'Private Owner','credits':9876})
  r=await self.client.get('/studio/');self.assertEqual(r.status,200)
  html=await r.text();self.assertNotIn('Открываю личный тест',html);self.assertNotIn('Прежнее меню',html)
  self.assertIn('as-content',html);self.assertNotIn('Private Owner',html)
  for path in ['/studio/home/app.js','/studio/home/style.css','/studio/shared/mobile.js','/studio/shared/mobile.css','/studio/templates/fonts.css','/studio/templates/fonts/017fab69b4be8ffc.woff2']:
   r=await self.client.get(path);self.assertEqual(r.status,200,path)
   self.assertEqual(r.headers['X-Content-Type-Options'],'nosniff')
  for path in ['/studio/profile','/studio/home/','/studio/forms/','/studio/templates/catalog.json','/studio/constructor/coffee.jpg','/studio/templates/assets/standard-space-0-hd-b78b5f1e45.webp']:
   self.assertEqual((await self.client.get(path)).status,401,path)
  self.studio.profile_callback.assert_not_awaited()
  self.assertEqual((await self.client.post('/studio/generate',json={})).status,401)
  await self.login()
  r=await self.client.get('/studio/profile');self.assertEqual((await r.json())['profile']['credits'],9876)
 async def test_example_pages_private_read_only_and_complete(self):
  for path in ['/studio/home/examples.js','/studio/home/examples/catalog.json','/studio/home/examples/essay-1.webp']:
   self.assertEqual((await self.client.get(path)).status,401)
  await self.login()
  r=await self.client.get('/studio/home/examples/catalog.json');self.assertEqual(r.status,200)
  examples=await r.json();self.assertEqual(set(examples),{'pres','word','excel'})
  for kind,count in [('pres',4),('word',3),('excel',1)]:
   self.assertEqual(len(examples[kind]['pages']),count)
   for page in examples[kind]['pages']+[{'src':examples[kind]['thumb']}]:
    self.assertTrue(page['src'].startswith('/studio/'))
    r=await self.client.get(page['src']);self.assertEqual(r.status,200,page['src'])
    content=await r.read();self.assertEqual(content[:4],b'RIFF');self.assertEqual(content[8:12],b'WEBP')
    self.assertNotIn('attachment',r.headers.get('Content-Disposition',''))
  self.action.assert_not_awaited();self.send.assert_not_awaited()
 async def test_valid_session_paths_and_csrf(self):
  await self.login()
  for path in ['/studio/templates/','/studio/constructor/app.js']:self.assertEqual((await self.client.get(path)).status,200)
  r=await self.client.post('/studio/generate',json={},headers={'Origin':'https://evil.test','X-Studio-Request':'1'});self.assertEqual(r.status,403)
  self.assertEqual((await self.client.post('/studio/session',json={},headers={'Origin':ORIGIN})).status,403)
  with patch('studio_private.Path.is_file',return_value=False):
   r=await self.client.get('/studio/constructor/coffee.jpg');self.assertEqual(r.status,200);self.assertGreater(len(await r.read()),1000)
  for session in self.studio.sessions.values():session['expires']=0
  self.assertEqual((await self.client.get('/studio/templates/')).status,401)
 async def test_template_choice_and_price_mode_forwarded(self):
  await self.login()
  for kind,style,key,expected in [('standard','strict','nature','nature'),('designer','soft','d16','designer_soft')]:
   r=await self.client.post('/studio/generate',headers=HEADERS,json={'kind':kind,'style':style,'key':key,'slides':6,'topic':'Киты','requestId':'test-'+kind});self.assertEqual(r.status,202)
   user,payload,rid=self.action.call_args.args;self.assertEqual(user['id'],PILOT_UID);self.assertEqual(payload['style'],expected)
   if kind=='designer':self.assertEqual(payload['studio_template'],'d16')
  r=await self.client.post('/studio/generate',headers=HEADERS,json={'kind':'designer','style':'soft','key':'d999','slides':6,'topic':'Киты'});self.assertEqual(r.status,400)
 def deck(self):return {'version':1,'confirmedCost':16,'pricingVersion':VERSION,'chat_id':123,'slides':[{'background':{'type':'solid','color':'#14293b'},'elements':[{'type':'text','text':'Проверка','x':40,'y':50,'w':900,'h':90,'size':40,'font':'Arial','color':'#e8dfcf'}]}]}
 async def test_export_delivery_deduplication_and_fixed_recipient(self):
  await self.login()
  for _ in range(2):
   r=await self.client.post('/studio/export',headers=HEADERS,json=self.deck());self.assertEqual(r.status,200,await r.text())
  self.assertEqual(self.users[PILOT_UID]['credits'],984)
  self.send.assert_awaited_once();uid,data=self.send.call_args.args;self.assertEqual(uid,PILOT_UID)
  p=Presentation(io.BytesIO(data));self.assertEqual(p.slides[0].shapes[0].text,'Проверка')
 async def test_both_testers_receive_own_identical_export_and_own_profile(self):
  self.studio.profile_callback=AsyncMock(side_effect=lambda user:{'name':str(user['id'])})
  for uid in sorted(PILOT_UIDS):
   await self.login(uid)
   r=await self.client.get('/studio/profile');self.assertEqual((await r.json())['profile']['name'],str(uid))
   self.assertEqual((await self.client.get('/studio/constructor/')).status,200)
   for _ in range(2):
    r=await self.client.post('/studio/export',headers=HEADERS,json=self.deck());self.assertEqual(r.status,200)
   self.assertEqual(self.send.call_args.args[0],uid)
  self.assertEqual([call.args[0] for call in self.send.await_args_list],sorted(PILOT_UIDS))
 async def test_second_tester_generation_and_actions_use_own_identity(self):
  await self.login(909828109)
  r=await self.client.post('/studio/generate',headers=HEADERS,json={'kind':'designer','style':'soft','key':'d16','slides':6,'topic':'Киты','chat_id':PILOT_UID})
  self.assertEqual(r.status,202);self.assertEqual(self.action.call_args.args[0]['id'],909828109)
  self.studio.menu_action=AsyncMock(return_value=web.json_response({'ok':True}))
  r=await self.client.post('/studio/action',headers=HEADERS,json={'action':'history','payload':{'user_id':PILOT_UID}})
  self.assertEqual(r.status,200);self.assertEqual(self.studio.menu_action.call_args.args[0]['id'],909828109)
 async def test_failed_delivery_retriable_and_busy(self):
  await self.login();self.studio.busy=True
  self.assertEqual((await self.client.post('/studio/export',headers=HEADERS,json={})).status,409)
  self.studio.busy=False;self.send.side_effect=RuntimeError('delivery')
  self.assertEqual((await self.client.post('/studio/export',headers=HEADERS,json=self.deck())).status,503)
  self.assertFalse(self.studio.busy);self.send.side_effect=None
  self.assertEqual(self.users[PILOT_UID]['credits'],1000)
  self.assertEqual((await self.client.post('/studio/export',headers=HEADERS,json=self.deck())).status,409)
 async def test_export_requires_exact_confirmed_price_and_version(self):
  await self.login()
  for changes in [{'confirmedCost':0},{'confirmedCost':15},{'confirmedCost':True},{'pricingVersion':'old'}]:
   r=await self.client.post('/studio/export',headers=HEADERS,json=self.deck()|changes)
   self.assertEqual(r.status,409)
  self.send.assert_not_awaited();self.assertEqual(self.users[PILOT_UID]['credits'],1000)
 async def test_export_over_limit_and_insufficient_balance(self):
  await self.login();d=self.deck();d['slides'][0]['elements']*=7
  self.assertEqual((await self.client.post('/studio/export',headers=HEADERS,json=d)).status,400)
  self.users[PILOT_UID]['credits']=15
  self.assertEqual((await self.client.post('/studio/export',headers=HEADERS,json=self.deck())).status,402)
  self.send.assert_not_awaited()
 def test_menu_target_only_owner(self):
  for uid in [None,123,str(PILOT_UID)]:self.assertEqual(menu_url(ORIGIN,uid,'old'),'old')
  for uid in PILOT_UIDS:self.assertEqual(menu_url(ORIGIN,uid,'old'),ORIGIN+'/studio/')
 def test_export_rejects_external_image_and_nonfinite_geometry(self):
  d=self.deck();d['slides'][0]['elements'][0].update(type='photo',raster='https://example.test/file')
  with self.assertRaises(ValueError):build_pptx(d)
  d['slides'][0]['elements'][0]['w']=float('nan')
  with self.assertRaises(ValueError):build_pptx(d)

class StudioTemplateIntegration(unittest.IsolatedAsyncioTestCase):
 async def test_every_designer_template_reaches_real_design_selection(self):
  import re
  from types import SimpleNamespace
  from unittest.mock import Mock
  from presentation_pricing import presentation_cost
  from presentation_profiles import choose_design
  start=AsyncMock();balance=Mock(return_value=True)
  fn=load_functions(['_handle_miniapp_payload'],{'ensure_answerable':lambda m:m,'user_lang':lambda uid:'ru','get_user':lambda uid:{},'_clip_text':lambda v,n:(v or '')[:n],'can_afford':balance,'presentation_cost':presentation_cost,'_start_presentation_request':start,'send_no_credits_notice':AsyncMock(),'re':re})['_handle_miniapp_payload']
  for uid in sorted(PILOT_UIDS):
   for n in range(1,51):
    for style in ('designer_soft','designer_strict'):
     key=f'd{n:02d}'
     await fn(SimpleNamespace(from_user=SimpleNamespace(id=uid)),object(),{'topic':'Кофейня','style':style,'slides':6,'studio_template':key},'gen_presentation')
     settings=start.call_args.kwargs['settings']
     self.assertEqual(choose_design({'topic':'Кофейня',**settings})['template_id'],key)
     self.assertEqual(balance.call_args.args,(uid,192))
  await fn(SimpleNamespace(from_user=SimpleNamespace(id=123)),object(),{'topic':'Кофейня','style':'designer_soft','slides':6,'studio_template':'d01'},'gen_presentation')
  self.assertNotIn('extra',start.call_args.kwargs['settings'])

class FullMenuTests(StudioTests):
 async def test_owner_profile_forms_and_action(self):
  self.studio.profile_callback=AsyncMock(return_value={'name':'Owner','credits':15,'history':[],'payments':[]})
  self.studio.menu_action=AsyncMock(side_effect=lambda *a:web.json_response({'ok':True},status=202))
  for path in ['/studio/home/','/studio/forms/','/studio/profile']:
   self.assertEqual((await self.client.get(path)).status,401)
  self.assertEqual((await self.client.post('/studio/action',json={})).status,401)
  await self.login()
  self.assertEqual((await self.client.get('/studio/home/')).status,200)
  r=await self.client.get('/studio/profile?user_id=123');self.assertEqual((await r.json())['profile']['credits'],15)
  self.assertEqual(self.studio.profile_callback.call_args.args[0]['id'],PILOT_UID)
  r=await self.client.get('/studio/forms/?screen=form-word');html=await r.text()
  self.assertIn("new URL(relativePath, window.location.origin + '/')",html);self.assertIn('<base href="/">',html);self.assertIn("showScreen(studioScreen)",html);self.assertIn("fetch('/studio/'",html)
  for action in ['gen_word','gen_excel','topup','history','help']:
   r=await self.client.post('/studio/action',headers=HEADERS,json={'action':action,'payload':{'topic':'test','user_id':123},'requestId':'id'});self.assertEqual(r.status,202)
   self.assertEqual(self.studio.menu_action.call_args.args[0]['id'],PILOT_UID)
  self.assertEqual((await self.client.post('/studio/action',json={'action':'gen_word'},headers={'Origin':'https://evil.test','X-Studio-Request':'1'})).status,403)
  self.assertEqual((await self.client.post('/studio/action',json={'action':'unknown'},headers=HEADERS)).status,400)
 def test_payment_data_filtered_no_secrets(self):
  from tempfile import TemporaryDirectory
  from pathlib import Path
  from studio_menu import recent_payments
  with TemporaryDirectory() as d:
   p=Path(d)/'payments.jsonl';p.write_text('\n'.join([json.dumps({'uid':123,'credits':999,'payment_id':'private'}),'{bad',json.dumps({'uid':str(PILOT_UID),'credits':20,'ts':'2026-10-23','payment_id':'private'})]))
   self.assertEqual(recent_payments(p,PILOT_UID),[{'credits':20,'ts':'2026-10-23'}])
 def test_menu_has_no_demo_balance_and_local_fonts(self):
  from pathlib import Path
  root=Path(__file__).resolve().parents[1]/'studio_ui/home'
  html=(root/'index.html').read_text();js=(root/'app.js').read_text()
  self.assertNotIn('120 кредитов',js);self.assertNotIn('const transactions=',js)
  self.assertIn('/studio/templates/fonts.css',html);self.assertIn('escape(data?.name',js)
