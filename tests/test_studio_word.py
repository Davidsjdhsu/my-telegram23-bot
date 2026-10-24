import asyncio, json, tempfile, time, unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from aiohttp import web, CookieJar
from aiohttp.test_utils import TestClient, TestServer
from studio_private import Studio, PILOT_UID, PILOT_UIDS
from studio_word import WordWorkflow, WordState, WordJob, atomic
from studio_word_catalog import catalog, normalize, price, SPECS, render_preview
import test_studio_private
from test_studio_private import HEADERS, ORIGIN
import test_office
from test_office import office_env
from studio_word_runtime import attach

class CatalogTests(unittest.TestCase):
 def test_slots_prices_and_scope(self):
  self.assertEqual(len(catalog()),36)
  self.assertIn('record_book',SPECS['notes']['fields']);self.assertNotIn('record_book',SPECS['essay']['fields'])
  for kind in SPECS:
   p=normalize(dict(kind=kind,topic='Тема',fields={},pages=6))
   self.assertEqual(price(p),10 if SPECS[kind]['fixed'] or kind=='resume' else 30)
  for p in [[],{'kind':'essay','topic':'A','fields':{'foo':'B'}},{'kind':'essay','topic':'A','pages':3},{'kind':'essay','topic':'A','pages':True},{'kind':'essay','topic':''},{'kind':'services_ru','content':'ignored'},{'kind':'services_ru','content_lang':'en'},{'kind':'dkp_car','fields':{'seller':'ООО Компания'}}]:
   with self.assertRaises(ValueError):normalize(p)
 def test_real_cover_fields_are_literal_and_not_ai_generated(self):
  from docx import Document
  with tempfile.TemporaryDirectory() as d:
   for kind in ('referat','coursework','essay','report','notes'):
    fields={'author':'Тестов Иван Сергеевич','school':'Учебный университет','teacher':'Петрова Анна Викторовна','city':'Казань','year':'2026'}
    path=Path(d)/(kind+'.docx');render_preview(path,dict(kind=kind,topic='Проверка полей',fields=fields,pages=6))
    doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for value in fields.values():self.assertIn(value,text,kind)
    self.assertIn('проверка полей',text.casefold())
 def test_resume_real_renderer(self):
  with tempfile.TemporaryDirectory() as d:
   path=Path(d)/'resume.docx';render_preview(path,dict(kind='resume',topic='Резюме',fields={'name':'Тестов Иван'},pages=2),office_env()['build_word'])
   from docx import Document
   self.assertEqual(Document(path).paragraphs[0].text,'Тестов Иван')

class WorkflowTests(unittest.IsolatedAsyncioTestCase):
 signed=test_studio_private.StudioTests.signed
 login=test_studio_private.StudioTests.login
 async def asyncSetUp(self):
  import hmac,hashlib
  from urllib.parse import parse_qsl
  from test_presentation_support import load_functions
  verify=load_functions(['verify_telegram_init_data','miniapp_verified_user'],{'BOT_TOKEN':'test-token','INIT_DATA_MAX_AGE_SECONDS':604800,'parse_qsl':parse_qsl,'hmac':hmac,'hashlib':hashlib,'json':json,'time':time})['miniapp_verified_user']
  self.tmp=tempfile.TemporaryDirectory();self.studio=Studio(verify,AsyncMock(),AsyncMock(),ORIGIN)
  self.send=AsyncMock();self.charge=AsyncMock(return_value=10);self.preview=AsyncMock(return_value=b'png')
  async def build(user,job,action):
   for ext in ('docx','pdf'):(job.directory/('source.'+ext)).write_bytes(b'file')
   await job.deliver(job.directory/'source.docx',job.directory/'source.pdf',10,'Документ')
  self.builder=AsyncMock(side_effect=build)
  self.flow=WordWorkflow(self.studio,self.tmp.name,self.builder,self.preview,self.send,self.charge,lambda *a:True)
  app=web.Application();self.flow.register(app);self.studio.register(app)
  self.client=TestClient(TestServer(app),cookie_jar=CookieJar(unsafe=True));await self.client.start_server();await self.login()
 async def asyncTearDown(self):
  await self.client.close()
  for archive in self.studio.archives:archive.close()
  self.tmp.cleanup()
 async def create(self,**kw):
  r=await self.client.post('/studio/word/orders',headers=HEADERS,json={'kind':'essay','topic':'Тема','pages':6,**kw});self.assertEqual(r.status,200,await r.text());return await r.json()
 async def confirm(self,order,**kw):
  return await self.client.post('/studio/word/orders/'+order['id']+'/confirm',headers=HEADERS,json={'confirmedCost':order['cost'],**kw})
 async def settle(self):
  tasks=[v[1] for v in self.flow.tasks.values()]
  await asyncio.gather(*tasks)
 async def status(self,o):return await (await self.client.get('/studio/word/orders/'+o['id'])).json()
 async def test_auth_ownership_csrf_and_server_price(self):
  order=await self.create(cost=1);self.assertEqual(order['cost'],30)
  self.assertEqual((await self.confirm(order,confirmedCost=1)).status,409)
  self.assertEqual((await self.client.post('/studio/word/orders',json={'kind':'essay'})).status,403)
  await self.login(909828109)
  self.assertEqual((await self.client.get('/studio/word/orders/'+order['id'])).status,404)
  self.assertEqual((await self.confirm(order)).status,404)
  self.builder.assert_not_awaited();self.charge.assert_not_awaited()
 async def test_success_repeated_click_restart_and_reopen(self):
  o=await self.create();await self.confirm(o);await self.confirm(o);await self.settle()
  self.assertEqual((await self.status(o))['status'],'completed')
  self.assertEqual([c.args[3] for c in self.send.await_args_list],['docx','pdf']);self.charge.assert_awaited_once();self.builder.assert_awaited_once()
  self.assertEqual((await self.create())['id'],o['id'])
  self.assertEqual((await (await self.client.get('/studio/word/draft')).json())['order']['status'],'completed')
  WordWorkflow(self.studio,self.tmp.name,self.builder,self.preview,self.send,self.charge,lambda *a:True)
  await self.confirm(o);self.charge.assert_awaited_once()
 async def test_transport_ambiguity_never_retries_or_charges(self):
  self.send.side_effect=[None,TimeoutError()]
  o=await self.create();await self.confirm(o);await self.settle()
  self.assertEqual((await self.status(o))['status'],'delivery_unknown');self.charge.assert_not_awaited()
  self.assertEqual((await self.confirm(o)).status,409)
  job=self.flow.get(PILOT_UID,o['id']);job.row['created']-=90000;job.save()
  r=await self.client.post('/studio/word/orders',headers=HEADERS,json={'kind':'essay','topic':'Другая тема','pages':6});self.assertEqual(r.status,409)
  self.assertTrue((job.directory/'order.json').exists())
 async def test_failed_build_keeps_fields_and_content_and_resumes(self):
  original=self.builder.side_effect
  async def fail(user,job,action):
   await WordState(job).update_data(saved='готовый текст');raise RuntimeError('fixture')
  self.builder.side_effect=fail;o=await self.create(fields={'author':'Тестов'});await self.confirm(o);await self.settle()
  self.assertEqual((await self.status(o))['status'],'interrupted');self.send.assert_not_awaited();self.charge.assert_not_awaited()
  job=self.flow.get(PILOT_UID,o['id']);self.assertEqual((await WordState(job).get_data())['saved'],'готовый текст');self.assertEqual(job.row['payload']['fields']['author'],'Тестов')
  self.builder.side_effect=original;await self.confirm(o);await self.settle();self.charge.assert_awaited_once()
 async def test_choice_needs_exact_accept_price(self):
  self.builder.side_effect=lambda user,job,action:job.choice(20,'Получилось четыре страницы')
  o=await self.create();await self.confirm(o);await self.settle();self.assertEqual((await self.status(o))['status'],'choice');self.send.assert_not_awaited()
  self.assertEqual((await self.confirm(o,action='accept')).status,409)
  self.assertEqual((await self.confirm(o,action='accept',confirmedCost=20)).status,200);await self.settle()
 async def test_draft_preview_cache_and_isolation(self):
  p={'kind':'notes','topic':'Память','pages':6,'fields':{'author':'Тестов'}}
  r=await self.client.post('/studio/word/draft',headers=HEADERS,json=p);self.assertEqual(r.status,200)
  for _ in range(2):self.assertEqual((await self.client.post('/studio/word/preview',headers=HEADERS,json=p)).status,200)
  self.preview.assert_awaited_once();self.builder.assert_not_awaited();self.charge.assert_not_awaited()
  await self.login(909828109);self.assertIsNone((await (await self.client.get('/studio/word/draft')).json())['draft'])
 async def test_crash_intent_is_held_and_safe_stage_is_resumable(self):
  a=await self.create();job=self.flow.get(PILOT_UID,a['id']);job.row.update(status='building');job.save()
  restored=WordWorkflow(self.studio,self.tmp.name,self.builder,self.preview,self.send,self.charge,lambda *a:True)
  self.assertEqual(restored.get(PILOT_UID,a['id']).row['status'],'interrupted')
  job.row.update(status='delivering',billing_pending=True);job.save()
  restored=WordWorkflow(self.studio,self.tmp.name,self.builder,self.preview,self.send,self.charge,lambda *a:True)
  self.assertEqual(restored.get(PILOT_UID,a['id']).row['status'],'delivery_unknown')

class MainIntegrationTests(unittest.IsolatedAsyncioTestCase):
 async def test_real_word_orchestration_uses_explicit_fields_and_durable_delivery(self):
  e,m,s=test_office.WordWorkflowTests().env();s.get_data.return_value.update(studio_word_fields={'org':'Тестовая организация'},topic='Точная тема')
  job=SimpleNamespace(stage=Mock(),deliver=AsyncMock(return_value={'delivered':True}))
  await e['word_build'](m,s,studio_job=job)
  job.deliver.assert_awaited_once();m.answer_document.assert_not_awaited();e['spend_credits'].assert_not_called()
  call=e['run_heavy'].await_args.args;self.assertEqual(call[2],'Точная тема');self.assertEqual(call[5],{'org':'Тестовая организация'})
  self.assertTrue(any('studio_word_content' in c.kwargs for c in s.update_data.await_args_list))
 async def test_runtime_payment_ledger_is_idempotent_and_failure_held(self):
  with tempfile.TemporaryDirectory() as d:
   user={'credits':100,'generations':0,'history':[]};save=AsyncMock(return_value=True)
   api={'USERS_FILE':str(Path(d)/'users.json'),'get_user':lambda uid:user,'can_afford':lambda uid,c: user['credits']>=c,'ADMIN_IDS':set(),'save_users_now':save,'note_success':Mock()}
   service=attach(Mock(),web.Application(),api)
   self.assertEqual(await service.charge(1,'order-one',10,'Тема'),10);self.assertEqual(await service.charge(1,'order-one',10,'Тема'),10)
   self.assertEqual(user['credits'],90);self.assertEqual(len(user['history']),1)
   user['studio_word_payments']['unresolved']={'status':'pending','charged':10}
   with self.assertRaises(RuntimeError):await service.charge(1,'unresolved',10,'Тема')
   self.assertEqual(user['credits'],90)
