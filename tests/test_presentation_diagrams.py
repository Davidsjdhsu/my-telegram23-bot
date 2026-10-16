import asyncio
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import presentation_diagrams as dg
import presentation_designer as pd
import presentation_quality as pq
import presentation_visuals as pv
import presentation_profiles as pp
import test_presentation_support as f


def diagram(kind='sequence',count=4):
    return {'kind':kind,'nodes':[{'label':'Этап '+str(i+1),'text':'Материал переходит на следующую стадию обработки.'} for i in range(count)]}


class DiagramTests(unittest.TestCase):
    def test_five_stage_water_diagram_fits_multiline_labels(self):
        labels=['Водозабор и крупный мусор','Коагуляция и осветление','Обеззараживание воды',
                'Подготовка к обеззараживанию','От станции до вашего дома']
        item={'title':'Путь воды','diagram':{'kind':'sequence','nodes':[
            {'label':label,'text':'Вода проходит обработку перед подачей в сеть.'} for label in labels]}}
        for style in (None,'designer_soft','designer_strict'):
            scene=dg.build_diagram_scene(item,8,8,style=style)
            labels_out=[n for n in scene['nodes'] if n.get('role')=='subheading']
            self.assertEqual(len(labels_out),5)
            self.assertTrue(any('\n' in n['text'] for n in labels_out))
            self.assertTrue(all(n['size']>=22 for n in labels_out))

    def test_invalid_nodes_and_kinds_rejected(self):
        for bad in ({'kind':'photo','nodes':[]},diagram(count=1),diagram(count=6),diagram('cycle',2),
                    {'kind':'sequence','nodes':[{'label':'Same','text':'Text'}]*2}):
            with self.assertRaises(ValueError):dg.validate_diagram(bad)
        for kind in ('sequence','cycle','structure'):self.assertEqual(dg.validate_diagram(diagram(kind))['kind'],kind)

    def test_technical_photo_cannot_pass_content_validation(self):
        c=f.make_content(3);c['slides'][0]['visual_kind']='technical_diagram'
        with self.assertRaises(ValueError):f.ps.validate_content(c,f.make_plan(3))
        c['slides'][0].update(diagram=diagram(),image_prompt='')
        result=f.ps.validate_content(c,f.make_plan(3));self.assertTrue(result['slides'][0]['diagram'])

    def test_request_is_enforced_for_both_styles_but_not_negation(self):
        c=f.make_content(3)
        for design in (False,True):
            issues=pq.local_issues(c,design,{'extra':'Добавь схему работы.'})
            self.assertTrue(any(x['problem']=='missing' for x in issues))
        for text in ('Без схем','Without diagrams','No diagram'):
            self.assertFalse(dg.diagram_requested({'extra':text}))

    def test_editor_sees_actual_diagram_text_and_return_edge(self):
        item={'number':2,'title':'Цикл','content':'Hidden paragraph','diagram':diagram('cycle')}
        review=pq.review_slide(item)
        self.assertNotIn('Hidden',review['visible_text'])
        self.assertEqual(review['relations'][-1],{'from':4,'to':1,'meaning':'next_stage'})
        item['diagram']['kind']='structure'
        self.assertTrue(all(x['meaning']=='part_of' for x in pq.review_slide(item)['relations']))

    def test_native_links_text_and_uploaded_image_survive_pptx(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'photo.jpg';Image.new('RGB',(800,600),'navy').save(path)
            asset={'path':str(path),'width':800,'height':600,'origin':'user','safe_frames':[]}
            for style in (None,'designer_soft','designer_strict'):
                for kind in ('sequence','cycle','structure'):
                    item={'title':'Как работает система','content':'Пояснение','diagram':diagram(kind,5),
                          '_designer_design':pp.choose_design({'topic':'Система','theme_name':style})}
                    scene=dg.build_diagram_scene(item,2,8,style=style,asset=asset)
                    text_nodes=[n for n in scene['nodes'] if n['type']=='text']
                    for i,a in enumerate(text_nodes):
                        for b in text_nodes[i+1:]:
                            x,y,w,h=a['box'];xx,yy,ww,hh=b['box']
                            self.assertFalse(min(x+w,xx+ww)>max(x,xx)+1 and min(y+h,yy+hh)>max(y,yy)+1)
                    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
                    pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
                    buf=io.BytesIO();prs.save(buf);buf.seek(0);slide=Presentation(buf).slides[0]
                    self.assertEqual(sum(x.name=='User photograph' for x in slide.shapes),1)
                    xml=slide._element.xml
                    arrows=xml.count('type="triangle"')
                    self.assertEqual(arrows,0 if kind=='structure' else 5 if kind=='cycle' else 4)
                    text=' '.join(x.text for x in slide.shapes if x.has_text_frame)
                    self.assertIn('Этап 5',text)

    def test_generated_technical_evidence_is_rejected(self):
        verdict=dict(relevant=True,has_text=False,has_chart=False,defective=False,requires_verified_source=True)
        self.assertFalse(pv.approved(verdict))


class WorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def test_format_repair_is_not_undone_by_old_editorial_snapshot(self):
        env,state,sent=f.BuildTests().build_env(4)
        original=f.make_content(4); repaired=copy.deepcopy(original)
        repaired['slides'][1]['content']='Короткая исправленная подпись.'
        env['ask_grok']=AsyncMock(side_effect=[json.dumps(original),json.dumps(original),json.dumps(repaired)])
        validator=env['validate_content']; calls=0
        def validate(*args):
            nonlocal calls
            calls+=1
            if calls==2:raise ValueError('Formatting repair required')
            return validator(*args)
        env['validate_content']=validate
        env['review_content']=AsyncMock(side_effect=[[{'slide':2,'problem':'language','fix':'Уточни формулировку.'}],[]])
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertEqual(len(sent),2)
        final=env['review_content'].await_args_list[-1].args[2]
        self.assertEqual(final['slides'][1]['content'],repaired['slides'][1]['content'])

    async def test_component_photo_review_batches_repairs_before_prose_review(self):
        c=f.make_content(4)
        for s in c['slides']:s['visual_kind']='photo'
        issues=[{'slide':s['number'],'problem':'image_mismatch','fix':'Replace component photo with technical_diagram.'} for s in c['slides']]
        model=AsyncMock(side_effect=[json.dumps({'issues':issues}),json.dumps({'issues':[]})])
        self.assertEqual(await pq.review_content({}, {}, c, model, json.loads),issues)
        self.assertEqual(model.await_count,2)

    async def test_editorial_repair_preserves_unflagged_slides(self):
        env,state,sent=f.BuildTests().build_env(4)
        first=f.make_content(4); revised=copy.deepcopy(first)
        revised['slides'][1]['content']='Unrequested new claim.'
        env['ask_grok']=AsyncMock(side_effect=[json.dumps(first),json.dumps(revised)])
        env['review_content']=AsyncMock(side_effect=[[{'slide':2,'problem':'language','fix':'Use a clearer verb here.'}],[]])
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertEqual(len(sent),2)
        second=env['review_content'].await_args_list[1].args[2]
        self.assertEqual(second['slides'][1]['content'],first['slides'][1]['content'])

    async def test_diagram_bypasses_paid_images_in_both_modes(self):
        for style in ('minimal','designer_soft','designer_strict'):
            env,state,sent=f.BuildTests().build_env(3);state.data['theme_name']=style
            c=f.make_content(3)
            for slide in c['slides']:
                slide.update(diagram=diagram(),visual_kind='technical_diagram',image_prompt='',designer={'layout':'auto','blocks':[]})
            env['ask_grok']=AsyncMock(return_value=json.dumps(c))
            await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
            self.assertEqual(len(sent),2,style)
            self.assertEqual(env['generate_image'].await_count,1) # cover only
            self.assertEqual(env['ask_presentation_vision'].await_count,1)
            env['spend_credits'].assert_called_once()
            saved=Presentation(io.BytesIO(sent[0][1]))
            self.assertTrue(any(s.name=='Diagram connection' for s in saved.slides[1].shapes))

    async def test_unintelligible_language_blocks_but_preferences_do_not(self):
        env,state,sent=f.BuildTests().build_env(3)
        env['review_content']=AsyncMock(return_value=[{'slide':2,'problem':'language','fix':'Назови конкретное действие вместо бессмысленного оборота.'}])
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertFalse(sent);env['generate_image'].assert_not_awaited();env['spend_credits'].assert_not_called()
