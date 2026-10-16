import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn
import presentation_visuals as pv
from presentation_designer import build_scene, render_scene
from presentation_style_direction import decorate_fallback
from presentation_profiles import choose_design
from types import SimpleNamespace
import presentation_support as ps
import test_presentation_support as fixtures


class OwnFrameFlowTests(unittest.IsolatedAsyncioTestCase):
    async def run_flow(self, directory, verdict, caption='слайд 2', context=''):
        source=Path(directory)/'source.png'
        Image.new('RGB',(800,600),'navy').save(source)
        before=source.read_bytes()
        slides=[{'title':f'Слайд {i}', 'content':'Фотография предмета.',
                 '_designer_style':'designer_soft', '_request_context':context,
                 '_skip_ai_image':True} for i in range(2)]
        inspector=AsyncMock(side_effect=verdict) if isinstance(verdict,Exception) else AsyncMock(return_value=json.dumps(verdict))
        generator=AsyncMock(return_value=False)
        assets,warnings=await pv.prepare_images('Предмет',slides,[str(source)],{str(source):caption},
                                                directory,generator,inspector,json.loads)
        generator.assert_not_awaited()
        self.assertEqual(source.read_bytes(),before)
        self.assertEqual(warnings['own_unreadable'],0)
        return assets,inspector

    async def test_explicit_slide_retained_and_frame_review_shared_with_mapping(self):
        with tempfile.TemporaryDirectory() as td:
            assets,inspector=await self.run_flow(td,{'slides':[1], 'preserve_full_frame':False,
                                                    'safe_frames':['circle','rounded','invented','circle',True]})
            self.assertIsNone(assets[0]);self.assertEqual(assets[1]['origin'],'user')
            self.assertEqual(assets[1]['safe_frames'],['circle','rounded'])
            self.assertTrue(assets[1]['frame_reviewed']);self.assertEqual(inspector.await_count,1)
            self.assertIn('треугольник вершиной вверх',inspector.call_args.args[1])

    async def test_protected_or_malformed_review_keeps_photo_and_explicit_slide(self):
        for verdict in [TimeoutError(),[],{}, {'preserve_full_frame':True,'safe_frames':['circle']},
                        {'preserve_full_frame':'false','safe_frames':['circle']},
                        {'preserve_full_frame':False,'safe_frames':'circle'}]:
            with tempfile.TemporaryDirectory() as td:
                assets,_=await self.run_flow(td,verdict)
                self.assertIsNone(assets[0]);self.assertEqual(assets[1]['origin'],'user')
                self.assertFalse(assets[1].get('frame_reviewed',False))
                self.assertFalse(assets[1].get('safe_frames'))

    async def test_no_crop_request_wins_over_classifier(self):
        for caption,context in [('слайд 2, не обрезать',''),('слайд 2','Все мои фото без обрезки')]:
            with tempfile.TemporaryDirectory() as td:
                assets,inspector=await self.run_flow(td,{'preserve_full_frame':False,'safe_frames':['circle']},caption,context)
                inspector.assert_not_awaited();self.assertFalse(assets[1].get('safe_frames'))

    async def test_semantic_assignment_also_returns_safe_shape(self):
        with tempfile.TemporaryDirectory() as td:
            assets,inspector=await self.run_flow(td,{'slides':[2], 'preserve_full_frame':False,'safe_frames':['square']},'Мой предмет')
            self.assertIsNone(assets[0]);self.assertEqual(assets[1]['safe_frames'],['square'])
            self.assertEqual(inspector.await_count,1)

    async def test_real_bot_builder_uses_uploaded_photo_shapes_without_regenerating_them(self):
        for style in ('designer_soft','designer_strict'):
            with tempfile.TemporaryDirectory() as td:
                env,state,sent=fixtures.BuildTests().build_env(3)
                state.data.update(theme_name=style,extra='',photo_mode='own')
                photos=[]
                for i,color in enumerate(('navy','green','orange')):
                    path=Path(td)/f'own-{i}.jpg';Image.new('RGB',(800,600),color).save(path);photos.append(str(path))
                state.data.update(user_photos=photos,pres_photo_captions={p:f'слайд {i+1}' for i,p in enumerate(photos)})
                state.data['pres_plan_key']=ps.plan_key(state.data)
                content=fixtures.make_content(3)
                for slide in content['slides']:slide['designer']={'layout':'auto','blocks':[]}
                env['ask_grok']=AsyncMock(return_value=json.dumps(content))
                env['ask_presentation_vision']=AsyncMock(return_value=json.dumps({'slides':[],
                    'preserve_full_frame':False,'safe_frames':['circle','square','hexagon','chamfer','triangle','rounded']}))
                env['generate_image']=AsyncMock(side_effect=AssertionError('Must keep uploaded photographs'))
                await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321),chat=SimpleNamespace(id=-987654321)),state)
                self.assertEqual(len(sent),2)
                deck=Presentation(io.BytesIO(sent[0][1]));presets=[]
                for slide in deck.slides:
                    pictures=[s for s in slide.shapes if s.name=='User photograph']
                    self.assertEqual(len(pictures),1)
                    presets.append(pictures[0]._element.spPr.find(qn('a:prstGeom')).get('prst'))
                self.assertTrue(any(p in ('ellipse','triangle','hexagon','snip2DiagRect') for p in presets))
                env['generate_image'].assert_not_awaited();env['spend_credits'].assert_called_once()
                self.assertEqual(env['ask_presentation_vision'].await_count,3)


class OwnFrameExportTests(unittest.TestCase):
    def test_approved_shapes_export_with_original_image_and_matching_outlines(self):
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'photo.jpg';Image.new('RGB',(800,600),'navy').save(source)
            for style in ['designer_soft','designer_strict']:
                for mask,number,preset in [('circle',2,'ellipse'),('square',4,'rect'),('triangle',3,'triangle'),
                                           ('hexagon',5,'hexagon'),('chamfer',6,'snip2DiagRect'),('rounded',1,'roundRect')]:
                    if style=='designer_strict' and mask in ('hexagon','chamfer','rounded'):continue
                    item={'title':'Фотография предмета','content':'Фотография предоставлена пользователем.',
                          '_designer_design':choose_design({'topic':'Тоскана','extra':'Дизайн 11','theme_name':style})}
                    asset=dict(path=str(source),width=800,height=600,origin='user',safe_frames=[mask],frame_reviewed=True)
                    scene=build_scene(item,asset,number,8,style=style)
                    self.assertIn('Decorative crop approved',scene['notes'])
                    prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
                    render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
                    stream=io.BytesIO();prs.save(stream);stream.seek(0)
                    slide=Presentation(stream).slides[0]
                    picture=next(s for s in slide.shapes if s.name=='User photograph')
                    self.assertEqual(picture.image.blob,source.read_bytes())
                    self.assertEqual(picture._element.spPr.find(qn('a:prstGeom')).get('prst'),preset)
                    self.assertEqual(len([s for s in slide.shapes if s.name=='Designer photo frame']),2)
                    for frame in [s for s in slide.shapes if s.name=='Designer photo frame']:
                        self.assertEqual(frame._element.spPr.find(qn('a:prstGeom')).get('prst'),preset)
                    if mask in ('circle','square','triangle','hexagon'):
                        self.assertEqual(picture.width,picture.height)
                        self.assertAlmostEqual(picture.crop_left,.125)
                    else:self.assertEqual(picture.crop_left+picture.crop_right+picture.crop_top+picture.crop_bottom,0)

    def test_custom_dense_layout_retains_text_and_image_bounds(self):
        item={'_designer_design':choose_design({'topic':'Тоскана','extra':'Дизайн 11'})}
        for approved in (False,True):
            asset=dict(origin='user',safe_frames=['circle'],frame_reviewed=approved)
            node=dict(type='image',box=[700,250,400,300],radius=0,crop=False)
            text=dict(type='text',role='body',color='#FFFFFF',text='Текст целиком',box=[64,250,500,300])
            scene=dict(background='#102030',nodes=[text,node],notes='Image: supplied by user. Full frame preserved.')
            decorate_fallback(scene,item,asset,2,'designer_soft',custom=True)
            self.assertEqual(text['text'],'Текст целиком');self.assertEqual(text['box'],[64,250,500,300])
            self.assertEqual(node['mask'],'circle' if approved else 'rect')
            self.assertEqual(node['crop'],approved)
            self.assertGreaterEqual(node['box'][0],700);self.assertLessEqual(node['box'][0]+node['box'][2],1100)
            self.assertIn('Decorative crop approved' if approved else 'Full frame preserved',scene['notes'])
