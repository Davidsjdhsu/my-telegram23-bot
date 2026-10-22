"""Exercise the bot builder with photos and editable diagrams in every ordinary theme."""
import ast
import copy
import io
import json
from pathlib import Path
from types import SimpleNamespace
import unittest
from pptx import Presentation
import test_presentation_support as f


def themes():
    tree = ast.parse((Path(__file__).resolve().parents[1] / 'main.py').read_text())
    return next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
                and any(isinstance(t, ast.Name) and t.id == 'THEMES' for t in n.targets))


class UniformBackgroundTests(unittest.IsolatedAsyncioTestCase):
    async def build(self, name, background='theme', custom=None):
        env, state, sent = f.BuildTests().build_env(4)
        palette = themes()
        env['THEMES'] = palette
        state.data.update(theme_name=name, extra='', custom_colors=custom)
        state.data['pres_plan']['design'] = {'background': background, 'accent': ''}
        state.data['pres_plan_key'] = f.ps.plan_key(state.data)
        content = f.make_content(4)
        content['slides'][1].update(visual_kind='technical_diagram', image_prompt='',
            diagram={'kind': 'sequence', 'nodes': [
                {'label': 'Источник', 'text': 'Панель получает солнечный свет.'},
                {'label': 'Результат', 'text': 'Панель вырабатывает электрический ток.'}]})
        env['ask_grok'].return_value = json.dumps(content)
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)), state)
        self.assertEqual(len(sent), 2)
        env['spend_credits'].assert_called_once_with(-987654321, 'presentation', amount=20)
        return Presentation(io.BytesIO(sent[0][1]))

    async def test_all_36_themes_keep_color_on_cover_photo_diagram_and_final(self):
        palette = themes()
        self.assertEqual(len(palette), 36)
        for name, colors in palette.items():
            with self.subTest(theme=name):
                deck = await self.build(name)
                self.assertEqual(len(deck.slides), 4)
                for slide in deck.slides:
                    self.assertEqual(tuple(slide.background.fill.fore_color.rgb), colors['bg'])
                    self.assertNotIn('<a:gradFill', slide._element.xml)
                # Both the picture and editable explanatory text survive serialization.
                self.assertTrue(any(s.shape_type == 13 for s in deck.slides[1].shapes))
                self.assertIn('Источник', ' '.join(s.text for s in deck.slides[2].shapes if s.has_text_frame))

    async def test_explicit_background_overrides_theme_for_every_slide(self):
        for background, expected in [('light', (250, 252, 250)), ('dark', (18, 24, 22))]:
            deck = await self.build('nature', background)
            for slide in deck.slides:
                self.assertEqual(tuple(slide.background.fill.fore_color.rgb), expected)

    async def test_imported_palette_keeps_its_background(self):
        custom = copy.deepcopy(themes()['tech'])
        custom.update(bg=(42, 25, 65), ink=(245, 240, 255), mid=(230, 220, 250))
        before = copy.deepcopy(custom)
        deck = await self.build('nature', custom=custom)
        self.assertEqual(custom, before)
        for slide in deck.slides:
            self.assertEqual(tuple(slide.background.fill.fore_color.rgb), custom['bg'])
