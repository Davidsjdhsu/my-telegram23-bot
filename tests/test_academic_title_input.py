"""Missing title data must be resolved before paid generation in every entry path."""
import unittest
from unittest.mock import AsyncMock

import academic_templates as academic
from test_office import WordWorkflowTests
from test_presentation_support import load_functions


class TitleInputTests(unittest.IsolatedAsyncioTestCase):
    def setup_state(self, data):
        env, m, state = WordWorkflowTests().env()
        state.get_data.side_effect = lambda: dict(data)
        async def update(**kwargs):
            data.update(kwargs)
        state.update_data.side_effect = update
        return env, m, state

    async def test_missing_fields_pause_before_model_job_delivery_and_charge(self):
        for kind in academic.KINDS:
            data = {'word_kind': kind, 'topic': 'Цифровые технологии\nВуз: Учебный вуз', 'word_pages': 6}
            env, m, state = self.setup_state(data)
            await env['word_build'](m, state)
            env['ask_grok'].assert_not_awaited()
            env['start_job'].assert_not_called()
            env['spend_credits'].assert_not_called()
            m.answer_document.assert_not_awaited()
            state.set_state.assert_awaited_once_with('Form:waiting_academic_meta')
            question = m.answer.call_args.args[0]
            self.assertNotIn('Учебное заведение: ', question)
            self.assertIn('Группа: ', question)
            self.assertIn('ФИО студента: ', question)

    async def test_details_append_to_request_and_resume_shared_builder(self):
        data = {'word_kind': 'essay', 'topic': 'Тема', 'extra': 'Без таблиц'}
        env, m, state = self.setup_state(data)
        resume = AsyncMock()
        env['word_build'] = resume
        handler = load_functions(['academic_title_details'], env)['academic_title_details']
        m.text = 'Вуз: Мой вуз\nФИО: Анна Петрова\nГруппа: ИС-21'
        await handler(m, state)
        self.assertTrue(data['extra'].startswith('Без таблиц'))
        self.assertEqual(academic.request_meta(data)['school'], 'Мой вуз')
        self.assertNotIn('school', academic.missing_title_fields(data))
        resume.assert_awaited_once_with(m, state)
        m.text = academic.LEAVE_TITLE_BLANKS
        await handler(m, state)
        self.assertEqual(data['academic_cover_allow_blanks'], academic.title_request_key(data))
        data['topic'] = 'Новая тема'
        self.assertNotEqual(data['academic_cover_allow_blanks'], academic.title_request_key(data))

    async def test_invalid_or_non_text_reply_does_not_start_generation(self):
        env, m, state = self.setup_state({'word_kind': 'essay'})
        env['word_build'] = AsyncMock()
        handler = load_functions(['academic_title_details'], env)['academic_title_details']
        for text in (None, 'просто текст', 'x' * 4001):
            m.text = text
            await handler(m, state)
        env['word_build'].assert_not_awaited()
        state.update_data.assert_not_awaited()

    async def test_explicit_blank_choice_reaches_generator_without_asking_again(self):
        data = {'word_kind': 'essay', 'word_pages': 6, 'topic': 'Тема'}
        data['academic_cover_allow_blanks'] = academic.title_request_key(data)
        env, m, state = self.setup_state(data)
        env['WORD_KIND_LABELS'] = {'essay': {'ru': 'Эссе'}}
        env['ask_grok'].side_effect = RuntimeError('offline fixture')
        await env['word_build'](m, state)
        env['ask_grok'].assert_awaited_once()
        env['start_job'].assert_called_once()
        state.set_state.assert_not_awaited()

    def test_cover_preserves_spacer_geometry_and_manual_fields(self):
        from docx import Document
        reference = Document(academic.blank_path('referat'))
        cover = academic._cover('referat', 'Тема', {})
        for i, paragraph in enumerate(reference.paragraphs):
            if not paragraph.text.strip():
                self.assertEqual(paragraph.paragraph_format.line_spacing,
                                 cover.paragraphs[i].paragraph_format.line_spacing)
                self.assertEqual([r.font.size for r in paragraph.runs],
                                 [r.font.size for r in cover.paragraphs[i].runs])
        text = '\n'.join(p.text for p in cover.paragraphs)
        for value in ('Учебное заведение:', 'ФИО студента:', 'Руководитель работы',
                      'Автор работы', 'Группа', 'Город', 'Год'):
            self.assertIn(value, text)
        self.assertNotIn('2026', text)

    def test_genres_have_their_own_slots_and_fields(self):
        expected = {
            'essay': ('Научный руководитель:', 'Выполнил: студент группы'),
            'report': ('Проверил:', 'Выполнил студент'),
            'notes': ('№ зачётной книжки:', 'Оценка:', 'Дата:', 'Подпись:'),
        }
        for kind, markers in expected.items():
            cover = academic._cover(kind, 'Моя тема', {'course': '2', 'record_book': '12345'})
            text = '\n'.join(p.text for p in cover.paragraphs)
            for marker in markers:
                self.assertIn(marker, text)
            self.assertNotIn('Автор работы', text)
            self.assertNotIn('[' , text)
            if kind != 'notes':
                self.assertNotIn('Подпись:', text)
                self.assertNotIn('Оценка:', text)
            if kind == 'report':
                self.assertIn('«МОЯ ТЕМА»', text)
            if kind == 'notes':
                self.assertIn('12345', text)

    def test_questions_follow_selected_cover_and_do_not_ask_unused_discipline(self):
        for kind in ('essay', 'report', 'coursework'):
            self.assertNotIn('discipline', academic.missing_title_fields({'word_kind': kind}))
        for kind in ('notes', 'report'):
            self.assertIn('course', academic.missing_title_fields({'word_kind': kind}))
            self.assertNotIn('course', academic.missing_title_fields({'word_kind': kind, 'extra': 'Курс: 2'}))
        parsed = academic.explicit_meta({}, 'Номер зачетной книжки: 123\nДолжность преподавателя: Доцент\nИнститут: Институт технологий')
        self.assertEqual(parsed['record_book'], '123')
        self.assertEqual(parsed['teacher_position'], 'Доцент')
        self.assertEqual(parsed['institute'], 'Институт технологий')

    def test_assets_have_distinct_provenance_and_unchanged_tsu_sources(self):
        import hashlib, json
        sources = json.loads((academic.ASSETS / 'sources.json').read_text())
        forms = {item['kind']: item for item in sources['forms']}
        for kind in academic.KINDS:
            self.assertEqual(hashlib.sha256(academic.blank_path(kind).read_bytes()).hexdigest(),
                             forms[kind]['sha256'])
        self.assertEqual(len({forms[k]['source'] for k in ('essay', 'notes', 'report')}), 3)
        self.assertEqual(forms['referat']['sha256'], 'c79f2e0ef6bbe4d42946c7fa531b34fc43becec8cb75f5614bddfd0543166ff2')
        self.assertEqual(forms['coursework']['sha256'], '37375b78a25a58517e990d7029393f62fea1b72ef09b6e9f9867759250876e70')
