import tempfile,unittest
from pathlib import Path
from docx import Document
import fixed_word_templates as fixed
import test_fixed_word_templates as existing

KINDS=set('dkp rent proxy statement act loan claim consent marriage_contract gift lawsuit alimony offer services employment work_act supply agency joint_activity nonresidential_rent cession nda self_employed warranty_letter'.split())

class CoverageTests(unittest.TestCase):
    def test_catalog_covers_remaining_menu_kinds(self):
        self.assertTrue(KINDS<=fixed.CATALOG.keys())
        for kind in KINDS:
            spec=fixed.CATALOG[kind]
            self.assertTrue(fixed.applies({'word_kind':kind},'ru'))
            self.assertFalse(fixed.applies({'word_kind':kind},'en'))
            self.assertFalse(fixed.applies({'word_kind':kind,'content_lang':'en'},'ru'))
            self.assertEqual(set(spec['fields']),set(spec['sample']))
            self.assertLess(len(spec['scope']+fixed.form_text(kind)),3500)
    def test_specialised_car_dispatch_and_study_are_preserved(self):
        self.assertFalse(fixed.applies({'word_kind':'dkp','topic':'Продать автомобиль'}))
        self.assertTrue(fixed.applies({'word_kind':'dkp','topic':'Продать велосипед'}))
        for k in ['dkp_car','dkp_realty','doc','referat','essay','coursework']:
            self.assertFalse(fixed.applies({'word_kind':k}))
    def test_notary_warning_is_in_downloaded_file(self):
        for kind in ['consent','marriage_contract','alimony']:
            spec=fixed.CATALOG[kind];d=Document(fixed.ASSETS/spec['file'])
            text='\n'.join(p.text for p in d.paragraphs)
            self.assertIn('ПРОЕКТ ДЛЯ НОТАРИУСА',text)
            self.assertIn('не удостоверен',text)
            self.assertNotIn('зарегистрировано в реестре',text.casefold())
    def test_identity_address_is_not_deal_subject(self):
        self.assertFalse(fixed.scope_issue({'word_kind':'dkp','topic':'Реквизиты первой стороны: Москва, квартира 5\nПредмет продажи и идентификаторы: Велосипед'}))
        self.assertTrue(fixed.scope_issue({'word_kind':'dkp','topic':'Предмет продажи и идентификаторы: Квартира'}))
    def test_scope_rejects_known_incompatible_variants(self):
        for kind,text in [('rent','Аренда автомобиля'),('proxy','Право распоряжаться недвижимостью'),('gift','Дарение квартиры'),('employment','Срочный договор'),('lawsuit','Арбитражный иск'),('agency','Действовать от своего имени'),('work_act','КС-2'),('cession','Уступка алиментов')]:
            self.assertTrue(fixed.scope_issue({'word_kind':kind,'topic':text}),(kind,text))
        self.assertFalse(fixed.scope_issue({'word_kind':'employment','topic':'Бессрочный договор без испытательного срока'}))
        self.assertFalse(fixed.scope_issue({'word_kind':'rent','topic':'Аренда офисного стула для офиса'}))
        self.assertTrue(fixed.scope_issue({'word_kind':'rent','topic':'Аренда офиса'}))
    def test_templates_do_not_assert_unconfirmed_payment(self):
        for kind in ['loan','act','work_act','warranty_letter']:
            text=' '.join(p.text for p in Document(fixed.ASSETS/fixed.CATALOG[kind]['file']).paragraphs)
            self.assertNotIn('деньги получил',text.casefold())
            self.assertNotIn('претензий не име',text.casefold())

class WorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def test_notary_warning_before_any_delivery_or_charge(self):
        for kind in ['consent','marriage_contract','alimony']:
            e,m,s=existing.WorkflowTests().env(kind)
            await e['word_build'](m,s)
            self.assertIn('ПРОЕКТ ДЛЯ НОТАРИУСА',str(m.answer.call_args_list))
            m.answer_document.assert_not_awaited();e['spend_credits'].assert_not_called()
    async def test_generic_sale_about_vehicle_still_uses_car_workflow(self):
        e,m,s=existing.WorkflowTests().env('dkp','Продаю автомобиль')
        await e['word_build'](m,s)
        self.assertIn('car_sale',s.data);self.assertNotIn('fixed_word',s.data)
        m.answer_document.assert_not_awaited()
    async def test_changed_scope_cannot_reuse_approved_fields(self):
        e,m,s=existing.WorkflowTests().env('gift')
        await e['word_build'](m,s)
        await s.update_data(extra='Дар и идентификаторы: квартира')
        await e['word_build'](m,s)
        m.answer_document.assert_not_awaited();e['spend_credits'].assert_not_called()
