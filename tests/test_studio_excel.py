import asyncio, json, tempfile, unittest
from pathlib import Path
from unittest.mock import AsyncMock, Mock
from aiohttp import web, CookieJar
from aiohttp.test_utils import TestClient, TestServer
from openpyxl import load_workbook
from studio_private import Studio, PILOT_UID
from studio_excel import ExcelWorkflow
from studio_excel_catalog import normalize
from studio_excel_runtime import attach, validate_file
import test_studio_word
from test_studio_private import HEADERS, ORIGIN
from test_office import office_env

ROWS=[{'name':'Краска','unit':'банка','qty':3,'price':3400,'category':'Материалы'},
      {'name':'Работы','unit':'м²','qty':24,'price':50000,'category':'Работы'}]
PAYLOAD=dict(kind='expense_estimate',title='Контрольная смета',rows=ROWS)

class InputTests(unittest.TestCase):
 def test_bad_values_and_missing_numbers_never_become_silent_zero(self):
  for value in [True,False,None,{},'NaN','inf',-1,'',1e12,'0.0001']:
   with self.subTest(value=value), self.assertRaises(ValueError):normalize({**PAYLOAD,'rows':[{**ROWS[0],'price':value}]})
  self.assertEqual(normalize({**PAYLOAD,'rows':[{**ROWS[0],'price':'1 234,50'}]})['rows'][0]['price'],1234.5)
  self.assertEqual(normalize({**PAYLOAD,'rows':[{**ROWS[0],'price':0}]})['rows'][0]['price'],0)
  self.assertEqual(normalize({**PAYLOAD,'rows':[{**ROWS[0],'price':''}]},draft=True)['rows'][0]['price'],'')
  for body in [[],{'kind':'price_list'}, {**PAYLOAD,'rows':ROWS*31},{**PAYLOAD,'style':[]},{**PAYLOAD,'palette':{}},{**PAYLOAD,'title':'a'*181}]:
   with self.assertRaises(ValueError):normalize(body)
 def test_real_workbook_all_palettes_and_literal_input(self):
  build=office_env()['build_excel_items']
  with tempfile.TemporaryDirectory() as directory:
   for style,palette in [('standard','copper')]+[('designer',p) for p in ('green','copper','emerald','lilac','burgundy')]:
    p=normalize({**PAYLOAD,'style':style,'palette':palette,'rows':[{**ROWS[0],'name':'=1+1'},ROWS[1]]})
    path=str(Path(directory)/'estimate.xlsx')
    build(path,p['title'],p['kind'],{},p['rows'],excel_style=p['style'],excel_palette=p['palette'],excel_protect=False)
    validate_file(path,p)
    wb=load_workbook(path);ws=wb.active;first=5 if style=='standard' else 9
    self.assertEqual(ws.cell(first,2).data_type,'s');self.assertFalse(ws.protection.sheet)
    self.assertEqual(load_workbook(path,data_only=True).active.cell(first+3 if style=='designer' else first+2,6).value,1210200)

class WorkflowTests(unittest.IsolatedAsyncioTestCase):
 signed=test_studio_word.WorkflowTests.signed
 login=test_studio_word.WorkflowTests.login
 async def asyncSetUp(self):
  import hmac,hashlib,time
  from urllib.parse import parse_qsl
  from test_presentation_support import load_functions
  verify=load_functions(['verify_telegram_init_data','miniapp_verified_user'],{'BOT_TOKEN':'test-token','INIT_DATA_MAX_AGE_SECONDS':604800,'parse_qsl':parse_qsl,'hmac':hmac,'hashlib':hashlib,'json':json,'time':time})['miniapp_verified_user']
  self.tmp=tempfile.TemporaryDirectory();self.studio=Studio(verify,AsyncMock(),AsyncMock(),ORIGIN)
  self.send=AsyncMock();self.charge=AsyncMock(return_value=5);self.parse=AsyncMock(return_value={'title':'Смета','rows':ROWS})
  self.afford=Mock(return_value=True)
  async def build(user,job):
   (job.directory/'document.xlsx').write_bytes(b'test')
   await self.flow.deliver(job)
  self.builder=AsyncMock(side_effect=build)
  self.flow=ExcelWorkflow(self.studio,self.tmp.name,self.builder,self.send,self.charge,self.afford,self.parse,5)
  app=web.Application();self.flow.register(app);self.studio.register(app)
  self.client=TestClient(TestServer(app),cookie_jar=CookieJar(unsafe=True));await self.client.start_server();await self.login()
 async def asyncTearDown(self):
  await self.client.close()
  for archive in self.studio.archives:archive.close()
  self.tmp.cleanup()
 async def create(self,**kw):
  r=await self.client.post('/studio/excel/orders',headers=HEADERS,json={**PAYLOAD,**kw});self.assertEqual(r.status,200,await r.text());return await r.json()
 async def confirm(self,o,**kw):return await self.client.post('/studio/excel/orders/'+o['id']+'/confirm',headers=HEADERS,json={'confirmedCost':o['cost'],**kw})
 async def settle(self):await asyncio.gather(*[v[1] for v in self.flow.tasks.values()])
 async def status(self,o):return await (await self.client.get('/studio/excel/orders/'+o['id'])).json()
 async def test_auth_owner_csrf_and_server_price(self):
  o=await self.create(cost=0);self.assertEqual(o['cost'],20)
  self.assertEqual((await self.confirm(o,confirmedCost=0)).status,409)
  self.assertEqual((await self.client.post('/studio/excel/orders',json=PAYLOAD)).status,403)
  await self.login(909828109);self.assertEqual((await self.client.get('/studio/excel/orders/'+o['id'])).status,404)
  self.builder.assert_not_awaited();self.charge.assert_not_awaited()
 async def test_double_click_restore_and_charge_after_send(self):
  events=[];self.send.side_effect=lambda *a:events.append('send');self.charge.side_effect=lambda *a:events.append('charge') or 5
  o=await self.create();await self.confirm(o);await self.confirm(o);await self.settle()
  self.assertEqual(events,['send','charge']);self.builder.assert_awaited_once()
  self.assertEqual((await self.status(o))['status'],'completed')
  self.assertEqual((await self.create())['id'],o['id'])
  self.assertEqual((await (await self.client.get('/studio/excel/draft')).json())['order']['status'],'completed')
  await self.confirm(o);self.charge.assert_awaited_once()
 async def test_send_timeout_never_resends_or_debits(self):
  self.send.side_effect=TimeoutError();o=await self.create();await self.confirm(o);await self.settle()
  self.assertEqual((await self.status(o))['status'],'delivery_unknown');self.charge.assert_not_awaited()
  self.assertEqual((await self.confirm(o)).status,409)
  job=self.flow.get(PILOT_UID,o['id']);job.row['created']-=90000;job.save()
  r=await self.client.post('/studio/excel/orders',headers=HEADERS,json={**PAYLOAD,'title':'Другая'});self.assertEqual(r.status,409)
 async def test_failure_keeps_source_and_resumes(self):
  original=self.builder.side_effect;self.builder.side_effect=RuntimeError('fixture')
  o=await self.create(source='исходный список');await self.confirm(o);await self.settle()
  self.assertEqual((await self.status(o))['status'],'interrupted');self.send.assert_not_awaited()
  self.assertEqual(self.flow.get(PILOT_UID,o['id']).row['payload']['source'],'исходный список')
  self.builder.side_effect=original;await self.confirm(o);await self.settle();self.charge.assert_awaited_once()
 async def test_insufficient_balance_does_not_start(self):
  o=await self.create();self.afford.return_value=False
  self.assertEqual((await self.confirm(o)).status,402);self.builder.assert_not_awaited()
 async def test_draft_and_import_are_private_cached_and_no_build(self):
  body={**PAYLOAD,'source':'Краска 3 по 3400; работы 24 по 50000'}
  self.assertEqual((await self.client.post('/studio/excel/draft',headers=HEADERS,json=body)).status,200)
  for _ in range(2):self.assertEqual((await self.client.post('/studio/excel/parse',headers=HEADERS,json=body)).status,200)
  self.parse.assert_awaited_once();self.charge.assert_not_awaited();self.builder.assert_not_awaited()
  await self.login(909828109);self.assertIsNone((await (await self.client.get('/studio/excel/draft')).json())['draft'])
 async def test_restart_intent_and_safe_build(self):
  o=await self.create();job=self.flow.get(PILOT_UID,o['id']);job.row.update(status='building');job.save()
  restored=ExcelWorkflow(self.studio,self.tmp.name,self.builder,self.send,self.charge,self.afford,self.parse,5)
  self.assertEqual(restored.get(PILOT_UID,o['id']).row['status'],'interrupted')
  job.row.update(status='delivering',sending='xlsx');job.save()
  restored=ExcelWorkflow(self.studio,self.tmp.name,self.builder,self.send,self.charge,self.afford,self.parse,5)
  self.assertEqual(restored.get(PILOT_UID,o['id']).row['status'],'delivery_unknown')

class RuntimeTests(unittest.IsolatedAsyncioTestCase):
 async def test_real_build_delivery_and_idempotent_charge(self):
  with tempfile.TemporaryDirectory() as d:
   user={'credits':20,'history':[]};send=AsyncMock();save=AsyncMock(return_value=True)
   async def heavy(fn,*args,**kw):return fn(*args,**kw)
   api=dict(USERS_FILE=str(Path(d)/'users.json'),get_user=lambda uid:user,can_afford=lambda uid,c:user['credits']>=c,
       ADMIN_IDS=set(),save_users_now=save,note_success=Mock(),CREDIT_COSTS={'excel':5},miniapp_user_record=Mock(),
       _miniapp_action_tasks={},start_job=Mock(return_value=(True,'')),finish_job=Mock(),run_heavy=heavy,
       build_excel_items=office_env()['build_excel_items'],bot=Mock(send_document=send),FSInputFile=lambda p,**kw:p,safe_filename=lambda *a,**kw:'Смета')
   service=attach(Mock(),web.Application(),api)
   from studio_word import WordJob
   token='a'*32;directory=service.root/token;directory.mkdir()
   job=WordJob(service,dict(id=token,uid=1,payload=normalize(PAYLOAD),cost=5,title='Смета',sent=[]));job.save()
   await service.build({'id':1},job)
   self.assertEqual(job.row['status'],'completed');self.assertEqual(user['credits'],15);send.assert_awaited_once()
   self.assertEqual(await service.charge(1,token,5,'Смета'),5);self.assertEqual(user['credits'],15)
   api['finish_job'].assert_called_once()
 async def test_pending_billing_blocks_retry_and_admin_free(self):
  with tempfile.TemporaryDirectory() as d:
   user={'credits':20,'history':[]};save=AsyncMock(return_value=True)
   api=dict(USERS_FILE=str(Path(d)/'users.json'),get_user=lambda uid:user,can_afford=lambda *a:True,ADMIN_IDS={1},save_users_now=save,note_success=Mock(),CREDIT_COSTS={'excel':5})
   service=attach(Mock(),web.Application(),api)
   self.assertEqual(await service.charge(1,'free',5,'Смета'),0);self.assertEqual(user['credits'],20)
   user['studio_excel_payments']['pending']={'status':'pending','charged':5}
   with self.assertRaises(RuntimeError):await service.charge(2,'pending',5,'Смета')
