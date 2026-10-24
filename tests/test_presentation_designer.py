import io
import json
import ast
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn
import presentation_designer as pd
import presentation_visuals as pv
import presentation_support as ps
import test_presentation_support as f

ROOT=Path(__file__).parents[1]

class DesignerTests(unittest.TestCase):
    def asset(self, directory, origin='ai', size=(1600,900)):
        path=Path(directory)/'test.png';Image.new('RGB',size,'navy').save(path)
        return {'path':str(path),'width':size[0],'height':size[1],'origin':origin}

    def test_explicit_style_only_and_custom_template_wins(self):
        self.assertFalse(pd.is_designer({'theme_name':'tech'}))
        self.assertTrue(pd.is_designer({'theme_name':'designer_soft'}))
        self.assertTrue(pd.is_designer({'theme_name':'designer_strict'}))
        self.assertFalse(pd.is_designer({'theme_name':'designer_soft','custom_colors':{'bg':[1,2,3]}}))

    def test_plan_explains_real_style_and_cache_changes(self):
        a={'theme_name':'designer_soft','slides':8};b={**a,'theme_name':'designer_strict'}
        self.assertNotEqual(ps.plan_key(a),ps.plan_key(b))
        self.assertIn('скруглённые',ps.plan_prompt(a,'ru'))
        self.assertIn('прямые границы',ps.plan_prompt(b,'ru'))

    def test_prompts_compose_cover_but_not_ordinary_photos(self):
        prompt=pv.image_prompt('Киты',{'title':'Киты','_designer_style':'designer_soft','_designer_cover':True})
        self.assertIn('RIGHT half',prompt);self.assertIn('Киты',prompt)
        self.assertNotIn('RIGHT half',pv.image_prompt('Киты',{'title':'Киты'}))

    def test_soft_and_strict_editable_geometry_survive_serialization(self):
        with tempfile.TemporaryDirectory() as td:
            a=self.asset(td)
            for style,shape in [('designer_soft','roundRect'),('designer_strict','rect')]:
                scene=pd.build_scene({'title':'Энергия на крыше','content':'Панели вырабатывают электричество.'},a,2,8,style=style)
                prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
                pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
                stream=io.BytesIO();prs.save(stream);stream.seek(0);s=Presentation(stream).slides[0]
                pic=next(x for x in s.shapes if x.shape_type==13)
                self.assertEqual(pic._element.spPr.find(qn('a:prstGeom')).get('prst'),shape)
                self.assertIn('Панели', ' '.join(x.text for x in s.shapes if x.has_text_frame))

    def test_user_frame_preserved_at_all_aspect_ratios(self):
        with tempfile.TemporaryDirectory() as td:
            for size in [(300,1200),(1600,350),(800,800)]:
                a=self.asset(td,'user',size)
                for number in (1,2):
                    scene=pd.build_scene({'title':'Фото пользователя','content':'Весь кадр.'},a,number,8)
                    n=next(x for x in scene['nodes'] if x['type']=='image')
                    self.assertFalse(n['crop']);self.assertEqual(n['radius'],0)
                    self.assertAlmostEqual(n['box'][2]/n['box'][3],size[0]/size[1])

    def test_dense_copy_keeps_all_words_and_photo(self):
        with tempfile.TemporaryDirectory() as td:
            body=('Проверенное предложение. '*19).strip()
            scene=pd.build_scene({'title':'Особенности использования солнечных панелей в разных климатических условиях','content':body},self.asset(td),2,8)
            actual=' '.join(x['text'] for x in scene['nodes'] if x.get('role')=='body')
            self.assertEqual(' '.join(actual.split()),body)
            self.assertEqual(sum(x['type']=='image' for x in scene['nodes']),1)

    def test_uploaded_photo_keeps_structured_content(self):
        with tempfile.TemporaryDirectory() as td:
            blocks=[{'heading':'Плюсы','text':'Факт А.'},{'heading':'Минусы','text':'Факт Б.'}]
            scene=pd.build_scene({'title':'Сравнение','content':'Два варианта','designer':{'layout':'comparison','blocks':blocks}},self.asset(td,'user'),2,8)
            txt=' '.join(x.get('text','') for x in scene['nodes'])
            for w in ('Плюсы','Минусы','Факт А.','Факт Б.'):self.assertIn(w,txt)

    def test_malformed_blocks_rejected(self):
        for d in [[],{'layout':'fake'},{'layout':'comparison','blocks':[]}]:
            with self.assertRaises(pd.DesignerValidationError):pd.validate_designer({'designer':d})

    def test_edited_chart_reaches_chart_renderer(self):
        chart={'categories':['A','B'],'series':[{'name':'Data','values':[2,4]}]}
        scene=pd.build_scene({'title':'Данные','content':'Из исходных данных.','chart':chart},None,2,3)
        prs=Presentation();render=Mock();pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene,render)
        self.assertEqual(render.call_args.args[5],chart)

    def test_fallback_font_metrics_are_conservative(self):
        metrics=pd.ConservativeMetrics(30)
        self.assertGreater(metrics.getlength('Большой кит'),0)
        self.assertGreater(metrics.getlength('汉字'),metrics.getlength('aa'))

class FlowTests(unittest.IsolatedAsyncioTestCase):
    async def test_no_ai_calls_for_schemes_but_uploaded_photo_kept(self):
        with tempfile.TemporaryDirectory() as td:
            slide={'title':'Схема','content':'Текст','_skip_ai_image':True};gen=AsyncMock();inspect=AsyncMock()
            assets,w=await pv.prepare_images('Тема',[slide],[],{},td,gen,inspect,json.loads)
            self.assertEqual(assets,[None]);self.assertEqual(w['ai_omitted'],0);gen.assert_not_awaited()
            a=DesignerTests().asset(td,'user')
            assets,w=await pv.prepare_images('Тема',[slide],[a['path']],{a['path']:'слайд 1'},td,gen,inspect,json.loads)
            self.assertEqual(assets[0]['origin'],'user');gen.assert_not_awaited()

    async def test_builder_billing_and_style_reset(self):
        for name in ('designer_soft','designer_strict','minimal'):
            env,state,sent=f.BuildTests().build_env(3)
            state.data['theme_name']=name;state.data['pres_plan_key']=ps.plan_key(state.data)
            content=f.make_content(3)
            for slide in content['slides']:slide['designer']={'layout':'auto','blocks':[]}
            env['ask_grok']=AsyncMock(return_value=json.dumps(content))
            message=SimpleNamespace(from_user=SimpleNamespace(id=-987654321),chat=SimpleNamespace(id=-987654321))
            await env['_build_presentation'](message,state)
            self.assertEqual(len(sent),2)
            prs=Presentation(io.BytesIO(sent[0][1]));self.assertEqual(len(prs.slides),3)
            names=[s.name for s in prs.slides[0].shapes]
            self.assertEqual(any(n.startswith('Designer') for n in names),name!='minimal')
            env['spend_credits'].assert_called_once_with(-987654321,'presentation',amount=48 if name=='minimal' else 96)
            self.assertEqual(state.data,{})

    async def test_invalid_designer_response_not_charged(self):
        env,state,sent=f.BuildTests().build_env(3);state.data['theme_name']='designer_soft';state.data['pres_plan_key']=ps.plan_key(state.data)
        data=f.make_content(3);data['slides'][0]['designer']={'layout':'bad'}
        env['ask_grok']=AsyncMock(return_value=json.dumps(data))
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321),chat=SimpleNamespace(id=-987654321)),state)
        self.assertEqual(sent,[]);env['spend_credits'].assert_not_called();self.assertEqual(env['ask_grok'].await_count,3)

    async def test_style_buttons_update_and_clear_old_choice(self):
        env=f.load_functions(['process_theme'],f.WorkflowTests().env(
            STYLE_BY_LABEL={'Мягкий':'designer_soft','Минимализм':'minimal'},
            TR={'btn_keep_style':{'ru':'Оставить'}},style_kb=lambda **k:None))
        state=f.State({'theme_name':'minimal','custom_colors':{'bg':[1,2,3]}})
        for label,key in [('Мягкий','designer_soft'),('Минимализм','minimal')]:
            await env['process_theme'](f.WorkflowTests().message(label),state)
            self.assertEqual(state.data['theme_name'],key)
            self.assertIsNone(state.data['custom_colors'])
