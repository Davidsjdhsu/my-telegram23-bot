import io
import tempfile
import unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn
from presentation_profiles import choose_design
from presentation_style_direction import build_style_scene
from presentation_designer import render_scene

class PhotoFrameTests(unittest.TestCase):
    def scene(self, asset, number=2, style='designer_strict'):
        item={'title':'Фотография','content':'Главный объект в кадре.',
              '_designer_design':choose_design({'topic':'Бутик-отель в Тоскане','theme_name':style,'extra':'Дизайн 11'})}
        return build_style_scene(item,asset,number,8,'ru',style)

    def test_uploaded_and_unreviewed_keep_entire_frame(self):
        with tempfile.TemporaryDirectory() as td:
            for size in ((1600,900),(400,1400),(2000,300)):
                p=Path(td)/'photo.png';Image.new('RGB',size,'navy').save(p)
                for origin,allowed in [('user',['circle','triangle','square']),('ai',[])]:
                    for style in ('designer_strict','designer_soft'):
                        a=dict(path=str(p),width=size[0],height=size[1],origin=origin,safe_frames=allowed)
                        s=self.scene(a,style=style)
                        n=next(n for n in s['nodes'] if n['type']=='image')
                        self.assertFalse(n['crop']);self.assertEqual(n['mask'],'rect')
                        self.assertAlmostEqual(n['box'][2]/n['box'][3],size[0]/size[1])
                        if origin=='user':self.assertEqual(n['radius'],0)

    def test_reviewed_masks_and_matching_frames_survive_pptx(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'photo.png';Image.new('RGB',(1600,900),'navy').save(p)
            cases=[('designer_strict',2,'circle','ellipse'),('designer_strict',3,'triangle','triangle'),
                   ('designer_strict',4,'square','rect'),('designer_soft',2,'circle','ellipse'),
                   ('designer_soft',3,'square','rect'),('designer_soft',4,'triangle','triangle'),
                   ('designer_soft',5,'hexagon','hexagon'),('designer_soft',6,'chamfer','snip2DiagRect')]
            for style,number,mask,preset in cases:
                a=dict(path=str(p),width=1600,height=900,origin='ai',safe_frames=[mask])
                scene=self.scene(a,number,style)
                prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
                render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
                stream=io.BytesIO();prs.save(stream);stream.seek(0)
                slide=Presentation(stream).slides[0]
                photo=next(s for s in slide.shapes if s.name=='AI illustration')
                self.assertEqual(photo._element.spPr.find(qn('a:prstGeom')).get('prst'),preset)
                if mask=='chamfer':
                    self.assertAlmostEqual(photo.width/photo.height,1600/900,places=4)
                    self.assertEqual(photo.crop_left,0)
                else:
                    self.assertEqual(photo.width,photo.height)
                    self.assertAlmostEqual(photo.crop_left,(1-900/1600)/2,places=4)
                frames=[s for s in slide.shapes if s.name=='Designer photo frame']
                self.assertEqual(len(frames),2)
                for frame in frames:
                    self.assertEqual(frame._element.spPr.find(qn('a:prstGeom')).get('prst'),preset)
                    self.assertLess(frame.left,photo.left)
                    self.assertGreater(frame.width,photo.width)

if __name__=='__main__':unittest.main()
