import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from test_office import office_env, COLORS
from excel_styles import PALETTES

ROWS=[{'name':'Зарплата','type':'income','amount':100000,'category':'Основное'},
      {'name':'Подработка','type':'доход','amount':20000,'category':'Основное'},
      {'name':'Аренда','type':'expense','amount':80000,'category':'Основное'}]

class BudgetTests(unittest.TestCase):
    def test_all_layouts_preserve_separate_flows_and_inputs(self):
        build=office_env()['build_excel_items']
        for kind in ['family_budget','project_budget']:
            for style,palette in [('standard','burgundy')]+[('designer',p) for p in PALETTES]:
                with self.subTest(kind=kind,style=style,palette=palette), tempfile.TemporaryDirectory() as td:
                    path=str(Path(td)/'test.xlsx')
                    build(path,'Бюджет',kind,COLORS,ROWS,excel_style=style,excel_palette=palette)
                    w=load_workbook(path);ws=w.active;cached=load_workbook(path,data_only=True).active
                    first=9 if style=='designer' else 5
                    inc=first+len(ROWS)+(1 if style=='designer' else 0)
                    result=inc+2;summary=result+4
                    self.assertEqual(cached[f'D{inc}'].value,120000)
                    self.assertEqual(cached[f'D{inc+1}'].value,80000)
                    self.assertEqual(cached[f'D{result}'].value,40000)
                    self.assertEqual([cached.cell(summary,c).value for c in (3,4,5)],[120000,80000,40000])
                    self.assertEqual(ws.cell(first-1,2).fill.fgColor.rgb[-6:],PALETTES[palette if style=='designer' else 'green'][0])
                    self.assertIn('SUMIF(',ws[f'D{inc}'].value)
                    self.assertIn('SUMIFS(',ws[f'C{summary}'].value)
                    self.assertIn(f'B{summary}',ws[f'C{summary}'].value)
                    self.assertFalse(ws[f'D{first}'].protection.locked)
                    self.assertTrue(ws[f'D{result}'].protection.locked)
                    note=next(c for row in ws for c in row if isinstance(c.value,str) and 'подсказку можно удалить' in c.value)
                    self.assertFalse(note.protection.locked)
                    self.assertEqual(len(ws._charts),1)
                    self.assertEqual(len(ws.data_validations.dataValidation),2)
                    formulas=[f for rules in ws.conditional_formatting._cf_rules.values() for rule in rules for f in rule.formula]
                    self.assertIn(f'AND(ISNUMBER(D{result}),D{result}<0)',formulas)
                    self.assertEqual(w.calculation.calcMode,'auto')

    def test_missing_amount_differs_from_zero(self):
        build=office_env()['build_excel_items']
        for amount,expected in [(None,'Заполните данные'),('', 'Заполните данные'),('не указано','Заполните данные'),(0,0)]:
            with tempfile.TemporaryDirectory() as td:
                p=str(Path(td)/'test.xlsx')
                build(p,'Тест','family_budget',COLORS,[dict(name='Аренда',type='расход',amount=amount)])
                ws=load_workbook(p,data_only=True).active
                self.assertEqual(ws['D8'].value,expected)
                self.assertEqual(ws['D5'].value,0 if amount==0 else None)

    def test_injected_labels_stay_text_and_invalid_types_fail(self):
        build=office_env()['build_excel_items']
        with tempfile.TemporaryDirectory() as td:
            p=str(Path(td)/'test.xlsx')
            build(p,'=1+1','family_budget',COLORS,[dict(name='=2+2',type='расход',amount=5,category='A*?"')])
            ws=load_workbook(p).active
            self.assertEqual(ws['A1'].data_type,'s');self.assertEqual(ws['B5'].data_type,'s')
            with self.assertRaises(ValueError):
                build(p,'Тест','family_budget',COLORS,[dict(name='X',type='Неизвестно',amount=5)])
            with self.assertRaises(ValueError):
                build(p,'Тест','family_budget',COLORS,[dict(name='X',type='доход',amount=float('nan'))])
