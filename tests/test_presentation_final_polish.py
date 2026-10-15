import asyncio
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock
from types import SimpleNamespace
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import presentation_visuals as pv
import presentation_support as ps
import presentation_designer as pd
from test_presentation_catalog import specimen
import test_presentation_support as base


class FinalPolishTests(unittest.TestCase):
    def test_cover_directions_removed_but_ai_topic_preserved(self):
        p=base.make_plan()
        p['slides'][0]['point']='Улей — дом пчелиной семьи; на обложке ИИ-иллюстрация улья среди цветов.'
        clean=ps.finalize_plan_design(p,{'theme_name':'minimal'})
        self.assertEqual(clean['slides'][0]['point'],'Улей — дом пчелиной семьи.')
        ps.validate_plan(clean,8)
        self.assertIn('на обложке',p['slides'][0]['point'])
        for text in ('ИИ-иллюстрации помогают художникам искать идеи.', 'На обложке книги указано имя автора.',
                     'На обложке книги изображён парусник.'):
            self.assertEqual(ps.clean_cover_point(text,'Тема'),text)
        self.assertEqual(ps.clean_cover_point('На обложке изображение улья.','Пчёлы'),'Пчёлы')

    def test_whole_deck_budget_not_slide_or_class_count(self):
        for text,expected in [
            ('6 слайдов для 9–11 классов. Добавь 3 подходящие ИИ-иллюстрации.',(3,True)),
            ('Добавь три картинки.',(3,False)),('Add 2 suitable AI images.',(2,True)),
            ('6 слайдов, 3 фотографии. Теперь две картинки.',(2,False)),
            ('Без изображений.',(0,False)),('Без ИИ-иллюстраций.',(0,True)),
            ('3 фото через ИИ',(3,True)),('6 слайдов для 3 класса',None),
            ('2–3 картинки',None),('по 2 картинки на каждом слайде',None)]:
            with self.subTest(text=text):self.assertEqual(pv.requested_image_budget(text),expected)

    def test_budget_retains_uploads_and_native_diagrams(self):
        slides=[{'title':str(i)} for i in range(6)]
        slides[3]['_skip_ai_image']=True;slides[5]['_skip_ai_image']=True
        pv.apply_image_budget(slides,[None]*6,'Добавь 3 ИИ-иллюстрации')
        self.assertEqual([i for i,s in enumerate(slides) if not s.get('_skip_ai_image')],[0,2,4])
        for request,expected in [('2 фотографии',1),('2 ИИ-иллюстрации',2),('Без изображений',0)]:
            ss=[{} for _ in range(5)];assets=[{'origin':'user'},None,None,None,None]
            pv.apply_image_budget(ss,assets,request)
            self.assertEqual(sum(not s.get('_skip_ai_image') and not a for s,a in zip(ss,assets)),expected)
            self.assertEqual(assets[0]['origin'],'user')

    def test_label_removed_and_complete_frame_serialized_all_renderers(self):
        env,_,_=base.BuildTests().build_env(3)
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'photo.png';Image.new('RGB',(1600,900),'gold').save(path)
            asset={'path':str(path),'width':1600,'height':900,'origin':'ai'}
            prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
            for kind in ('catalog','profile','legacy','ordinary'):
                item,n=specimen('d01','photo_bottom')
                if kind=='profile':item['_designer_design'].pop('template_id')
                if kind=='legacy':item.pop('_designer_design')
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                if kind=='ordinary':pv.render_slide(slide,item,asset,env['THEMES']['minimal'],n,12,'ru',env['slide_background'],Mock())
                else:
                    scene=pd.build_scene(item,asset,n,12)
                    pd.render_scene(slide,scene,Mock())
                visible=' '.join(s.text for s in slide.shapes if s.has_text_frame)
                self.assertNotIn('ИИ-иллюстрация',visible)
                self.assertIn('AI-generated',slide.notes_slide.notes_text_frame.text)
            out=io.BytesIO();prs.save(out);out.seek(0);loaded=Presentation(out)
            for slide in loaded.slides:
                for pic in [s for s in slide.shapes if s.shape_type==13]:
                    self.assertEqual(pic.crop_left+pic.crop_right+pic.crop_top+pic.crop_bottom,0)
                    self.assertAlmostEqual(pic.width/pic.height,1600/900,places=5)


class ActualBuildPolishTests(unittest.IsolatedAsyncioTestCase):
    async def test_actual_builder_respects_budget_before_paid_calls(self):
        for style in ('minimal','designer_soft'):
            with self.subTest(style=style):
                env,state,sent=base.BuildTests().build_env(6)
                content=base.make_content(6)
                for slide in content['slides']:slide['designer']={'layout':'photo','blocks':[]}
                env['ask_grok'].return_value=json.dumps(content)
                state.data['theme_name']=style
                state.data['extra']='Добавь 3 подходящие ИИ-иллюстрации.'
                state.data['pres_plan']['slides'][0]['point']='Свет превращается в ток; на обложке ИИ-иллюстрация панели.'
                state.data['pres_plan_key']=ps.plan_key(state.data)
                await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321),chat=SimpleNamespace(id=-987654321)),state)
                self.assertEqual(len(sent),2)
                self.assertEqual(env['generate_image'].await_count,3)
                prs=Presentation(io.BytesIO(sent[0][1]))
                self.assertEqual(len(prs.slides),6)
                self.assertEqual(sum(s.shape_type==13 for sl in prs.slides for s in sl.shapes),3)
                visible=' '.join(s.text for sl in prs.slides for s in sl.shapes if s.has_text_frame)
                self.assertNotIn('ИИ-иллюстрация',visible)
                self.assertNotIn('на обложке',visible)
                if style=='designer_soft':self.assertIn('Свет превращается в ток.',visible)
                env['spend_credits'].assert_called_once()
