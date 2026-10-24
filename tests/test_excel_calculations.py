import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from excel_calculations import normalize_calculation, calculate, formula
from studio_excel_catalog import normalize
from studio_excel_runtime import validate_file
from test_office import office_env

ROW = dict(name='Товар', qty=3, price=500, unit='шт.', category='Материалы')

class CalculationTests(unittest.TestCase):
    def test_known_results_and_strict_grammar(self):
        for spec, expected in [({},1500), ({'mode':'discount','percent':10},1350),
                               ({'mode':'markup','percent':20},1800)]:
            self.assertEqual(calculate(ROW,spec or None),expected)
        for op, expected in [('multiply',1500),('divide',.006),('add',503),('subtract',-497)]:
            spec=dict(mode='custom',left='qty',right='price',op=op)
            self.assertAlmostEqual(calculate(ROW,spec),expected)
        for value in [[], {'mode':[]}, {'mode':'discount','percent':101}, {'mode':'markup','percent':True},
                      {'mode':'custom','left':'qty','right':'price','op':'=WEBSERVICE("x")'},
                      {'mode':'custom','left':'F5','right':'price','op':'add'},
                      {'mode':'custom','left':'qty','right':'number','right_value':0,'op':'divide'}]:
            with self.subTest(value=value), self.assertRaises(ValueError): normalize_calculation(value)

    def test_division_and_limits_preserve_draft_but_block_build(self):
        c=dict(mode='custom',left='price',right='qty',op='divide')
        p=dict(title='Смета',rows=[{**ROW,'qty':0}],calculation=c)
        self.assertEqual(normalize(p,draft=True)['rows'][0]['qty'],0)
        with self.assertRaisesRegex(ValueError,'Строка 1'): normalize(p)
        p['rows'][0]['qty']=''
        self.assertIsNone(calculate(p['rows'][0],c))
        p=dict(title='Смета',rows=[{**ROW,'qty':1000000,'price':1000000000}],calculation={'mode':'discount','percent':100})
        self.assertEqual(calculate(normalize(p)['rows'][0],p['calculation']),0)
        p['calculation']={'mode':'cost'}
        with self.assertRaises(ValueError):normalize(p)

    def test_actual_xlsx_formula_drivers_and_caches_all_modes_styles(self):
        build=office_env()['build_excel_items']
        specs=[{'mode':'cost'}, {'mode':'discount','percent':10}, {'mode':'markup','percent':20}]
        specs += [dict(mode='custom',left='price',right='qty',op=op) for op in ('multiply','divide','add','subtract')]
        specs += [dict(mode='custom',left='number',left_value=2,right='number',right_value=3,op='multiply')]
        with tempfile.TemporaryDirectory() as directory:
            for style in ('standard','designer'):
                for spec in specs:
                    with self.subTest(style=style,spec=spec):
                        payload=normalize(dict(title='Проверка',style=style,rows=[ROW,{**ROW,'qty':6}],calculation=spec))
                        path=str(Path(directory)/'test.xlsx')
                        build(path,payload['title'],'expense_estimate',{},payload['rows'],excel_style=style,excel_calculation=spec)
                        validate_file(path,payload)
                        wb=load_workbook(path); cached=load_workbook(path,data_only=True)
                        first=5 if style=='standard' else 9
                        self.assertEqual(wb.active.cell(first,6).value,formula(first,spec))
                        self.assertAlmostEqual(cached.active.cell(first,6).value,calculate(ROW,spec))
                        self.assertFalse(wb.active.protection.sheet)
                        if spec['mode']=='cost':self.assertEqual(len(wb.sheetnames),1)
                        else:
                            self.assertIn('Расчёт',wb.sheetnames)
                            self.assertFalse(wb['Расчёт'].protection.sheet)
                            if spec['mode'] in ('discount','markup'): self.assertEqual(wb['Расчёт']['B3'].value,spec['percent'])
                        if spec['mode']=='custom':self.assertNotIn('₽',wb.active.cell(first,6).number_format)
                        wb.close();cached.close()

    def test_calculation_is_in_saved_payload_and_changes_identity(self):
        base=normalize(dict(title='Смета',rows=[ROW]))
        changed=normalize({**base,'calculation':dict(mode='discount',percent='10,5')})
        self.assertNotEqual(base,changed)
        self.assertEqual(changed['calculation']['percent'],10.5)

if __name__=='__main__':unittest.main()
