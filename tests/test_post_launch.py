"""Regressions observed during the public launch acceptance run."""
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import academic_templates as academic
import test_office as office
import test_document_intake as intake_tests
from test_presentation_support import load_functions
from docx import Document

class CoverChoiceTests(unittest.IsolatedAsyncioTestCase):
    def test_explicit_choice_and_later_correction(self):
        for request in ('Эссе без титульного листа', 'Не делай титульный лист',
                        'Титульный не нужен', 'Убери титульник', 'Essay without a title page',
                        'Essay without title page', 'No cover page'):
            with self.subTest(request=request):
                self.assertFalse(academic.wants_title_page({'topic': request}))
        self.assertTrue(academic.wants_title_page({'topic': 'Эссе о технологиях'}))
        self.assertTrue(academic.wants_title_page({'topic': 'Без титульного', 'extra': 'Добавь титульный лист'}))
        self.assertFalse(academic.wants_title_page({'topic': 'С титульным', 'extra': 'Без титульного листа'}))

    def test_render_without_cover_preserves_body_and_correct_front_count(self):
        for kind in academic.KINDS:
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as tmp:
                p=Path(tmp)/'no-cover.docx'
                office.office_env()['build_word'](p,'Как технологии помогают учиться',
                    [{'title':'Введение','content':'Учебный текст сохраняется.'}],kind,{},
                    extra_data={'include_title':False,'requested_pages':2})
                doc=Document(p);text='\n'.join(x.text for x in doc.paragraphs)
                self.assertIn('Учебный текст сохраняется.',text)
                self.assertNotIn('ФИО',text)
                self.assertNotIn('Руководитель',text)
                self.assertFalse(doc.sections[0].different_first_page_header_footer)
                self.assertEqual(academic.front_matter_pages(kind,False),int(kind in ('referat','coursework')))
                if kind in ('referat','coursework'):
                    toc=next(x for x in doc.paragraphs if x.text=='ОГЛАВЛЕНИЕ')
                    self.assertFalse(toc.paragraph_format.page_break_before)
                else:
                    self.assertFalse(any(x.paragraph_format.page_break_before for x in doc.paragraphs))

    async def test_no_title_fields_requested_and_no_extra_page_billed(self):
        env,m,state=office.WordWorkflowTests().env()
        state.get_data.return_value={'price_version':env['PRICE_VERSION'],'word_kind':'essay',
            'word_pages':2,'topic':'Эссе о технологиях без титульного листа'}
        env['WORD_KIND_LABELS']={'essay':{'ru':'Эссе'}}
        env['word_cost']=lambda n,kind='doc':n*5
        env['ask_grok'].return_value='{"title":"Технологии","sections":[{"title":"Тема","content":"Текст"}]}'
        async def prepare(value,requested,*args,**kw):
            self.assertEqual(requested,2)
            await kw['build'](value,'/tmp/fixture-no-cover.docx')
            return value,2
        env['prepare_word_files']=AsyncMock(side_effect=prepare)
        await env['word_build'](m,state)
        state.set_state.assert_not_awaited()
        self.assertEqual(m.answer_document.await_count,2)
        self.assertFalse(env['run_heavy'].call_args.kwargs['extra_data']['include_title'])
        self.assertEqual(env['spend_credits'].call_args.kwargs['amount'],10)
        self.assertIn('титульный лист НЕ нужен',env['ask_grok'].call_args.args[0])

class ExcelQuoteTests(unittest.IsolatedAsyncioTestCase):
    async def test_each_chat_entry_records_displayed_tariff(self):
        for name in ('excel_topic','excel_data','excel_startup_data','excel_extra'):
            env=load_functions([name],{'user_lang':lambda _:'ru','tr':lambda key,*a:key,
                'is_thin_input':lambda _:False,'excel_confirm_kb':lambda _:'buttons',
                'Form':SimpleNamespace(waiting_excel_confirm='confirm')})
            state=intake_tests.State();state.data={'excel_kind':'expense_estimate'}
            m=SimpleNamespace(text='Краска 3 по 100',from_user=SimpleNamespace(id=1),answer=AsyncMock())
            await env[name](m,state)
            self.assertEqual(state.data['price_version'],env['PRICE_VERSION'])
            self.assertEqual(m.answer.await_count,1)
            self.assertIn('Стоимость:',m.answer.call_args.args[0])

    async def test_free_text_quote_records_tariff_and_build_does_not_repeat_it(self):
        fixture=intake_tests.IntakeFlowTests();env=fixture.env();state=intake_tests.State()
        draft={'fmt':'excel','kind':'expense_estimate','text':intake_tests.REQUEST,'mode':'user','style':'standard'}
        await env['advance_document_intake'](fixture.msg(intake_tests.REQUEST),state,draft,'ru')
        self.assertEqual(state.data['price_version'],env['PRICE_VERSION'])
        start=Mock(return_value=(False,'busy fixture'))
        build=load_functions(['excel_build'],{'user_lang':lambda _:'ru','can_afford':lambda *a:True,
            'get_user':lambda _:{},'start_job':start})['excel_build']
        m=fixture.msg('Собрать')
        await build(m,state)
        start.assert_called_once()
        self.assertEqual(m.answer.call_args.args[0],'busy fixture')

    async def test_old_quote_still_requires_fresh_confirmation(self):
        env=load_functions(['excel_build'],{'user_lang':lambda _:'ru',
            'Form':SimpleNamespace(waiting_excel_confirm='confirm'),'excel_confirm_kb':lambda _:'buttons',
            'start_job':Mock()})
        state=intake_tests.State();state.data={'price_version':'old','excel_kind':'expense_estimate'}
        m=SimpleNamespace(from_user=SimpleNamespace(id=1),answer=AsyncMock())
        await env['excel_build'](m,state)
        env['start_job'].assert_not_called()
        self.assertIn('Подтвердите',m.answer.call_args.args[0])

class VoiceAndConcurrencyTests(unittest.IsolatedAsyncioTestCase):
    async def test_voice_success_empty_failure_and_cancel_cleanup(self):
        import asyncio, os
        for outcome in ('success','empty','failed','cancelled'):
            with self.subTest(outcome=outcome):
                paths=[]
                async def download(remote,path):
                    paths.append(Path(path));Path(path).write_bytes(b'fixture audio')
                provider=AsyncMock(return_value=SimpleNamespace(text='  Сколько будет 7 умножить на 8?  ' if outcome=='success' else ''))
                if outcome=='failed': provider.side_effect=RuntimeError('test service failure')
                if outcome=='cancelled': provider.side_effect=asyncio.CancelledError()
                bot=SimpleNamespace(get_file=AsyncMock(return_value=SimpleNamespace(file_path='remote.ogg')),download_file=AsyncMock(side_effect=download))
                env=load_functions(['transcribe_voice'],dict(whisper_client=SimpleNamespace(audio=SimpleNamespace(transcriptions=SimpleNamespace(create=provider))),
                    bot=bot,MAX_VOICE_SECONDS=120,os=os,random=SimpleNamespace(randint=lambda *a:7781)))
                if outcome=='cancelled':
                    with self.assertRaises(asyncio.CancelledError):
                        await env['transcribe_voice'](SimpleNamespace(file_id='test',duration=5),-992233)
                else:
                    result=await env['transcribe_voice'](SimpleNamespace(file_id='test',duration=5),-992233)
                    self.assertEqual(result,'Сколько будет 7 умножить на 8?' if outcome=='success' else None)
                self.assertTrue(paths)
                self.assertTrue(all(not p.exists() for p in paths))
                provider.reset_mock();bot.download_file.reset_mock()
                self.assertIsNone(await env['transcribe_voice'](SimpleNamespace(file_id='test',duration=121),-992233))
                provider.assert_not_awaited();bot.download_file.assert_not_awaited()

    async def test_concurrent_accounts_and_retries_have_one_debit_each(self):
        import asyncio,copy
        from order_billing import charge_callback
        users={uid:{'credits':100,'history':[]} for uid in range(1,11)}
        disk={}
        async def persist():
            await asyncio.sleep(0)
            disk.clear();disk.update(copy.deepcopy(users));return True
        api=dict(get_user=lambda uid:users[uid],save_users_now=persist,
            ADMIN_IDS=set(),can_afford=lambda uid,c:users[uid]['credits']>=c,note_success=Mock())
        charge=charge_callback(api,'studio_excel_payments','Excel')
        await asyncio.gather(*(charge(uid,'same-local-order-id',20,'Таблица') for uid in users for _ in range(3)))
        for uid,u in users.items():
            self.assertEqual(u['credits'],80)
            self.assertEqual(len(u['wallet_transactions']),1)
            self.assertEqual(disk[uid]['credits'],80)
