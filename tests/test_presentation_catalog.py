import copy
import io
import tempfile
import unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import presentation_catalog as pc
import presentation_profiles as pp
import presentation_designer as pd
import presentation_support as ps
import test_presentation_support as base


def specimen(key, comp):
    t=pc.TEMPLATES[key]
    d=pp.choose_design({'topic':'Пчёлы: жизнь улья','extra':f'Дизайн {t["number"]}'})
    number=1 if comp.startswith('cover') else 4
    if comp.startswith('cover'):
        d['variant']=((comp=='cover_band')-t['sequence'])%2
    elif comp.startswith('photo'):
        d['variant']=(pc.COMPOSITIONS.index(comp)-2-(number-2)-t['sequence'])%4
    item={'title':'Как устроена жизнь улья','content':'Рабочие пчёлы ухаживают за расплодом, строят соты и собирают корм.',
          'point':'Пчелиная семья и её повседневная работа.', '_designer_design':d}
    blocks=[{'heading':'Матка','text':'Откладывает яйца.'},{'heading':'Рабочие','text':'Строят соты и собирают корм.'},
            {'heading':'Трутни','text':'Участвуют в размножении.'}]
    if comp=='process':
        item['title']='Как получается мёд'
        blocks=[{'heading':'Сбор','text':'Пчела собирает нектар с цветков.'},
                {'heading':'Передача','text':'В улье нектар передают другим пчёлам.'},
                {'heading':'Созревание','text':'Нектар перерабатывают и удаляют лишнюю воду.'},
                {'heading':'Хранение','text':'Созревший мёд хранят в сотах.'}]
    elif comp=='comparison':
        item['title']='Две роли в пчелиной семье'
        blocks=blocks[:2]
    elif comp=='conclusion':
        item['title']='Что важно запомнить'
    elif comp=='chart':
        item['title']='Учебный пример наблюдений'
        item['content']='Условные данные для проверки графика. Это не результаты исследования.'
        item['chart']={'chart_type':'bar','title':'Условный пример','categories':['А','Б','В'],
                       'series':[{'name':'Наблюдения','values':[2,4,3]}]}
    if comp in ('sections','process','comparison','conclusion'):
        item['designer']={'layout':comp,'blocks':blocks}
        if comp=='comparison':
            item['designer']['rows']=[{'criterion':'Роль','left':'Размножение','right':'Сбор корма'},
                                      {'criterion':'Место','left':'В улье','right':'В улье и снаружи'},
                                      {'criterion':'Численность','left':'Обычно одна','right':'Много особей'}]
    elif comp=='statement':item['designer']={'layout':'statement','blocks':[]}
    return item,number


class CatalogTests(unittest.TestCase):
    def test_fifty_named_kits_and_twelve_compositions(self):
        self.assertEqual(len(pc.TEMPLATES),50)
        self.assertEqual(len(set(t['name'] for t in pc.TEMPLATES.values())),50)
        for t in pc.TEMPLATES.values():self.assertEqual(len(set(t['compositions'])),12)
        # Designs differ in geometry AND typography, not just colour names.
        sig={(t['system'],t['margin'],t['split'],t['font'],t['body_font'],t['title_size']) for t in pc.TEMPLATES.values()}
        self.assertEqual(len(sig),50)

    def test_subject_precedes_audience_and_explicit_choice_wins(self):
        for topic,key in [('Пчёлы: жизнь улья','d01'),('Океан и киты','d03'),('Космос','d06'),
                          ('Бутик-отель в Тоскане','d11'),('Кофейня','d16'),('Роботы на складе','d21')]:
            with self.subTest(topic=topic):
                self.assertEqual(pp.choose_design({'topic':topic,'audience':'школьники 7 класса'})['template_id'],key)
        for key,t in pc.TEMPLATES.items():
            self.assertEqual(pp.choose_design({'topic':'Пчёлы','extra':f'Шаблон №{t["number"]}'})['template_id'],key)
        self.assertEqual(pp.choose_design({'topic':'Пчёлы','extra':'Дизайн 6, затем дизайн 12'})['template_id'],'d12')
        self.assertEqual(pp.choose_design({'topic':'Пчёлы','extra':'Технологический стиль'})['family'],'technology')

    def test_selection_stable_and_validation_preserves_kit(self):
        data={'topic':'Пчёлы','theme_name':'designer_soft'}
        p=ps.finalize_plan_design(base.make_plan(),data)
        self.assertEqual(ps.validate_plan(p,8)['design'],p['design'])
        self.assertEqual(pp.choose_design(data),pp.choose_design(data))
        for value in ('d99',None,[],{}) :
            q=copy.deepcopy(p);q['design']['template_id']=value
            with self.assertRaises((ps.PresentationValidationError,TypeError)):ps.validate_plan(q,8)

    def test_all_six_hundred_compositions_keep_text_images_and_valid_xml(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'image.png';Image.new('RGB',(1200,800),'#789765').save(path)
            asset={'path':str(path),'width':1200,'height':800,'origin':'ai'}
            for key in pc.TEMPLATES:
                prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
                geometry=set()
                for comp in pc.COMPOSITIONS:
                    with self.subTest(key=key,composition=comp):
                        item,n=specimen(key,comp)
                        photo=asset if comp.startswith(('cover','photo')) else None
                        scene=pd.build_scene(item,photo,n,12)
                        self.assertEqual(scene['composition'],comp)
                        geometry.add(tuple((x['type'],tuple(x['box'])) for x in scene['nodes']))
                        texts=' '.join(x['text'] for x in scene['nodes'] if x['type']=='text')
                        self.assertIn(' '.join(item['title'].split()),' '.join(texts.split()))
                        if photo:self.assertEqual(sum(x['type']=='image' for x in scene['nodes']),1)
                        if comp not in ('chart','comparison'):
                            for b in item.get('designer',{}).get('blocks',[]):
                                self.assertIn(' '.join(b['text'].split()),' '.join(texts.split()))
                        # Title/body/image regions cannot unintentionally overlap; full-bleed cover is intentional.
                        relevant=[x for x in scene['nodes'] if x['type'] in ('text','image')]
                        if not(comp.startswith('cover') and pc.TEMPLATES[key]['system']=='cinema'):
                            for i,a in enumerate(relevant):
                                for b in relevant[i+1:]:
                                    ax,ay,aw,ah=a['box'];bx,by,bw,bh=b['box']
                                    self.assertFalse(min(ax+aw,bx+bw)>max(ax,bx)+1 and min(ay+ah,by+bh)>max(ay,by)+1,(key,comp,a,b))
                        pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene,lambda *a:None)
                self.assertEqual(len(geometry),12,key)
                out=io.BytesIO();prs.save(out);out.seek(0)
                reopened=Presentation(out)
                self.assertEqual(len(reopened.slides),12)
                self.assertIsNotNone(reopened._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}notesMasterIdLst'))

    def test_long_text_and_user_images_preserved(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'portrait.png';Image.new('RGB',(450,1200)).save(path)
            asset={'path':str(path),'width':450,'height':1200,'origin':'user'}
            for key in pc.TEMPLATES:
                item,n=specimen(key,'photo_left')
                item['title']='Подробное объяснение работы системы и взаимодействия её участников'
                item['content']=('Полное описание работы системы. '*14).strip()
                scene=pd.build_scene(item,asset,n,12)
                copytext=' '.join(x['text'] for x in scene['nodes'] if x.get('role')=='body')
                self.assertEqual(' '.join(copytext.split()),item['content'])
                self.assertFalse(next(x for x in scene['nodes'] if x['type']=='image')['crop'])

    def test_live_builder_uses_catalog_without_extra_calls(self):
        import asyncio,json
        from types import SimpleNamespace
        env,state,sent=base.BuildTests().build_env(3)
        state.data.update(theme_name='designer_soft',topic='Пчёлы',extra='Дизайн 1')
        env['_create_presentation_plan'].return_value['audience']='школьники'
        content=base.make_content(3)
        for s in content['slides']:s['designer']={'layout':'photo','blocks':[]}
        env['ask_grok'].return_value=json.dumps(content)
        asyncio.run(env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state))
        self.assertEqual(len(sent),2)
        result=Presentation(io.BytesIO(sent[0][1]))
        for slide in result.slides:self.assertIn('d01 Медовый атлас',slide.notes_slide.notes_text_frame.text)
        self.assertEqual(env['generate_image'].await_count,3)
        env['spend_credits'].assert_called_once()
