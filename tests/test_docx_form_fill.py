import tempfile, unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx_form_fill import replace_fields, adapt_filled_layout

class FormFillTests(unittest.TestCase):
    def test_split_runs_multiline_multiple_tokens_preserve_format_and_break(self):
        d=Document();p=d.add_paragraph();p.add_run('До [Им').bold=True;p.add_run('я] и [Адрес]. После.').italic=True
        br=OxmlElement('w:br');br.set(qn('w:type'),'page');p.runs[-1]._r.append(br)
        replace_fields(d,{'[Имя]':'Анна\nИванова','[Адрес]':'Москва\nУлица.'})
        self.assertEqual(p.text,'До Анна\nИванова и Москва\nУлица. После.')
        self.assertTrue(p.runs[0].bold);self.assertTrue(p.runs[1].italic)
        self.assertEqual(len(d.element.xpath('.//w:br[@w:type="page"]')),1)
    def test_same_run_multiple_multiline_fields_not_interpreted_twice(self):
        d=Document();p=d.add_paragraph('[А] / [Б] / [А]')
        replace_fields(d,{'[А]':'[Б]\nпервая','[Б]':'вторая\nстрока'})
        self.assertEqual(p.text,'[Б]\nпервая / вторая\nстрока / [Б]\nпервая')
    def test_nested_tables_and_footer(self):
        d=Document();t=d.add_table(rows=1,cols=1);nested=t.cell(0,0).add_table(rows=1,cols=1)
        nested.cell(0,0).text='[Поле]';d.sections[0].footer.paragraphs[0].text='[Поле]'
        replace_fields(d,{'[Поле]':'Значение'})
        self.assertEqual(nested.cell(0,0).text,'Значение');self.assertEqual(d.sections[0].footer.paragraphs[0].text,'Значение')
    def test_long_row_can_expand_and_split(self):
        d=Document();t=d.add_table(rows=1,cols=1);t.cell(0,0).text='длинные реквизиты '*100
        pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'));h=OxmlElement('w:trHeight');h.set(qn('w:hRule'),'exact');pr.append(h)
        adapt_filled_layout(d)
        self.assertIsNone(pr.find(qn('w:cantSplit')));self.assertEqual(h.get(qn('w:hRule')),'atLeast')
