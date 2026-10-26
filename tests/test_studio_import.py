import io
import time
import zipfile
import unittest
from unittest.mock import AsyncMock
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from aiohttp import web,CookieJar
from aiohttp.test_utils import TestClient,TestServer
from studio_import import import_pptx
from studio_export import build_pptx
from studio_private import Studio,COOKIE

def fixture():
 p=Presentation();p.slide_width=11430000;p.slide_height=6429375
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('14293B')
 t=s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2));t.text_frame.paragraphs[0].text='Свой текст';t.text_frame.paragraphs[0].font.size=Pt(32);t.text_frame.paragraphs[0].font.name='Georgia';t.text_frame.paragraphs[0].font.color.rgb=RGBColor.from_string('EEE7DC')
 im=Image.new('RGB',(80,40),'red');b=io.BytesIO();im.save(b,format='PNG');s.shapes.add_picture(io.BytesIO(b.getvalue()),Inches(8),Inches(2),Inches(3),Inches(2))
 o=io.BytesIO();p.save(o);return o.getvalue()

def mutate_zip(raw,name,content):
 b=io.BytesIO()
 with zipfile.ZipFile(io.BytesIO(raw)) as src,zipfile.ZipFile(b,'w') as dst:
  for f in src.infolist():
   if f.filename!=name:dst.writestr(f,src.read(f))
  dst.writestr(name,content)
 return b.getvalue()

class ImportTests(unittest.TestCase):
 def test_import_edit_export_roundtrip(self):
  raw=fixture();r=import_pptx(raw,'Тест.pptx');d=r['deck'];s=d['slides'][0]
  self.assertEqual(d['name'],'Тест');self.assertEqual(len(s['elements']),2)
  self.assertEqual(s['background']['color'],'#14293B')
  t=s['elements'][0];self.assertEqual(t['color'],'#EEE7DC');self.assertEqual(t['font'],'Georgia')
  t['text']='Исправленный текст';s['elements'][1]['raster']=s['elements'][1]['src']
  p=Presentation(io.BytesIO(build_pptx(d)))
  self.assertIn('Исправленный текст',p.slides[0].shapes[0].text)
  self.assertEqual(len(p.slides[0].shapes),2);self.assertEqual(r['price']['cost'],10)
 def test_unsupported_objects_reported_per_slide(self):
  p=Presentation(io.BytesIO(fixture()));p.slides[0].shapes.add_table(2,2,0,0,1000000,1000000)
  p.slides[0].shapes.add_shape(MSO_SHAPE.OVAL,0,0,1000000,1000000)
  b=io.BytesIO();p.save(b);r=import_pptx(b.getvalue())
  self.assertTrue(any('таблицы (1)' in w and 'фигуры' in w for w in r['warnings']))
  self.assertEqual(r['deck']['importWarnings'],r['warnings'])
 def test_bad_file_external_payload_and_macro_rejected(self):
  for raw in (b'not pptx',b'x'*10000001,mutate_zip(fixture(),'../escape','x'),mutate_zip(fixture(),'ppt/vbaProject.bin',b'x'),mutate_zip(fixture(),'ppt/test.xml','<!DOCTYPE x><x/>'),mutate_zip(fixture(),'ppt/_rels/bad.rels','<Relationships><Relationship TargetMode="External" Type="image" Target="http://private.invalid"/></Relationships>')):
   with self.assertRaises(ValueError):import_pptx(raw)
 def test_complex_slide_not_silently_empty(self):
  p=Presentation(io.BytesIO(fixture()));s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_table(2,2,0,0,1000000,1000000);b=io.BytesIO();p.save(b)
  with self.assertRaisesRegex(ValueError,'Слайд 2'):import_pptx(b.getvalue())
 def test_existing_limits_checked_before_open(self):
  p=Presentation(io.BytesIO(fixture()));s=p.slides[0]
  for i in range(6):s.shapes.add_textbox(0,0,1000000,1000000).text='Ещё'
  b=io.BytesIO();p.save(b)
  with self.assertRaisesRegex(ValueError,'6 текстовых'):import_pptx(b.getvalue())

class ImportEndpointTests(unittest.IsolatedAsyncioTestCase):
 async def asyncSetUp(self):
  self.send=AsyncMock();self.action=AsyncMock();self.studio=Studio(None,self.send,self.action,'http://example.test')
  self.studio.sessions['test']={'expires':time.monotonic()+300,'user':{'id':465823470}}
  app=web.Application();self.studio.register(app);self.client=TestClient(TestServer(app),cookie_jar=CookieJar(unsafe=True));await self.client.start_server()
 async def asyncTearDown(self):
  await self.client.close()
  for z in self.studio.archives:z.close()
 async def test_auth_csrf_free_import(self):
  h={'Origin':'http://example.test','X-Studio-Request':'1'}
  self.assertEqual((await self.client.post('/studio/import-pptx',data=fixture(),headers=h)).status,401)
  self.client.session.cookie_jar.update_cookies({COOKIE:'test'})
  self.assertEqual((await self.client.post('/studio/import-pptx',data=fixture())).status,403)
  r=await self.client.post('/studio/import-pptx',data=fixture(),headers=h);self.assertEqual(r.status,200);self.assertEqual(len((await r.json())['deck']['slides']),1)
  self.send.assert_not_awaited();self.action.assert_not_awaited();self.assertFalse(self.studio.import_busy)
 async def test_invalid_upload_returns_readable_error_and_unlocks(self):
  self.client.session.cookie_jar.update_cookies({COOKIE:'test'})
  r=await self.client.post('/studio/import-pptx',data=b'bad',headers={'Origin':'http://example.test','X-Studio-Request':'1'})
  self.assertEqual(r.status,400);self.assertIn('PPTX',(await r.json())['error']);self.assertFalse(self.studio.import_busy)
