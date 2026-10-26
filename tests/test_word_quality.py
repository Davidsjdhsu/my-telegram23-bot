"""Regressions for delivered Word defects and bounded page-aware recovery."""
import asyncio
import ast
import copy
import json
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest
from unittest.mock import AsyncMock, Mock
from docx import Document
from test_office import office_env
import test_office
from word_quality import (clean_word_meta, merge_word_expansion, prepare_word_files,
                          word_generation_schema, WORD_EDITORIAL_GUIDANCE,
                          readable_word_paragraphs, word_review_findings, WORD_SYSTEM, WORD_REVIEW_SYSTEM,
                          word_section_budgets, expansion_word_target, needs_word_revision)


class WordMetadataTests(unittest.TestCase):
    def test_optional_placeholders_removed_but_values_and_legal_fields_preserved(self):
        meta = {'org': '[учебное заведение]', 'author': 'Анна Петрова', 'school': '[название учебного заведения]'}
        self.assertEqual(clean_word_meta(meta, 'doc'), {'org':'', 'author':'Анна Петрова', 'school':''})
        self.assertEqual(clean_word_meta(meta, 'statement'), meta)
        self.assertEqual(meta['org'], '[учебное заведение]')
        self.assertEqual(clean_word_meta({'org':'[Лаборатория Солнце]'}, 'doc')['org'], '[Лаборатория Солнце]')

    def test_generated_schema_has_no_fake_institution_or_date(self):
        schema = '{"org":"Министерство образования","school":"[учебное заведение]","year":"2026"}'
        self.assertEqual(json.loads(word_generation_schema('referat',schema)), {'org':'','school':'','year':''})
        self.assertEqual(word_generation_schema('statement',schema), schema)

    def test_plain_document_omits_placeholder_and_retains_real_organization(self):
        env = office_env()
        with tempfile.TemporaryDirectory() as td:
            path = str(Path(td)/'sample.docx')
            for org in ('[учебное заведение]', 'Мастерская Север'):
                env['build_word'](path,'Хранение файлов',[{'title':'Сравнение','content':'Текст'}],meta={'org':org})
                paragraphs = Document(path).paragraphs
                text='\n'.join(p.text for p in paragraphs)
                self.assertNotIn('[учебное заведение]',text)
                if org=='Мастерская Север': self.assertIn(org,text)
                title=next(p for p in paragraphs if p.text=='Хранение файлов')
                self.assertEqual(title.style.name,'Title')
                self.assertNotIn('pBdr',title.style.element.xml)

    def test_study_cover_omits_unknown_details_and_template_keeps_blanks(self):
        env=office_env()
        with tempfile.TemporaryDirectory() as td:
            path=str(Path(td)/'sample.docx')
            for kind in ('referat','report','essay','coursework'):
                env['build_word'](path,'Хранение файлов',[{'title':'Введение','content':'Текст'}],kind,
                                  {'school':'[учебное заведение]','author':'Анна Петрова'})
                text='\n'.join(p.text for p in Document(path).paragraphs)
                self.assertIn('Анна Петрова',text)
                self.assertNotIn('[',text)
                self.assertNotIn('Министерство',text)
            env['build_word'](path,'Шаблон',[], 'referat',{},extra_data={'template':True})
            self.assertIn('[Название учебного заведения]','\n'.join(p.text for p in Document(path).paragraphs))

    def test_invalid_revision_is_not_accepted(self):
        original={'sections':[{'title':'А','content':'Текст','table':{'headers':['X'],'rows':[['42']]}}]}
        for expanded in (None, [], {'sections':None}, {'sections':[{'title':'Б','content':'Другой'}]}):
            self.assertIs(merge_word_expansion(original,expanded), original)


class WordPaginationTests(unittest.IsolatedAsyncioTestCase):
    async def run_case(self, candidate_pages=2, failure=None, requested=2, initial_pages=1):
        with tempfile.TemporaryDirectory() as td:
            docx,pdf=Path(td)/'original.docx',Path(td)/'original.pdf'
            original={'title':'Тема','meta':{'author':'Анна'},'sections':[{'title':'А','content':'initial','table':{'headers':['X'],'rows':[['42']]}}]}
            built=[]
            async def build(value,path):
                built.append(copy.deepcopy(value))
                pages=initial_pages if value['sections'][0]['content']=='initial' else candidate_pages
                Path(path).write_text(str(pages))
            async def convert(source,target):
                if failure=='conversion' and 'revised' in source: raise RuntimeError('fixture')
                Path(target).write_text(Path(source).read_text())
            async def expand(value,actual,wanted,findings):
                if failure=='model': raise RuntimeError('fixture')
                if failure=='cancel': raise asyncio.CancelledError()
                if failure=='malformed': return {'sections':None}
                return {'title':'Malicious replacement','meta':{'author':'Other'},'sections':[{'title':'А','content':'expanded','table':{'rows':[['99']]}}]}
            revision=AsyncMock(side_effect=expand)
            content,count=await prepare_word_files(original,requested,str(docx),str(pdf),build=build,convert=convert,
                                                  expand=revision,measure=lambda p:int(Path(p).read_text()))
            return dict(content=content,count=count,docx=docx.read_text(),pdf=pdf.read_text(),calls=revision.await_count,built=built)

    async def test_two_page_order_really_gets_one_bounded_revision(self):
        r=await self.run_case()
        self.assertEqual((r['count'],r['calls'],r['docx'],r['pdf']),(2,1,'2','2'))
        self.assertEqual(r['content']['title'],'Тема')
        self.assertEqual(r['content']['meta'],{'author':'Анна'})
        self.assertEqual(r['content']['sections'][0]['table']['rows'],[['42']])

    async def test_no_extra_model_request_when_length_is_met(self):
        r=await self.run_case(initial_pages=2)
        self.assertEqual(r['calls'],0)

    async def test_optional_model_conversion_and_json_failures_keep_valid_pair(self):
        for failure in ('model','conversion','malformed'):
            r=await self.run_case(failure=failure)
            self.assertEqual((r['count'],r['docx'],r['pdf']),(1,'1','1'))
            self.assertEqual(r['calls'],1)

    async def test_shorter_or_oversized_revision_cannot_replace_original(self):
        for pages in (1,10):
            r=await self.run_case(candidate_pages=pages)
            self.assertEqual((r['count'],r['docx'],r['pdf']),(1,'1','1'))

    async def test_cancellation_propagates(self):
        with self.assertRaises(asyncio.CancelledError): await self.run_case(failure='cancel')


class WordContentReviewTests(unittest.IsolatedAsyncioTestCase):
    def test_editor_budgets_can_pass_preservation_gate_for_long_introduction(self):
        original = {'sections': [
            {'title': title, 'content': ' '.join(['слово'] * count)}
            for title, count in [('Введение', 62), ('Различия', 76), ('Условия', 86),
                                 ('Пример', 86), ('Заключение', 57)]]}
        budgets = word_section_budgets(original, 560)
        self.assertEqual(sum(b['target_words'] for b in budgets), 560)
        revised = {'sections': [{'title': b['title'], 'content': ' '.join(['пояснение'] * b['target_words'])}
                                for b in budgets]}
        self.assertIsNot(merge_word_expansion(original, revised), original)
        self.assertTrue(all(b['target_words'] >= len(s['content'].split())
                            for s, b in zip(original['sections'], budgets)))

    def test_editor_budgets_preserve_existing_reference_text_without_expansion(self):
        original = {'sections': [{'title': 'Источники', 'content': 'Источник указан пользователем'},
                                 {'title': 'Тема', 'content': 'Подробное объяснение темы'}]}
        self.assertEqual([b['target_words'] for b in word_section_budgets(original, 100)], [3, 97])
        self.assertEqual([b['target_words'] for b in word_section_budgets(original, 2)], [3, 3])

    def test_prose_split_preserves_all_words_numbers_and_tables(self):
        text = ('Студент И. И. Иванов сохранил 3.14 МБ данных на компьютере. ' * 16).strip()
        split = readable_word_paragraphs(text)
        self.assertGreater(len(split.splitlines()), 1)
        self.assertEqual(text.split(), split.split())
        self.assertNotIn('И.\n', split)
        for special in ('| ' + text + ' |', '1. ' + text, '- ' + text, '`' + text + '`'):
            self.assertEqual(readable_word_paragraphs(special), special)

    def test_builder_splits_informational_prose_without_touching_legal_text(self):
        e = office_env()
        content = ('Это предложение описывает учебный пример выбора устройства для хранения файлов. ' * 15).strip()
        with tempfile.TemporaryDirectory() as td:
            path = str(Path(td) / 'sample.docx')
            e['build_word'](path, 'Хранение файлов', [{'title': 'Выбор', 'content': content}])
            paras = [p.text for p in Document(path).paragraphs if p.text.startswith('Это предложение')]
            self.assertGreater(len(paras), 1)
            self.assertEqual(' '.join(paras).split(), content.split())
            e['build_word'](path, 'Договор', [{'title': 'Условия', 'content': content}], 'dkp')
            self.assertIn(content, [p.text for p in Document(path).paragraphs])

    def test_review_signal_distinguishes_conditions_and_detects_sparse_last_page(self):
        value = {'sections': [{'content': 'Без связи файлы становятся недоступны. Если файл не загружен, он недоступен.'}]}
        found = word_review_findings(value, 2, 2, {'last_page_fill': .2})
        self.assertEqual(found['categorical_claims'], 1)
        self.assertTrue(found['sparse_last_page'])
        self.assertFalse(found['short_pages'])
        self.assertFalse(word_review_findings(value, 1, 1, {'last_page_fill': .2})['sparse_last_page'])
        self.assertTrue(word_review_findings(value, 3, 2, {'last_page_fill': .2})['sparse_last_page'])

    def test_sparse_middle_page_and_oversized_order_are_reviewed_without_cover_padding(self):
        findings = word_review_findings({'sections': []}, 7, 6,
            {'body_page_fills': [.8, .3, .9, .8, .8], 'last_page_fill': .8, 'front_matter_pages': 2})
        self.assertEqual(findings['sparse_body_pages'], [4])
        self.assertTrue(findings['excess_pages'])
        self.assertTrue(needs_word_revision(findings))
        clean = word_review_findings({'sections': []}, 6, 6,
            {'body_page_fills': [.9, .9, .9, .7], 'last_page_fill': .7, 'front_matter_pages': 2})
        self.assertFalse(needs_word_revision(clean))

    def test_measured_density_bounds_expansion_instead_of_doubling_live_example(self):
        value = {'sections': [{'title': 'Тема', 'content': ' '.join(['слово'] * 577)}]}
        target = expansion_word_target(value, 6, 6, front_matter_pages=2,
            layout={'body_page_fills': [.994, .977, .956, .221]})
        self.assertGreater(target, 577)
        self.assertLess(target, 750)
        oversized = expansion_word_target(value, 8, 6, front_matter_pages=2,
            layout={'body_page_fills': [.9] * 6})
        budgets = word_section_budgets(value, oversized, allow_shorten=True)
        candidate = {'sections': [{'title': 'Тема', 'content': ' '.join(['слово'] * budgets[0]['target_words'])}]}
        self.assertLess(budgets[0]['target_words'], 577)
        self.assertIsNot(merge_word_expansion(value, candidate), value)

    def test_revision_cannot_erase_section_shorten_dramatically_or_lose_numbers(self):
        old = {'sections': [{'title': 'А', 'content': 'Храним 3.14 МБ данных и 42 учебных документа отдельно.'}]}
        for bad in ('', 'Коротко.', 'Храним 3.15 МБ данных и 42 учебных документа отдельно.'):
            self.assertIs(merge_word_expansion(old, {'sections': [{'title': 'А', 'content': bad}]}), old)

    def test_inline_markdown_table_cannot_be_rewritten_by_editor(self):
        old = {'sections': [{'title': 'А', 'content': 'Пояснение\n| Способ | Условие |\n| Облако | Нужен интернет |'}]}
        new = copy.deepcopy(old)
        new['sections'][0]['content'] = new['sections'][0]['content'].replace('Нужен интернет', 'Не нужен интернет')
        self.assertIs(merge_word_expansion(old, new), old)

    async def sparse_case(self, new_fill, failure=None):
        original = {'sections': [{'title': 'Тема', 'content': 'Исходное пояснение условий выбора.'}]}
        revised = {'sections': [{'title': 'Тема', 'content': 'Уточнённое пояснение условий выбора с подробным примером.'}]}
        with tempfile.TemporaryDirectory() as td:
            docx, pdf = Path(td) / 'sample.docx', Path(td) / 'sample.pdf'
            async def build(value, path): Path(path).write_text(value['sections'][0]['content'])
            async def convert(source, target): Path(target).write_text(Path(source).read_text())
            def inspect(path):
                if failure == 'layout': raise RuntimeError('fixture')
                return {'last_page_fill': new_fill if 'revised' in str(path) else .2}
            revision = AsyncMock(return_value=revised)
            if failure == 'model': revision.side_effect = RuntimeError('fixture')
            result, pages = await prepare_word_files(original, 2, str(docx), str(pdf), build=build, convert=convert,
                measure=lambda _: 2, inspect=inspect, expand=revision)
            self.assertEqual(docx.read_text(), pdf.read_text())
            return result, pages, revision

    async def test_sparse_two_page_document_can_improve_without_adding_third_page(self):
        value, pages, revision = await self.sparse_case(.75)
        self.assertTrue(value['sections'][0]['content'].startswith('Уточнённое'))
        self.assertEqual(pages, 2)
        self.assertEqual(revision.await_count, 1)
        self.assertTrue(revision.await_args.args[3]['sparse_last_page'])

    async def test_revision_cannot_replace_file_with_emptier_or_unimproved_ending(self):
        for fill in (.1, .2):
            value, _, revision = await self.sparse_case(fill)
            self.assertTrue(value['sections'][0]['content'].startswith('Исходное'))
            self.assertEqual(revision.await_count, 1)

    async def test_optional_layout_or_editor_failure_does_not_block_delivery(self):
        for failure in ('layout', 'model'):
            value, pages, _ = await self.sparse_case(.75, failure)
            self.assertTrue(value['sections'][0]['content'].startswith('Исходное'))
            self.assertEqual(pages, 2)

    async def test_middle_page_improvement_can_replace_same_length_document(self):
        original = {'sections': [{'title': 'Тема', 'content': 'Исходное пояснение условий.'}]}
        updated = {'sections': [{'title': 'Тема', 'content': 'Уточнённое пояснение условий и пример.'}]}
        with tempfile.TemporaryDirectory() as td:
            docx, pdf = Path(td)/'a.docx', Path(td)/'a.pdf'
            async def build(value, path): Path(path).write_text(value['sections'][0]['content'])
            async def convert(source, target): Path(target).write_text(Path(source).read_text())
            def inspect(path):
                return {'last_page_fill': .8, 'body_page_fills': [.8 if 'revised' in str(path) else .2, .8],
                        'front_matter_pages': 2}
            revision = AsyncMock(return_value=updated)
            value, pages = await prepare_word_files(original, 4, str(docx), str(pdf), build=build, convert=convert,
                measure=lambda _:4, inspect=inspect, expand=revision)
            self.assertEqual(value, updated)
            self.assertEqual(pages, 4)
            self.assertEqual(docx.read_text(), pdf.read_text())
            self.assertEqual(revision.await_count, 1)


class WordReviewRequestTests(unittest.IsolatedAsyncioTestCase):
    def request_env(self):
        source = ast.parse((Path(__file__).resolve().parents[1] / 'main.py').read_text())
        function = next(n for n in source.body if isinstance(n, ast.AsyncFunctionDef) and n.name == 'ask_grok')
        reply = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"sections": []}'), finish_reason='stop')])
        create = AsyncMock(return_value=reply)
        client = Mock(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
        client.with_options.return_value = client
        ns = dict(client=client, PRESENTATION_SYSTEM='presentation', REVIEW_SYSTEM='review',
                  WORD_REVIEW_SYSTEM=WORD_REVIEW_SYSTEM, GROK_ERROR_PREFIX='ERROR:',
                  asyncio=SimpleNamespace(sleep=AsyncMock()))
        exec(compile(ast.Module(body=[function], type_ignores=[]), 'main.py', 'exec'), ns)
        return ns, client, create

    async def test_word_editor_uses_bounded_structured_request_and_keeps_author_model(self):
        ns, client, create = self.request_env()
        await ns['ask_grok']('Review', system_prompt=WORD_REVIEW_SYSTEM, temperature=.15)
        client.with_options.assert_called_once_with(timeout=60., max_retries=0)
        call = create.await_args.kwargs
        self.assertEqual(call['model'], 'grok-4.6')
        self.assertEqual(call['response_format'], {'type':'json_object'})
        self.assertEqual(call['messages'][0]['content'], WORD_REVIEW_SYSTEM)
        await ns['ask_grok']('Write', system_prompt=WORD_SYSTEM)
        self.assertEqual(create.await_args.kwargs['model'], 'grok-4.3')
        self.assertNotIn('response_format', create.await_args.kwargs)
        await ns['ask_grok']('Slides', system_prompt='presentation')
        self.assertEqual(create.await_args.kwargs['model'], 'grok-4.6')

    async def test_word_editor_does_not_retry_truncated_json_or_loop_on_network_failure(self):
        ns, client, create = self.request_env()
        create.return_value.choices[0].finish_reason = 'length'
        result = await ns['ask_grok']('Review', system_prompt=WORD_REVIEW_SYSTEM)
        self.assertTrue(result.startswith('ERROR:'))
        self.assertEqual(create.await_count, 1)
        create.reset_mock()
        create.side_effect = TimeoutError('fixture')
        result = await ns['ask_grok']('Review', system_prompt=WORD_REVIEW_SYSTEM)
        self.assertTrue(result.startswith('ERROR:'))
        self.assertEqual(create.await_count, 2)

    async def test_full_academic_author_and_growth_use_bounded_json_but_drafts_keep_existing_model(self):
        ns, client, create = self.request_env()
        for system in (WORD_SYSTEM, WORD_SYSTEM + ' Верни additions.'):
            await ns['ask_grok']('Full academic text', system_prompt=system, document_json=True)
            client.with_options.assert_called_with(timeout=90., max_retries=0)
            self.assertEqual(create.await_args.kwargs['model'], 'grok-4.6')
            self.assertEqual(create.await_args.kwargs['response_format'], {'type': 'json_object'})
        await ns['ask_grok']('Outline', system_prompt=WORD_SYSTEM)
        self.assertEqual(create.await_args.kwargs['model'], 'grok-4.3')
        await ns['ask_grok']('Focused review', system_prompt=WORD_REVIEW_SYSTEM + ' Return index.')
        self.assertEqual(create.await_args.kwargs['model'], 'grok-4.6')


class WordDeliveryPageTests(unittest.IsolatedAsyncioTestCase):
    def setup_env(self, actual):
        e,m,s=test_office.WordWorkflowTests().env()
        async def prepare(*args, **kwargs):
            return await prepare_word_files(*args, **kwargs, measure=lambda _:actual)
        e['prepare_word_files']=prepare
        return e,m,s

    async def test_short_result_disclosed_and_only_delivered_pages_charged(self):
        e,m,s=self.setup_env(1)
        await e['word_build'](m,s)
        self.assertEqual(m.answer_document.await_count,2)
        e['spend_credits'].assert_called_once_with(-9901,'word',amount=10, operation_id=__import__('unittest.mock',fromlist=['ANY']).ANY)
        self.assertTrue(any('1 стр.' in str(c) and '10 кредитов' in str(c) for c in m.answer.await_args_list))
        self.assertEqual(e['ask_grok'].await_count,2)

    async def test_extra_pages_do_not_increase_charge(self):
        e,m,s=self.setup_env(3)
        await e['word_build'](m,s)
        e['spend_credits'].assert_called_once_with(-9901,'word',amount=20, operation_id=__import__('unittest.mock',fromlist=['ANY']).ANY)
        self.assertEqual(e['ask_grok'].await_count,2)

    async def test_plain_document_prompt_has_editorial_structure_without_legal_language(self):
        e,m,s=self.setup_env(2)
        await e['word_build'](m,s)
        prompt=e['ask_grok'].await_args.args[0]
        self.assertIn(WORD_EDITORIAL_GUIDANCE,prompt)
        self.assertNotIn('точно по формулировкам ГК РФ',prompt)
        self.assertNotIn('детекторы ИИ-текста',prompt)

    async def test_word_generation_and_revision_use_document_system_instead_of_slide_persona(self):
        e,m,s=self.setup_env(1)
        await e['word_build'](m,s)
        self.assertEqual(e['ask_grok'].await_count, 2)
        for call, system in zip(e['ask_grok'].await_args_list, (WORD_SYSTEM, WORD_REVIEW_SYSTEM)):
            self.assertEqual(call.kwargs['system_prompt'], system)
            self.assertLessEqual(call.kwargs['temperature'], .25)
        self.assertIn('target_words', e['ask_grok'].await_args_list[-1].args[0])

    async def test_uploaded_word_also_uses_document_system(self):
        e,m,s=self.setup_env(2)
        await e['build_word_from_upload'](m,s,'Учебные материалы','Сравни способы хранения')
        self.assertEqual(e['ask_grok'].await_args.kwargs['system_prompt'], WORD_SYSTEM)
        self.assertEqual(m.answer_document.await_count, 1)
