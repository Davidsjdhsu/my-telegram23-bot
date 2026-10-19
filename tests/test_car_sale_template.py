import ast
import asyncio
import copy
import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from zipfile import ZipFile
from docx import Document
import car_sale_template as car
import test_office

class State:
    def __init__(self,data):self.data=data;self.current=None
    async def get_data(self):return copy.deepcopy(self.data)
    async def update_data(self,**kw):self.data.update(copy.deepcopy(kw))
    async def set_state(self,v):self.current=v
    async def clear(self):self.data={};self.current=None

class CarFieldsTests(unittest.TestCase):
    def test_scope(self):
        self.assertTrue(car.applies({'word_kind':'dkp_car'}))
        self.assertTrue(car.applies({'word_kind':'dkp','topic':'Продаю автомобиль'}))
        self.assertFalse(car.applies({'word_kind':'dkp','topic':'Продажа квартиры'}))
        self.assertFalse(car.applies({'word_kind':'referat','topic':'Продажа автомобилей'}))
        self.assertFalse(car.applies({'word_kind':'dkp_car','content_lang':'en'}))
        self.assertFalse(car.applies({'word_kind':'dkp','topic':'Автор продаёт книгу'}))
    def test_scope_rejects_explicit_pledge_representative_and_installments(self):
        for text in ['Продавец ООО Ромашка', 'Продавец ИП Петров', 'продажа по доверенности', 'оплата в рассрочку', 'машина в залоге']:
            self.assertTrue(car.scope_issue({'topic':text}))
        for text in ['Автомобиль не в залоге', 'без залога, без рассрочки', 'Залог: нет']:
            self.assertFalse(car.scope_issue({'topic':text}))
    def test_rejects_invented_values_and_preserves_exact_details(self):
        source='ФИО продавца: Петров Александр\nЦена: 1 234 500\nVIN: ABC123'
        clean=car.validate_fields({'seller_name':'Петров Александр','price':'1234500','vin':'ABC123','buyer_name':'Иванов','unknown':'ABC123','received_amount':123,'date':'[дата]'},source)
        self.assertEqual(clean,{'seller_name':'Петров Александр','vin':'ABC123'})
    def test_form_corrections_and_empty_values(self):
        self.assertEqual(car.parse_labeled('Цена: 100\nЦена: 200\nVIN: '),{'price':'200','vin':''})
    def test_fixed_document_fills_without_rewriting_or_losing_breaks(self):
        with tempfile.TemporaryDirectory() as td:
            blank=Path(td)/'blank.docx';filled=Path(td)/'filled.docx'
            car.render(blank);self.assertEqual(blank.read_bytes(),car.TEMPLATE.read_bytes())
            car.render(filled,{'seller_name':'Петров Александр Константинович','price':'1 200 000'})
            d=Document(filled);before=Document(car.TEMPLATE)
            texts='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
            self.assertIn('Петров Александр Константинович',texts)
            self.assertIn('[ФИО покупателя полностью]',texts)
            self.assertIn('1 200 000',texts)
            for a,b in zip(before.paragraphs,d.paragraphs):
                if '[' not in a.text:self.assertEqual(a.text,b.text)
            self.assertEqual(d.element.xml.count('w:type="page"'),before.element.xml.count('w:type="page"'))
            with ZipFile(filled) as z:
                self.assertIsNone(z.testzip())
                self.assertFalse(any(n.startswith('word/media/') for n in z.namelist()))
                xml=''.join(z.read(n).decode() for n in z.namelist() if n.endswith('.xml'))
                self.assertNotIn('auto.ru',xml.lower());self.assertNotIn('авто.ру',xml.lower())
                self.assertNotIn('TargetMode="External"',''.join(z.read(n).decode() for n in z.namelist() if n.endswith('.rels')))
    def test_currency_payment_and_terminal_punctuation(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'filled.docx'
            for unit in ('рублей', 'руб.', '₽'):
                car.render(path, {'price':'850 000 '+unit, 'price_words':'восемьсот пятьдесят тысяч рублей', 'payment_method':'банковский перевод', 'extra':'Условие.', 'received_amount':'850 000 рублей', 'balance':'0 руб.'})
                text='\n'.join(p.text for p in Document(path).paragraphs)
                self.assertIn('850 000 (восемьсот пятьдесят тысяч) рублей.',text)
                self.assertIn('Способ оплаты: банковский перевод.',text)
                self.assertIn('Дополнительные условия: Условие. Изменения',text)
                self.assertNotIn('рублей рублей',text)
                self.assertNotIn('..',text)

    def test_all_entry_routes_call_shared_workflow(self):
        source=(Path(__file__).resolve().parents[1]/'main.py').read_text()
        tree=ast.parse(source);funcs={n.name:ast.get_source_segment(source,n) for n in tree.body if isinstance(n,ast.AsyncFunctionDef)}
        self.assertIn('await build_car_sale(m, state, globals())',funcs['word_build'])
        self.assertIn('await preview_car_sale(m, state, globals())',funcs['word_build_draft'])
        self.assertIn('await preview_car_sale(m, state, globals())',funcs['word_extra'])
        self.assertIn('await run_heavy(render_car_sale, tpl_path)',funcs['word_mode_template'])
        self.assertIn('await word_build(m, state)',funcs['_handle_miniapp_payload'])

class CarWorkflowTests(unittest.IsolatedAsyncioTestCase):
    def env(self,data=None):
        e,m,_=test_office.WordWorkflowTests().env()
        e['Form']=SimpleNamespace(waiting_word_confirm='confirm',waiting_word_extra='extra')
        e['cancel_kb']=lambda *a:None
        e['word_cost']=lambda p:5*p
        e['send_no_credits_notice']=AsyncMock()
        async def heavy(fn,*a,**kw):return fn(*a,**kw)
        e['run_heavy']=heavy
        async def prepare(content,pages,docx,pdf,**kw):
            await kw['build'](content,docx);Path(pdf).write_bytes(b'%PDF-fixture');return content,2
        e['prepare_word_files']=AsyncMock(side_effect=prepare)
        e['ask_grok']=AsyncMock(return_value=json.dumps({'fields':{'seller_name':'Петров','price':'1200000'}}))
        state=State(data or {'word_kind':'dkp_car','topic':'Продавец Петров. Цена 1200000.'})
        return e,m,state
    async def cleanup(self,s):
        directory=s.data.get('car_sale',{}).get('directory')
        if directory:shutil.rmtree(directory,ignore_errors=True)
    async def test_first_call_requires_review_no_delivery_or_charge(self):
        e,m,s=self.env();await car.build(m,s,e)
        self.assertEqual(s.current,'confirm');e['spend_credits'].assert_not_called();m.answer_document.assert_not_awaited()
        self.assertIn('Не заполнено:',''.join(str(c) for c in m.answer.call_args_list))
    async def test_confirmed_delivery_order_and_fixed_price(self):
        e,m,s=self.env();await car.build(m,s,e)
        events=[]
        m.answer_document.side_effect=lambda *a,**kw:events.append('send')
        e['spend_credits'].side_effect=lambda *a,**kw:events.append('charge')
        await car.build(m,s,e)
        self.assertEqual(events,['send','send','charge']);e['ask_grok'].assert_awaited_once()
        e['spend_credits'].assert_called_once_with(-9901,'word',amount=10)
        self.assertEqual(s.data,{})
    async def test_pdf_send_failure_resume_skips_docx_and_generation(self):
        e,m,s=self.env();await car.build(m,s,e)
        m.answer_document.side_effect=[None,RuntimeError('offline')]
        await car.build(m,s,e);e['spend_credits'].assert_not_called()
        self.assertTrue(s.data['car_sale']['sent_docx']);self.assertFalse(s.data['car_sale'].get('sent_pdf'))
        m.answer_document.side_effect=None
        await car.build(m,s,e)
        self.assertEqual(m.answer_document.await_count,3);e['prepare_word_files'].assert_awaited_once()
        e['ask_grok'].assert_awaited_once();e['spend_credits'].assert_called_once()
    async def test_render_failure_no_charge_and_retry_uses_cached_fields(self):
        e,m,s=self.env();await car.build(m,s,e)
        e['prepare_word_files'].side_effect=RuntimeError('pdf')
        await car.build(m,s,e);e['spend_credits'].assert_not_called();m.answer_document.assert_not_awaited()
        self.assertEqual(s.current,'confirm');e['ask_grok'].assert_awaited_once();e['finish_job'].assert_called_once()
        await self.cleanup(s)
    async def test_final_message_failure_no_double_charge_or_resend(self):
        e,m,s=self.env();await car.build(m,s,e)
        async def answer(text,**kw):
            if text.startswith('Готово.'):raise RuntimeError('offline')
        m.answer.side_effect=answer
        await car.build(m,s,e);e['spend_credits'].assert_called_once()
        m.answer.side_effect=None
        await car.build(m,s,e);e['spend_credits'].assert_called_once()
        self.assertEqual(m.answer_document.await_count,2)
    async def test_insufficient_balance_after_preview_does_not_start(self):
        e,m,s=self.env();await car.build(m,s,e);e['can_afford']=lambda *a:False
        await car.build(m,s,e);e['start_job'].assert_not_called();e['spend_credits'].assert_not_called()
    async def test_extraction_failure_keeps_input_and_requests_fields(self):
        e,m,s=self.env();e['ask_grok'].side_effect=RuntimeError('offline')
        await car.build(m,s,e);self.assertEqual(s.current,'extra');self.assertIn('topic',s.data)
        self.assertNotIn('car_sale',s.data);e['spend_credits'].assert_not_called()
    async def test_explicit_correction_preserves_other_fields_no_model_call(self):
        e,m,s=self.env();await car.preview(m,s,e)
        await s.update_data(extra='Цена: 1500000\nVIN: ABCDEF')
        await car.preview(m,s,e)
        self.assertEqual(s.data['car_sale']['fields'],{'seller_name':'Петров','price':'1500000','vin':'ABCDEF'})
        e['ask_grok'].assert_awaited_once()
        await s.update_data(extra='Цена: 1500000\nVIN: ABCDEF\nVIN: ')
        await car.preview(m,s,e);self.assertNotIn('vin',s.data['car_sale']['fields'])
    async def test_new_request_does_not_reuse_old_data(self):
        e,m,s=self.env();await car.preview(m,s,e)
        await s.update_data(topic='Марка и модель: Lada',extra='')
        await car.build(m,s,e)
        self.assertEqual(s.data['car_sale']['fields'],{'model':'Lada'})
        m.answer_document.assert_not_awaited()
    async def test_actual_word_build_dispatches_template(self):
        e,m,s=self.env();await e['word_build'](m,s)
        self.assertEqual(s.current,'confirm');e['spend_credits'].assert_not_called()
        await e['word_build'](m,s);self.assertEqual(m.answer_document.await_count,2)

    async def test_out_of_scope_never_generates_or_charges(self):
        e,m,s=self.env({'word_kind':'dkp_car','topic':'Автомобиль в залоге'})
        await car.build(m,s,e)
        e['ask_grok'].assert_not_awaited();m.answer_document.assert_not_awaited();e['spend_credits'].assert_not_called()
        self.assertIsNone(s.data['car_sale'])
