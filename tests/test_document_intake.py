import asyncio
import re
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock
import document_intake as intake
from test_presentation_support import load_functions

REQUEST='Смета обустройства мастерской. Валюта — рубли.\nКраска: 3 банки по 3400 рублей.\nГрунтовка: 2 канистры по 1100 рублей.\nРабочий стол: 2 штуки по 14500 рублей.\nРабочее кресло: 2 штуки по 7200 рублей.\nИспользуй только эти позиции и цены.'

class State:
    def __init__(self): self.data={};self.current=None
    async def get_data(self): return dict(self.data)
    async def update_data(self,**kw): self.data.update(kw)
    async def set_state(self,value): self.current=value
    async def clear(self): self.data={};self.current=None


class SignalsTests(unittest.TestCase):
    def test_supplied_data_and_short_requests(self):
        for text,fmt in [(REQUEST,'excel'),('Краска: 3 банки по 3400\nСтол: 2 по 14500','excel'),('Нужен реферат про облака','word'),('Сделай эссе на 6 страниц','word'),('презентация про рыб','presentation'),('сделай презу про космос','presentation')]:
            with self.subTest(text=text):self.assertEqual(intake.suggest_format(text),fmt)
        self.assertEqual(intake.source_mode(REQUEST,'excel'),'user')
        self.assertEqual(intake.source_mode('Составь смету. Ничего не добавляй','excel'),'user')
        self.assertIsNone(intake.source_mode('Бюджет на 6 месяцев','excel'))
        self.assertEqual(intake.source_mode('Придумай условный пример сметы','excel'),'ai')
        self.assertEqual(intake.source_mode('Напиши эссе на 6 страниц','word'),'ai')

    def test_questions_and_negative_intents_stay_in_chat(self):
        for text in ['Что такое Excel?','Как сделать презентацию?','Почему нужна таблица?','Мне не нужна презентация','Объясни смету: краска 3 по 3400','Реши задачу: автобус проехал 100 км','Привет','Я не хочу файл Word']:
            with self.subTest(text=text):self.assertIsNone(intake.suggest_format(text))
        self.assertTrue(intake.needs_topic('Сделай Word пожалуйста'))
        self.assertFalse(intake.needs_topic('Сделай презентацию про рыб'))


class IntakeFlowTests(unittest.IsolatedAsyncioTestCase):
    def env(self):
        from excel_styles import STYLE_LABELS,PALETTE_LABELS
        kinds={'Смета расходов':'expense_estimate','Семейный бюджет':'family_budget','Бюджет проекта':'project_budget','Прайс-лист':'price_list','Учебный расчёт':'calc_table','Финансовая модель':'startup_model'}
        # AST extraction keeps these tests independent of live bot tokens and accounts.
        guesses=load_functions(['guess_word_kind','guess_excel_kind'],dict(WORD_KIND_KEYWORDS=[(('эссе',),'essay'),(('договор оказания услуг',),'services')],
            EXCEL_KIND_KEYWORDS=[(('смета',),'expense_estimate')],DKP_CONTEXT_KEYWORDS=(),_STUDY_WORK_RE=re.compile('эссе'),re=re))
        e=dict(re=re,user_lang=lambda _: 'ru',_intake_busy=set(),INTAKE_EXCEL_KINDS=kinds,
            get_user=lambda _:dict(lang_chosen=True),chat_rate_limited=lambda _:False,
            prefers_chat_reply=lambda t:bool(intake.QUESTION.search(t) or intake.DECLINE.search(t)),
            looks_like_document_request=lambda _:False,may_request_document=lambda _:False,
            classify_uncertain_intent=AsyncMock(return_value=None),respond_in_chat=AsyncMock(),
            bot=SimpleNamespace(send_message=AsyncMock()),tr=lambda key,*a,**k:key,
            Form=SimpleNamespace(waiting_document_intake='intake',waiting_excel_confirm='excel_confirm',waiting_word_size='word_size'),
            KeyboardButton=lambda **kw:kw,ReplyKeyboardMarkup=lambda **kw:kw,
            EXCEL_KIND_DESC={'expense_estimate':'смета расходов','calc_table':'учебный расчёт'},WORD_KIND_LABELS={'essay':{'ru':'Эссе'},'doc':{'ru':'Документ'},'services':{'ru':'Договор услуг'}},WORD_KIND_DESC={'essay':'эссе','doc':'документ','services':'договор услуг'},
            EXCEL_STYLES=STYLE_LABELS,EXCEL_PALETTES=PALETTE_LABELS,CREDIT_COSTS={'excel':5},
            excel_confirm_kb=lambda _: 'excel_confirm_buttons',main_kb=lambda *a,**k:'main',
            can_afford=lambda *a:True,word_cost=lambda p:p*5,normalize_word_pages=int,WORD_DEFAULT_PAGES=2,
            extract_page_count=lambda t:int(re.search(r'(\d+) страниц',t)[1]) if re.search(r'(\d+) страниц',t) else None,
            WORD_PAGE_CHOICE_KINDS={'doc','essay'},word_after_input=AsyncMock(),word_build_draft=AsyncMock(),
            extract_slide_count=lambda t:8,extract_style=lambda t:None,presentation_cost=lambda *a:40,
            _start_presentation_request=AsyncMock(),send_no_credits_notice=AsyncMock(),
            excel_build=AsyncMock(),word_build=AsyncMock(),spend_credits=AsyncMock())
        e.update({k:guesses[k] for k in ['guess_word_kind','guess_excel_kind']})
        return load_functions(['intake_kb','offer_document_intake','intake_question','advance_document_intake','document_intake_reply','handle_free_text_request'],e)

    def msg(self,text):return SimpleNamespace(text=text,from_user=SimpleNamespace(id=11),answer=AsyncMock())

    async def reply(self,e,s,text):await e['document_intake_reply'](self.msg(text),s)

    async def test_original_estimate_never_builds_before_final_confirmation(self):
        e=self.env();s=State();m=self.msg(REQUEST)
        await e['handle_free_text_request'](m,s,REQUEST)
        self.assertEqual(s.current,'intake');self.assertIn('Правильно понял',m.answer.call_args.args[0])
        await self.reply(e,s,intake.YES)
        self.assertEqual(s.data['document_intake']['step'],'style')
        await self.reply(e,s,intake.STANDARD)
        self.assertEqual(s.current,'excel_confirm');self.assertEqual(s.data['excel_topic'],REQUEST)
        self.assertEqual(s.data['excel_mode'],'user');self.assertFalse(s.data['excel_protect'])
        e['excel_build'].assert_not_awaited();e['spend_credits'].assert_not_awaited()

    async def test_ambiguous_table_asks_kind_and_source_then_preserves_new_data(self):
        e=self.env();s=State()
        await e['offer_document_intake'](self.msg(''),s,'Создай Excel для мастерской','excel','ru')
        await self.reply(e,s,intake.YES);self.assertEqual(s.data['document_intake']['step'],'kind')
        await self.reply(e,s,'Смета расходов');self.assertEqual(s.data['document_intake']['step'],'source')
        await self.reply(e,s,intake.OWN);self.assertEqual(s.data['document_intake']['step'],'details')
        await self.reply(e,s,'Краска: 3 банки по 3400\nСтол: 2 по 14500')
        await self.reply(e,s,intake.DESIGNER);await self.reply(e,s,'Графит и медь')
        self.assertEqual(s.data['excel_palette'],'copper');self.assertIn('для мастерской',s.data['excel_topic'])
        self.assertIn('14500',s.data['excel_topic']);self.assertEqual(s.data['excel_mode'],'user')

    async def test_known_settings_are_not_asked_again(self):
        e=self.env();s=State();text=REQUEST+'\nДизайнерский, графит и медь.'
        await e['offer_document_intake'](self.msg(''),s,text,'excel','ru')
        await self.reply(e,s,intake.YES);self.assertEqual(s.current,'excel_confirm')
        self.assertEqual(s.data['excel_style'],'designer');self.assertEqual(s.data['excel_palette'],'copper')

    async def test_word_known_pages_proceeds_only_to_draft(self):
        e=self.env();s=State();text='Напиши эссе про технологии на 6 страниц'
        await e['handle_free_text_request'](self.msg(text),s,text)
        e['word_build_draft'].assert_not_awaited()
        await self.reply(e,s,intake.YES)
        e['word_after_input'].assert_not_awaited();e['word_build_draft'].assert_awaited_once()
        self.assertEqual(s.data['word_pages'],6);self.assertEqual(s.data['word_kind'],'essay')
        e['word_build'].assert_not_awaited()

    async def test_word_missing_pages_uses_existing_questions_and_contract_uses_template(self):
        for text,kind in [('Напиши эссе про технологии','essay'),('Составь договор оказания услуг для ремонта','services')]:
            e=self.env();s=State();await e['offer_document_intake'](self.msg(''),s,text,'word','ru')
            await self.reply(e,s,intake.YES);e['word_after_input'].assert_awaited_once()
            self.assertEqual(s.data['word_kind'],kind);self.assertEqual(s.data['topic'],text)
            e['word_build'].assert_not_awaited()

    async def test_switch_format_retains_original_and_presentation_uses_shared_flow(self):
        e=self.env();s=State();await e['offer_document_intake'](self.msg(''),s,REQUEST,'excel','ru')
        await self.reply(e,s,intake.FORMATS['presentation'])
        args=e['_start_presentation_request'].call_args
        self.assertEqual(args.args[2],REQUEST);self.assertEqual(args.kwargs['settings']['user_text'],REQUEST)
        self.assertEqual(args.kwargs['settings']['mode'],'user')

    async def test_chat_and_cancel_do_not_generate(self):
        e=self.env();s=State();await e['offer_document_intake'](self.msg(''),s,REQUEST,'excel','ru')
        await self.reply(e,s,intake.CHAT)
        self.assertEqual(s.data,{});self.assertEqual(e['respond_in_chat'].call_args.args[1],REQUEST)
        e['excel_build'].assert_not_awaited()
        await self.reply(e,s,intake.YES);e['word_build'].assert_not_awaited()

    async def test_miniapp_chat_reply_uses_same_pending_request(self):
        e=self.env();s=State();await e['handle_free_text_request'](self.msg(REQUEST),s,REQUEST)
        await e['handle_free_text_request'](self.msg(''),s,'да')
        self.assertEqual(s.data['document_intake']['step'],'style')
        self.assertEqual(s.data['document_intake']['text'],REQUEST)

    async def test_low_balance_retains_request(self):
        for fmt,text in [('word','Напиши эссе на 6 страниц'),('presentation','Сделай презентацию про китов'),('excel',REQUEST)]:
            e=self.env();e['can_afford']=lambda *a:False;s=State()
            await e['offer_document_intake'](self.msg(''),s,text,fmt,'ru');await self.reply(e,s,intake.YES)
            self.assertEqual(s.data['document_intake']['text'],text);e['send_no_credits_notice'].assert_awaited_once()
            e['word_build_draft'].assert_not_awaited();e['_start_presentation_request'].assert_not_awaited()

    async def test_duplicate_confirmation_does_not_start_twice(self):
        e=self.env();s=State();await e['offer_document_intake'](self.msg(''),s,'Презентация про рыб','presentation','ru')
        entered=asyncio.Event();release=asyncio.Event()
        async def start(*a,**kw):entered.set();await release.wait()
        e['_start_presentation_request']=AsyncMock(side_effect=start)
        task=asyncio.create_task(self.reply(e,s,intake.YES));await entered.wait()
        await self.reply(e,s,intake.YES);release.set();await task
        e['_start_presentation_request'].assert_awaited_once()

class ProductionRegressionTests(unittest.IsolatedAsyncioTestCase):
    async def test_four_real_rows_survive_ai_default_and_both_protection_modes(self):
        import json, os
        from datetime import datetime
        from openpyxl import load_workbook
        from unittest.mock import Mock
        from test_office import office_env, COLORS
        rows=[dict(name='Краска',qty=3,price=3400),dict(name='Грунтовка',qty=2,price=1100),
              dict(name='Рабочий стол',qty=2,price=14500),dict(name='Рабочее кресло',qty=2,price=7200)]
        for protect in [False,True]:
            s=State();s.data=dict(excel_kind='expense_estimate',excel_topic=REQUEST,excel_mode='ai',excel_protect=protect)
            ask=AsyncMock(return_value=json.dumps({'title':'Мастерская','rows':rows}));order=[]
            async def run(fn,*a,**kw):return fn(*a,**kw)
            async def send(path,**kw):
                w=load_workbook(path);v=load_workbook(path,data_only=True).active
                self.assertEqual(w.active.protection.sheet,protect)
                self.assertEqual([v[f'B{r}'].value for r in range(5,9)],[r['name'] for r in rows])
                self.assertEqual(v['F9'].value,55800)
                self.assertNotIn('Условный пример',v['A2'].value)
                self.assertFalse(w.active['D5'].protection.locked)
                self.assertTrue(w.active['F5'].protection.locked)
                order.append('send')
            m=SimpleNamespace(from_user=SimpleNamespace(id=-903090),answer=AsyncMock(),answer_document=send)
            user={'generations':0,'history':[]}
            e=load_functions(['excel_build'],dict(user_lang=lambda _:'ru',get_user=lambda _:user,
                start_job=lambda _: (True,''),finish_job=Mock(),tr=lambda *a,**kw:'notice',pick_theme=lambda _:('default',COLORS),
                grok_json_lang_instruction=lambda _: '',content_gen_lang=lambda *a,**kw:'ru',ask_grok=ask,extract_json=json.loads,
                run_heavy=run,build_excel_items=office_env()['build_excel_items'],
                EXCEL_ROW_SCHEMA={'expense_estimate':{'fields':['name','qty','price'],'example':'{}'}},
                EXCEL_MIN_ROWS_AI={'expense_estimate':14},EXCEL_KIND_DESC={'expense_estimate':'Смета'},
                datetime=datetime,os=os,safe_filename=lambda *a,**kw:'test',FSInputFile=lambda p,**kw:p,
                spend_credits=lambda *a:order.append('charge'),note_success=Mock(),signature_line=AsyncMock(return_value=''),main_kb=lambda *a,**kw:None))
            await e['excel_build'](m,s)
            self.assertEqual(order,['send','charge']);ask.assert_awaited_once()
            self.assertIn('ничего не досочиняя',ask.call_args.args[0]);self.assertNotIn('МИНИМУМ',ask.call_args.args[0])

    async def test_explanations_never_offer_document_or_call_classifier(self):
        e=IntakeFlowTests().env();s=State()
        for text in ['Почему нужна таблица Excel с данными о доходах и расходах за все месяцы?',
                     'Мне не нужна презентация, просто расскажи про формулы и числа в Excel подробно']:
            await e['handle_free_text_request'](SimpleNamespace(text=text,from_user=SimpleNamespace(id=11)),s,text)
        e['classify_uncertain_intent'].assert_not_awaited();self.assertEqual(e['respond_in_chat'].await_count,2)
        self.assertEqual(s.data,{})

    async def test_all_existing_legal_types_map_to_their_own_templates(self):
        import ast
        from pathlib import Path
        tree=ast.parse((Path(__file__).resolve().parents[1]/'main.py').read_text())
        names={'WORD_KIND_KEYWORDS','EXCEL_KIND_KEYWORDS','DKP_CONTEXT_KEYWORDS','DKP_CAR_RE','DKP_REALTY_RE','_STUDY_WORK_RE'}
        ns={'re':re}
        nodes=[n for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id in names for t in n.targets)]
        exec(compile(ast.Module(body=nodes,type_ignores=[]),'test','exec'),ns)
        f=load_functions(['guess_word_kind'],ns)['guess_word_kind']
        for text,kind in [('акт оказанных услуг','services_act_ru'),('акт выполненных работ','work_act'),
            ('договор аренды квартиры','rent_flat_ru'),('договор аренды нежилого помещения','nonresidential_rent'),
            ('договор с самозанятым','self_employed'),('претензия о замене товара','claim_replace_ru'),
            ('агентский договор','agency'),('договор простого товарищества','joint_activity'),
            ('договор уступки требования','cession'),('соглашение о конфиденциальности','nda'),
            ('соглашение об уплате алиментов','alimony'),('гарантийное письмо','warranty_letter'),
            ('акт приёма-передачи имущества','act'),('договор возмездного оказания услуг','services')]:
            with self.subTest(text=text):
                self.assertEqual(f(text),kind)
                self.assertEqual(intake.suggest_format('Составь '+text),'word')
