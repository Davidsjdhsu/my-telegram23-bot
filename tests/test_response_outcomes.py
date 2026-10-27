import asyncio
import json
import unittest
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock, Mock
import chat_allowance as allowance
from response_outcomes import Reply, GenerationStopped, parse_reply, billable, ensure_generation
from presentation_support import validate_plan, validate_content
from test_presentation_support import make_plan, make_content, load_functions
import test_chat_support as cs
from test_chat_support import main_namespace
from chat_support import build_chat_messages


def rejected_plan():
    p = make_plan()
    p['title'] = 'Запрос отклонён'
    p['design_summary'] = 'План не готовится.'
    for slide in p['slides']:
        slide.update(title='Отказ', point='Содержание по этой формулировке не планируется.')
    return p

class GenerationTests(unittest.TestCase):
    def test_real_incident_rejected_before_build(self):
        with self.assertRaises(GenerationStopped):
            validate_plan(rejected_plan(), 8)

    def test_structured_refusal_and_clarification_need_no_slides(self):
        for status in ('refusal', 'clarification'):
            with self.assertRaises(GenerationStopped): validate_plan({'status': status}, 8)
            with self.assertRaises(GenerationStopped): validate_content({'status': status}, make_plan())

    def test_legacy_body_refusal_and_restored_content_rejected(self):
        content = make_content()
        content['slides'][0]['content'] = 'По этой формулировке презентацию составлять нельзя.'
        with self.assertRaises(GenerationStopped): validate_content(content, make_plan())
        with self.assertRaises(GenerationStopped): ensure_generation(content)

    def test_educational_adult_relationships_and_refusal_as_topic_are_valid(self):
        for title in ('Сексуальные отношения взрослых: согласие и здоровье', 'Отказ от курения', 'Как сказать нет'):
            p = make_plan(); p['title'] = title
            self.assertEqual(validate_plan(p, 8)['title'], title)
            self.assertTrue(validate_content(make_content(), p))

    def test_semantic_result_not_the_word_no(self):
        for text in ('Нет, 2 + 2 = 4, а не 5.', 'Отказ от курения снижает риски для здоровья.', 'Нет, я ИИ и не обижаюсь.', 'Фраза «Разговор окончен» означает завершение беседы.', 'Пример отказа: «Не могу выполнить этот запрос».'):
            self.assertTrue(billable(parse_reply(json.dumps({'status':'answer','text':text}))))
        for status in ('refusal', 'clarification'):
            self.assertFalse(billable(parse_reply(json.dumps({'status':status,'text':'Уточните цель.'}))))
        self.assertFalse(billable(parse_reply(json.dumps({'status':'answer','text':'Разговор окончен.'}))))

    def test_invalid_envelope_cannot_be_billed_as_answer(self):
        for raw in ('Разговор окончен.', '{}', '{"status":"success","text":"OK"}', '{"status":"answer","text":null}'):
            with self.assertRaises(ValueError): parse_reply(raw)

    def test_old_argument_not_replayed_as_assistant_instruction(self):
        messages = build_chat_messages('Объясни согласие в отношениях взрослых', history=[
            {'role':'assistant', 'content':'Нет, отвечать не обязан. Разговор окончен.'}])
        self.assertFalse(any(m['role']=='assistant' for m in messages))
        self.assertIn('Образовательные темы', messages[0]['content'])
        self.assertIn('Грубая лексика', messages[0]['content'])

class HandlerTests(unittest.IsolatedAsyncioTestCase):
    def message(self, n):
        return NS(from_user=NS(id=1), chat=NS(id=1), message_id=n, voice=None, answer=AsyncMock())

    async def test_refusals_and_clarifications_do_not_use_free_or_paid_allowance(self):
        for paid in (False, True):
            ns, users = cs.IntegrationTests().chat_ns(); user=users[1]
            if paid:
                user['chat_allowance']={'day':allowance.day(), 'receipts':{str(i):'chat' for i in range(10)}}
                allowance.consent(user, True)
            before = allowance.status(user)['remaining']
            for i, reply in enumerate((Reply('Не могу помочь с этим запросом.', 'refusal'),
                                       Reply('Какую задачу решить?', 'clarification'), 'Разговор окончен.')):
                ns['ask_grok_chat']=AsyncMock(return_value=reply)
                await ns['respond_in_chat'](self.message(i+1),'Вопрос','ru')
            self.assertEqual(user['credits'],100)
            self.assertEqual(allowance.status(user)['remaining'], before)
            self.assertFalse(user.get('wallet_transactions'))
            self.assertFalse(user.get('_chat_access_active'))
            self.assertFalse(any('Разговор окончен' in text for _,text,_ in ns['bot'].sent))

    async def test_delayed_memory_does_not_block_next_paid_reply(self):
        ns, users=cs.IntegrationTests().chat_ns();user=users[1]
        user['chat_allowance']={'day':allowance.day(), 'receipts':{str(i):'chat' for i in range(10)}}
        allowance.consent(user,True)
        started, release=asyncio.Event(),asyncio.Event()
        async def memory(uid): started.set(); await release.wait()
        ns['update_long_term_memory']=memory
        ns['ask_grok_chat']=AsyncMock(return_value=Reply('Полезный ответ.'))
        await ns['respond_in_chat'](self.message(1),'Первый вопрос','ru')
        await asyncio.wait_for(started.wait(),1)
        try:
            self.assertFalse(ns['_chat_gate'].active)
            self.assertFalse(user.get('_chat_access_active'))
            await ns['respond_in_chat'](self.message(2),'Второй вопрос','ru')
            self.assertEqual(user['credits'],99)
            await ns['respond_in_chat'](self.message(2),'Второй вопрос','ru')
            self.assertEqual(user['credits'],99)
            self.assertEqual(ns['ask_grok_chat'].await_count,2)
        finally:
            release.set(); await asyncio.sleep(0.01)

    async def test_refused_plan_is_not_retried_as_formatting_error(self):
        ns=load_functions(['_create_presentation_plan'], dict(
            grok_json_lang_instruction=lambda x:'',content_gen_lang=lambda *a:'ru',
            ask_grok=AsyncMock(return_value='{"status":"refusal"}'),
            grok_failed=lambda x:False,extract_json=json.loads))
        with self.assertRaises(GenerationStopped):
            await ns['_create_presentation_plan']({'slides':8}, 'ru')
        ns['ask_grok'].assert_awaited_once()

    async def test_plan_stop_clears_old_confirmation_and_asks_topic(self):
        ns=load_functions(['process_slides'],dict(
            presentation_cost=lambda *a:80,user_lang=lambda uid:'ru',extract_slide_count=lambda text:None,
            tr=lambda *a:'working',cancel_kb=lambda lang:'cancel',
            _create_presentation_plan=AsyncMock(side_effect=GenerationStopped()),
            Form=NS(waiting_topic='topic',waiting_user_text='text')))
        state=NS(get_data=AsyncMock(return_value={'slides':8}),update_data=AsyncMock(),set_state=AsyncMock())
        message=self.message(1);message.text='8 слайдов'
        await ns['process_slides'](message,state)
        self.assertIn('не списаны',message.answer.await_args.args[0])
        state.set_state.assert_awaited_once_with('topic')
        state.update_data.assert_awaited_once_with(pres_plan=None,pres_plan_key=None,sample=None)

if __name__=='__main__': unittest.main()

class FileDeliveryTests(unittest.IsolatedAsyncioTestCase):
    async def test_rejected_content_stops_actual_builder_before_images_files_and_charge(self):
        import test_presentation_support as fixtures
        env,state,sent=fixtures.BuildTests().build_env(3)
        env['Form'].waiting_topic='topic'
        env['ask_grok']=AsyncMock(return_value='{"status":"refusal"}')
        await env['_build_presentation'](NS(from_user=NS(id=-987654321)),state)
        env['ask_grok'].assert_awaited_once()
        env['generate_image'].assert_not_awaited()
        env['spend_credits'].assert_not_called()
        self.assertEqual(sent,[])
        self.assertIsNone(state.data['pres_plan'])
        self.assertEqual(state.current,'topic')

    async def test_cached_refusal_file_cannot_be_delivered_or_charged(self):
        import tempfile
        from pathlib import Path
        from pptx import Presentation
        from presentation_orders import PresentationOrders
        with tempfile.TemporaryDirectory() as tmp:
            send,charge,notify=AsyncMock(),AsyncMock(),AsyncMock()
            orders=PresentationOrders(tmp,send,charge,lambda *a:True,notify)
            data={'topic':'test','slides':3}
            job=orders.get(1,data,30)
            deck=Presentation()
            for title,body in [('Запрос отклонён',''),('Отказ','Содержание по этой формулировке не планируется.'),('Отказ','Не готовится.')]:
                slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=title;slide.placeholders[1].text=body
            deck.save(job.directory/'document.pptx');(job.directory/'document.pdf').write_bytes(b'fixture')
            job.row.update(files_ready=True);job.save()
            result=await orders.run(1,data,30,AsyncMock())
            self.assertEqual(result['status'],'rejected')
            send.assert_not_awaited();charge.assert_not_awaited()
            self.assertIn('не списаны',notify.await_args.args[1])
