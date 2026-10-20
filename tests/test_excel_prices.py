import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from excel_styles import PALETTES
from test_office import office_env, COLORS

ROWS = [
    dict(sku='0007',name='Диагностика',unit='услуга',price=1250.55,discount=10,note='Согласовать время',category='Сервис'),
    dict(sku='0008',name='Настройка',unit='услуга',price=0,discount=None,note='',category='Сервис'),
    dict(sku='0009',name='Ремонт',unit='услуга',price=None,discount=20,note='По согласованию',category='Сервис'),
]

class PriceListTests(unittest.TestCase):
    def build(self, path, rows, **kw):
        office_env()['build_excel_items'](str(path),'Прайс-лист','price_list',COLORS,rows,**kw)

    def test_six_styles_preserve_prices_skus_discount_formulas_and_sorting(self):
        for style,palette in [('standard','copper')]+[('designer',p) for p in PALETTES]:
            with self.subTest(style=style,palette=palette), tempfile.TemporaryDirectory() as td:
                p=Path(td)/'prices.xlsx';self.build(p,ROWS,excel_style=style,excel_palette=palette)
                w=load_workbook(p);s=w.active;c=load_workbook(p,data_only=True).active
                h=6 if style=='designer' else 4;first=h+1
                cols={x.value:x.column for x in s[h] if x.value}
                self.assertEqual(s.cell(first,cols['Артикул']).value,'0007')
                self.assertEqual(s.cell(first,cols['Цена, ₽']).value,1250.55)
                self.assertEqual(s.cell(first+1,cols['Цена, ₽']).value,0)
                self.assertIsNone(s.cell(first+2,cols['Цена, ₽']).value)
                self.assertEqual(c.cell(first,cols['Цена со скидкой, ₽']).value,1125.5)
                self.assertEqual(c.cell(first+1,cols['Цена со скидкой, ₽']).value,0)
                self.assertEqual(c.cell(first+2,cols['Цена со скидкой, ₽']).value,'Укажите цену')
                self.assertEqual(s.cell(first,cols['Цена со скидкой, ₽']).data_type,'f')
                self.assertFalse(s.protection.sheet)
                self.assertEqual(len(s.tables),1);self.assertEqual(len(s._charts),0)
                self.assertEqual(s.cell(h,cols['Наименование']).fill.fgColor.rgb[-6:],PALETTES[palette if style=='designer' else 'green'][0])
                self.assertNotIn('Кол-во',cols);self.assertNotIn('Сумма, ₽',cols)
                self.assertFalse(any('SUM(' in x.value for row in s for x in row if x.data_type=='f'))
                self.assertEqual(w.calculation.calcMode,'auto')
                self.assertEqual(s.print_title_rows,f'${h}:${h}')

    def test_minimal_catalogue_has_no_invented_columns_or_discounts(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'prices.xlsx';self.build(p,[dict(name='=1+2',price=10),dict(name='Нет цены',price=None)])
            s=load_workbook(p).active
            self.assertEqual([c.value for c in s[4]],['Наименование','Цена, ₽'])
            self.assertEqual(s['A5'].data_type,'s')
            self.assertIsNone(s['B6'].value)
            self.assertFalse(any(c.data_type=='f' for row in s for c in row))

    def test_invalid_prices_discounts_and_nonfinite_numbers_are_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            for field,value in [('price',-1),('price',float('inf')),('discount',-1),('discount',101),('discount',float('nan')),('discount','неясно')]:
                with self.subTest(field=field,value=value), self.assertRaises(ValueError):
                    self.build(Path(td)/'bad.xlsx',[dict(name='Товар',price=10,**{field:value})] if field!='price' else [dict(name='Товар',price=value)])

    def test_long_labels_have_wrap_and_sufficient_height(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'long.xlsx';self.build(p,[dict(sku='00000000000001',name='Настройка оборудования '*8,price=1200000000,note='Подробное пояснение '*10)])
            s=load_workbook(p).active
            self.assertGreater(s.row_dimensions[5].height,80)
            self.assertTrue(s['B5'].alignment.wrap_text)
