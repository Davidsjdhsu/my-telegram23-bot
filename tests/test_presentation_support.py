import ast
import copy
import io
import json
import random
import re
import textwrap
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import presentation_copy as pc
import presentation_support as ps
import presentation_visuals as pv
import presentation_designer as pd
import presentation_pdf as pp
import presentation_recovery as pr
import tempfile


def make_plan(n=8):
    return {'title': 'Солнечная энергетика', 'audience': '9–11 классы',
        'design_summary': 'Светлый фон, зелёные акценты',
        'design': {'background': 'light', 'accent': '#238452'},
        'slides': [{'number': i, 'title': f'Тема {i}', 'point': f'Главная мысль {i}'} for i in range(1,n+1)]}


def make_content(n=8):
    return {'title': 'Другое название', 'slides': [
        {'number': i, 'title': f'Изменённый заголовок {i}', 'content': 'Свет передаёт энергию. Панель вырабатывает ток.',
         'image_prompt': 'A solar panel', 'chart': None} for i in range(2,n+1)]}


def load_functions(names, env=None):
    tree = ast.parse((Path(__file__).resolve().parents[1]/'main.py').read_text())
    nodes = []
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name in names:
            node.decorator_list = []
            nodes.append(node)
    if len(nodes) != len(names):
        raise AssertionError(set(names)-{n.name for n in nodes})
    ns = {k: getattr(ps,k) for k in dir(ps) if not k.startswith('__')}
    ns.update({k: getattr(pc,k) for k in ('final_process_requested','personal_photo_story','needs_photo_context','photo_context_message','PHOTO_ONLY_LABELS','photo_sources')})
    ns.update({k: getattr(pv,k) for k in ('prepare_own_images','prepare_images', 'render_slide', 'visual_quality_notice', 'require_images', 'ImageDeliveryError', 'image_failure_notice', 'photo_caption_hint')})
    ns.update({k: getattr(pd,k) for k in ('is_designer','render_designer_slide','content_instruction','validate_designer','build_scene','plan_designer_images','DesignerValidationError','designer_copy')})
    ns.update({k: getattr(pp,k) for k in ('ensure_pdf_converter','convert_presentation_pdf','PresentationPDFError','pdf_failure_notice')})
    ns.update({k:getattr(pr,k) for k in ('PresentationRecovery','missing_images','recovery_notice','PARTIAL_LABELS','RESUME_LABELS')})
    ns.update(env or {})
    module = ast.Module(body=[ast.ImportFrom(module='__future__', names=[ast.alias(name='annotations')],level=0)] + nodes, type_ignores=[])
    exec(compile(ast.fix_missing_locations(module),'main.py','exec'),ns)
    return ns


class PlanValidationTests(unittest.TestCase):
    def test_total_counts(self):
        for n in [3,8,12,16,30]:
            self.assertEqual(len(ps.validate_plan(make_plan(n),n)['slides']),n)
        for n in [True,2,31,'eight',8.5,None]:
            with self.assertRaises(ps.PresentationValidationError): ps.total_slides(n)

    def test_wrong_plan_count_rejected(self):
        with self.assertRaises(ps.PresentationValidationError): ps.validate_plan(make_plan(9),8)

    def test_duplicate_numbers_rejected(self):
        p=make_plan(); p['slides'][2]['number']=2
        with self.assertRaises(ps.PresentationValidationError): ps.validate_plan(p,8)

    def test_invalid_color_rejected(self):
        for accent in ['#fff','green','url(secret)',None]:
            p=make_plan(); p['design']['accent']=accent
            with self.assertRaises(ps.PresentationValidationError): ps.validate_plan(p,8)

    def test_wrong_shapes_rejected(self):
        for p in [None,[],{}, {'slides': [None]*8}]:
            with self.assertRaises(ps.PresentationValidationError): ps.validate_plan(p,8)

    def test_empty_or_overlong_text_rejected(self):
        for v in ['', 'x'*181, False]:
            p=make_plan(); p['slides'][3]['point']=v
            with self.assertRaises(ps.PresentationValidationError): ps.validate_plan(p,8)

    def test_content_excludes_cover_and_preserves_titles(self):
        p=ps.validate_plan(make_plan(),8)
        c=ps.validate_content(make_content(),p)
        self.assertEqual(c['title'],p['title'])
        self.assertEqual([x['title'] for x in c['slides']], [x['title'] for x in p['slides'][1:]])
        self.assertEqual(len(c['slides'])+1,8)

    def test_extra_missing_and_reordered_body_rejected(self):
        p=ps.validate_plan(make_plan(),8)
        for c in [make_content(9),make_content(7), {'slides': list(reversed(make_content()['slides']))}]:
            with self.assertRaises(ps.PresentationValidationError): ps.validate_content(c,p)

    def test_extra_paragraphs_are_preserved_by_two_block_layouts(self):
        c=make_content();c['slides'][0]['content']='Первое.\nВторое.\n\nТретье.\nЧетвёртое.'
        text=ps.validate_content(c,make_plan())['slides'][0]['content']
        self.assertEqual(len(text.split('\n\n')),2)
        for word in ['Первое.','Второе.','Третье.','Четвёртое.']: self.assertIn(word,text)

    def test_empty_content_rejected(self):
        c=make_content(); c['slides'][0]['content']=' '
        with self.assertRaises(ps.PresentationValidationError): ps.validate_content(c,make_plan())

    def test_palette_does_not_mutate_shared_theme(self):
        original={'bg':(1,2,3),'ink':(4,5,6),'line':(0,0,0),'heading_font':'Calibri'}
        before=copy.deepcopy(original)
        result=ps.apply_design(original,ps.validate_plan(make_plan(),8))
        self.assertEqual(original,before)
        self.assertEqual(result['line'],(35,132,82))
        self.assertEqual(result['bg'],(250,252,250))
        p=make_plan();p['design']={'background':'theme','accent':''}
        self.assertEqual(ps.apply_design(original,p),original)

    def test_cache_invalidated_when_requirements_change(self):
        data={'topic':'Солнце','extra':'зелёный','slides':8,'theme_name':'minimal'}
        key=ps.plan_key(data)
        for field,value in [('topic','Космос'),('extra','синий'),('slides',12),('theme_name','tech')]:
            self.assertNotEqual(ps.plan_key({**data,field:value}),key)
        self.assertEqual(ps.plan_key({**data,'user_photos':['file']}),key)

    def test_preview_lists_every_slide(self):
        text=ps.format_plan(ps.validate_plan(make_plan(),8))
        self.assertIn('включая обложку: 8',text)
        for i in range(1,9): self.assertIn(f'{i}. ',text)
        self.assertIn('9–11 классы',text)
        self.assertIn('зелёные акценты',text)


class State:
    def __init__(self,data): self.data=copy.deepcopy(data); self.current=None
    async def get_data(self): return copy.deepcopy(self.data)
    async def update_data(self,**kw): self.data.update(kw)
    async def set_state(self,value): self.current=value
    async def clear(self): self.data={}; self.current=None


class WorkflowTests(unittest.IsolatedAsyncioTestCase):
    def env(self, **extra):
        env={'user_lang':lambda uid:'ru', 'can_afford':lambda *a:True,
            'presentation_cost':lambda n:n*5, 'send_no_credits_notice':AsyncMock(),
            'process_slides':AsyncMock(), 'extract_slide_count':lambda t:int(re.search(r'(\d+) слайд',t)[1]) if re.search(r'(\d+) слайд',t) else None,
            'extract_style':lambda t:'minimal' if 'минимал' in t else None,
            'extract_photo_mode':lambda t:'own' if 'свои фото' in t else None,
            'Form':SimpleNamespace(waiting_pres_clarify='clarify',waiting_confirm='confirm',waiting_slides='slides'),
            're':re,
            'content_gen_lang':lambda d,l:l,'grok_json_lang_instruction':lambda l:f' Язык {l}',
            'grok_failed':lambda s:s.startswith('ERROR'), 'extract_json':json.loads,
            'ask_grok':AsyncMock(), 'tr':lambda key,lang,**kw:key+'\n'+kw.get('sample',''),
            'slides_kb':lambda lang:'slides','confirm_kb':lambda lang:'confirm',
            'cancel_kb':lambda lang:'cancel', 'THEME_LABELS_I18N':{},'THEME_LABELS':{}}
        env.update(extra);return env

    def message(self,text=''):
        return SimpleNamespace(text=text,from_user=SimpleNamespace(id=123),answer=AsyncMock())

    async def test_clarification_preserves_original_count_and_photos(self):
        env=load_functions(['pres_clarify_handler'],self.env())
        state=State({'pres_topic':'Презентация 12 слайдов, свои фото, для 9 класса'})
        await env['pres_clarify_handler'](self.message('минимализм, светлый фон, зелёные акценты'),state)
        self.assertEqual(state.data['slides'],12)
        self.assertEqual(state.data['photo_mode'],'own')
        self.assertIn('зелёные акценты',state.data['extra'])
        env['process_slides'].assert_awaited_once()

    async def test_explicit_clarification_overrides_count(self):
        env=load_functions(['pres_clarify_handler'],self.env())
        state=State({'pres_topic':'Презентация 12 слайдов'})
        await env['pres_clarify_handler'](self.message('8 слайдов'),state)
        self.assertEqual(state.data['slides'],8)

    async def test_multilingual_slide_button_keeps_number(self):
        for text in ['12 Folien', '16 diapositives', '12 张幻灯片']:
            n=int(re.search(r'\d+',text)[0])
            env=load_functions(['process_slides'],self.env(
                _create_presentation_plan=AsyncMock(return_value=ps.validate_plan(make_plan(n),n)),
                _send_presentation_draft=AsyncMock()))
            state=State({})
            await env['process_slides'](self.message(text),state)
            self.assertEqual(state.data['slides'],n)

    async def test_failed_initial_plan_can_retry_with_slide_buttons(self):
        env=load_functions(['process_slides'],self.env(
            _create_presentation_plan=AsyncMock(side_effect=ps.PresentationValidationError()),
            _send_presentation_draft=AsyncMock()))
        state=State({'slides':12})
        await env['process_slides'](self.message(),state)
        self.assertEqual(state.current,'slides')

    async def test_bad_plan_retry_is_bounded(self):
        model=AsyncMock(side_effect=[json.dumps(make_plan(9)),json.dumps(make_plan(8))])
        env=load_functions(['_create_presentation_plan'],self.env(ask_grok=model))
        result=await env['_create_presentation_plan']({'slides':8},'ru')
        self.assertEqual(len(result['slides']),8); self.assertEqual(model.await_count,2)
        model=AsyncMock(return_value='not json')
        env=load_functions(['_create_presentation_plan'],self.env(ask_grok=model))
        with self.assertRaises(ps.PresentationValidationError): await env['_create_presentation_plan']({'slides':8},'ru')
        self.assertEqual(model.await_count,2)

    async def test_failed_extra_keeps_confirmed_plan(self):
        env=load_functions(['process_extra'],self.env(_create_presentation_plan=AsyncMock(side_effect=ps.PresentationValidationError()),_send_presentation_draft=AsyncMock()))
        state=State({'extra':'старое','slides':8,'sample':'старый план','pres_plan':make_plan()})
        before=copy.deepcopy(state.data)
        await env['process_extra'](self.message('12 слайдов, новый запрос'),state)
        self.assertEqual(state.data,before)
        self.assertEqual(state.current,'confirm')

    async def test_long_draft_sent_as_text_with_buttons_only_at_end(self):
        env=load_functions(['_send_presentation_draft'],self.env())
        m=self.message()
        await env['_send_presentation_draft'](m,'\n\n'.join(['🌞'*175]*30),'buttons')
        self.assertGreater(m.answer.await_count,1)
        for call in m.answer.await_args_list:
            self.assertLessEqual(len(call.args[0].encode('utf-16-le'))//2,3500)
            self.assertIsNone(call.kwargs['parse_mode'])
        self.assertEqual(m.answer.await_args_list[-1].kwargs['reply_markup'],'buttons')
        for call in m.answer.await_args_list[:-1]: self.assertIsNone(call.kwargs['reply_markup'])





class BuildTests(unittest.IsolatedAsyncioTestCase):
    def build_env(self, count=8, bad_content=False):
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.lib.pagesizes import A4
        from datetime import datetime
        import os
        p=ps.validate_plan(make_plan(count),count)
        data={'slides':count,'topic':'Солнечная энергетика','extra':'Светлый фон, зелёные акценты','theme_name':'minimal','mode':'ai'}
        data.update(pres_plan=p,pres_plan_key=ps.plan_key(data))
        sent=[]
        async def send_document(chat_id,file,**kw): sent.append((file.filename,Path(file.path).read_bytes()))
        async def run_heavy(fn,*args): return fn(*args)
        theme={'bg':(255,255,255),'ink':(15,15,15),'mid':(70,70,70),'mute':(120,120,120),'line':(0,0,0),'photo':'natural photo','heading_font':'Calibri'}
        model=AsyncMock(return_value=json.dumps(make_content(count+1 if bad_content else count)))
        async def generate_fixture(prompt, path, **kwargs):
            from PIL import Image
            import hashlib
            color=tuple(hashlib.sha256(path.encode()).digest()[:3])
            Image.new('RGB',(640,480),color).save(path)
            return True
        async def fixture_pdf(source, target, count):
            # Delivery tests isolate conversion. Real rendering has its own integration check.
            prs=Presentation(source)
            pdf=canvas.Canvas(target,pagesize=(float(prs.slide_width)/12700,float(prs.slide_height)/12700))
            for slide in prs.slides:
                y=480
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        pdf.drawString(30,y,' '.join(shape.text.split()))
                        y-=20
                pdf.showPage()
            pdf.save()
            return target
        recovery_root=tempfile.TemporaryDirectory(prefix='alix-recovery-test-')
        env=WorkflowTests().env(ask_grok=model)
        env.update(USERS_FILE=str(Path(recovery_root.name)/'users.json'),
                   _recovery_root=recovery_root, presentation_recovery_kb=lambda lang: 'recovery')
        env.update({'local_issues':lambda *a:[], 'review_content':AsyncMock(return_value=[]), 'DeckQualityError':__import__('presentation_quality').DeckQualityError,'Presentation':Presentation,'Inches':Inches,'Pt':Pt,'RGBColor':RGBColor,
            'MSO_SHAPE':MSO_SHAPE,'PP_ALIGN':PP_ALIGN,'MSO_ANCHOR':MSO_ANCHOR,
            'canvas':canvas,'pdfmetrics':pdfmetrics,'A4':A4,'datetime':datetime,'os':os,
            'random':random.Random(7),'re':re,'textwrap':textwrap,'THEMES':{'minimal':theme,'default':theme},
            'pick_theme':lambda t:('minimal',theme),'start_job':Mock(return_value=(True,None)),
            'finish_job':Mock(),'get_user':lambda uid:{'generations':0,'history':[]},
            'tempfile':tempfile,'ask_presentation_vision':AsyncMock(return_value=json.dumps({'relevant':True,'has_text':False,'has_chart':False,'defective':False})), 'add_chart':Mock(),
            'ensure_pdf_converter':Mock(return_value='soffice'), 'convert_presentation_pdf':AsyncMock(side_effect=fixture_pdf),
            'run_heavy':run_heavy,'generate_image':AsyncMock(side_effect=generate_fixture),
            'register_pdf_fonts':lambda:('Helvetica','Helvetica-Bold'),
            'spend_credits':Mock(),'note_success':Mock(),'signature_line':AsyncMock(return_value=''),
            'main_kb':lambda *a,**kw:None,'FSInputFile':lambda path,filename:SimpleNamespace(path=path,filename=filename),
            'bot':SimpleNamespace(send_message=AsyncMock(),send_document=send_document),
            '_create_presentation_plan':AsyncMock(return_value=p)})
        names=['_build_presentation','slide_background','rect','txt','fill_slide_text','slide_paragraphs',
               '_estimate_wrapped_lines','_relative_luminance','_blend_toward','safe_line','invert_colors',
               'sandwich_colors','card_fill_for','text_on_fill','wrap_lines','safe_filename']
        return load_functions(names,env),State(data),sent

    async def test_real_pptx_exact_total_and_colors_and_charge(self):
        from pptx import Presentation
        for count in (3,8,30):
            with self.subTest(count=count):
                env,state,sent=self.build_env(count)
                m=SimpleNamespace(from_user=SimpleNamespace(id=-987654321),chat=SimpleNamespace(id=-987654321))
                await env['_build_presentation'](m,state)
                self.assertEqual(len(sent),2)
                pptx=Presentation(io.BytesIO(sent[0][1]))
                self.assertEqual(len(pptx.slides),count)
                for slide in pptx.slides:
                    stops=slide.shapes[0].fill.gradient_stops
                    self.assertEqual(tuple(stops[0].color.rgb),(250,252,250))
                env['spend_credits'].assert_called_once_with(-987654321,'presentation',amount=count*5)
                env['finish_job'].assert_called_once()
                env['_create_presentation_plan'].assert_not_awaited()
                self.assertEqual(state.data,{})
                # Keep a reproducible local assembly fixture, never sent to Telegram.
                if count==8:
                    artifact=Path(__file__).resolve().parents[2]/'presentation-review'/'stage1-assembly.pptx'
                    artifact.parent.mkdir(exist_ok=True)
                    artifact.write_bytes(sent[0][1])

    async def test_invalid_count_stops_before_images_or_charge(self):
        env,state,sent=self.build_env(8,bad_content=True)
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)),state)
        self.assertEqual(sent,[])
        self.assertEqual(env['ask_grok'].await_count,3)
        env['generate_image'].assert_not_awaited()
        env['spend_credits'].assert_not_called()
        env['finish_job'].assert_called_once()

    async def test_content_repair_then_build(self):
        env,state,sent=self.build_env(8)
        env['ask_grok'].side_effect=['{}',json.dumps(make_content())]
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)),state)
        self.assertEqual(len(sent),2)
        self.assertEqual(env['ask_grok'].await_count,2)
        env['spend_credits'].assert_called_once()

    async def test_upload_without_plan_is_validated(self):
        env,state,sent=self.build_env(3)
        del state.data['pres_plan']
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-987654321)),state)
        self.assertEqual(len(sent),2)
        env['_create_presentation_plan'].assert_awaited_once()


if __name__ == "__main__": unittest.main()
