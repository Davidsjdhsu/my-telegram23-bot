import copy
import io
import json
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock
from pptx import Presentation
from pptx.util import Inches
import presentation_copy as pc
import presentation_quality as pq
import presentation_designer as pd
import presentation_profiles as pp
import test_presentation_support as f


def steps():
    return [{'heading':'Приём','text':'Записываем неисправность.'},
            {'heading':'Осмотр','text':'Мастер проверяет велосипед.'},
            {'heading':'Согласование','text':'Клиент утверждает ремонт.'}]

class SemanticTests(unittest.IsolatedAsyncioTestCase):
    def test_explicit_final_process_request(self):
        for text in ['На последнем слайде кратко покажи весь процесс.',
                     'В конце добавь схему круговорота воды.',
                     'Итоговая схема этапов ремонта.',
                     'On the final slide show a process diagram.']:
            self.assertTrue(pc.final_process_requested({'extra':text}),text)
        for text in ['В конце покажи фотографии.', 'Итоговая схема расположения мебели.',
                     'В конце не добавляй схему процесса.', 'Final slide without a process diagram.']:
            self.assertFalse(pc.final_process_requested({'extra':text}),text)

    def test_process_edges_and_comparison_rows_reach_editor(self):
        item={'number':2,'title':'Процесс','content':'Скрытый текст',
              'designer':{'layout':'process','blocks':steps()}}
        observed=pq.review_slide(item)
        self.assertEqual(observed['relations'],[{'from':1,'to':2,'meaning':'next_stage'},
                                                {'from':2,'to':3,'meaning':'next_stage'}])
        self.assertEqual(observed['blocks'],steps())
        self.assertNotIn('Скрытый',observed['visible_text'])
        item['designer']['layout']='sections'
        self.assertEqual(pq.review_slide(item)['relations'],[])
        item['designer']={'layout':'comparison','blocks':[{'heading':'A','text':'hidden'},{'heading':'B','text':'hidden'}],
                          'rows':[{'criterion':'Условия','left':'Может','right':'Зависит от'}]}
        observed=pq.review_slide(item)
        self.assertEqual(observed['rows'],item['designer']['rows']);self.assertEqual(observed['blocks'],[])
        self.assertNotIn('hidden',json.dumps(observed))

    def test_final_process_not_forced_into_conclusion(self):
        content={'title':'Ремонт','slides':[{'number':2,'title':'Порядок','content':'Ремонт.',
                     'designer':{'layout':'process','blocks':steps()}}]}
        data={'extra':'На последнем слайде покажи весь процесс.'}
        self.assertEqual(pq.local_issues(content,True,data),[])
        content['slides'][0]['designer']={'layout':'conclusion','blocks':steps()[:2]}
        self.assertEqual(pq.local_issues(content,True,data)[0]['problem'],'missing')

    async def test_logic_error_blocks_images_and_billing_after_bounded_repairs(self):
        env,state,sent=f.BuildTests().build_env(3)
        issue={'slide':2,'problem':'logic','fix':'Альтернативу нельзя связывать стрелкой как следующий этап.'}
        self.assertEqual(pq.parse_review({'issues':[issue]},{2}),[issue])
        env['review_content']=AsyncMock(return_value=[issue])
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertFalse(sent);env['generate_image'].assert_not_awaited();env['spend_credits'].assert_not_called()
        self.assertEqual(env['ask_grok'].await_count,3)

    async def test_builder_passes_final_diagram_request_to_writer_and_gate(self):
        env,state,sent=f.BuildTests().build_env(3)
        state.data.update(theme_name='designer_soft',extra='На последнем слайде покажи весь процесс.')
        c=f.make_content(3)
        for s in c['slides']:s['designer']={'layout':'process','blocks':steps()}
        env['ask_grok']=AsyncMock(return_value=json.dumps(c))
        seen=[]
        env['local_issues']=lambda content,designer,data: seen.append(data) or []
        await env['_build_presentation'](SimpleNamespace(from_user=SimpleNamespace(id=-1)),state)
        self.assertIn('Последний слайд должен быть process',env['ask_grok'].await_args.args[0])
        self.assertTrue(pc.final_process_requested(seen[0]));self.assertEqual(len(sent),2)

    def test_two_criteria_render_without_blank_third_row(self):
        d={'layout':'comparison','blocks':[{'heading':'Дома'},{'heading':'В мастерской'}],
           'rows':[{'criterion':'Инструменты','left':'Свои','right':'Оборудование мастера'},
                   {'criterion':'Опыт','left':'Опыт владельца','right':'Навыки специалиста'}]}
        for design in [{**pp.choose_design({'topic':'Ремонт'}),'family':family} for family in pp.PROFILES]+[pp.choose_design({'topic':'Ремонт велосипедов','theme_name':'designer_soft'})]:
            item={'title':'Где ремонтировать','content':'Сравнение условий.','designer':copy.deepcopy(d)}
            if design:item['_designer_design']=design
            self.assertEqual(len(pd.validate_designer(item)['rows']),2)
            scene=pd.build_scene(item,None,3,4)
            table=next(n for n in scene['nodes'] if n['type']=='table')
            self.assertEqual(len(table['cells']),3);self.assertEqual(len(table['heights']),3)
            prs=Presentation();prs.slide_width=Inches(1280/96);prs.slide_height=Inches(7.5)
            pd.render_scene(prs.slides.add_slide(prs.slide_layouts[6]),scene)
            b=io.BytesIO();prs.save(b);b.seek(0)
            saved=Presentation(b)
            self.assertEqual(len(next(s.table for s in saved.slides[0].shapes if s.has_table).rows),3)

    def test_repetition_uses_displayed_blocks_not_hidden_prose(self):
        a={'number':2,'title':'Один','content':'Неиспользуемая общая заготовка.',
           'designer':{'layout':'sections','blocks':steps()[:2]}}
        b={'number':3,'title':'Два','content':a['content'],
           'designer':{'layout':'sections','blocks':[{'heading':'Полив','text':'Влажность зависит от растения.'},
                                                      {'heading':'Свет','text':'Листьям нужен свет.'}]}}
        self.assertEqual(pq.local_issues({'slides':[a,b]}),[])
        b['designer']=copy.deepcopy(a['designer'])
        self.assertEqual(pq.local_issues({'slides':[a,b]})[0]['problem'],'repetition')
