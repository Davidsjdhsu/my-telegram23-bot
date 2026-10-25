import copy
import io
import unittest
from pptx import Presentation
from studio_palette_preview import scene_for, render_preview, ITEMS
from presentation_designer import build_scene, render_scene
from presentation_profiles import choose_design
from presentation_support import plan_key, finalize_plan_design, validate_plan
from presentation_diagrams import diagram_relations

class DesignerOptionsTests(unittest.TestCase):
 def test_scene_variants_preserve_semantic_nodes_edges_and_bounds(self):
  for style in ('strict','soft'):
   for layout in ('classic','cards','steps'):
    for kind in ('sequence','cycle','structure'):
     for count in (3,4,5):
      item=copy.deepcopy(ITEMS[3]);item['diagram']={'kind':kind,'nodes':[{'label':f'Этап номер {i+1}','text':'Вещество изменяется, после чего поступает на следующий этап.'} for i in range(count)]}
      item['_designer_design']=choose_design({'extra':'Шаблон №16','studio_palette':'tech','studio_design_options':{'background_mode':'gradient','diagram_layout':layout}})
      scene=build_scene(item,None,4,6,'ru','designer_'+style)
      texts=[n['text'].replace('\n',' ') for n in scene['nodes'] if n['type']=='text']
      for n in item['diagram']['nodes']:
       self.assertTrue(any(n['label'] in t for t in texts));self.assertTrue(any(n['text'] in t for t in texts))
      arrows=[n for n in scene['nodes'] if n['type']=='diagram_edge' and n['arrow']]
      self.assertEqual(len(arrows),0 if kind=='structure' else count if kind=='cycle' else count-1)
      self.assertEqual(len([n for n in scene['nodes'] if n.get('role')=='background_gradient']),1)
 def test_saved_options_survive_plan_validation_and_repair(self):
  data={'topic':'Кофе','slides':4,'theme_name':'designer_strict','studio_design_options':{'background_mode':'gradient','diagram_layout':'cards'}}
  plan={'title':'Кофе','audience':'Все','design_summary':'Оформление','design':{'background':'light','accent':''},'slides':[{'title':'Кофе','number':i+1,'point':'Как устроена кофейня'} for i in range(4)]}
  result=finalize_plan_design(validate_plan(plan,4),data)
  self.assertEqual(result['design']['diagram_layout'],'cards');self.assertEqual(result['design']['background_mode'],'gradient')
  self.assertNotEqual(plan_key(data),plan_key({**data,'studio_design_options':{}}))
 def test_native_export_keeps_gradient_and_editable_text(self):
  from pptx.util import Inches
  for layout in ('classic','cards','steps'):
   scene=scene_for('strict','d42','wine',3,'gradient',layout)
   p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
   slide=p.slides.add_slide(p.slide_layouts[6]);render_scene(slide,scene)
   out=io.BytesIO();p.save(out);loaded=Presentation(io.BytesIO(out.getvalue()))
   xml=loaded.slides[0]._element.xml
   self.assertIn('gradFill',xml);self.assertIn('Приготовление',xml)
 def test_empty_palette_supported_for_new_options(self):
  for style in ('strict','soft'):
   for page in range(4):self.assertIn('<svg',render_preview(style,'d16','',page,'gradient','cards'))
