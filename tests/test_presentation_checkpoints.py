import copy
import io
import json
import os
from pathlib import Path
import tempfile
import time
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock

from pptx import Presentation
import test_presentation_support as f
import presentation_recovery as recovery

MESSAGE = SimpleNamespace(from_user=SimpleNamespace(id=-987654321))


class CheckpointTests(unittest.IsolatedAsyncioTestCase):
    async def test_consented_missing_cover_renders_in_all_three_styles(self):
        for style in ('minimal', 'designer_strict', 'designer_soft'):
            with self.subTest(style=style):
                env, state, sent = f.BuildTests().build_env(8)
                state.data['theme_name'] = style
                state.data['pres_plan_key'] = f.ps.plan_key(state.data)
                content = f.make_content(8)
                for slide in content['slides']:
                    slide['designer'] = {'layout':'auto', 'blocks':[]}
                env['ask_grok'] = AsyncMock(return_value=json.dumps(content))
                original = env['generate_image']
                async def fail_cover(prompt, path, **kw):
                    if 'generated-0-' in path:
                        return False
                    return await original(prompt, path, **kw)
                env['generate_image'] = fail_cover
                await env['_build_presentation'](MESSAGE, state)
                self.assertEqual(sent, [])
                await env['_build_presentation'](MESSAGE, state, allow_missing_images=True)
                self.assertEqual(len(sent), 2)
                deck = Presentation(io.BytesIO(sent[0][1]))
                self.assertEqual(len(deck.slides), 8)
                self.assertFalse(any(s.shape_type == 13 and s.name != 'Designer texture' for s in deck.slides[0].shapes))
                self.assertEqual(env['ask_grok'].await_count, 1)
                env['spend_credits'].assert_called_once()

    async def failed_order(self):
        env, state, sent = f.BuildTests().build_env(3)
        original = env['generate_image']
        async def one_failure(prompt, path, **kw):
            if 'generated-1-' in path:
                return False
            return await original(prompt, path, **kw)
        env['generate_image'] = AsyncMock(side_effect=one_failure)
        await env['_build_presentation'](MESSAGE, state)
        self.assertEqual(sent, [])
        env['spend_credits'].assert_not_called()
        self.assertIn('2', env['bot'].send_message.await_args.args[1])
        return env, state, sent, original

    async def test_resume_only_missing_picture_without_reauthoring_or_rechecking(self):
        env, state, sent, generate = await self.failed_order()
        author_calls = env['ask_grok'].await_count
        editorial_calls = env['review_content'].await_count
        env['generate_image'] = AsyncMock(side_effect=generate)
        env['ask_presentation_vision'].reset_mock()
        await env['_build_presentation'](MESSAGE, state)
        self.assertEqual(env['ask_grok'].await_count, author_calls)
        self.assertEqual(env['review_content'].await_count, editorial_calls)
        self.assertEqual(env['generate_image'].await_count, 1)
        self.assertEqual(env['ask_presentation_vision'].await_count, 1)
        self.assertEqual(len(sent), 2)
        self.assertEqual([sum(s.shape_type == 13 for s in slide.shapes)
                          for slide in Presentation(io.BytesIO(sent[0][1])).slides], [1, 1, 1])
        env['spend_credits'].assert_called_once()
        self.assertFalse(list(Path(env['USERS_FILE']).parent.glob('presentation-recovery/*/checkpoint.json')))

    async def test_explicit_partial_consent_retains_text_and_accepted_images_without_api_calls(self):
        env, state, sent, _ = await self.failed_order()
        for name in ['generate_image', 'ask_presentation_vision', 'ask_grok', 'review_content']:
            env[name] = AsyncMock(side_effect=AssertionError('No external calls allowed'))
        await env['_build_presentation'](MESSAGE, state, allow_missing_images=True)
        self.assertEqual(len(sent), 2)
        deck = Presentation(io.BytesIO(sent[0][1]))
        self.assertEqual(len(deck.slides), 3)
        self.assertEqual([sum(s.shape_type == 13 for s in slide.shapes) for slide in deck.slides], [1, 0, 1])
        for slide in deck.slides:
            self.assertTrue(any(s.has_text_frame and s.text for s in slide.shapes))
        env['spend_credits'].assert_called_once_with(MESSAGE.from_user.id, 'presentation', amount=15)

    async def test_second_failure_keeps_successful_images_and_remains_unbilled(self):
        env, state, sent, _ = await self.failed_order()
        env['generate_image'] = AsyncMock(return_value=False)
        await env['_build_presentation'](MESSAGE, state)
        self.assertEqual(env['generate_image'].await_count, 1)
        self.assertEqual(env['ask_grok'].await_count, 1)
        self.assertEqual(sent, [])
        env['spend_credits'].assert_not_called()
        self.assertEqual(state.current, 'confirm')

    async def test_changed_topic_cannot_reuse_partial_consent(self):
        env, state, sent, _ = await self.failed_order()
        state.data['topic'] = 'Fish, a completely different presentation'
        env['ask_grok'].reset_mock()
        env['generate_image'].reset_mock()
        await env['_build_presentation'](MESSAGE, state, allow_missing_images=True)
        self.assertEqual(sent, [])
        env['ask_grok'].assert_not_awaited()
        env['generate_image'].assert_not_awaited()
        env['spend_credits'].assert_not_called()

    async def test_no_checkpoint_cannot_silently_deliver_partial(self):
        env, state, sent = f.BuildTests().build_env(3)
        await env['_build_presentation'](MESSAGE, state, allow_missing_images=True)
        self.assertEqual(sent, [])
        env['ask_grok'].assert_not_awaited()
        env['spend_credits'].assert_not_called()
        env['finish_job'].assert_called_once()

    async def test_checkpoint_survives_new_reader_but_expired_or_corrupt_files_are_rejected(self):
        env, state, _, _ = await self.failed_order()
        root = Path(env['USERS_FILE']).parent / 'presentation-recovery'
        reader = recovery.PresentationRecovery(MESSAGE.from_user.id, state.data, root)
        with tempfile.TemporaryDirectory() as td:
            value = reader.load(td)
            self.assertEqual(value['missing'], [2])
            checkpoint = reader.directory / 'checkpoint.json'
            os.utime(checkpoint, (time.time()-recovery.TTL-1,)*2)
            self.assertIsNone(reader.load(td))
            os.utime(checkpoint, None)
            (reader.directory / 'asset-0.jpg').write_bytes(b'corrupt')
            self.assertIsNone(reader.load(td))

    async def test_other_user_or_changed_style_cannot_read_order(self):
        env, state, _, _ = await self.failed_order()
        root = Path(env['USERS_FILE']).parent / 'presentation-recovery'
        with tempfile.TemporaryDirectory() as td:
            self.assertIsNone(recovery.PresentationRecovery(333, state.data, root).load(td))
            changed = dict(state.data, theme_name='designer_soft')
            self.assertIsNone(recovery.PresentationRecovery(MESSAGE.from_user.id, changed, root).load(td))

    async def test_bot_buttons_route_to_actual_builder_with_explicit_consent_only(self):
        build = AsyncMock()
        env = f.load_functions(['resume_presentation_images','accept_presentation_missing_images'], {'_build_presentation':build})
        await env['resume_presentation_images'](MESSAGE, 'state')
        build.assert_awaited_with(MESSAGE, 'state')
        await env['accept_presentation_missing_images'](MESSAGE, 'state')
        build.assert_awaited_with(MESSAGE, 'state', allow_missing_images=True)
