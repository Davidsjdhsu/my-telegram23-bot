import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

from openpyxl import load_workbook
from excel_study import prepare_rows, calculated_values, StudyValidationError, MISSING, INVALID
from excel_styles import PALETTES
from test_office import office_env, COLORS
from test_presentation_support import load_functions


MOTION = [
    dict(id='s', name='Путь', value=120, unit='м'),
    dict(id='t', name='Время', value=15, unit='с'),
    dict(id='v', name='Скорость', formula='s/t', unit='м/с', note='Равномерное движение'),
    dict(id='kmh', name='Скорость в км/ч', formula='v*3.6', unit='км/ч'),
]


class StudyTests(unittest.TestCase):
    def test_motion_in_six_styles_and_real_formula_caches(self):
        build = office_env()['build_excel_items']
        for style, palette in [('standard', 'copper')] + [('designer', p) for p in PALETTES]:
            with self.subTest(style=style,palette=palette), tempfile.TemporaryDirectory() as td:
                p = Path(td)/'study.xlsx'
                build(str(p),'Движение','calc_table',COLORS,MOTION,excel_style=style,excel_palette=palette)
                w = load_workbook(p); s = w.active; c = load_workbook(p,data_only=True).active
                h, start = (6,2) if style=='designer' else (4,1)
                self.assertEqual(s.cell(h,start).fill.fgColor.rgb[-6:], PALETTES[palette if style=='designer' else 'green'][0])
                self.assertEqual(c.cell(h+3,start+2).value,8)
                self.assertAlmostEqual(c.cell(h+4,start+2).value,28.8)
                self.assertEqual(s.cell(h+3,start+2).data_type,'f')
                self.assertFalse(s.cell(h+1,start+2).protection.locked)
                self.assertTrue(s.cell(h+3,start+2).protection.locked)
                self.assertTrue(s.protection.sheet)
                self.assertEqual(len(s._charts),0)
                self.assertEqual(s.print_title_rows,f'${h}:${h}')
                self.assertEqual(w.calculation.calcMode,'auto')
                self.assertIn('v = s/t',s.cell(h+3,start+4).value)
                self.assertEqual(s.cell(h+3,start+3).value,'м/с')
                note=next(x for row in s for x in row if isinstance(x.value,str) and 'Эту подсказку' in x.value)
                self.assertFalse(note.protection.locked)

    def test_statistics_units_and_percentages_are_explicit(self):
        rows = [dict(id='a',value=10),dict(id='b',value=20),dict(id='c',value=30),
                dict(id='mean',formula='AVERAGE(a,b,c)'),dict(id='total',formula='SUM(a,b,c)'),
                dict(id='share',formula='a/total*100',unit='%'),dict(id='range1',formula='MAX(a,b,c)-MIN(a,b,c)'),
                dict(id='root',formula='SQRT(ABS(-a**2))')]
        v = calculated_values(prepare_rows(rows))
        self.assertEqual(v['mean'],20);self.assertEqual(v['total'],60)
        self.assertAlmostEqual(v['share'],100/6);self.assertEqual(v['range1'],20)
        self.assertEqual(v['root'],10)

    def test_missing_zero_and_invalid_results_propagate_without_fabricated_zero(self):
        for t, expected in [(None,MISSING),(0,INVALID),(-15,-8),(15,8)]:
            rows = [dict(r) for r in MOTION];rows[1]['value']=t
            values=calculated_values(prepare_rows(rows))
            self.assertEqual(values['v'],expected)
            if isinstance(expected,str):self.assertEqual(values['kmh'],expected)
        rows=[dict(id='a',value=None),dict(id='b',value=2),dict(id='mean',formula='AVERAGE(a,b)')]
        self.assertEqual(calculated_values(prepare_rows(rows))['mean'],MISSING)
        self.assertEqual(calculated_values(prepare_rows([dict(id='a',value=0),dict(id='b',formula='a*2')]))['b'],0)
        self.assertEqual(calculated_values(prepare_rows([dict(id='a',value=-1),dict(id='b',formula='SQRT(a)')]))['b'],INVALID)

    def test_unsafe_unknown_and_cyclic_formulas_never_become_excel_code(self):
        for expr in ['=a+1','A1+B1','__import__("os")','a.__class__','[a]','a[0]',
                     'WEBSERVICE("https://example.com")','SUM(a,other)','SUM()','SQRT(a,a)',
                     'a**a','a**999','a if a else 1','True','SUM(a, x=1)','"text"']:
            with self.subTest(expr=expr),self.assertRaises(StudyValidationError):
                prepare_rows([dict(id='a',value=1),dict(id='b',formula=expr)])
        for rows in [[dict(id='a',formula='a+1')],[dict(id='a',formula='b'),dict(id='b',formula='a')],
                     [dict(id='a',value=1),dict(id='a',value=2)], [dict(id='a',value=float('nan'))],
                     [dict(id='a',value='abc')],[dict(id='a',value=True)]]:
            with self.assertRaises(StudyValidationError):prepare_rows(rows)

    def test_forward_references_and_literal_text_and_tiny_values(self):
        rows=[dict(id='v',formula='s/t',name='=1+2',unit='=HYPERLINK("x")'),
              dict(id='t',value=2),dict(id='s',value=1e-12)]
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'tiny.xlsx'
            office_env()['build_excel_items'](str(p),'=1+1','calc_table',COLORS,rows)
            w=load_workbook(p).active;c=load_workbook(p,data_only=True).active
            self.assertEqual(c['C5'].value,5e-13)
            self.assertEqual(w['B5'].data_type,'s');self.assertEqual(w['D5'].data_type,'s')
            self.assertEqual(w['A1'].data_type,'s')
            rules=[r for cf in w.conditional_formatting for r in w.conditional_formatting[cf]]
            self.assertTrue(any(r.dxf and r.dxf.numFmt and 'E+00' in r.dxf.numFmt.formatCode for r in rules))


class StudyFailureFlowTests(unittest.IsolatedAsyncioTestCase):
    async def test_invalid_calculation_keeps_request_and_never_sends_or_charges(self):
        data=dict(excel_kind='calc_table',excel_topic='Вычислить скорость',excel_mode='user')
        state=SimpleNamespace(get_data=AsyncMock(return_value=data),set_state=AsyncMock(),clear=AsyncMock())
        msg=SimpleNamespace(from_user=SimpleNamespace(id=123),answer=AsyncMock(),answer_document=AsyncMock())
        spend=Mock(); finish=Mock()
        env=load_functions(['excel_build'],dict(user_lang=lambda _: 'ru',get_user=lambda _: {},
            start_job=lambda _: (True,''),tr=lambda *a:'Сборка',pick_theme=lambda _:('default',COLORS),
            grok_json_lang_instruction=lambda _: '',content_gen_lang=lambda *a,**k:'ru',
            EXCEL_ROW_SCHEMA={'calc_table':{'fields':['id','name','value','unit','formula','note'],'example':'{}'}},
            EXCEL_MIN_ROWS_AI={'calc_table':0},EXCEL_KIND_DESC={'calc_table':'Учебная таблица'},
            ask_grok=AsyncMock(return_value=json.dumps({'unsupported':True,'rows':[]})),extract_json=json.loads,
            spend_credits=spend,finish_job=finish,Form=SimpleNamespace(waiting_excel_confirm='confirm'),
            excel_confirm_kb=lambda _:None))
        await env['excel_build'](msg,state)
        spend.assert_not_called();msg.answer_document.assert_not_awaited();state.clear.assert_not_awaited()
        state.set_state.assert_awaited_with('confirm');finish.assert_called_once()
        self.assertIn('не списаны',msg.answer.call_args.args[0])
