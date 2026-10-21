import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from datetime import datetime
import os

from openpyxl import load_workbook
from test_office import office_env, COLORS
from test_presentation_support import load_functions
from excel_styles import PALETTES, STYLE_LABELS, PALETTE_LABELS, normalize_excel_appearance

ROWS = [
    {'name': 'Краска для стен', 'unit': 'банка', 'qty': 3, 'price': 3400, 'category': 'Материалы'},
    {'name': 'Подготовка стен', 'unit': 'м²', 'qty': 24, 'price': 50000, 'category': 'Работы'},
    {'name': 'Покраска стен', 'unit': 'м²', 'qty': 24, 'price': 450, 'category': 'Работы'},
]

class ExcelAppearanceTests(unittest.TestCase):
    def test_styles_keep_real_formulas_inputs_and_removable_note(self):
        build = office_env()['build_excel_items']
        for style, palette in [('standard', 'burgundy')] + [('designer', p) for p in PALETTES]:
            with self.subTest(style=style, palette=palette), tempfile.TemporaryDirectory() as td:
                p = str(Path(td) / 'estimate.xlsx')
                build(p, 'Смета', 'expense_estimate', COLORS, ROWS, excel_style=style, excel_palette=palette)
                ws = load_workbook(p).active
                cached = load_workbook(p, data_only=True).active
                header = 4 if style == 'standard' else 8
                first = header + 1
                total = first + 3 + (1 if style == 'designer' else 0)
                self.assertEqual(ws.cell(first, 6).value, f'=D{first}*E{first}')
                self.assertEqual(cached.cell(total, 6).value, 1221000)
                self.assertEqual(ws.cell(header, 2).fill.fgColor.rgb[-6:], PALETTES[palette if style == 'designer' else 'green'][0])
                self.assertEqual(ws.cell(header, 2).font.color.rgb[-6:], 'FFFFFF')
                self.assertFalse(ws.cell(first, 4).protection.locked)
                self.assertTrue(ws.cell(first, 6).protection.locked)
                note = next(c for row in ws for c in row if isinstance(c.value, str) and 'Эту подсказку' in c.value)
                self.assertFalse(note.protection.locked)
                self.assertTrue(ws.data_validations.dataValidation[0].showErrorMessage)
                self.assertEqual(ws.calculation if hasattr(ws, 'calculation') else load_workbook(p).calculation.calcMode, 'auto')

    def test_normalization(self):
        for invalid in [None, 'bad', {}, []]:
            self.assertEqual(normalize_excel_appearance(invalid, invalid), ('standard','green'))
        self.assertEqual(normalize_excel_appearance('standard','copper'), ('standard','green'))
        self.assertEqual(normalize_excel_appearance('designer','copper'), ('designer','copper'))

    def test_labels_are_literal_and_categories_escape_patterns(self):
        with tempfile.TemporaryDirectory() as td:
            p = str(Path(td)/'labels.xlsx')
            office_env()['build_excel_items'](p, '=1+1', 'expense_estimate', COLORS,
                [{'name':'=1+2','qty':0,'price':12.5,'category':'A*?"'}, {'name':'B','qty':1.5,'price':12.5,'category':'AB'}])
            ws = load_workbook(p).active
            self.assertEqual(ws['A1'].data_type, 's')
            self.assertEqual(ws['B5'].data_type, 's')
            formulas = [c.value for row in ws for c in row if c.data_type == 'f']
            self.assertTrue(any('=A~*~?""' in f for f in formulas))
            self.assertEqual(load_workbook(p,data_only=True).active['F7'].value, 18.75)


class ExcelStyleFlowTests(unittest.IsolatedAsyncioTestCase):
    async def test_miniapp_settings_reach_file_and_charge_after_send(self):
        for kind, style, palette, fail_send, protect in [(kind,style,palette,failed,protect)
                for kind in ['expense_estimate','family_budget','project_budget','price_list','calc_table']
                for style,palette,failed in [('standard','copper',False), ('designer','burgundy',False), ('designer','copper',True)] for protect in [False, True, 'false']]:
            state_data = {}
            async def update(**kwargs): state_data.update(kwargs)
            state = SimpleNamespace(update_data=update, get_data=AsyncMock(side_effect=lambda: state_data), clear=AsyncMock())
            user = {'generations':0,'history':[]}
            order = []
            async def send(file, **kwargs):
                ws = load_workbook(file).active
                self.assertEqual(ws.protection.sheet, protect is True)
                header = (6 if kind in ('price_list','calc_table') else 8) if style == 'designer' else 4
                expected = PALETTES[palette if style == 'designer' else 'green'][0]
                self.assertEqual(ws.cell(header,2).fill.fgColor.rgb[-6:], expected)
                order.append('send')
                if fail_send: raise RuntimeError('simulated Telegram failure')
            msg = SimpleNamespace(from_user=SimpleNamespace(id=-903020),answer=AsyncMock(),answer_document=send)
            async def run(fn,*args,**kwargs): return fn(*args,**kwargs)
            e = dict(ensure_answerable=lambda m:m,user_lang=lambda _: 'ru',get_user=lambda _:user,
                _clip_text=lambda v,n:(v or '')[:n],can_afford=lambda *a:True,CREDIT_COSTS={'excel':5},
                normalize_excel_appearance=normalize_excel_appearance,start_job=lambda _: (True,''),
                finish_job=Mock(),tr=lambda key,*a,**kw:key,pick_theme=lambda _:('default',COLORS),
                grok_json_lang_instruction=lambda _: '',content_gen_lang=lambda *a,**kw:'ru',
                EXCEL_ROW_SCHEMA={kind:{'fields':['name','qty','price','category'],'example':'{}'}},
                EXCEL_MIN_ROWS_AI={'expense_estimate':12},EXCEL_KIND_DESC={'expense_estimate':'Смета'},
                ask_grok=AsyncMock(return_value=json.dumps({'title':'Проверка','rows':[{'id':'s','name':'Путь','value':100,'unit':'м'}, {'id':'t','name':'Время','value':20,'unit':'с'}, {'id':'v','name':'Скорость','formula':'s/t','unit':'м/с'}] if kind=='calc_table' else ROWS if kind in ('expense_estimate','price_list') else
                    [{'name':'Доход','type':'income','amount':1000,'category':'Проект'},
                     {'name':'Расход','type':'expense','amount':300,'category':'Проект'}]})),extract_json=json.loads,
                run_heavy=run,build_excel_items=office_env()['build_excel_items'],datetime=datetime,
                safe_filename=lambda *a,**k:'estimate',FSInputFile=lambda path,**kw:path,
                spend_credits=lambda *a:order.append('charge'),note_success=Mock(),os=os,
                signature_line=AsyncMock(return_value=''),main_kb=lambda *a,**kw:None)
            env=load_functions(['excel_build','_handle_miniapp_payload'],e)
            await env['_handle_miniapp_payload'](msg,state,{'kind':kind,'content':'Данные',
                'mode':'user','excel_style':style,'excel_palette':palette,'excel_protect':protect},'gen_excel')
            self.assertEqual(order,['send'] if fail_send else ['send','charge'])
            self.assertEqual(user['generations'],0 if fail_send else 1)
            if fail_send: os.remove('/tmp/xl_-903020.xlsx')

    async def test_chat_buttons_select_style_and_palette(self):
        data = {}
        async def update(**kw):data.update(kw)
        state=SimpleNamespace(update_data=update,set_state=AsyncMock(),get_data=AsyncMock(side_effect=lambda:data))
        msg=SimpleNamespace(text='Дизайнерский',from_user=SimpleNamespace(id=1),answer=AsyncMock())
        form=SimpleNamespace(waiting_excel_mode='mode',waiting_excel_palette='palette')
        env=load_functions(['excel_style_choice','excel_palette_choice'],dict(user_lang=lambda _:'ru',
            EXCEL_STYLES=STYLE_LABELS,EXCEL_PALETTES=PALETTE_LABELS,Form=form,
            excel_appearance_kb=lambda *a:None,excel_mode_kb=lambda *a:None,tr=lambda *a:'Next'))
        await env['excel_style_choice'](msg,state)
        state.set_state.assert_awaited_with('palette')
        msg.text='Графит и медь'
        await env['excel_palette_choice'](msg,state)
        self.assertEqual(data,{'excel_style':'designer','excel_palette':'copper'})
        msg.text='Стандартный'
        await env['excel_style_choice'](msg,state)
        self.assertEqual(data,{'excel_style':'standard','excel_palette':'green'})
        state.set_state.assert_awaited_with('mode')
