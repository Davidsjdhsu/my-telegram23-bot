"""Editable family/project budgets with separate cash flows and live sign styling."""
import math
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.workbook.properties import CalcProperties
from excel_styles import PALETTES, normalize_excel_appearance

TYPE_MAP = {'доход':'Доход','доходы':'Доход','income':'Доход',
            'расход':'Расход','расходы':'Расход','expense':'Расход','expenses':'Расход'}
MONEY = '#,##0.00;(#,##0.00);"—"'


def build_budget(path, title, kind, rows, *, style='standard', palette='green', subtitle=None, number=None):
    if kind not in ('family_budget', 'project_budget') or not rows:
        raise ValueError('Нужен семейный бюджет или бюджет проекта с позициями')
    number = number or (lambda value, default=None: default if value in (None, '') else float(value))
    style, palette = normalize_excel_appearance(style, palette)
    designer = style == 'designer'
    dark, accent, soft, stripe = PALETTES[palette]
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    ws = wb.active
    ws.title = 'Бюджет'
    ws.sheet_properties.tabColor = dark
    ws.sheet_view.showGridLines = not designer
    for col, width in {'A':3 if designer else 6,'B':47,'C':15,'D':30,'E':32}.items():
        ws.column_dimensions[col].width = width
    edge = Side(style='thin', color='808080')
    grid = Border(left=edge, right=edge, top=edge, bottom=edge)
    cache, categories = {}, OrderedDict()
    start_col, title_row, header = (2,2,8) if designer else (1,1,4)
    result_label = 'Остаток' if kind == 'family_budget' else 'Результат бюджета'

    def text(row, col, value):
        c = ws.cell(row, col, str(value or ''))
        c.data_type = 's'
        return c

    def band(row, first, last, fill='FFFFFF', color='000000', size=11, bold=False):
        for col in range(first, last + 1):
            c = ws.cell(row, col)
            c.fill = PatternFill('solid', fgColor=fill)
            c.font = Font(name='Arial', size=size, color=color, bold=bold)
            c.alignment = Alignment(vertical='center', wrap_text=True)

    def sign_style(cell):
        ws.conditional_formatting.add(cell, FormulaRule(formula=[f'AND(ISNUMBER({cell}),{cell}<0)'],
            font=Font(color='9C0006',bold=True),fill=PatternFill('solid',fgColor='FCE4D6')))
        ws.conditional_formatting.add(cell, FormulaRule(formula=[f'AND(ISNUMBER({cell}),{cell}>=0)'],
            font=Font(color='217346',bold=True),fill=PatternFill('solid',fgColor='E7F1EA')))

    for r, value, size in [(title_row,title,17 if designer else 12), (title_row+1,subtitle or 'Суммы в рублях',10)]:
        ws.merge_cells(start_row=r,start_column=start_col,end_row=r,end_column=5)
        band(r,start_col,5,color=dark if r == title_row and designer else '000000',size=size,bold=r==title_row)
        text(r,start_col,value)
        ws.row_dimensions[r].height=max(32 if r==title_row and designer else 24, math.ceil(len(str(value))/85)*20)
    headers=['Статья','Тип','Сумма, ₽','Категория']
    if not designer: headers.insert(0,'№')
    band(header,start_col,5,dark,'FFFFFF',11,True)
    ws.row_dimensions[header].height=30 if designer else 25
    for col,label in enumerate(headers,start_col):
        c=text(header,col,label)
        c.alignment=Alignment(horizontal='center',vertical='center')
        c.border=grid if not designer else Border(right=Side(style='thin',color='FFFFFF'))
    first=header+1
    totals={'Доход':0.0,'Расход':0.0}
    complete=True
    for i,row in enumerate(rows):
        r=first+i
        row_type=TYPE_MAP.get(str(row.get('type') or '').strip().lower())
        if row_type is None: raise ValueError('Неизвестный тип строки бюджета')
        amount=number(row.get('amount'),None)
        if amount is not None and not math.isfinite(amount): raise ValueError('Некорректная сумма бюджета')
        complete=complete and amount is not None
        category=str(row.get('category') or '').strip() or 'Прочее'
        category_totals=categories.setdefault(category,{'Доход':0.0,'Расход':0.0})
        if amount is not None:
            totals[row_type]+=amount
            category_totals[row_type]+=amount
        band(r,start_col,5,stripe if designer and i%2 else 'FFFFFF')
        if not designer: ws.cell(r,1,i+1)
        text(r,2,row.get('name'));text(r,3,row_type);text(r,5,category)
        c=ws.cell(r,4,amount)
        c.number_format=MONEY
        c.fill=PatternFill('solid',fgColor=soft)
        c.font=Font(name='Arial',size=11,color='1767B0' if designer else '000000')
        c.alignment=Alignment(horizontal='right',vertical='center')
        ws.cell(r,3).alignment=Alignment(horizontal='center',vertical='center')
        ws.cell(r,5).alignment=Alignment(horizontal='left',vertical='center',wrap_text=True,indent=1)
        if not designer:
            for c in ws[r][:5]: c.border=grid
        ws.row_dimensions[r].height=max(30 if designer else 22,
            max(math.ceil(len(str(row.get('name') or ''))/43),math.ceil(len(category)/28))*16+6)
    if not all(math.isfinite(v) for v in totals.values()):
        raise ValueError('Слишком большая сумма бюджета')
    last=first+len(rows)-1
    amounts=f'D{first}:D{last}'
    types=f'C{first}:C{last}'
    cat_range=f'E{first}:E{last}'
    # Pasted invalid values can bypass validation. Do not silently undercount them.
    missing=f'COUNT({amounts})<>ROWS({amounts})'
    invalid=f'COUNTIF({types},"Доход")+COUNTIF({types},"Расход")<>ROWS({types})'
    def guarded(formula):
        return f'=IF(OR({missing},{invalid}),"Заполните данные",{formula})'
    inc_row=last+(2 if designer else 1)
    exp_row,res_row=inc_row+1,inc_row+2
    for r,label,formula,value in [
        (inc_row,'Итого доходы',guarded(f'SUMIF({types},"Доход",{amounts})'),totals['Доход']),
        (exp_row,'Итого расходы',guarded(f'SUMIF({types},"Расход",{amounts})'),totals['Расход']),
        (res_row,result_label,f'=IF(AND(ISNUMBER(D{inc_row}),ISNUMBER(D{exp_row})),D{inc_row}-D{exp_row},"Заполните данные")',totals['Доход']-totals['Расход'])]:
        band(r,start_col,5,soft if designer else dark,dark if designer else 'FFFFFF',11,True)
        ws.merge_cells(start_row=r,start_column=start_col,end_row=r,end_column=3)
        text(r,start_col,label)
        if designer: ws.merge_cells(f'D{r}:E{r}')
        c=ws.cell(r,4,formula);c.number_format=MONEY;c.alignment=Alignment(horizontal='right',vertical='center')
        ws.row_dimensions[r].height=32 if designer else 25
        cache[f'D{r}']=value if complete else 'Заполните данные'
    sign_style(f'D{res_row}')
    if designer:
        ws.merge_cells('B5:C5');ws.merge_cells('D5:E5')
        band(5,2,5,color=dark,size=17,bold=True)
        text(5,2,result_label)
        ws['B5'].font=Font(name='Arial',size=12,color=dark)
        ws['D5']=f'=D{res_row}';ws['D5'].number_format=MONEY
        ws['D5'].alignment=Alignment(horizontal='right',vertical='center')
        ws.row_dimensions[5].height=40
        cache['D5']=cache[f'D{res_row}'];sign_style('D5')
    summary= res_row+3
    for col,label in enumerate(['Категория','Доходы, ₽','Расходы, ₽','Разница, ₽'],2): text(summary,col,label)
    band(summary,2,5,dark,'FFFFFF',10,True)
    ws.row_dimensions[summary].height=30
    for offset,(category,values) in enumerate(categories.items(),1):
        r=summary+offset
        band(r,2,5,stripe if designer and offset%2==0 else 'FFFFFF')
        text(r,2,category)
        criterion=f'"="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(B{r},"~","~~"),"*","~*"),"?","~?")'
        for col,row_type in [(3,'Доход'),(4,'Расход')]:
            letter='C' if col==3 else 'D'
            ws.cell(r,col,guarded(f'SUMIFS({amounts},{cat_range},{criterion},{types},"{row_type}")'))
            cache[f'{letter}{r}']=values[row_type] if complete else 'Заполните данные'
        ws.cell(r,5,f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),C{r}-D{r},"Заполните данные")')
        cache[f'E{r}']=values['Доход']-values['Расход'] if complete else 'Заполните данные'
        for col in range(3,6):
            c=ws.cell(r,col);c.number_format=MONEY;c.alignment=Alignment(horizontal='right',vertical='center')
        ws.row_dimensions[r].height=max(24,math.ceil(len(category)/43)*16+6)
        sign_style(f'E{r}')
    # The type column also carries category income totals; reserve room for large sums.
    ws.column_dimensions['C'].width=27
    note_row=summary+len(categories)+2
    ws.merge_cells(start_row=note_row,start_column=start_col,end_row=note_row,end_column=5)
    text(note_row,start_col,'Суммы можно менять в цветных ячейках. Тип — «Доход» или «Расход». Пустая сумма требует заполнения. Эту подсказку можно удалить.')
    band(note_row,start_col,5,color='66726B',size=10)
    ws.row_dimensions[note_row].height=32
    nums=DataValidation(type='decimal',operator='between',formula1=-1e15,formula2=1e15,
        allow_blank=True,showErrorMessage=True,errorStyle='stop',error='Введите число.',errorTitle='Неверная сумма')
    ws.add_data_validation(nums);nums.add(amounts)
    choices=DataValidation(type='list',formula1='"Доход,Расход"',allow_blank=False,showErrorMessage=True,
        errorStyle='stop',error='Выберите «Доход» или «Расход».')
    ws.add_data_validation(choices);choices.add(types)
    ws.freeze_panes=f'C{first}';ws.auto_filter.ref=f'{"B" if designer else "A"}{header}:E{last}'
    chart=BarChart();chart.title='Доходы и расходы';chart.y_axis.title='₽';chart.legend=None
    chart.add_data(Reference(ws,min_col=4,min_row=inc_row,max_row=exp_row),titles_from_data=False)
    chart.set_categories(Reference(ws,min_col=start_col,min_row=inc_row,max_row=exp_row))
    chart.series[0].graphicalProperties.solidFill=dark
    chart.height=9;chart.width=17;ws.add_chart(chart,'G3')
    for row in ws:
        for c in row:c.protection=Protection(locked=c.data_type=='f')
    ws.protection.sheet=True
    ws.protection.formatCells=ws.protection.formatColumns=ws.protection.formatRows=False
    ws.protection.autoFilter=False
    ws.protection.selectLockedCells=ws.protection.selectUnlockedCells=False
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='portrait' if designer else 'landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    # Keep short budgets together when printing, without shrinking long ledgers.
    if len(rows) <= 10 and len(categories) <= 8:
        ws.page_setup.fitToHeight=1
    ws.page_margins=PageMargins(left=0.3,right=0.3,top=0.3,bottom=0.3,header=0.1,footer=0.1)
    ws.print_title_rows=f'{header}:{header}';ws.print_area=f'A1:E{note_row}'
    wb.save(path)
    return cache
