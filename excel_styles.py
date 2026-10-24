"""Reviewed estimate layouts. Other workbook kinds keep their existing builders."""
import math
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

PALETTES = {
    'green': ('217346', '33A064', 'E7F1EA', 'F4F8F5'),
    'copper': ('343B40', 'AE6E48', 'F2E5DA', 'FAF6F2'),
    'emerald': ('174C43', 'B59A61', 'EDE6D6', 'F5F7F2'),
    'lilac': ('292E55', '8170B3', 'ECE8F5', 'F6F5FA'),
    'burgundy': ('6A2E43', 'B9828F', 'F1E1E5', 'FCF7F3'),
}
STYLE_LABELS = {'standard': 'Стандартный', 'designer': 'Дизайнерский'}
PALETTE_LABELS = dict(zip(PALETTES, ['Классический зелёный', 'Графит и медь',
    'Изумруд и песочный', 'Чернильный и сирень', 'Бордовый и кремовый']))


def normalize_excel_appearance(style=None, palette=None):
    style = style if isinstance(style, str) and style in STYLE_LABELS else 'standard'
    return style, palette if style == 'designer' and isinstance(palette, str) and palette in PALETTES else 'green'


def build_estimate(path, title, rows, *, style='standard', palette='green', subtitle=None, number=None, protect=False, calculation=None):
    """Save a real XLSX and return caches for the existing preview-cache writer."""
    from excel_calculations import normalize_calculation, calculate, formula, description
    calculation = normalize_calculation(calculation)
    custom = calculation['mode'] == 'custom'
    adjusted = calculation['mode'] != 'cost'
    if not rows:
        raise ValueError('В смете нет позиций')
    number = number or (lambda value, default=0: default if value is None else float(value))
    style, palette = normalize_excel_appearance(style, palette)
    designer = style == 'designer'
    dark, accent, soft, stripe = PALETTES[palette]
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    ws = wb.active
    ws.title = 'Смета'
    ws.sheet_view.showGridLines = not designer
    ws.sheet_properties.tabColor = dark
    cache, categories = {}, OrderedDict()
    any_categories = any(str(row.get('category') or '').strip() for row in rows)
    widths = {'A': 3 if designer else 6, 'B': 39, 'C': 12, 'D': 12, 'E': 22, 'F': 26, 'G': 22}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    thin = Side(style='thin', color='808080')
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)

    def text(row, col, value):
        cell = ws.cell(row, col, str(value or ''))
        cell.data_type = 's'  # User labels must never become Excel formulas.
        return cell

    def band(row, first, last, fill, color='000000', size=11, bold=False):
        for col in range(first, last + 1):
            c = ws.cell(row, col)
            c.font = Font(name='Arial', size=size, color=color, bold=bold)
            c.fill = PatternFill('solid', fgColor=fill)
            c.alignment = Alignment(vertical='center', wrap_text=True)

    title_row, header = (2, 8) if designer else (1, 4)
    first_col = 2 if designer else 1
    ws.merge_cells(start_row=title_row, start_column=first_col, end_row=title_row, end_column=7)
    text(title_row, first_col, title)
    band(title_row, first_col, 7, 'FFFFFF', dark if designer else '000000', 17 if designer else 12, True)
    ws.row_dimensions[title_row].height = max(32 if designer else 24, math.ceil(len(str(title)) / 90) * 22)
    ws.merge_cells(start_row=title_row + 1, start_column=first_col, end_row=title_row + 1, end_column=7)
    text(title_row + 1, first_col, subtitle or 'Суммы в рублях')
    band(title_row + 1, first_col, 7, 'FFFFFF', '66726B', 10)
    ws.row_dimensions[title_row + 1].height = max(22, math.ceil(len(str(subtitle or '')) / 110) * 16)
    headers = ['Позиция', 'Ед.', 'Кол-во', 'Цена, ₽', 'Сумма, ₽', 'Раздел']
    if custom: headers[4] = 'Результат'
    if not designer:
        headers.insert(0, '№')
    for col, label in enumerate(headers, first_col):
        text(header, col, label)
    band(header, first_col, 7, dark, 'FFFFFF', 11, True)
    ws.row_dimensions[header].height = 30 if designer else 25
    for c in ws[header][first_col - 1:7]:
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = grid if not designer else Border(right=Side(style='thin', color='FFFFFF'))
    first = header + 1
    total = 0.0
    for i, row in enumerate(rows):
        r = first + i
        qty, price = number(row.get('qty'), 1), number(row.get('price'), 0)
        if not all(math.isfinite(v) for v in (qty, price)):
            raise ValueError('Число в смете не является конечным')
        category = str(row.get('category') or '').strip() or ('Прочее' if any_categories else '')
        name, unit = str(row.get('name') or ''), str(row.get('unit') or '')
        band(r, first_col, 7, stripe if designer and i % 2 else 'FFFFFF')
        if not designer:
            ws.cell(r, 1, i + 1)
        text(r, 2, name); text(r, 3, unit); text(r, 7, category)
        ws.cell(r, 4, qty); ws.cell(r, 5, price)
        ws.cell(r, 6, formula(r, calculation))
        value = calculate({'qty': qty, 'price': price}, calculation)
        if not math.isfinite(value):
            raise ValueError('Слишком большая сумма в смете')
        cache[f'F{r}'] = value
        total += value
        if category:
            categories[category] = categories.get(category, 0) + value
        for c in ws[r][first_col - 1:7]:
            if not designer:
                c.border = grid
            if c.column in (4, 5):
                c.fill = PatternFill('solid', fgColor=soft)
                c.font = Font(name='Arial', size=11, color='1767B0' if designer else '000000')
            if c.column in (4, 5, 6):
                c.alignment = Alignment(horizontal='right', vertical='center')
                c.number_format = '#,##0.###' if c.column == 4 else '#,##0.00;(#,##0.00);"—"'
            elif c.column in (1, 3):
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            elif designer and c.column == 7:
                c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
        ws.row_dimensions[r].height = max(30 if designer else 22,
            max(math.ceil(len(name) / 36), math.ceil(len(unit) / 10), math.ceil(len(category) / 20)) * 16 + 6)
    last = first + len(rows) - 1
    total_row = last + (2 if designer else 1)
    band(total_row, first_col, 7, soft if designer else dark, dark if designer else 'FFFFFF', 12 if designer else 11, True)
    text(total_row, 2, 'Сумма результатов' if custom else 'Итого по смете')
    guard = f'COUNT(F{first}:F{last})={len(rows)}'
    ws.cell(total_row, 6, f'=IF({guard},SUM(F{first}:F{last}),"")' if custom else f'=SUM(F{first}:F{last})')
    ws.cell(total_row, 6).number_format = '#,##0.00" ₽"'
    ws.cell(total_row, 6).alignment = Alignment(horizontal='right', vertical='center')
    cache[f'F{total_row}'] = total
    ws.row_dimensions[total_row].height = 35 if designer else 25
    if designer:
        ws.merge_cells(f'F{total_row}:G{total_row}')
        text(5, 2, 'Сумма результатов' if custom else 'Общая стоимость')
        ws.merge_cells('F5:G5')
        ws.cell(5, 6, f'=IF(ISNUMBER(F{total_row}),F{total_row},"")' if custom else f'=F{total_row}')
        cache['F5'] = total
        band(5, 2, 7, 'FFFFFF', dark, 17, True)
        ws['B5'].font = Font(name='Arial', size=12, color=dark)
        ws['F5'].number_format = '#,##0.00" ₽"'
        ws['F5'].alignment = Alignment(horizontal='right', vertical='center')
        ws.row_dimensions[5].height = 40
    else:
        for c in ws[total_row][:7]:
            c.border = grid
    summary_row = total_row + 2
    for offset, (category, value) in enumerate(categories.items()):
        r = summary_row + offset
        text(r, 2, category)
        # SUMIF interprets *, ? and ~ as patterns; categories are literal labels.
        criterion = '=' + category.replace('~', '~~').replace('*', '~*').replace('?', '~?')
        criterion = criterion.replace('"', '""')
        expression = f'SUMIF(G{first}:G{last},"{criterion}",F{first}:F{last})'
        ws.cell(r, 6, f'=IF({guard},{expression},"")' if custom else '=' + expression)
        cache[f'F{r}'] = value
        band(r, 2, 7, 'FFFFFF', dark if designer else '000000', 11, designer)
        ws['F' + str(r)].number_format = '#,##0.00" ₽"'
        ws['F' + str(r)].alignment = Alignment(horizontal='right', vertical='center')
        ws.row_dimensions[r].height = max(23, math.ceil(len(category) / 36) * 16 + 6)
        if designer:
            ws.merge_cells(f'F{r}:G{r}')
    note_row = summary_row + len(categories) + 1
    ws.merge_cells(start_row=note_row, start_column=first_col, end_row=note_row, end_column=7)
    text(note_row, first_col, 'Количество и цену можно менять в цветных ячейках. Эту подсказку можно удалить.')
    band(note_row, first_col, 7, 'FFFFFF', '66726B', 10)
    ws.row_dimensions[note_row].height = 26
    validation = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0,
        allow_blank=True, showErrorMessage=True, errorTitle='Нужно число', error='Введите число от 0.', errorStyle='stop')
    ws.add_data_validation(validation)
    validation.add(f'D{first}:E{last}')
    ws.freeze_panes = f'D{first}'
    ws.auto_filter.ref = f'{"B" if designer else "A"}{header}:G{last}'
    for row in ws:
        for cell in row:
            cell.protection = Protection(locked=cell.data_type == 'f')
    ws.protection.sheet = True
    ws.protection.formatCells = ws.protection.formatColumns = ws.protection.formatRows = False
    ws.protection.autoFilter = False
    ws.protection.selectLockedCells = ws.protection.selectUnlockedCells = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait' if designer else 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = f'{header}:{header}'
    ws.print_area = f'A1:G{note_row}'
    if adjusted:
        # Keep the approved main-sheet geometry; editable drivers get their own sheet.
        settings = wb.create_sheet('Расчёт')
        settings.column_dimensions['A'].width = 34
        settings.column_dimensions['B'].width = 75
        if calculation['mode'] in ('discount', 'markup'):
            sign, label = ('−', 'Скидка') if calculation['mode']=='discount' else ('+', 'Наценка')
            expression_label = f'Количество × Цена × (1 {sign} {label} / 100)'
        else:
            from excel_calculations import LABELS, SYMBOLS
            left = 'Первое число' if calculation['left']=='number' else LABELS[calculation['left']]
            right = 'Второе число' if calculation['right']=='number' else LABELS[calculation['right']]
            expression_label = f'{left} {SYMBOLS[calculation["op"]]} {right}'
        settings.append(['Расчёт для всех строк', expression_label])
        settings.append(['Результат', 'Число без единицы; проверьте смысл выбранных величин.' if custom else 'Сумма в рублях'])
        if calculation['mode'] in ('discount', 'markup'):
            settings.cell(3, 1, 'Скидка, %' if calculation['mode'] == 'discount' else 'Наценка, %')
            settings.cell(3, 2, calculation['percent'])
            check = DataValidation(type='decimal', operator='between', formula1=0,
                formula2=100 if calculation['mode']=='discount' else 1000, showErrorMessage=True,
                error='Проверьте допустимый процент.', allow_blank=False)
            settings.add_data_validation(check); check.add('B3')
        for side, r in [('left', 4), ('right', 5)]:
            if calculation.get(side) == 'number':
                settings.cell(r, 1, 'Первое число' if side == 'left' else 'Второе число')
                settings.cell(r, 2, calculation[side + '_value'])
        settings.cell(7, 1, 'Как пользоваться')
        settings.cell(7, 2, 'Меняйте числа в цветных ячейках. Расчёт применяется ко всем строкам сметы. Пустой результат при делении означает нулевой делитель; исправьте его, чтобы появился итог.')
        for row in settings:
            for cell in row:
                cell.font = Font(name='Arial', size=11, color='000000')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                if cell.row in (3, 4, 5) and cell.column == 2 and cell.value is not None:
                    cell.fill = PatternFill('solid', fgColor=soft)
        settings.row_dimensions[1].height = 36
        settings.row_dimensions[2].height = 36
        settings.row_dimensions[7].height = 72
        settings.print_area = 'A1:B7'
        settings.sheet_properties.pageSetUpPr.fitToPage = True
        settings.page_setup.fitToWidth = 1
        settings.page_setup.fitToHeight = 1
        # Keep the expression beside the main table as well.
        ws.cell(note_row, first_col, 'Расчёт: ' + expression_label + '. Настройки — на листе «Расчёт».')
        ws.row_dimensions[note_row].height = 42
    if custom:
        for row in ws:
            for cell in row:
                if cell.column == 6: cell.number_format = '#,##0.###;(#,##0.###);"0"'
    apply_formula_protection(wb, protect)
    wb.save(path)
    return cache


def apply_formula_protection(workbook, enabled=False):
    """Opt-in accidental-edit protection; keep explicit input overrides editable."""
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                cell.protection = Protection(
                    locked=cell.data_type == 'f' and cell.protection.locked,
                    hidden=False,
                )
        sheet.protection.sheet = enabled is True
        sheet.protection.selectLockedCells = sheet.protection.selectUnlockedCells = False
        sheet.protection.formatCells = sheet.protection.formatColumns = sheet.protection.formatRows = False
        sheet.protection.autoFilter = sheet.protection.sort = False
