"""Coursework calculations compiled from a small, non-executable expression language.

Model output never becomes raw Excel code. Only arithmetic, named row references,
and the functions below are accepted. This validates calculation mechanics, not
the scientific appropriateness of a model-selected equation or unit.
"""
import ast
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.workbook.properties import CalcProperties
from excel_styles import apply_formula_protection, PALETTES, normalize_excel_appearance


class StudyValidationError(ValueError):
    """Invalid structured calculation; do not deliver or charge for it."""


MISSING = 'Заполните данные'
INVALID = 'Проверьте расчёт'
FUNCTIONS = {'SUM', 'AVERAGE', 'MIN', 'MAX', 'SQRT', 'ABS'}
CALC_INSTRUCTIONS = """
Учебная таблица должна решать указанную задачу, а не автоматически считать доли.
Для calc_table нет минимального числа строк: не добавляй лишние данные ради объёма.
Сохраняй числа пользователя точно, включая явные нули и отрицательные числа.
Неизвестное исходное значение: value:null, formula:""; не заменяй его нулём.
В режиме своих данных не придумывай результаты экспериментов и исходные числа.
В режиме примера можно дать условные исходные числа, но обязательно отмечай пример.
Каждая строка: id (уникальное короткое латинское обозначение, например mass или t1),
name (полное название), value (число или null), unit (единица), formula, note.
У исходных данных formula:""; у расчётных строк value:null и formula с выражением.
Формулы используют id других строк и только + - * / **, скобки, SUM, AVERAGE,
MIN, MAX, SQRT, ABS. У ** показатель — целое число от -12 до 12.
Примеры: distance/time; AVERAGE(t1,t2,t3); mass*speed**2/2; count/total*100.
Нельзя использовать адреса Excel, диапазоны, знак = в начале или иные функции.
Не записывай вычисленный ответ как исходное число: все результаты — формулы.
Не создавай циклы. Указывай единицы измерения; преобразование единиц оформляй
отдельной формулой (например cm/100 в метры), не меняя исходные числа.
Для процентов используй результат от 0 до 100 и unit:"%". В note кратко поясняй
метод расчёта и допущения. Не складывай разные величины/единицы; не добавляй общий
итог или доли, если они не нужны. Доли опроса допустимы только для совместимых
категорий; повторные измерения можно усреднять. Не выдавай условные данные за опыт.
Если задача выходит за эти операции, не подменяй её простым расчётом: верни
{"unsupported":true,"title":"Расчёт требует уточнения","rows":[]}.
"""


def _finite(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise StudyValidationError('Значение должно быть числом')
    if not math.isfinite(value) or abs(value) > 1e100:
        raise StudyValidationError('Число вне допустимого диапазона')
    return float(value)


def _expression(source, ids):
    if not isinstance(source, str) or len(source) > 500:
        raise StudyValidationError('Слишком длинная формула')
    try:
        node = ast.parse(source.strip(), mode='eval').body
    except (SyntaxError, ValueError, RecursionError):
        raise StudyValidationError('Некорректная формула') from None
    if sum(1 for _ in ast.walk(node)) > 128:
        raise StudyValidationError('Слишком сложная формула')
    refs = set()

    def check(n):
        if isinstance(n, ast.Constant):
            _finite(n.value)
        elif isinstance(n, ast.Name):
            if n.id not in ids:
                raise StudyValidationError('Неизвестное обозначение в формуле')
            refs.add(n.id)
        elif isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.UAdd, ast.USub)):
            check(n.operand)
        elif isinstance(n, ast.BinOp) and isinstance(n.op, (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow)):
            check(n.left)
            check(n.right)
            if isinstance(n.op, ast.Pow):
                exponent = n.right
                sign = 1
                if isinstance(exponent, ast.UnaryOp) and isinstance(exponent.op, (ast.USub, ast.UAdd)):
                    sign = -1 if isinstance(exponent.op, ast.USub) else 1
                    exponent = exponent.operand
                if not isinstance(exponent, ast.Constant) or type(exponent.value) is not int or not -12 <= sign*exponent.value <= 12:
                    raise StudyValidationError('Степень должна быть целой от -12 до 12')
        elif isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id in FUNCTIONS:
            if n.keywords or not 1 <= len(n.args) <= 64 or (n.func.id in ('SQRT', 'ABS') and len(n.args) != 1):
                raise StudyValidationError('Некорректные аргументы функции')
            for arg in n.args:
                check(arg)
        else:
            raise StudyValidationError('Недопустимая операция в формуле')
    check(node)
    return node, refs


def prepare_rows(rows, number=None):
    if not isinstance(rows, list) or not 1 <= len(rows) <= 200:
        raise StudyValidationError('Нужны от 1 до 200 показателей')
    result, ids = [], set()
    for i, row in enumerate(rows, 1):
        if not isinstance(row, dict):
            raise StudyValidationError('Некорректная строка')
        ident = row.get('id') or f'x{i}'
        if not isinstance(ident, str) or not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,19}', ident) or ident.upper() in FUNCTIONS or ident in ids:
            raise StudyValidationError('Некорректные или повторные обозначения')
        ids.add(ident)
        formula = row.get('formula') or ''
        value = row.get('value')
        if formula:
            value = None  # A model-provided answer never replaces the calculation.
        elif value not in (None, ''):
            if isinstance(value, bool):
                raise StudyValidationError('Значение должно быть числом')
            try:
                value = number(value, None) if number else float(str(value).replace(',', '.'))
            except (TypeError, ValueError, OverflowError):
                raise StudyValidationError('Значение должно быть числом') from None
            value = _finite(value)
        else:
            value = None
        result.append(dict(id=ident, name=str(row.get('name') or ident), unit=str(row.get('unit') or ''),
                           note=str(row.get('note') or ''), value=value, formula=formula))
    for row in result:
        row['node'], row['refs'] = _expression(row['formula'], ids) if row['formula'] else (None, set())
    by_id = {row['id']: row for row in result}
    done, active = set(), set()

    def visit(ident):
        if ident in active:
            raise StudyValidationError('Циклическая зависимость')
        if ident in done:
            return
        active.add(ident)
        for ref in by_id[ident]['refs']:
            visit(ref)
        active.remove(ident)
        done.add(ident)
    for ident in ids:
        visit(ident)
    return result


def _compile(node, addresses):
    if isinstance(node, ast.Constant):
        return str(node.value)
    if isinstance(node, ast.Name):
        return addresses[node.id]
    if isinstance(node, ast.UnaryOp):
        return '(' + ('-' if isinstance(node.op, ast.USub) else '+') + _compile(node.operand, addresses) + ')'
    if isinstance(node, ast.BinOp):
        op = {ast.Add: '+', ast.Sub: '-', ast.Mult: '*', ast.Div: '/', ast.Pow: '^'}[type(node.op)]
        return '(' + _compile(node.left, addresses) + op + _compile(node.right, addresses) + ')'
    return node.func.id + '(' + ','.join(_compile(arg, addresses) for arg in node.args) + ')'


def _evaluate(node, values):
    if isinstance(node, ast.Constant):
        return float(node.value)
    if isinstance(node, ast.Name):
        return values[node.id]
    if isinstance(node, ast.UnaryOp):
        return (-1 if isinstance(node.op, ast.USub) else 1)*_evaluate(node.operand, values)
    if isinstance(node, ast.BinOp):
        a, b = _evaluate(node.left, values), _evaluate(node.right, values)
        if isinstance(node.op, ast.Add): return a+b
        if isinstance(node.op, ast.Sub): return a-b
        if isinstance(node.op, ast.Mult): return a*b
        if isinstance(node.op, ast.Div): return a/b
        return a**b
    args = [_evaluate(arg, values) for arg in node.args]
    return {'SUM': sum, 'AVERAGE': lambda a: sum(a)/len(a), 'MIN': min, 'MAX': max,
            'SQRT': lambda a: math.sqrt(a[0]), 'ABS': lambda a: abs(a[0])}[node.func.id](args)


def calculated_values(rows):
    by_id = {row['id']: row for row in rows}
    values = {}

    def resolve(ident):
        if ident in values: return values[ident]
        row = by_id[ident]
        if row['node'] is None:
            value = MISSING if row['value'] is None else row['value']
        else:
            deps = [resolve(ref) for ref in row['refs']]
            if INVALID in deps:
                value = INVALID
            elif MISSING in deps:
                value = MISSING
            else:
                try:
                    value = _evaluate(row['node'], values)
                    if not isinstance(value, (float, int)) or not math.isfinite(value) or abs(value) > 1e100:
                        value = INVALID
                except (ArithmeticError, ValueError):
                    value = INVALID
        values[ident] = value
        return value
    for row in rows: resolve(row['id'])
    return values


def build_study(path, title, rows, *, style='standard', palette='green', subtitle=None, number=None, protect=False):
    rows = prepare_rows(rows, number)
    values = calculated_values(rows)
    style, palette = normalize_excel_appearance(style, palette)
    designer = style == 'designer'
    dark, accent, soft, stripe = PALETTES[palette]
    wb = Workbook()
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    ws = wb.active
    ws.title = 'Расчёт'
    ws.sheet_view.showGridLines = not designer
    ws.sheet_properties.tabColor = dark
    start, title_row, header = (2, 2, 6) if designer else (1, 1, 4)
    widths = [12, 32, 27, 12, 44]
    labels = ['Код', 'Показатель', 'Значение', 'Ед.', 'Расчёт / пояснение']
    end, first, last = start+4, header+1, header+len(rows)
    val_col = get_column_letter(start+2)
    addresses = {row['id']: f'${val_col}${r}' for r, row in enumerate(rows, first)}
    border = Border(**{k: Side(style='thin', color='A2B3A9') for k in ('left','right','top','bottom')})
    cache = {}

    def text(r, c, value):
        cell = ws.cell(r, c, str(value))
        cell.data_type = 's'
        return cell

    if designer: ws.column_dimensions['A'].width = 3
    for r, value, size in [(title_row, title, 17 if designer else 13),
                            (title_row+1, subtitle or 'Исходные данные и результаты расчёта', 10)]:
        ws.merge_cells(start_row=r, start_column=start, end_row=r, end_column=end)
        c = text(r, start, value)
        c.font = Font(name='Arial', size=size, bold=r==title_row, color=dark if r==title_row else '66726B')
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = max(30, math.ceil(len(str(value))/95)*18)
    for col, (label, width) in enumerate(zip(labels, widths), start):
        ws.column_dimensions[get_column_letter(col)].width = width
        c = text(header, col, label)
        c.fill = PatternFill('solid', fgColor=dark)
        c.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border
    ws.row_dimensions[header].height = 30
    for r, row in enumerate(rows, first):
        method = (row['id'] + ' = ' + row['formula']) if row['node'] else ('Исходные данные' if row['value'] is not None else 'Укажите исходное значение')
        if row['note']: method += '\n' + row['note']
        row_values = [row['id'], row['name'], '', row['unit'], method]
        height = 30 if designer else 26
        for cnum, (value, width) in enumerate(zip(row_values, widths), start):
            c = text(r, cnum, value)
            c.font = Font(name='Arial', size=11, color='222222')
            c.fill = PatternFill('solid', fgColor=stripe if designer and (r-first)%2 else 'FFFFFF')
            c.border = border if not designer else Border(bottom=Side(style='hair', color='D9E2DC'))
            c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
            # Formula explanation must remain consistent with the locked formula.
            c.protection = Protection(locked=True)
            height = max(height, sum(max(1, math.ceil(len(line)/(width-3))) for line in str(value).split('\n'))*16+8)
        cell = ws.cell(r, start+2)
        # General preserves tiny scientific values instead of displaying a rounded zero.
        cell.number_format = 'General'
        cell.alignment = Alignment(horizontal='right', vertical='center', shrink_to_fit=True)
        if row['node']:
            cell.number_format = '#,##0.00####'
            expr = _compile(row['node'], addresses)
            # Validate all dependencies, even when SUM/AVERAGE would silently skip text.
            refs = [addresses[ref] for ref in sorted(row['refs'])]
            result = f'IFERROR(IF(ABS({expr})>1E100,"{INVALID}",{expr}),"{INVALID}")'
            if refs:
                invalid = ','.join(f'{ref}="{INVALID}"' for ref in refs)
                valid = ','.join(f'ISNUMBER({ref})' for ref in refs)
                result = f'IF(OR({invalid}),"{INVALID}",IF(AND({valid}),{result},"{MISSING}"))'
            cell.value = '=' + result
            cell.font = Font(name='Arial', size=11, bold=True, color=dark)
            cache[cell.coordinate] = values[row['id']]
        else:
            cell.value = row['value']
            cell.data_type = 'n'
            cell.fill = PatternFill('solid', fgColor=soft)
            cell.font = Font(name='Arial', size=11, color='1767B0' if designer else '000000')
            cell.protection = Protection(locked=False)
            dv = DataValidation(type='decimal', operator='between', formula1=-1e100, formula2=1e100,
                                allow_blank=True, showErrorMessage=True, errorStyle='stop', error='Введите число.')
            ws.add_data_validation(dv)
            dv.add(cell)
            ws.conditional_formatting.add(cell.coordinate, FormulaRule(formula=[f'NOT(ISNUMBER({cell.coordinate}))'],
                fill=PatternFill('solid',fgColor='FFF2CC')))
        ws.row_dimensions[r].height = min(height, 409)
    ws.conditional_formatting.add(f'{val_col}{first}:{val_col}{last}', FormulaRule(
        formula=[f'ISTEXT({val_col}{first})'], font=Font(color='9C0006', bold=True), fill=PatternFill('solid',fgColor='FCE4D6')))
    ref = f'{val_col}{first}'
    ws.conditional_formatting.add(f'{ref}:{val_col}{last}', Rule(type='expression',
        formula=[f'AND(ISNUMBER({ref}),{ref}<>0,OR(ABS({ref})<0.000001,ABS({ref})>=1E12))'],
        dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=164, formatCode='0.######E+00'))))
    note_row = last+2
    ws.merge_cells(start_row=note_row, start_column=start, end_row=note_row, end_column=end)
    c = text(note_row, start, 'Светлые цветные ячейки — исходные числа: меняйте их, результаты пересчитаются. '
             'Коды связывают показатели с формулами. Пустое значение не равно нулю. Единицы и метод указаны рядом. ' +
             ('Формулы защищены от случайной замены; защиту листа можно снять без пароля. ' if protect else 'Защита выключена: формулы тоже можно менять. ') + 'Эту подсказку можно удалить.')
    c.font = Font(name='Arial', size=10, color='66726B')
    c.alignment = Alignment(wrap_text=True, vertical='center')
    c.protection = Protection(locked=False)
    ws.row_dimensions[note_row].height = 60
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.freeze_panes = f'{val_col}{first}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=.3, right=.3, top=.3, bottom=.3, header=.1, footer=.1)
    ws.print_title_rows = f'{header}:{header}'
    ws.print_area = f'A1:{get_column_letter(end)}{note_row}'
    apply_formula_protection(wb, protect)
    wb.save(path)
    return cache
