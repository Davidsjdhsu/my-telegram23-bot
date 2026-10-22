"""Portable native PPTX export for the private manual constructor (no AI calls)."""
import base64
import io
import math
import re

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

SCALE = 9525  # one CSS pixel at 96 dpi
FONTS = {'Arial', 'Georgia', 'Trebuchet MS', 'Verdana', 'Times New Roman'}


def number(v, low, high):
    if isinstance(v, bool) or not isinstance(v, (int,float)) or not math.isfinite(v) or not low <= v <= high:
        raise ValueError('number')
    return v


def color(v, none=False):
    if none and v == 'none':
        return None
    if not isinstance(v,str) or not re.fullmatch(r'#[0-9a-fA-F]{6}',v):
        raise ValueError('color')
    return v[1:]


def rgb(v):
    return RGBColor.from_string(color(v))


def alpha(element, opacity):
    for c in element.iter():
        if c.tag.endswith('}srgbClr'):
            a = OxmlElement('a:alpha')
            a.set('val',str(round(opacity*100000)))
            c.append(a)


def paint(fill, value, opacity=1):
    if value == 'none':
        fill.background()
    else:
        fill.solid()
        fill.fore_color.rgb = rgb(value)
        alpha(fill._xPr, opacity)


def text_shape(slide, e, x, y, w, h, diagram=False):
    shape = slide.shapes.add_textbox(x,y,w,h)
    shape.rotation = number(e.get('rotation',0),-360,360)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=False
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE if diagram else MSO_ANCHOR.TOP
    lines = e.get('lines',str(e.get('text','')).split('\n'))
    if not isinstance(lines,list) or not 1<=len(lines)<=100 or any(not isinstance(v,str) or len(v)>3000 for v in lines):
        raise ValueError('lines')
    font=e.get('font','Arial')
    if font not in FONTS:
        raise ValueError('font')
    size=number(e.get('size',24),1,200)*.75
    opacity=number(e.get('opacity',1),0,1)
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line
        p.space_before=p.space_after=Pt(0)
        p.line_spacing=1.0
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(e.get('align','center' if diagram else 'left'),PP_ALIGN.LEFT)
        p.font.name=font
        p.font.size=Pt(size)
        p.font.bold=e.get('bold') is True
        p.font.italic=e.get('italic') is True
        p.font.color.rgb=rgb(e.get('color','#ffffff'))
        alpha(p._p,opacity)
    return shape


def position(e, relative=False):
    return [round(number(e.get(k), *bounds)*SCALE) for k,bounds in (
        ('x',(-2400,2400) if relative else (-1200,2400)),
        ('y',(-1350,1350)),('w',(0.1,2400)),('h',(0.1,1350)))]


def build_pptx(body):
    if not isinstance(body,dict) or body.get('version')!=1:
        raise ValueError('version')
    slides=body.get('slides')
    if not isinstance(slides,list) or not 1<=len(slides)<=20:
        raise ValueError('slides')
    p=Presentation()
    p.slide_width=1200*SCALE
    p.slide_height=675*SCALE
    total_pixels=0
    for s in slides:
        slide=p.slides.add_slide(p.slide_layouts[6])
        b=s['background']
        if b.get('type')=='gradient':
            color(b['color']);color(b['second'])
            fill=slide.background.fill
            fill.gradient()
            for stop,c in zip(fill.gradient_stops,[b['color'],b['second']]):
                stop.color.rgb=rgb(c)
            # CSS/SVG angle runs clockwise; DrawingML angle also runs clockwise.
            grad=fill._xPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}gradFill')
            lin=grad.find('{http://schemas.openxmlformats.org/drawingml/2006/main}lin')
            lin.set('ang',str(round(number(b.get('angle',0),-360,360)%360*60000)))
        elif b.get('type')=='solid':
            paint(slide.background.fill,b['color'])
        else:
            raise ValueError('background')
        elements=s.get('elements')
        if not isinstance(elements,list) or len(elements)>50:
            raise ValueError('elements')
        for e in elements:
            x,y,w,h=position(e)
            opacity=number(e.get('opacity',1),0,1)
            rotation=number(e.get('rotation',0),-360,360)
            if e['type']=='text':
                text_shape(slide,e,x,y,w,h)
            elif e['type']=='diagram':
                # Same client geometry used in the visual editor. Bounds, primitive
                # types and text are validated independently; no SVG/HTML execution.
                parts=e.get('parts')
                if not isinstance(parts,list) or not 1<=len(parts)<=100:
                    raise ValueError('parts')
                for part in parts:
                    px,py,pw,ph=position(part,True)
                    geom=part['geometry']
                    if geom=='textbox':
                        text_shape(slide,{**part,'opacity':opacity},x+px,y+py,pw,ph,True)
                    else:
                        shapes={'rect':MSO_SHAPE.RECTANGLE,'roundRect':MSO_SHAPE.ROUNDED_RECTANGLE,'ellipse':MSO_SHAPE.OVAL,'rightArrow':MSO_SHAPE.RIGHT_ARROW}
                        if geom not in shapes:
                            raise ValueError('shape')
                        obj=slide.shapes.add_shape(shapes[geom],x+px,y+py,pw,ph)
                        obj.rotation=number(part.get('rotation',0),-360,360)
                        paint(obj.fill,part['fill'],opacity)
                        paint(obj.line.fill,part.get('stroke','none'),opacity)
                        obj.line.width=Pt(1.125)
                        if geom=='roundRect':
                            obj.adjustments[0]=min(.5,12*SCALE/min(pw,ph))
            elif e['type'] in ('photo','pattern'):
                raster=e.get('raster','')
                prefix='data:image/png;base64,'
                if not isinstance(raster,str) or not raster.startswith(prefix) or len(raster)>8_000_000:
                    raise ValueError('image')
                raw=base64.b64decode(raster[len(prefix):],validate=True)
                with Image.open(io.BytesIO(raw)) as im:
                    pixels=im.width*im.height
                    total_pixels+=pixels
                    if im.format!='PNG' or pixels>10_000_000 or total_pixels>40_000_000:
                        raise ValueError('image dimensions')
                    im.verify()
                obj=slide.shapes.add_picture(io.BytesIO(raw),x,y,w,h)
                obj.rotation=rotation
            else:
                raise ValueError('type')
    out=io.BytesIO()
    p.save(out)
    return out.getvalue()
